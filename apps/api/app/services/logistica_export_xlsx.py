"""Excel do painel Logística (botão "Excel" das abas de marketplace).

Quem decide O QUE vai pra planilha é a tela: os filtros do painel (Mostrar
tudo, sub-aba DBA × Envio próprio da Amazon, conta, status Bling, período,
busca) vivem no front, e o texto de cada coluna também (Status Plataforma em
PT, "Destino: …" do Envio próprio, resumo do chamado…). A tela manda as linhas
já filtradas e formatadas, na ordem em que aparecem; aqui só escrevemos o .xlsx
— assim o Excel é sempre igual ao que o operador está vendo (Vinicius 06/10).
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_DIA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
# Nome de aba do Excel: até 31 caracteres e sem []:*?/\
_ABA_PROIBIDO = re.compile(r"[\[\]:*?/\\]")
_LARGURA_MIN = 8
_LARGURA_MAX = 60


def _dia(valor: str) -> date | None:
    m = _DIA_ISO.match(valor)
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def nome_aba(aba: str) -> str:
    nome = _ABA_PROIBIDO.sub(" ", aba or "").strip()[:31].strip()
    return nome or "Logística"


def montar_xlsx(
    aba: str,
    colunas: Sequence[tuple[str, str]],
    linhas: Sequence[Sequence[str | None]],
) -> bytes:
    """`colunas` = (título, tipo); tipo "data" recebe "YYYY-MM-DD" e vira data
    de verdade no Excel (ordena e filtra por período); o resto é texto. Número
    de pedido fica TEXTO de propósito: "2000018792949872" virava 2,00002E+15."""
    wb = Workbook()
    ws = wb.active
    ws.title = nome_aba(aba)

    ws.append([titulo for titulo, _ in colunas])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F5F9")
        cell.alignment = Alignment(vertical="center")
    larguras = [len(titulo) for titulo, _ in colunas]

    quebra = Alignment(wrap_text=True, vertical="top")
    topo = Alignment(vertical="top")
    for i, linha in enumerate(linhas, start=2):
        for j, (_, tipo) in enumerate(colunas, start=1):
            bruto = linha[j - 1] if j - 1 < len(linha) else None
            texto = ILLEGAL_CHARACTERS_RE.sub("", bruto or "").strip()
            if not texto:
                continue
            cell = ws.cell(row=i, column=j)
            dia = _dia(texto) if tipo == "data" else None
            if dia is not None:
                cell.value = dia
                cell.number_format = "DD/MM/YYYY"
                larguras[j - 1] = max(larguras[j - 1], 10)
            else:
                cell.value = texto
                # O openpyxl grava como FÓRMULA todo texto que começa com "=".
                cell.data_type = "s"
                larguras[j - 1] = max(larguras[j - 1], *(len(p) for p in texto.split("\n")))
            cell.alignment = quebra if "\n" in texto else topo

    for j, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, _LARGURA_MIN), _LARGURA_MAX)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(len(linhas) + 1, 1)}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
