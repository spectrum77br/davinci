"""Relatório do dia do robô de denúncia (Ouvidoria › Denúncia › Robô › Ocorrências).

Vinicius, 05/10/2026: "precisamos criar… um relatório no final do dia, mostrando quantos anúncios
ele achou, quantos denunciou na loja, quantos abriu reclamação na Anatel etc." — no DaVinci, dentro
das Ocorrências, dia do calendário (0h–24h) e "tem que sair em Excel".

Duas fontes:

- **números** — a cópia do banco do mini (`denuncia_*`, linha inteira em `dados`). As datas do mini
  são texto na hora de Brasília ("2026-10-05 02:19:23"; o SEI com "T"), então o dia é o prefixo
  "AAAA-MM-DD". Anúncio novo = `visto_primeiro`; denúncia = `criado_em`; Anatel =
  `sei_peticionado_em` (várias linhas do mesmo processo = uma loja); resposta = `resultado_em`
  (fora o "Aguardando" que a denúncia ganha ao nascer); saiu do ar = `saiu_em`; print = prova
  "Captura no ato"; conferido = `verificacoes.ts`. Depois da meia-noite as contas do dia são
  congeladas (`numeros`): resultado que muda depois não mexe no dia que passou.
- **anotações** — o DaVinci só guarda o resumo de agora do robô, então a cada resumo (60 s) o dia
  anota os passos que rodaram, as ocorrências que apareceram e os buracos sem notícia do mini.

Dia de antes de 05/10 não tem anotações: o relatório dele sai só com os números (abre sob demanda e
congela).
"""

from __future__ import annotations

import copy
from collections import Counter, defaultdict
from collections.abc import Iterable
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.denuncia import (
    DenunciaAnuncio,
    DenunciaDenuncia,
    DenunciaProva,
    DenunciaRelatorio,
    DenunciaVerificacao,
)
from app.services.denuncia_robo import FUSO, SEM_NOTICIA, _ordem_nome, _quando, montar_painel

SITES = ("Mercado Livre", "Shopee", "TikTok Shop", "Amazon")
GRUPO_NOME = {"GRUPO 1": "Nosso", "GRUPO 2": "Diversos"}
GRUPOS_FORA = ("DESCARTADO", "FORA DE ESCOPO")
# resposta da plataforma → coluna do resumo
RESPOSTA_TIPO = {
    "Anúncio removido": "removidos",
    "Procedente": "removidos",
    "Provável procedente": "removidos",
    "Improcedente": "recusados",
}
RESPOSTAS = ("removidos", "recusados", "sem_resposta", "outras")
_SITUACAO_NOME = {
    "Enviada": "enviados (sem leitura do andamento ainda)",
    "Recebida": "recebidos por uma unidade da Anatel",
    "Em tratamento": "na fiscalização",
    "Respondida — analisar": "Anatel respondeu / concluiu",
    "Exigência": "Anatel pede complemento",
}
# 05/10: status_anatel dos processos do SEI (o sistema do mini grava com a leitura do passo 1);
# "Enviada" = ainda sem leitura
SITUACOES_ANATEL = ("Enviada", "Recebida", "Em tratamento", "Respondida — analisar", "Exigência")
# dias pra trás que o fechamento confere (worker de hora em hora + a cada restart)
DIAS_FECHAR = 7


def _dia(v: Any) -> str:
    """ "2026-10-05 02:19:23" / "2026-10-05T14:02:10" → "2026-10-05"."""
    return str(v or "")[:10]


def _hora(v: Any) -> str:
    s = str(v or "")
    return s[11:16] if len(s) >= 16 else ""


def _grupo(g: Any) -> str:
    return GRUPO_NOME.get(str(g or ""), "Outros")


def _site(canal: Any) -> str:
    c = str(canal or "")
    return "Anatel" if c.startswith("Anatel") else c


def _sim(v: Any) -> bool:
    try:
        return int(v or 0) == 1
    except (TypeError, ValueError):
        return False


def _resposta_tipo(resultado: str) -> str:
    if resultado.startswith("Sem resposta"):
        return "sem_resposta"
    return RESPOSTA_TIPO.get(resultado, "outras")


def _anuncio_linha(a: DenunciaAnuncio | None, aid: str | None) -> dict:
    d = (a.dados if a else None) or {}
    return {
        "anuncio_id": aid or (a.id if a else None),
        "site": (a.marketplace if a else None) or d.get("marketplace") or "",
        "loja": (a.loja if a else None) or d.get("loja") or "",
        "titulo": (a.titulo if a else None) or d.get("titulo") or "",
        "grupo": _grupo(a.grupo if a else None),
        "vendas": int(a.vendas or 0) if a else 0,
        "url": d.get("url") or "",
    }


# ─────────────────────────────────────────────────────────────── números (cópia do banco)


async def calcular_numeros(session: AsyncSession, dia: date) -> dict:
    """As contas do dia, da cópia do banco. Listas inteiras: o Excel mostra linha por linha."""
    d = dia.isoformat()
    J = DenunciaAnuncio.dados  # noqa: N806

    # 1. anúncios novos (fora os de teste)
    novos = (
        (
            await session.execute(
                select(DenunciaAnuncio).where(J["visto_primeiro"].astext.like(f"{d}%"))
            )
        )
        .scalars()
        .all()
    )
    achou_lista, proprias, descartados = [], 0, 0
    achou_site: dict[str, Counter] = defaultdict(Counter)
    for a in novos:
        x = a.dados or {}
        if _sim(x.get("teste")):
            continue
        if a.propria:
            proprias += 1
            continue
        if str(a.grupo or "") in GRUPOS_FORA or _sim(x.get("fora_escopo")):
            descartados += 1
            continue
        linha = _anuncio_linha(a, a.id) | {"hora": _hora(a.visto_primeiro), "preco": x.get("preco")}
        achou_lista.append(linha)
        achou_site[linha["site"]][linha["grupo"]] += 1

    # 2. denúncias criadas no dia (lojas e Anatel), respostas registradas no dia
    D = DenunciaDenuncia.dados  # noqa: N806
    criadas = (
        (
            await session.execute(
                select(DenunciaDenuncia).where(
                    or_(
                        D["criado_em"].astext.like(f"{d}%"),
                        and_(D["criado_em"].astext.is_(None), DenunciaDenuncia.data == d),
                        D["sei_peticionado_em"].astext.like(f"{d}%"),
                    )
                )
            )
        )
        .scalars()
        .all()
    )
    respondidas = (
        (
            await session.execute(
                select(DenunciaDenuncia).where(
                    D["resultado_em"].astext.like(f"{d}%"),
                    DenunciaDenuncia.resultado.is_not(None),
                    DenunciaDenuncia.resultado != "Aguardando",
                )
            )
        )
        .scalars()
        .all()
    )
    saiu = (
        (
            await session.execute(
                select(DenunciaAnuncio).where(
                    J["saiu_em"].astext.like(f"{d}%"), DenunciaAnuncio.propria == 0
                )
            )
        )
        .scalars()
        .all()
    )

    ids = {x.anuncio_id for x in (*criadas, *respondidas) if x.anuncio_id}
    anuncios: dict[str, DenunciaAnuncio] = {}
    if ids:
        anuncios = {
            a.id: a
            for a in (
                await session.execute(select(DenunciaAnuncio).where(DenunciaAnuncio.id.in_(ids)))
            )
            .scalars()
            .all()
        }

    den_lista, de_novo, replicas, consumidor = [], 0, 0, 0
    den_site: dict[str, Counter] = defaultdict(Counter)
    processos: dict[str, dict] = {}
    for x in criadas:
        dd = x.dados or {}
        if _sim(dd.get("teste")):
            continue
        canal = str(x.canal or "")
        if canal == "Anatel SEI":
            if not _dia(dd.get("sei_peticionado_em")) == d:
                continue
            chave = dd.get("sei_processo") or f"sem-numero-{x.id}"
            p = processos.get(chave)
            if p is None:
                a = _anuncio_linha(anuncios.get(x.anuncio_id), x.anuncio_id)
                p = processos[chave] = {
                    "processo": dd.get("sei_processo") or "",
                    "recibo": dd.get("sei_recibo") or "",
                    "hora": _hora(dd.get("sei_peticionado_em")),
                    "loja": a["loja"],
                    "site": a["site"],
                    "grupo": a["grupo"],
                    "anuncios": 0,
                    "anuncio_ids": [],
                }
            p["anuncios"] += 1
            p["anuncio_ids"].append(x.anuncio_id)
            continue
        if _dia(dd.get("criado_em") or x.data) != d:
            continue
        if canal == "Anatel":
            consumidor += 1
            continue
        if canal not in SITES:
            continue
        a = _anuncio_linha(anuncios.get(x.anuncio_id), x.anuncio_id)
        tentativa = int(dd.get("tentativa") or 1)
        # 05/10: a Réplica Denúncias Diversos (outra empresa, citando o processo SEI) marca a obs
        replica = str(dd.get("obs") or "").startswith("RÉPLICA")
        replicas += replica
        de_novo += tentativa > 1 and not replica
        den_lista.append(
            a
            | {
                "site": canal,
                "hora": _hora(dd.get("criado_em")),
                "tentativa": tentativa,
                "protocolo": x.protocolo or "",
                "replica": str(dd.get("obs") or "").split(" · cita")[0] if replica else "",
            }
        )
        den_site[canal][a["grupo"]] += 1

    resp_lista = []
    resp_site: dict[str, Counter] = defaultdict(Counter)
    for x in respondidas:
        dd = x.dados or {}
        if _sim(dd.get("teste")):
            continue
        res = str(x.resultado or "")
        tipo = _resposta_tipo(res)
        site = _site(x.canal)
        a = _anuncio_linha(anuncios.get(x.anuncio_id), x.anuncio_id)
        resp_lista.append(
            a
            | {
                "site": site,
                "resultado": res,
                "tipo": tipo,
                "hora": _hora(dd.get("resultado_em")),
                "denunciado_em": _dia(dd.get("criado_em") or x.data),
            }
        )
        resp_site[site][tipo] += 1

    sairam = []
    for a in saiu:
        x = a.dados or {}
        if _sim(x.get("teste")):
            continue
        sairam.append(_anuncio_linha(a, a.id) | {"hora": _hora(x.get("saiu_em"))})

    # 3. prints e conferências
    P = DenunciaProva.dados  # noqa: N806
    provas = (
        await session.execute(
            select(DenunciaProva.anuncio_id, DenunciaProva.tipo).where(
                P["enviado_em"].astext.like(f"{d}%")
            )
        )
    ).all()
    capturas = [p for p in provas if (p.tipo or "") == "Captura no ato"]
    V = DenunciaVerificacao.dados  # noqa: N806
    verif = (
        await session.execute(
            select(DenunciaVerificacao.anuncio_id, V["situacao"].astext).where(
                V["ts"].astext.like(f"{d}%")
            )
        )
    ).all()
    fora = {v[0] for v in verif if v[1] and v[1] != "ativo"}

    # 4. (05/10) respostas da Anatel: o passo 1 lê o andamento dos processos no SEI e o sistema
    # do mini grava status_anatel (+ área e a data do andamento em status_anatel_em) nas
    # denúncias "Anatel SEI".
    # Situação de TODOS os processos (foto do dia) e os que a Anatel mexeu no dia.
    sei = (
        await session.execute(
            select(DenunciaDenuncia.anuncio_id, DenunciaDenuncia.dados).where(
                DenunciaDenuncia.canal == "Anatel SEI"
            )
        )
    ).all()
    por_proc: dict[str, dict] = {}
    for aid, dd in sei:
        dd = dd or {}
        proc = dd.get("sei_processo") or dd.get("protocolo")
        if not proc or _sim(dd.get("teste")):
            continue
        p = por_proc.setdefault(
            proc,
            {
                "processo": proc,
                "situacao": dd.get("status_anatel") or "Enviada",
                "area": dd.get("status_anatel_area") or "",
                "desde": dd.get("status_anatel_em") or "",
                "anuncio_id": aid,
            },
        )
        if (dd.get("status_anatel_em") or "") > p["desde"]:
            p.update(
                situacao=dd.get("status_anatel") or p["situacao"],
                area=dd.get("status_anatel_area") or p["area"],
                desde=dd["status_anatel_em"],
            )
    situacao_anatel = Counter(p["situacao"] for p in por_proc.values())
    movidos = [p for p in por_proc.values() if _dia(p["desde"]) == d]
    if movidos:
        ids_mov = {p["anuncio_id"] for p in movidos if p["anuncio_id"]} - set(anuncios)
        if ids_mov:
            anuncios.update(
                {
                    a.id: a
                    for a in (
                        await session.execute(
                            select(DenunciaAnuncio).where(DenunciaAnuncio.id.in_(ids_mov))
                        )
                    )
                    .scalars()
                    .all()
                }
            )
    movimentos = sorted(
        (
            {
                "processo": p["processo"],
                "situacao": p["situacao"],
                "area": p["area"],
                "loja": _anuncio_linha(anuncios.get(p["anuncio_id"]), p["anuncio_id"])["loja"],
                "site": _anuncio_linha(anuncios.get(p["anuncio_id"]), p["anuncio_id"])["site"],
            }
            for p in movidos
        ),
        key=lambda m: (m["situacao"], m["processo"]),
    )

    def _por_site(c: dict[str, Counter], chaves: Iterable[str]) -> dict:
        return {s: {k: int(c[s][k]) for k in chaves} for s in sorted(c)}

    return {
        "achou": {
            "total": len(achou_lista),
            "lojas_proprias": proprias,
            "descartados": descartados,
            "por_site": _por_site(achou_site, ("Nosso", "Diversos", "Outros")),
            "lista": sorted(
                achou_lista, key=lambda r: (r["site"], r["loja"], r["anuncio_id"] or "")
            ),
        },
        "denunciou": {
            "total": len(den_lista),
            "de_novo": de_novo,
            "replicas": replicas,
            "por_site": _por_site(den_site, ("Nosso", "Diversos", "Outros")),
            "lista": sorted(den_lista, key=lambda r: (r["hora"], r["site"])),
        },
        "anatel": {
            "lojas": len(processos),
            "anuncios": sum(p["anuncios"] for p in processos.values()),
            "consumidor": consumidor,
            "processos": sorted(processos.values(), key=lambda p: (p["hora"], p["processo"])),
            # 05/10: situação de todos os processos no SEI + os que a Anatel mexeu no dia
            "situacao": {k: int(situacao_anatel.get(k, 0)) for k in SITUACOES_ANATEL},
            "movimentos": movimentos,
        },
        "respostas": {
            "total": len(resp_lista),
            "por_site": _por_site(resp_site, RESPOSTAS),
            "lista": sorted(resp_lista, key=lambda r: (r["site"], r["tipo"], r["loja"])),
        },
        "sairam": {
            "total": len(sairam),
            "lista": sorted(sairam, key=lambda r: (r["site"], r["loja"])),
        },
        "prints": {
            "capturas": len(capturas),
            "anuncios": len({p.anuncio_id for p in capturas}),
            "registros": sum(1 for p in provas if (p.tipo or "") == "Registro da denúncia"),
        },
        "conferidos": {"anuncios": len({v[0] for v in verif}), "fora_do_ar": len(fora)},
    }


# ─────────────────────────────────────────────────────────────── anotações (resumo de 60 s)


def anotar(
    anot: dict, resumo: dict, painel: dict, agora: datetime, dia: date, so_passos: bool = False
) -> dict:
    """Junta no dia o que o resumo do robô mostra agora. `so_passos`: dia de ontem (passo que
    começou antes da meia-noite e só terminou depois)."""
    anot = copy.deepcopy(anot or {})
    agora = agora.astimezone(FUSO)
    passos = anot.setdefault("passos", {})
    for it in resumo.get("itens") or []:
        if not isinstance(it, dict) or not str(it.get("chave") or "").startswith("tarefa_"):
            continue
        x = it.get("dados") or {}
        comeco = _quando(x.get("inicio")) or _quando(x.get("pedido_em"))
        if (comeco.astimezone(FUSO).date() if comeco else agora.date()) != dia:
            continue
        acao = str(x.get("acao") or "")
        passos[str(it["chave"])] = {
            "acao": acao,
            "nome": _ordem_nome(acao, x.get("nome"))[1],
            "ordem": _ordem_nome(acao, x.get("nome"))[0],
            "status": x.get("status") or "fila",
            "feito": x.get("feito") or "",
            "pedido_em": x.get("pedido_em"),
            "inicio": x.get("inicio"),
            "fim": x.get("fim"),
            "erro": (x.get("erro") or "")[:500],
        }
    if so_passos:
        return anot

    ocs = anot.setdefault("ocorrencias", {})
    agora_iso = agora.isoformat(timespec="seconds")
    for o in painel.get("ocorrencias") or []:
        q = _quando(o.get("quando"))
        if o.get("origem") == "robo" and q and q.astimezone(FUSO).date() != dia:
            continue  # o robô guarda 24 h: a de ontem já foi anotada ontem
        ja = ocs.get(o["chave"])
        if ja:
            ja["ultima"] = agora_iso
            ja["detalhe"] = o.get("detalhe") or ja.get("detalhe") or ""
            continue
        ocs[o["chave"]] = {
            "titulo": o.get("titulo") or "",
            "detalhe": o.get("detalhe") or "",
            "o_que_fazer": o.get("o_que_fazer") or "",
            "tipo": o.get("tipo") or "aviso",
            "primeira": (q.astimezone(FUSO).isoformat(timespec="seconds") if q else agora_iso),
            "ultima": agora_iso,
        }

    ultimo = _quando(anot.get("ultimo_contato"))
    if ultimo and agora - ultimo > SEM_NOTICIA:
        anot.setdefault("sem_noticia", []).append(
            {"de": ultimo.astimezone(FUSO).isoformat(timespec="seconds"), "ate": agora_iso}
        )
    anot["ultimo_contato"] = agora_iso
    return anot


async def _linha(session: AsyncSession, dia: date) -> DenunciaRelatorio:
    """A linha do dia, criada se não existe (o worker e a tela podem chegar juntos)."""
    await session.execute(
        pg_insert(DenunciaRelatorio).values(dia=dia, anotacoes={}).on_conflict_do_nothing()
    )
    row = (
        await session.execute(
            select(DenunciaRelatorio).where(DenunciaRelatorio.dia == dia).with_for_update()
        )
    ).scalar_one()
    return row


async def anotar_resumo(
    session: AsyncSession, resumo: dict, agora: datetime, agenda: dict[str, dict] | None = None
) -> None:
    """Chamado a cada resumo do mini (sync/robo). Não commita."""
    painel = montar_painel(resumo, agora, agora, agenda=agenda)
    hoje = agora.astimezone(FUSO).date()
    ontem = hoje - timedelta(days=1)
    tem_ontem = any(
        (
            c := _quando(
                (it.get("dados") or {}).get("inicio") or (it.get("dados") or {}).get("pedido_em")
            )
        )
        and c.astimezone(FUSO).date() == ontem
        for it in resumo.get("itens") or []
        if isinstance(it, dict) and str(it.get("chave") or "").startswith("tarefa_")
    )
    if tem_ontem:
        row = await _linha(session, ontem)
        row.anotacoes = anotar(row.anotacoes, resumo, painel, agora, ontem, so_passos=True)
    row = await _linha(session, hoje)
    anot = row.anotacoes or {}
    if not anot.get("ultimo_contato"):
        # primeiro resumo do dia: o buraco sem notícia pode ter começado ontem (mini desligado)
        prev = await session.get(DenunciaRelatorio, ontem)
        if prev and (prev.anotacoes or {}).get("ultimo_contato"):
            anot = {**anot, "ultimo_contato": prev.anotacoes["ultimo_contato"]}
    row.anotacoes = anotar(anot, resumo, painel, agora, hoje)


# ─────────────────────────────────────────────────────────────── o relatório


def _passos(anot: dict) -> list[dict]:
    out = []
    for p in (anot.get("passos") or {}).values():
        status = p.get("feito") or p.get("status") or ""
        situ = {
            "concluido": "feito",
            "concluida": "feito",
            "cedeu": "parou pra outro passo",
            "erro": "erro",
            "rodando": "rodando",
            "fila": "na fila",
        }.get(status, status)
        inicio, fim = _quando(p.get("inicio")), _quando(p.get("fim"))
        erro = p.get("erro") or ""  # no passo que cedeu o mini põe só "cedeu"
        out.append(
            {
                "nome": p.get("nome") or p.get("acao") or "",
                "ordem": p.get("ordem"),
                "inicio": p.get("inicio"),
                "fim": p.get("fim"),
                "situacao": situ,
                "erro": "" if erro in ("cedeu", "concluido", "erro") else erro,
                "minutos": int((fim - inicio).total_seconds() // 60) if inicio and fim else None,
            }
        )
    return sorted(out, key=lambda p: str(p["inicio"] or ""))


def _ocorrencias(anot: dict) -> list[dict]:
    """Agrupa por título (o "perfil 50 fechava ao abrir" aparece dezenas de vezes num dia)."""
    grupos: dict[tuple[str, str], dict] = {}
    for o in (anot.get("ocorrencias") or {}).values():
        k = (o.get("titulo") or "", o.get("tipo") or "aviso")
        g = grupos.get(k)
        if g is None:
            grupos[k] = {
                "titulo": k[0],
                "tipo": k[1],
                "vezes": 1,
                "primeira": o.get("primeira"),
                "ultima": o.get("ultima"),
                "detalhe": o.get("detalhe") or "",
                "o_que_fazer": o.get("o_que_fazer") or "",
            }
            continue
        g["vezes"] += 1
        g["primeira"] = min(str(g["primeira"] or ""), str(o.get("primeira") or "")) or None
        if str(o.get("ultima") or "") >= str(g["ultima"] or ""):
            g["ultima"], g["detalhe"] = o.get("ultima"), o.get("detalhe") or g["detalhe"]
    return sorted(grupos.values(), key=lambda g: (g["tipo"] != "pessoa", str(g["primeira"] or "")))


def montar(dia: date, numeros: dict, row: DenunciaRelatorio | None, agora: datetime) -> dict:
    anot = (row.anotacoes if row else None) or {}
    return {
        "dia": dia.isoformat(),
        "parcial": dia >= agora.astimezone(FUSO).date(),
        "fechado_em": row.fechado_em if row else None,
        "lido_em": row.lido_em if row else None,
        "lido_por": row.lido_por if row else None,
        "numeros": numeros,
        "passos": _passos(anot),
        "ocorrencias": _ocorrencias(anot),
        "sem_noticia": anot.get("sem_noticia") or [],
        # sem nenhum resumo do mini no dia: os passos/ocorrências não foram anotados
        "anotado": bool(anot.get("ultimo_contato") or anot.get("passos")),
    }


def manchete(numeros: dict | None) -> dict:
    """Os quatro números da linha nas Ocorrências e na lista de relatórios."""
    n = numeros or {}
    resp = (n.get("respostas") or {}).get("por_site") or {}
    return {
        "achou": (n.get("achou") or {}).get("total", 0),
        "denunciou": (n.get("denunciou") or {}).get("total", 0),
        "anatel": (n.get("anatel") or {}).get("lojas", 0),
        "removidos": sum(int(v.get("removidos") or 0) for v in resp.values()),
    }


async def relatorio(session: AsyncSession, dia: date, agora: datetime) -> dict:
    """Hoje: contas ao vivo. Dia que passou: as congeladas (congela agora se ainda não)."""
    hoje = agora.astimezone(FUSO).date()
    if dia >= hoje:
        return montar(
            dia,
            await calcular_numeros(session, dia),
            await session.get(DenunciaRelatorio, dia),
            agora,
        )
    row = await session.get(DenunciaRelatorio, dia)
    if row is None or row.numeros is None:
        row = await fechar(session, dia, agora)
    return montar(dia, row.numeros or {}, row, agora)


async def fechar(session: AsyncSession, dia: date, agora: datetime) -> DenunciaRelatorio:
    row = await _linha(session, dia)
    if row.numeros is None:
        row.numeros = await calcular_numeros(session, dia)
        row.fechado_em = agora
    await session.flush()
    return row


async def fechar_pendentes(session: AsyncSession, agora: datetime) -> list[date]:
    """Worker (de hora em hora e a cada restart): congela ontem — mesmo sem notícia do mini, pra
    o relatório aparecer todo dia — e qualquer dia anotado dos últimos 7 que ficou aberto."""
    hoje = agora.astimezone(FUSO).date()
    ontem = hoje - timedelta(days=1)
    abertos = set(
        (
            await session.execute(
                select(DenunciaRelatorio.dia).where(
                    DenunciaRelatorio.numeros.is_(None),
                    DenunciaRelatorio.dia < hoje,
                    DenunciaRelatorio.dia >= hoje - timedelta(days=DIAS_FECHAR),
                )
            )
        )
        .scalars()
        .all()
    )
    if (await session.get(DenunciaRelatorio, ontem)) is None:
        abertos.add(ontem)
    feitos = []
    for dia in sorted(abertos):
        await fechar(session, dia, agora)
        feitos.append(dia)
    return feitos


async def pendentes(session: AsyncSession, agora: datetime) -> list[dict]:
    """Relatórios fechados e ainda não lidos (últimos 7 dias) — viram linha nas Ocorrências."""
    hoje = agora.astimezone(FUSO).date()
    rows = (
        (
            await session.execute(
                select(DenunciaRelatorio)
                .where(
                    DenunciaRelatorio.fechado_em.is_not(None),
                    DenunciaRelatorio.lido_em.is_(None),
                    DenunciaRelatorio.dia >= hoje - timedelta(days=DIAS_FECHAR),
                )
                .order_by(DenunciaRelatorio.dia.desc())
            )
        )
        .scalars()
        .all()
    )
    return [
        {"dia": r.dia.isoformat(), "fechado_em": r.fechado_em, **manchete(r.numeros)} for r in rows
    ]


# ─────────────────────────────────────────────────────────────── Excel

_FONTE = "Arial"
_TITULO = Font(name=_FONTE, bold=True, size=13)
_CAB_FILL = PatternFill("solid", fgColor="1F3864")
_CAB_FONT = Font(name=_FONTE, bold=True, color="FFFFFF", size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
_NORMAL = Font(name=_FONTE, size=10)
_NEGRITO = Font(name=_FONTE, size=10, bold=True)
_fino = Side(style="thin", color="BFBFBF")
_BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
_RESP_NOME = {
    "removidos": "removidos / procedentes",
    "recusados": "recusados (improcedente)",
    "sem_resposta": "sem resposta (prazo vencido)",
    "outras": "outras respostas",
}


def _br(d: str) -> str:
    return f"{d[8:10]}/{d[5:7]}/{d[:4]}" if len(d) >= 10 else d


def _hm(v: Any) -> str:
    q = _quando(v)
    return q.astimezone(FUSO).strftime("%d/%m %H:%M") if q else ""


def _tabela(
    ws,
    linha: int,
    cabecalho: list[str],
    linhas: Iterable[list],
    larguras: list[int] | None = None,
    total: list | None = None,
) -> int:
    for j, h in enumerate(cabecalho, 1):
        c = ws.cell(linha, j, h)
        c.font, c.fill, c.border = _CAB_FONT, _CAB_FILL, _BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n = linha
    for n, valores in enumerate(linhas, linha + 1):
        for j, v in enumerate(valores, 1):
            c = ws.cell(n, j, v)
            c.font, c.border = _NORMAL, _BORDA
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    if total is not None:
        n += 1
        for j, v in enumerate(total, 1):
            c = ws.cell(n, j, v)
            c.font, c.fill, c.border = _NEGRITO, _TOTAL_FILL, _BORDA
    for j, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if linha == 1:
        ws.freeze_panes = ws.cell(2, 1)
    return n + 2


def excel(rel: dict) -> BytesIO:
    n = rel["numeros"] or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    dia = _br(rel["dia"])
    ws.cell(
        1,
        1,
        f"Robô de Denúncia — relatório de {dia}"
        + (" (parcial, dia em andamento)" if rel["parcial"] else ""),
    ).font = _TITULO
    ws.cell(
        2,
        1,
        "Dia do calendário (0h–24h, Brasília). Anúncios novos sem lojas próprias, capas e peças.",
    ).font = _NORMAL

    sites = [*SITES]
    ach = (n.get("achou") or {}).get("por_site") or {}
    den = (n.get("denunciou") or {}).get("por_site") or {}
    resp = (n.get("respostas") or {}).get("por_site") or {}
    linhas = []
    for titulo, fonte, chave in (
        ("Anúncios novos — Nosso", ach, "Nosso"),
        ("Anúncios novos — Diversos", ach, "Diversos"),
        ("Anúncios novos — outros", ach, "Outros"),
        ("Denúncias nas lojas — Nosso", den, "Nosso"),
        ("Denúncias nas lojas — Diversos", den, "Diversos"),
        ("Denúncias nas lojas — outros", den, "Outros"),
        *((f"Respostas: {_RESP_NOME[k]}", resp, k) for k in RESPOSTAS),
    ):
        valores = [int((fonte.get(s) or {}).get(chave) or 0) for s in sites]
        if titulo.endswith("outros") and not sum(valores):
            continue
        extra = int((resp.get("Anatel") or {}).get(chave) or 0) if fonte is resp else ""
        linhas.append([titulo, *valores, extra, sum(valores) + (extra or 0)])
    linha = _tabela(
        ws,
        4,
        ["", "Mercado Livre", "Shopee", "TikTok", "Amazon", "Anatel", "Total"],
        linhas,
        [38, 14, 12, 12, 12, 12, 12],
    )

    an = n.get("anatel") or {}
    pr = n.get("prints") or {}
    co = n.get("conferidos") or {}
    soltos = [
        ["Anatel (SEI): lojas peticionadas", an.get("lojas", 0)],
        ["Anatel (SEI): anúncios nas petições", an.get("anuncios", 0)],
        ["Anatel Consumidor: reclamações", an.get("consumidor", 0)],
        ["Réplicas Denúncias Diversos (outra empresa, citando o processo SEI)",
         (n.get("denunciou") or {}).get("replicas", 0)],
        *(
            [f"Anatel (SEI): processos — {_SITUACAO_NOME[k]}", (an.get("situacao") or {}).get(k, 0)]
            for k in SITUACOES_ANATEL
            if an.get("situacao")
        ),
        ["Anatel (SEI): processos que a Anatel mexeu no dia", len(an.get("movimentos") or [])],
        ["Anúncios que saíram do ar", (n.get("sairam") or {}).get("total", 0)],
        ["Prints de anúncio (capturas)", pr.get("capturas", 0)],
        ["Anúncios conferidos (ativo/inativo)", co.get("anuncios", 0)],
        ["  …desses, fora do ar", co.get("fora_do_ar", 0)],
        [
            "Anúncios novos de lojas próprias (fora da conta)",
            (n.get("achou") or {}).get("lojas_proprias", 0),
        ],
        ["Anúncios novos descartados (capa, peça…)", (n.get("achou") or {}).get("descartados", 0)],
    ]
    linha = _tabela(ws, linha, ["Outros números", "Quantidade"], soltos)
    if rel["sem_noticia"]:
        _tabela(
            ws,
            linha,
            ["Mac mini sem notícia", "De", "Até"],
            [["sem notícia", _hm(b["de"]), _hm(b["ate"])] for b in rel["sem_noticia"]],
        )

    def aba(nome: str, cab: list[str], rows: list[list], larg: list[int]) -> None:
        w = wb.create_sheet(nome)
        _tabela(w, 1, cab, rows or [["(nada no dia)"]], larg)
        w.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{max(len(rows), 1) + 1}"

    aba(
        "Anúncios novos",
        ["Hora", "Site", "Loja", "Anúncio", "Título", "Certificado", "Vendas", "Preço", "Link"],
        [
            [
                r["hora"],
                r["site"],
                r["loja"],
                r["anuncio_id"],
                r["titulo"],
                r["grupo"],
                r["vendas"],
                r.get("preco"),
                r["url"],
            ]
            for r in (n.get("achou") or {}).get("lista") or []
        ],
        [7, 14, 24, 18, 60, 12, 9, 10, 40],
    )
    aba(
        "Denúncias nas lojas",
        ["Hora", "Site", "Loja", "Anúncio", "Título", "Certificado", "Vez", "Protocolo", "Réplica"],
        [
            [
                r["hora"],
                r["site"],
                r["loja"],
                r["anuncio_id"],
                r["titulo"],
                r["grupo"],
                r["tentativa"],
                r["protocolo"],
                r.get("replica") or "",
            ]
            for r in (n.get("denunciou") or {}).get("lista") or []
        ],
        [7, 14, 24, 18, 60, 12, 6, 18, 48],
    )
    aba(
        "Anatel",
        ["Hora", "Processo SEI", "Loja", "Site", "Certificado", "Anúncios", "Ids dos anúncios"],
        [
            [
                p["hora"],
                p["processo"],
                p["loja"],
                p["site"],
                p["grupo"],
                p["anuncios"],
                ", ".join(p["anuncio_ids"]),
            ]
            for p in an.get("processos") or []
        ],
        [7, 24, 26, 14, 12, 9, 50],
    )
    # 05/10: o que a Anatel fez nos processos (lido no SEI pelo passo 1)
    w_an = wb["Anatel"]
    _tabela(
        w_an,
        w_an.max_row + 3,
        ["Processo SEI", "Situação na Anatel", "Unidade", "Loja", "Site"],
        [
            [
                m["processo"],
                _SITUACAO_NOME.get(m["situacao"], m["situacao"]),
                m["area"],
                m["loja"],
                m["site"],
            ]
            for m in an.get("movimentos") or []
        ]
        or [["(a Anatel não mexeu em nenhum processo no dia)"]],
    )
    aba(
        "Respostas",
        ["Hora", "Site", "Resposta", "Loja", "Anúncio", "Título", "Certificado", "Denunciado em"],
        [
            [
                r["hora"],
                r["site"],
                r["resultado"],
                r["loja"],
                r["anuncio_id"],
                r["titulo"],
                r["grupo"],
                _br(r["denunciado_em"]),
            ]
            for r in (n.get("respostas") or {}).get("lista") or []
        ],
        [7, 14, 28, 24, 18, 60, 12, 13],
    )
    aba(
        "Saíram do ar",
        ["Hora", "Site", "Loja", "Anúncio", "Título", "Certificado", "Vendas"],
        [
            [r["hora"], r["site"], r["loja"], r["anuncio_id"], r["titulo"], r["grupo"], r["vendas"]]
            for r in (n.get("sairam") or {}).get("lista") or []
        ],
        [7, 14, 24, 18, 60, 12, 9],
    )
    w = wb.create_sheet("Robô")
    linha = _tabela(
        w,
        1,
        ["Início", "Fim", "Minutos", "Passo", "Como terminou", "Erro"],
        [
            [_hm(p["inicio"]), _hm(p["fim"]), p["minutos"], p["nome"], p["situacao"], p["erro"]]
            for p in rel["passos"]
        ]
        or [["(sem passos anotados)"]],
        [13, 13, 9, 40, 22, 70],
    )
    _tabela(
        w,
        linha,
        ["Primeira vez", "Última vez", "Vezes", "Ocorrência", "Tipo", "Detalhe"],
        [
            [
                _hm(o["primeira"]),
                _hm(o["ultima"]),
                o["vezes"],
                o["titulo"],
                "precisa de pessoa" if o["tipo"] == "pessoa" else "aviso",
                o["detalhe"],
            ]
            for o in rel["ocorrencias"]
        ]
        or [["(nenhuma ocorrência anotada)"]],
    )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
