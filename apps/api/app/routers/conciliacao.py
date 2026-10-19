"""Conciliação — passo 1 (14/10/2026): sondagem só-leitura das fontes do
Mercado Livre / Mercado Pago pra uma loja.

`GET /api/conciliacao/sonda?conta=jlas2` responde, numa chamada só e sem
gravar nada, se a chave de acesso que o DaVinci já tem da loja:

* abre a **porta do extrato** no Mercado Pago (`/v1/account/release_report`
  — o "Relatório de liberações"), e o que já existe lá (configuração,
  arquivos gerados, uma amostra das linhas do último arquivo pronto);
* lê a **fatura do ML por período** (`/billing/integration/periods/key/…`:
  períodos, resumo, detalhes do mês atual e do anterior, pagamentos da
  fatura, movimentações do grupo MP).

Cada chamada volta com status, tempo e uma amostra recortada do corpo, pra
decidir por onde o extrato entra no painel (API ou importação do CSV) antes
de desenhar tabela. Avaliação completa em `docs/conciliacao-avaliacao.md`.

Só GET. O token nunca aparece na resposta. A renovação do token, se o ML
pedir, passa pelo caminho blindado do Atendimento (trava + commit próprio),
porque o refresh token do ML é de uso único.
"""

from __future__ import annotations

import asyncio
import gzip
import math
import re
import time
from collections import Counter, defaultdict
from collections.abc import Awaitable, Callable
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any
from uuid import UUID

import httpx
import structlog
from fastapi import (
    APIRouter,
    Body,
    Depends,
    File,
    Form,
    Header,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.datastructures import UploadFile as UploadFileBruto

from app.db import get_session
from app.deps.auth import require_permission_any
from app.models import Integration, User
from app.services import logistica_meli as lm
from app.services.atendimento.clientes import cliente_da_integracao
from app.services.conciliacao_extrato import (
    MP_API_BASE,
    baixar_relatorio,
    garantir_config,
    listar_relatorios,
    mp_get,
    pedir_relatorio,
)
from app.services.conciliacao_jobs import (
    RodadaEmAndamentoError,
    loja_por_conta,
    lojas_ativas,
    trava_da_loja,
)
from app.services.conciliacao_mp_csv import ler_csv_mp
from app.worker_pool import get_arq_pool

logger = structlog.get_logger()
router = APIRouter(prefix="/api/conciliacao", tags=["conciliacao"])

_PERM = require_permission_any(("margem", "logistica"), "view")
# 16/10, revisão: toda escrita da Conciliação é permissão de Conciliação (ou Margem) em
# editar — quem só tem Logística não mexe na config do relatório do MP pela sonda, e
# quem só tem Conciliação não toma 403 no "Sincronizar"
_PERM_PAINEL = require_permission_any(("conciliacao", "margem"), "view")
_PERM_PAINEL_EDIT = require_permission_any(("conciliacao", "margem"), "edit")
_AMOSTRA = 5
_LINHAS_CSV = 12
_MAX_TEXTO = 1500


def _chave(d: date) -> str:
    """Chave do período de faturamento do ML: sempre o primeiro dia do mês."""
    return d.strftime("%Y-%m-01")


def _mes_anterior(d: date) -> date:
    primeiro = d.replace(day=1)
    if primeiro.month == 1:
        return primeiro.replace(year=primeiro.year - 1, month=12)
    return primeiro.replace(month=primeiro.month - 1)


def _recorta(obj: Any) -> Any:
    """Amostra do corpo: listas ficam com os primeiros itens, textos encurtam."""
    if isinstance(obj, list):
        return {"total_itens": len(obj), "primeiros": [_recorta(x) for x in obj[:_AMOSTRA]]}
    if isinstance(obj, dict):
        out: dict[str, Any] = {}
        for k, v in obj.items():
            if isinstance(v, list) and len(v) > _AMOSTRA:
                out[k] = {"total_itens": len(v), "primeiros": [_recorta(x) for x in v[:_AMOSTRA]]}
            else:
                out[k] = _recorta(v)
        return out
    if isinstance(obj, str) and len(obj) > _MAX_TEXTO:
        return obj[:_MAX_TEXTO] + "…"
    return obj


def _amostra_csv(texto: str) -> dict[str, Any]:
    """Cabeçalho + primeiras linhas de um CSV do Mercado Pago (separador ; ou ,)."""
    cab, rows, sep = ler_csv_mp(texto)
    return {
        "total_linhas": len(rows),
        "separador": sep,
        "colunas": cab,
        "primeiras": rows[:_LINHAS_CSV],
    }


async def _chamada(
    nome: str, url: str, fn: Callable[[], Awaitable[httpx.Response]], *, csv_texto: bool = False
) -> dict[str, Any]:
    t0 = time.monotonic()
    item: dict[str, Any] = {"nome": nome, "url": url}
    try:
        r = await fn()
    except Exception as e:  # noqa: BLE001 — a sondagem registra e segue
        item.update(
            ok=False,
            erro=f"{type(e).__name__}: {str(e)[:300]}",
            ms=int((time.monotonic() - t0) * 1000),
        )
        return item
    item.update(
        status=r.status_code, ok=r.status_code == 200, ms=int((time.monotonic() - t0) * 1000)
    )
    ctype = r.headers.get("content-type", "")
    if csv_texto and r.status_code == 200 and "json" not in ctype:
        item["amostra"] = _amostra_csv(r.text)
        return item
    try:
        item["amostra"] = _recorta(r.json())
    except ValueError:
        item["amostra"] = (r.text or "")[:_MAX_TEXTO]
    return item


@router.get("/sonda")
async def sonda(
    conta: Annotated[
        str, Query(min_length=1, description="nome da conta ML na tela Lojas, ex. jlas2")
    ],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    mes: Annotated[
        str | None, Query(description="AAAA-MM do período a ler; padrão = mês atual")
    ] = None,
    parte: Annotated[
        int | None,
        Query(
            ge=1,
            le=3,
            description="1 = chave + fatura do mês atual; 2 = mês anterior, bonificações, "
            "grupo MP, pagamentos da fatura; 3 = extrato do Mercado Pago. Sem valor = tudo. "
            "A API de faturamento do ML aceita 5 chamadas por minuto (14/10), daí as partes.",
        ),
    ] = None,
) -> dict[str, Any]:
    integ = await lm._ml_integration_for_conta(session, conta)
    if integ is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "conta_ml_nao_encontrada", "conta": conta}
        )
    cliente = await cliente_da_integracao(integ)

    # 16/10, revisão: mês inválido (2026-13) é 422, não 500; o padrão é o mês em BRT
    atual, _fim = _mes_range(mes)
    anterior = _mes_anterior(atual)
    k_atual, k_ant = _chave(atual), _chave(anterior)

    chamadas: list[dict[str, Any]] = []

    async def ml(nome: str, path: str, params: dict | None = None) -> dict[str, Any]:
        item = await _chamada(nome, path, lambda: cliente._request("GET", path, params=params))
        chamadas.append(item)
        return item

    async def mp(
        nome: str, path: str, params: dict | None = None, *, csv_texto: bool = False
    ) -> dict[str, Any]:
        item = await _chamada(
            nome,
            f"{MP_API_BASE}{path}",
            lambda: mp_get(cliente, path, params),
            csv_texto=csv_texto,
        )
        chamadas.append(item)
        return item

    vazio: dict[str, Any] = {}
    eu = det_atual = det_ant = cfg = lista = vazio
    arquivo: str | None = None
    pagina = {
        "document_type": "BILL",
        "limit": _AMOSTRA,
        "from_id": 0,
        "sort_by": "ID",
        "order_by": "ASC",
    }

    if parte in (None, 1):
        # 1. A chave funciona? (quem é a loja) + fatura do ML do mês atual
        eu = await ml("ml_quem_sou", "/users/me")
        await ml(
            "ml_periodos",
            "/billing/integration/monthly/periods",
            {"group": "ML", "document_type": "BILL", "limit": 3},
        )
        await ml(
            "ml_resumo_mes_atual",
            f"/billing/integration/periods/key/{k_atual}/summary/details",
            {"group": "ML", "document_type": "BILL"},
        )
        det_atual = await ml(
            "ml_detalhes_mes_atual",
            f"/billing/integration/periods/key/{k_atual}/group/ML/details",
            pagina,
        )

    if parte in (None, 2):
        # 2. Mês anterior (fechado), bonificações, grupo MP e pagamentos da fatura
        det_ant = await ml(
            "ml_detalhes_mes_anterior",
            f"/billing/integration/periods/key/{k_ant}/group/ML/details",
            pagina,
        )
        await ml(
            "ml_bonificacoes_mes_anterior",
            f"/billing/integration/periods/key/{k_ant}/group/ML/details",
            {**pagina, "detail_type": "BONUS"},
        )
        await ml(
            "mp_grupo_detalhes_mes_atual",
            f"/billing/integration/periods/key/{k_atual}/group/MP/details",
            pagina,
        )
        await ml(
            "ml_pagamentos_da_fatura_mes_anterior",
            f"/billing/integration/periods/key/{k_ant}/group/ML/payment/details",
            {"limit": 3},
        )

    if parte in (None, 3):
        # 3. Porta do extrato no Mercado Pago (a dúvida que define o desenho)
        cfg = await mp("mp_extrato_config", "/v1/account/release_report/config")
        lista = await mp("mp_extrato_lista", "/v1/account/release_report/list")
        amostra = lista.get("amostra")
        itens = (
            amostra.get("primeiros")
            if isinstance(amostra, dict)
            else (amostra if isinstance(amostra, list) else None)
        )
        if lista.get("ok") and isinstance(itens, list):
            prontos = [x for x in itens if isinstance(x, dict) and x.get("file_name")]
            if prontos:
                arquivo = str(prontos[0]["file_name"])
        if arquivo:
            await mp(
                "mp_extrato_arquivo_mais_recente",
                f"/v1/account/release_report/{arquivo}",
                csv_texto=True,
            )
        await mp("mp_todas_transacoes_lista", "/v1/account/settlement_report/list")

    def _tem_resultados(item: dict[str, Any]) -> bool:
        a = item.get("amostra")
        if not isinstance(a, dict):
            return False
        res = a.get("results")
        if isinstance(res, dict):
            return bool(res.get("total_itens"))
        return bool(res)

    conclusao = {
        "chave_da_loja_funciona": bool(eu.get("ok")),
        "extrato_mp_abre_com_a_chave_do_ml": bool(lista.get("ok")),
        "extrato_mp_tem_config": bool(cfg.get("ok")),
        "extrato_mp_tem_arquivo_pronto": bool(arquivo),
        "fatura_mes_atual_tem_detalhes": _tem_resultados(det_atual),
        "fatura_mes_anterior_tem_detalhes": _tem_resultados(det_ant),
    }
    # O log da API é o único canal que dá pra ler em produção sem login: além
    # das conclusões, deixa o status de cada chamada e as colunas do extrato.
    por_nome = {c["nome"]: c.get("status") or c.get("erro") for c in chamadas}
    arquivo_amostra = next(
        (c.get("amostra") for c in chamadas if c["nome"] == "mp_extrato_arquivo_mais_recente"),
        None,
    )
    logger.info(
        "conciliacao_sonda",
        conta=conta,
        integration_id=str(integ.id),
        chamadas=por_nome,
        colunas_extrato=(arquivo_amostra or {}).get("colunas")
        if isinstance(arquivo_amostra, dict)
        else None,
        **conclusao,
    )
    return {
        "conta": conta,
        "integration_id": str(integ.id),
        "periodos": {"atual": k_atual, "anterior": k_ant},
        "quando": datetime.now(UTC).isoformat(),
        "conclusao": conclusao,
        "chamadas": chamadas,
    }


# --- Extrato do Mercado Pago (Relatório de liberações) ----------------------
# 14/10/2026, com OK do Vinicius: duas escritas de configuração na conta da loja
# (criar a configuração do relatório e pedir um período) e a leitura do arquivo.
# A rotina diária de verdade está em `services/conciliacao_extrato.py`; estes
# endpoints são a sondagem manual e o resumo de um arquivo pronto.


def resumir_extrato(texto: str) -> dict[str, Any]:
    """Resumo de um CSV do Relatório de liberações: tipos de movimento, somas,
    quantas linhas têm número de pedido, e uma amostra por tipo. Não grava."""
    from app.services.conciliacao_mp_csv import dinheiro

    cab, rows, sep = ler_csv_mp(texto)
    if not cab:
        return {"total_linhas": 0}
    por_tipo: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "linhas": 0,
            "credito": Decimal("0"),
            "debito": Decimal("0"),
            "com_pedido": 0,
            "sem_pedido": 0,
        }
    )
    registros: Counter[str] = Counter()
    amostra: dict[str, list[dict[str, Any]]] = defaultdict(list)
    com_pedido = sem_pedido = 0
    campos = [
        c
        for c in (
            "DATE",
            "SOURCE_ID",
            "EXTERNAL_REFERENCE",
            "RECORD_TYPE",
            "DESCRIPTION",
            "GROSS_AMOUNT",
            "MP_FEE_AMOUNT",
            "FINANCING_FEE_AMOUNT",
            "SHIPPING_FEE_AMOUNT",
            "COUPON_AMOUNT",
            "NET_CREDIT_AMOUNT",
            "NET_DEBIT_AMOUNT",
            "BALANCE_AMOUNT",
            "ORDER_ID",
            "PACK_ID",
            "SHIPPING_ID",
            "INSTALLMENTS",
            "PAYMENT_METHOD",
            "PRODUCT_SKU",
            "METADATA",
        )
        if c in cab
    ]
    for r in rows:
        tipo = (r.get("DESCRIPTION") or "").strip()
        registros[(r.get("RECORD_TYPE") or "").strip()] += 1
        t = por_tipo[tipo]
        t["linhas"] += 1
        t["credito"] += dinheiro(r.get("NET_CREDIT_AMOUNT"))
        t["debito"] += dinheiro(r.get("NET_DEBIT_AMOUNT"))
        tem_pedido = bool((r.get("ORDER_ID") or "").strip() or (r.get("PACK_ID") or "").strip())
        t["com_pedido" if tem_pedido else "sem_pedido"] += 1
        if tem_pedido:
            com_pedido += 1
        else:
            sem_pedido += 1
        if len(amostra[tipo]) < 3:
            amostra[tipo].append({c: (r.get(c) or "")[:120] for c in campos})
    return {
        "total_linhas": len(rows),
        "separador": sep,
        "colunas": cab,
        "registros": dict(registros),
        "por_tipo": {
            k: {**v, "credito": str(v["credito"]), "debito": str(v["debito"])}
            for k, v in por_tipo.items()
        },
        "com_pedido": com_pedido,
        "sem_pedido": sem_pedido,
        "credito_total": str(sum((v["credito"] for v in por_tipo.values()), Decimal("0"))),
        "debito_total": str(sum((v["debito"] for v in por_tipo.values()), Decimal("0"))),
        "saldo_primeira_linha": rows[0].get("BALANCE_AMOUNT") if rows else None,
        "saldo_ultima_linha": rows[-1].get("BALANCE_AMOUNT") if rows else None,
        "amostra_por_tipo": dict(amostra),
    }


async def _cliente_da_conta(session: AsyncSession, conta: str) -> Any:
    integ = await lm._ml_integration_for_conta(session, conta)
    if integ is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "conta_ml_nao_encontrada", "conta": conta}
        )
    return await cliente_da_integracao(integ)


@router.get("/sonda/ml-conta")
async def sonda_ml_conta(
    conta: Annotated[str, Query(min_length=1)],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
) -> dict[str, Any]:
    """Só leitura (16/10, lucas mei 403 PolicyAgent na fatura): com quem o token do ML
    foi emitido (`/users/me`, sem dados pessoais) e o que a fatura responde."""
    from app.services.conciliacao_fatura import ler_resumo

    cliente = await _cliente_da_conta(session, conta)
    saida: dict[str, Any] = {"conta": conta}
    r = await cliente._request("GET", "/users/me")
    try:
        corpo = r.json()
    except ValueError:
        corpo = {"texto": (r.text or "")[:300]}
    if isinstance(corpo, dict):
        corpo = {
            k: corpo.get(k)
            for k in (
                "id",
                "nickname",
                "user_type",
                "tags",
                "site_status",
                "seller_experience",
                "registration_date",
                "points",
            )
        }
    saida["usuario"] = {"http": r.status_code, **({"dados": corpo})}
    try:
        resumo = await ler_resumo(cliente, "2026-10-01")
        saida["fatura_outubro"] = {"ok": True, "chaves": sorted(resumo)[:12]}
    except Exception as e:  # noqa: BLE001
        saida["fatura_outubro"] = {"ok": False, "erro": str(e)[:400]}
    return saida


@router.post("/sonda/extrato/pedir")
async def sonda_extrato_pedir(
    conta: Annotated[str, Query(min_length=1)],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
    de: Annotated[date, Query(description="primeiro dia (BRT)")],
    ate: Annotated[date, Query(description="último dia (BRT), inclusive")],
    _body: Annotated[dict | None, Body()] = None,
) -> dict[str, Any]:
    """Cria a configuração do Relatório de liberações se não existir e pede o
    relatório do período. Nada de dinheiro, nada de pedido, nada gravado no DaVinci."""
    cliente = await _cliente_da_conta(session, conta)
    config = await garantir_config(cliente)
    pedido = await pedir_relatorio(cliente, de, ate)
    logger.info(
        "conciliacao_sonda_extrato_pedido",
        conta=conta,
        config=config.get("status"),
        http=pedido["status"],
    )
    return {"conta": conta, "config": config, "periodo": pedido["janela"], "pedido": pedido}


@router.get("/sonda/extrato")
async def sonda_extrato_ler(
    conta: Annotated[str, Query(min_length=1)],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    bruto: Annotated[
        int, Query(ge=0, le=20, description="devolve também as N primeiras linhas cruas do CSV")
    ] = 0,
) -> dict[str, Any]:
    """Lista os relatórios da conta e, se o mais novo estiver pronto, baixa o
    CSV e devolve um resumo. Só leitura."""
    cliente = await _cliente_da_conta(session, conta)
    relatorios = await listar_relatorios(cliente)
    pronto = next((x for x in relatorios if x.get("file_name")), None)
    resumo: dict[str, Any] | None = None
    if pronto:
        texto = await baixar_relatorio(cliente, str(pronto["file_name"]))
        if texto is not None:
            resumo = resumir_extrato(texto)
            if bruto:
                linhas = texto.splitlines()
                cheia = next((x for x in linhas[1:] if '"{' in x or '"[' in x), None)
                resumo["linhas_cruas"] = linhas[: bruto + 1] + ([cheia] if cheia else [])
    logger.info(
        "conciliacao_sonda_extrato_lido",
        conta=conta,
        relatorios=len(relatorios),
        pronto=bool(pronto),
        total_linhas=(resumo or {}).get("total_linhas"),
    )
    return {"conta": conta, "relatorios": relatorios, "pronto": pronto, "resumo": resumo}


# --- Painel: status das lojas ligadas e "sincronizar agora" -------------------


@router.get("/status")
async def status_lojas(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
) -> dict[str, Any]:
    """Por loja ligada: últimas leituras, movimentos por origem/tipo e pedidos por
    situação e etapa. É o painel "cru" da fase 2a — a tela vem na 2b."""
    from sqlalchemy import func

    from app.models import ConciliacaoMovimento, ConciliacaoPedido, ConciliacaoSync

    saida: list[dict[str, Any]] = []
    for loja in await lojas_ativas(session):
        por_tipo = (
            await session.execute(
                select(
                    ConciliacaoMovimento.origem,
                    ConciliacaoMovimento.tipo,
                    func.count(),
                    func.sum(ConciliacaoMovimento.credito),
                    func.sum(ConciliacaoMovimento.debito),
                )
                .where(ConciliacaoMovimento.integration_id == loja.integration_id)
                .group_by(ConciliacaoMovimento.origem, ConciliacaoMovimento.tipo)
                .order_by(ConciliacaoMovimento.origem, ConciliacaoMovimento.tipo)
            )
        ).all()
        por_situacao = (
            await session.execute(
                select(ConciliacaoPedido.situacao, func.count())
                .where(ConciliacaoPedido.integration_id == loja.integration_id)
                .group_by(ConciliacaoPedido.situacao)
            )
        ).all()
        por_etapa = (
            await session.execute(
                select(ConciliacaoPedido.etapa, func.count())
                .where(ConciliacaoPedido.integration_id == loja.integration_id)
                .group_by(ConciliacaoPedido.etapa)
            )
        ).all()
        syncs = (
            (
                await session.execute(
                    select(ConciliacaoSync)
                    .where(ConciliacaoSync.integration_id == loja.integration_id)
                    .order_by(ConciliacaoSync.origem, ConciliacaoSync.chave)
                )
            )
            .scalars()
            .all()
        )
        sub_tipos = (
            await session.execute(
                select(
                    ConciliacaoMovimento.sub_tipo,
                    ConciliacaoMovimento.tipo,
                    func.count(),
                    func.sum(ConciliacaoMovimento.debito),
                    func.sum(ConciliacaoMovimento.credito),
                )
                .where(
                    ConciliacaoMovimento.integration_id == loja.integration_id,
                    ConciliacaoMovimento.origem == "ml_fatura",
                )
                .group_by(ConciliacaoMovimento.sub_tipo, ConciliacaoMovimento.tipo)
                .order_by(func.count().desc())
            )
        ).all()
        sem_dono_tipos = (
            await session.execute(
                select(ConciliacaoMovimento.origem, ConciliacaoMovimento.tipo, func.count())
                .where(
                    ConciliacaoMovimento.integration_id == loja.integration_id,
                    ConciliacaoMovimento.origem == "mp_extrato",
                    ConciliacaoMovimento.pedido_bling.is_(None),
                    ConciliacaoMovimento.classificacao.is_(None),
                )
                .group_by(ConciliacaoMovimento.origem, ConciliacaoMovimento.tipo)
            )
        ).all()
        sem_pedido = (
            await session.execute(
                select(func.count()).where(
                    ConciliacaoMovimento.integration_id == loja.integration_id,
                    ConciliacaoMovimento.origem == "mp_extrato",
                    ConciliacaoMovimento.pedido_bling.is_(None),
                    ConciliacaoMovimento.classificacao.is_(None),
                )
            )
        ).scalar()
        saida.append(
            {
                "conta": loja.conta,
                "integration_id": str(loja.integration_id),
                "inicio": loja.inicio.isoformat(),
                "ultimo_extrato_em": loja.ultimo_extrato_em,
                "ultimo_fatura_em": loja.ultimo_fatura_em,
                "ultimo_recalculo_em": loja.ultimo_recalculo_em,
                "ultimo_erro": loja.ultimo_erro,
                "movimentos": [
                    {
                        "origem": o,
                        "tipo": t,
                        "linhas": n,
                        "credito": str(c or 0),
                        "debito": str(d or 0),
                    }
                    for o, t, n, c, d in por_tipo
                ],
                "movimentos_sem_dono": sem_pedido,
                "sem_dono_por_tipo": [
                    {"origem": o, "tipo": t, "linhas": n} for o, t, n in sem_dono_tipos
                ],
                "fatura_por_sub_tipo": [
                    {
                        "sub_tipo": st,
                        "tipo": t,
                        "linhas": n,
                        "debito": str(d or 0),
                        "credito": str(c or 0),
                    }
                    for st, t, n, d, c in sub_tipos
                ],
                "pedidos_por_situacao": dict(por_situacao),
                "pedidos_por_etapa": dict(por_etapa),
                "leituras": [
                    {
                        "origem": x.origem,
                        "chave": x.chave,
                        "status": x.status,
                        "linhas": x.linhas,
                        "lido_em": x.lido_em,
                        "erro": x.erro,
                        "last_id": (x.detalhe or {}).get("last_id"),
                    }
                    for x in syncs
                ],
            }
        )
    return {"lojas": saida}


@router.post("/lojas/{conta}/sincronizar")
async def sincronizar_agora(
    conta: str,
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    """Enfileira uma rodada completa da loja (fatura + extrato + recálculo) no worker."""
    loja = await loja_por_conta(session, conta)
    if loja is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "loja_nao_ligada", "conta": conta}
        )
    pool = await get_arq_pool()
    # O arq recusa um job com o MESMO id enquanto o resultado do anterior ainda
    # está guardado (1 h): o id leva o minuto, e a trava por loja em
    # `sincronizar_loja` é quem evita duas rodadas ao mesmo tempo.
    job = await pool.enqueue_job(
        "conciliacao_sincronizar_loja",
        str(loja.integration_id),
        _job_id=f"conciliacao_sincronizar_loja:{loja.integration_id}:{int(time.time() // 60)}",
    )
    return {
        "conta": loja.conta,
        "job_id": job.job_id if job else None,
        "enfileirado": job is not None,
    }


# --- Painel: lojas, pedidos (lista + totais), Raio-X, marcações e Excel ----------
# Passo 2b (14/10/2026, "Ok" do Vinicius): a tela. Lê `conciliacao_pedidos`
# (reconstruída pelo recálculo) e `conciliacao_movimentos`; grava só as marcas
# da pessoa (observação, conferido, finalizar à mão com motivo).

from datetime import timedelta  # noqa: E402
from io import BytesIO  # noqa: E402
from zoneinfo import ZoneInfo  # noqa: E402

from fastapi.responses import StreamingResponse  # noqa: E402
from pydantic import BaseModel, Field  # noqa: E402
from sqlalchemy import and_ as sa_and  # noqa: E402
from sqlalchemy import case as sa_case  # noqa: E402
from sqlalchemy import func, literal_column, or_  # noqa: E402

from app.models import ConciliacaoLoja, ConciliacaoMovimento, ConciliacaoPedido  # noqa: E402
from app.models.conciliacao import ORIGEM_TIKTOK, ORIGENS_EXTRATO  # noqa: E402
from app.services.conciliacao_tiktok import AGUARDANDO, aguardando_repasse  # noqa: E402

_BRT = ZoneInfo("America/Sao_Paulo")

SITUACAO_LABELS = {
    "bateu": "Bateu",
    "aguardando": "Aguardando liberação",
    "vencido": "Não liberado (vencido)",
    "recebido_a_menos": "Recebido a menos",
    "recebido_a_mais": "Recebido a mais",
    "bling_diferente": "Bling ≠ venda",
    "cancelado": "Cancelado",
    "estornado_depois": "Estornado depois",
    "sem_venda": "Venda não lida",
    "falta_frete": "Falta frete (Melhor Envio)",
    "fora_da_loja": "Saiu da loja",
    "em_disputa": "Em disputa (dinheiro retido)",
}
ETAPA_LABELS = {
    "aguardando": "aguardando",
    "conciliado": "conciliado",
    "pendencia": "pendência",
    "finalizado": "finalizado",
    "reaberto": "reaberto",
}
_CAMPOS_PEDIDO = (
    "pedido_bling",
    "bling_id",
    "numeroloja",
    "pack_id",
    "conta",
    "situacao_bling",
    "bling_total",
    "bling_produtos",
    "bling_valorbase",
    "bling_custofrete",
    "bling_taxacomissao",
    "venda_valor",
    "venda_tarifa",
    "venda_frete",
    "venda_cupom",
    "venda_afiliados",
    "venda_frete_pago",
    "venda_estorno",
    "venda_liquido",
    "venda_status",
    "fatura_tarifa",
    "fatura_parcelamento",
    "fatura_frete",
    "fatura_bonificacao",
    "fatura_outros",
    "liberado_valor",
    "saidas_depois",
    "recebido",
    "custo",
    "reembolso",
    "sobrou",
    "margem",
    "dif_bling_venda",
    "dif_previsto_recebido",
    "situacao",
    "motivo",
    "etapa",
    "etapa_por",
    "etapa_motivo",
    "conferido_por",
    "observacao",
)
_DATAS_PEDIDO = (
    "data_pedido",
    "previsao_liberacao",
    "liberado_em",
    "ultimo_movimento_em",
    "etapa_em",
    "bateu_em",
    "conferido_em",
    "recalculado_em",
)


def _num(v: Any) -> float | None:
    return None if v is None else float(v)


def _pedido_out(p: ConciliacaoPedido) -> dict[str, Any]:
    out: dict[str, Any] = {"id": str(p.id), "integration_id": str(p.integration_id)}
    for c in _CAMPOS_PEDIDO:
        v = getattr(p, c)
        out[c] = _num(v) if isinstance(v, Decimal) else v
    for c in _DATAS_PEDIDO:
        v = getattr(p, c)
        out[c] = v.isoformat() if v else None
    out["situacao_label"] = SITUACAO_LABELS.get(p.situacao, p.situacao)
    out["etapa_label"] = ETAPA_LABELS.get(p.etapa, p.etapa)
    out["conferido"] = bool(p.conferido_em)
    out["venda_extras"] = _extras_da_venda(p)
    return out


def _extras_da_venda(p: ConciliacaoPedido) -> float | None:
    """Tudo que a plataforma tira antes de liberar além da tarifa e do frete (cupom da
    loja, afiliados, parcelamento, outros): valor − tarifa − frete − líquido previsto.
    Com isso a linha da grade fecha sozinha (15/10, Vinicius: "Extras")."""
    if p.situacao in ("cancelado", "sem_venda"):
        return None  # sem venda com valor, "extras" seria o valor cheio
    if None in (p.venda_valor, p.venda_tarifa, p.venda_frete, p.venda_liquido):
        return None
    extras = p.venda_valor - p.venda_tarifa - p.venda_frete - p.venda_liquido
    return float(extras.quantize(Decimal("0.01")))


def _movimento_out(m: ConciliacaoMovimento) -> dict[str, Any]:
    return {
        "id": str(m.id),
        "origem": m.origem,
        "data": m.data.isoformat() if m.data else None,
        "tipo": m.tipo,
        "descricao": m.descricao,
        "sub_tipo": m.sub_tipo,
        "pagamento_id": m.pagamento_id,
        "pedido_marketplace": m.pedido_marketplace,
        "pack_id": m.pack_id,
        "pedido_bling": m.pedido_bling,
        "bruto": _num(m.bruto),
        "tarifa": _num(m.tarifa),
        "parcelamento": _num(m.parcelamento),
        "frete": _num(m.frete),
        "cupom": _num(m.cupom),
        "credito": _num(m.credito),
        "debito": _num(m.debito),
        "saldo_apos": _num(m.saldo_apos),
        "descontado_da_venda": m.descontado_da_venda,
        "sku": m.sku,
        "classificacao": m.classificacao,
        "data_aprovacao": m.data_aprovacao.isoformat() if m.data_aprovacao else None,
        # A linha ORIGINAL (a linha do relatório do Mercado Pago ou o detalhe da fatura
        # do ML), pra quem quiser conferir na fonte — 15/10, pedido do Vinicius.
        "original": _original_do_movimento(m),
    }


def _original_do_movimento(m: ConciliacaoMovimento) -> dict[str, Any]:
    raw = m.raw if isinstance(m.raw, dict) else {}
    if m.origem in ORIGENS_EXTRATO:
        # só as colunas preenchidas, na ordem do relatório
        return {k: v for k, v in raw.items() if v not in (None, "", "0.00", "0", "{}", "[]")}
    ci = raw.get("charge_info") or {}
    vendas = raw.get("sales_info") or []
    venda = vendas[0] if isinstance(vendas, list) and vendas else {}
    return {
        "detail_id": ci.get("detail_id"),
        "creation_date_time": ci.get("creation_date_time"),
        "detail_type": ci.get("detail_type"),
        "detail_sub_type": ci.get("detail_sub_type"),
        "detail_amount": ci.get("detail_amount"),
        "transaction_detail": ci.get("transaction_detail"),
        "debited_from_operation": ci.get("debited_from_operation"),
        "charge_bonified_id": ci.get("charge_bonified_id"),
        "order_id": venda.get("order_id") if isinstance(venda, dict) else None,
        "operation_id": venda.get("operation_id") if isinstance(venda, dict) else None,
        "_periodo": raw.get("_periodo"),
    }


async def _loja_ou_404(
    session: AsyncSession, conta: str | None, *, escrita: bool = False
) -> ConciliacaoLoja:
    """Loja ligada pelo nome. Leitura sem `conta` cai na primeira loja (padrão da tela);
    rota de ESCRITA exige a loja (16/10, revisão: fechar/reabrir dia, marcar pedido e o
    CSV da KFA sem `loja` iam parar na primeira loja da lista)."""
    lojas = await lojas_ativas(session)
    if not lojas:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "nenhuma_loja_ligada"})
    if not (conta or "").strip():
        if escrita:
            raise HTTPException(
                422, detail={"code": "informe_a_loja", "mensagem": "informe a loja"}
            )
        return lojas[0]
    loja = next((x for x in lojas if x.conta.strip().lower() == conta.strip().lower()), None)
    if loja is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "loja_nao_ligada", "conta": conta}
        )
    return loja


def _filtros(
    loja: ConciliacaoLoja,
    *,
    de: date | None,
    ate: date | None,
    data: str,
    situacao: str | None,
    etapa: str | None,
    q: str | None,
) -> list[Any]:
    conds: list[Any] = [ConciliacaoPedido.integration_id == loja.integration_id]
    col = ConciliacaoPedido.liberado_em if data == "liberacao" else ConciliacaoPedido.data_pedido
    if de:
        conds.append(col >= datetime.combine(de, datetime.min.time(), tzinfo=_BRT))
    if ate:
        conds.append(
            col < datetime.combine(ate, datetime.min.time(), tzinfo=_BRT) + timedelta(days=1)
        )
    if situacao:
        conds.append(
            ConciliacaoPedido.situacao.in_([x.strip() for x in situacao.split(",") if x.strip()])
        )
    if etapa:
        conds.append(
            ConciliacaoPedido.etapa.in_([x.strip() for x in etapa.split(",") if x.strip()])
        )
    if q:
        termo = f"%{q.strip()}%"
        conds.append(
            or_(
                ConciliacaoPedido.pedido_bling.ilike(termo),
                ConciliacaoPedido.numeroloja.ilike(termo),
                ConciliacaoPedido.pack_id.ilike(termo),
            )
        )
    return conds


@router.get("/lojas")
async def listar_lojas(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
) -> dict[str, Any]:
    return {
        "lojas": [
            {
                "conta": loja.conta,
                "integration_id": str(loja.integration_id),
                "platform": getattr(loja.platform, "value", loja.platform),
                "inicio": loja.inicio.isoformat(),
                "ultimo_extrato_em": loja.ultimo_extrato_em.isoformat()
                if loja.ultimo_extrato_em
                else None,
                "ultimo_fatura_em": loja.ultimo_fatura_em.isoformat()
                if loja.ultimo_fatura_em
                else None,
                "ultimo_recalculo_em": loja.ultimo_recalculo_em.isoformat()
                if loja.ultimo_recalculo_em
                else None,
                "ultimo_erro": loja.ultimo_erro,
                "modo_extrato": loja.modo_extrato,
            }
            for loja in await lojas_ativas(session)
        ]
    }


# 16/10, revisão: tetos do envio de arquivo (bytes) — o corpo inteiro do script e o
# arquivo em si, lido em pedaços
_UPLOAD_TETO_CORPO = 20 * 1024 * 1024
_UPLOAD_TETO_ARQUIVO = 5_000_000
_UPLOAD_PEDACO = 1 << 20


async def _ler_arquivo_com_teto(arquivo: UploadFileBruto) -> bytes:
    """Lê em pedaços e desiste no teto (413) em vez de carregar tudo na memória."""
    partes: list[bytes] = []
    lido = 0
    while True:
        pedaco = await arquivo.read(_UPLOAD_PEDACO)
        if not pedaco:
            break
        lido += len(pedaco)
        if lido > _UPLOAD_TETO_ARQUIVO:
            raise HTTPException(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail={"code": "arquivo_grande"}
            )
        partes.append(pedaco)
    return b"".join(partes)


async def _travar_loja(session: AsyncSession, integration_id: UUID) -> None:
    """Clique que recalcula entra na trava da loja (16/10, revisão): a rodada no meio do
    caminho responde 409 em vez de 500 ou de sobrescrever o vínculo da pessoa."""
    try:
        await trava_da_loja(session, integration_id)
    except RodadaEmAndamentoError as e:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={
                "code": "rodada_em_andamento",
                "mensagem": "rodada em andamento, tente em instantes",
            },
        ) from e


async def _amazon_upload_corpo(
    session: AsyncSession, loja: str, arquivo: UploadFileBruto
) -> dict[str, Any]:
    from app.services.conciliacao_amazon_csv import ingerir_csv_amazon

    alvo = await _loja_ou_404(session, loja, escrita=True)
    plataforma = getattr(alvo.platform, "value", str(alvo.platform))
    if plataforma != "amazon" or alvo.modo_extrato != "upload":
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={"code": "loja_nao_e_upload", "mensagem": "esta loja é lida pela API"},
        )
    dados = await _ler_arquivo_com_teto(arquivo)
    try:
        texto = dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = dados.decode("latin-1")
    await _travar_loja(session, alvo.integration_id)
    try:
        stats = await ingerir_csv_amazon(session, alvo, texto)
    except ValueError as e:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": "csv_invalido", "mensagem": str(e)},
        ) from e
    recalculo = await recalcular_loja(session, alvo)
    await session.commit()
    return {"loja": alvo.conta, "arquivo": arquivo.filename, **stats, "recalculo": recalculo}


@router.post("/amazon/upload")
async def amazon_upload(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
    loja: Annotated[str, Form()],
    arquivo: Annotated[UploadFile, File()],
) -> dict[str, Any]:
    """Transações da Amazon por arquivo (loja com `modo_extrato = upload`, 16/10: KFA).
    O CSV é o de Pagamentos › Visualizar transações › Fazer download, com "Todos os
    tipos de conta" e o período escolhido (no máximo 600 transações por arquivo)."""
    return await _amazon_upload_corpo(session, loja, arquivo)


@router.post("/amazon/upload/script")
async def amazon_upload_pelo_script(
    request: Request,
    session: Annotated[AsyncSession, Depends(get_session)],
    x_upload_chave: Annotated[str | None, Header()] = None,
    x_agent_token: Annotated[str | None, Header()] = None,
) -> dict[str, Any]:
    """O mesmo envio, feito pelo script do Mac (sem ninguém logado): a chave gerada em
    "Gerar chave pro script" vai no header X-Upload-Chave. O executor do Mac mini usa o
    X-Agent-Token dele (o mesmo da Logística).

    16/10, revisão: a chave é conferida ANTES de ler o corpo (rota sem login não lê o
    arquivo de quem não tem a chave) e o corpo tem teto (`_UPLOAD_TETO_CORPO` → 413).
    Por isso o formulário é lido à mão (`request.form()`) depois da chave."""
    import secrets as _secrets

    from app.config import get_settings
    from app.services.conciliacao_amazon_csv import chave_upload_valida

    esperado = get_settings().marketing_agent_token
    pelo_agente = bool(
        esperado and x_agent_token and _secrets.compare_digest(x_agent_token, esperado)
    )
    if not pelo_agente and not await chave_upload_valida(session, x_upload_chave):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, detail={"code": "chave_invalida"})
    try:
        tamanho = int(request.headers.get("content-length") or 0)
    except ValueError:
        tamanho = 0
    if tamanho > _UPLOAD_TETO_CORPO:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail={"code": "arquivo_grande"}
        )
    form = await request.form()
    loja = form.get("loja")
    arquivo = form.get("arquivo")
    if not isinstance(loja, str) or not isinstance(arquivo, UploadFileBruto):
        raise HTTPException(
            422, detail={"code": "formulario_invalido", "mensagem": "informe a loja e o arquivo"}
        )
    return await _amazon_upload_corpo(session, loja, arquivo)


@router.post("/lojas/{conta}/buscar-transacoes")
async def amazon_buscar_transacoes(
    conta: str,
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    """Botão ⟳ da loja por arquivo (16/10): pede ao robô do Mac mini pra baixar o CSV
    da Amazon e enviar. Fica na fila do robô da Logística (ação amazon_transacoes_csv)."""
    from app.services import logistica_robo
    from app.services.conciliacao_amazon_csv import enfileirar_busca_csv

    alvo = await _loja_ou_404(session, conta, escrita=True)
    if alvo.modo_extrato != "upload":
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={"code": "loja_nao_e_upload", "mensagem": "esta loja é lida pela API"},
        )
    cmd, novo = await enfileirar_busca_csv(session, alvo, user_id=user.id)
    await session.commit()
    visto = await logistica_robo.ultimo_heartbeat(session)
    return {
        "conta": alvo.conta,
        "comando_id": str(cmd.id),
        "enfileirado": novo,
        "robo_visto_em": visto.last_seen_at.isoformat() if visto and visto.last_seen_at else None,
    }


@router.get("/amazon/upload/chave")
async def amazon_upload_chave_situacao(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    from app.services.conciliacao_amazon_csv import situacao_chave_upload

    return await situacao_chave_upload(session)


@router.post("/amazon/upload/chave")
async def amazon_upload_chave_gerar(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    """Gera (ou troca) a chave do script. Aparece UMA vez; a anterior deixa de valer."""
    from app.services.conciliacao_amazon_csv import gerar_chave_upload

    chave = await gerar_chave_upload(session, por=getattr(user, "name", None) or str(user.id))
    await session.commit()
    return {"chave": chave}


@router.get("/pedidos")
async def listar_pedidos(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    de: date | None = None,
    ate: date | None = None,
    data: Annotated[str, Query(pattern="^(pedido|liberacao)$")] = "pedido",
    situacao: str | None = None,
    etapa: str | None = None,
    q: str | None = None,
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=500)] = 100,
) -> dict[str, Any]:
    """Uma linha por pedido do Bling + os totais do filtro (os cartões e o rodapé)."""
    alvo = await _loja_ou_404(session, loja)
    conds = _filtros(alvo, de=de, ate=ate, data=data, situacao=situacao, etapa=etapa, q=q)
    base = select(ConciliacaoPedido).where(*conds)
    total = (await session.execute(select(func.count()).select_from(base.subquery()))).scalar() or 0
    rows = (
        (
            await session.execute(
                base.order_by(
                    ConciliacaoPedido.data_pedido.desc().nullslast(),
                    ConciliacaoPedido.pedido_bling.desc(),
                )
                .offset((page - 1) * size)
                .limit(size)
            )
        )
        .scalars()
        .all()
    )
    ped = ConciliacaoPedido
    somas = (
        await session.execute(
            select(
                func.count(),
                func.sum(ped.venda_valor),
                func.sum(ped.venda_tarifa),
                func.sum(ped.venda_frete),
                func.sum(ped.venda_liquido),
                func.sum(ped.recebido),
                func.sum(ped.dif_previsto_recebido),
                func.sum(ped.custo),
                func.sum(ped.reembolso),
                func.sum(ped.sobrou),
                # "liberado" = o dinheiro saiu do ML e FICOU. No cancelado por
                # reclamação o ML libera e tira no mesmo dia: não conta aqui.
                func.count(
                    sa_case((ped.liberado_em.isnot(None) & (ped.situacao != "cancelado"), 1))
                ),
            ).where(*conds)
        )
    ).one()
    por_situacao = dict(
        (
            await session.execute(
                select(ped.situacao, func.count()).where(*conds).group_by(ped.situacao)
            )
        ).all()
    )
    por_etapa = dict(
        (
            await session.execute(select(ped.etapa, func.count()).where(*conds).group_by(ped.etapa))
        ).all()
    )
    custo = somas[7] or Decimal("0")
    sobrou = somas[9] or Decimal("0")
    return {
        "loja": alvo.conta,
        "total": total,
        "page": page,
        "size": size,
        "itens": [_pedido_out(p) for p in rows],
        "resumo": {
            "pedidos": somas[0],
            "vendido": _num(somas[1] or 0),
            "tarifas_frete": _num((somas[2] or 0) + (somas[3] or 0)),
            "previsto": _num(somas[4] or 0),
            "recebido": _num(somas[5] or 0),
            "diferenca": _num(somas[6] or 0),
            "custo": _num(custo),
            "reembolso": _num(somas[8] or 0),
            "sobrou": _num(sobrou),
            "margem": _num((sobrou / custo).quantize(Decimal("0.0001"))) if custo > 0 else None,
            "liberados": somas[10],
            "por_situacao": por_situacao,
            "por_etapa": por_etapa,
        },
    }


async def _pedido_ou_404(
    session: AsyncSession, loja: ConciliacaoLoja, pedido: str
) -> ConciliacaoPedido:
    p = (
        await session.execute(
            select(ConciliacaoPedido).where(
                ConciliacaoPedido.integration_id == loja.integration_id,
                ConciliacaoPedido.pedido_bling == pedido,
            )
        )
    ).scalar_one_or_none()
    if p is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "pedido_nao_conciliado", "pedido": pedido}
        )
    return p


@router.get("/pedidos/{pedido}")
async def raio_x(
    pedido: str,
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
) -> dict[str, Any]:
    """As quatro versões do dinheiro lado a lado + cada movimento do pedido."""
    alvo = await _loja_ou_404(session, loja)
    p = await _pedido_ou_404(session, alvo, pedido)
    chaves = {x for x in (p.numeroloja, p.pack_id) if x}
    movs = (
        (
            await session.execute(
                select(ConciliacaoMovimento)
                .where(
                    ConciliacaoMovimento.integration_id == alvo.integration_id,
                    or_(
                        ConciliacaoMovimento.pedido_bling == p.pedido_bling,
                        ConciliacaoMovimento.pedido_marketplace.in_(chaves or {"-"}),
                        ConciliacaoMovimento.pack_id.in_(chaves or {"-"}),
                    ),
                )
                .order_by(ConciliacaoMovimento.data, ConciliacaoMovimento.origem)
            )
        )
        .scalars()
        .all()
    )
    liberacoes = [m for m in movs if m.origem in ORIGENS_EXTRATO and m.tipo == "venda_liberada"]

    def _soma(attr: str, lista: list[ConciliacaoMovimento]) -> float | None:
        vals = [getattr(m, attr) for m in lista if getattr(m, attr) is not None]
        return float(abs(sum(vals, Decimal("0")))) if vals else None

    # Frete do comprador: o ML põe dentro do GROSS do pagamento o que o comprador pagou
    # de frete, e desconta o frete INTEIRO (SHIPPING_FEE); a venda guarda só a parte da
    # loja. Sem isso o Raio-X mostrava 24,45 × 55,44 e parecia erro (300415, 15/10).
    bruto_extrato = _soma("bruto", liberacoes)
    venda_valor = _num(p.venda_valor)
    frete_comprador = (
        round(bruto_extrato - venda_valor, 2)
        if (
            bruto_extrato is not None
            and venda_valor is not None
            and bruto_extrato - venda_valor > 0.005
        )
        else 0.0
    )
    frete_extrato = _soma("frete", liberacoes)
    venda_frete = _num(p.venda_frete)
    versoes = {
        "bling": {
            "valor": _num(p.bling_total),
            "tarifa": _num(p.bling_taxacomissao),
            "frete": _num(p.bling_custofrete),
            "liquido": _num(
                (p.bling_valorbase or Decimal("0"))
                - (p.bling_custofrete or Decimal("0"))
                - (p.bling_taxacomissao or Decimal("0"))
            )
            if p.bling_valorbase is not None
            else None,
        },
        "venda": {
            "valor": venda_valor,
            "frete_comprador": frete_comprador,
            "bruto": round(venda_valor + frete_comprador, 2) if venda_valor is not None else None,
            "tarifa": _num(p.venda_tarifa),
            "frete": venda_frete,
            "frete_cobrado": round(venda_frete + frete_comprador, 2)
            if venda_frete is not None
            else None,
            "cupom": _num(p.venda_cupom),
            "estorno": _num(p.venda_estorno),
            "liquido": _num(p.venda_liquido),
            "status": p.venda_status,
        },
        "fatura": {
            "tarifa": _num(p.fatura_tarifa),
            "parcelamento": _num(p.fatura_parcelamento),
            "frete": _num(p.fatura_frete),
            "bonificacao": _num(p.fatura_bonificacao),
            "outros": _num(p.fatura_outros),
        },
        "extrato": {
            "valor": bruto_extrato,
            "frete_comprador": frete_comprador if bruto_extrato is not None else None,
            "produtos": round(bruto_extrato - frete_comprador, 2)
            if bruto_extrato is not None
            else None,
            "tarifa": _soma("tarifa", liberacoes),
            "parcelamento": _soma("parcelamento", liberacoes),
            "frete": frete_extrato,
            "frete_liquido": round(frete_extrato - frete_comprador, 2)
            if frete_extrato is not None
            else None,
            "cupom": _soma("cupom", liberacoes),
            "liquido": _num(p.liberado_valor),
            "saidas_depois": _num(p.saidas_depois),
            "recebido": _num(p.recebido),
        },
    }
    # reserva retida do pedido (19/10): o par é conferido dentro dos movimentos do pedido
    saldo_reservas: dict[str, Decimal] = defaultdict(Decimal)
    for m in movs:
        if m.tipo in TIPOS_RESERVA_SET and m.pagamento_id:
            saldo_reservas[str(m.pagamento_id)] += (m.credito or Decimal("0")) - (
                m.debito or Decimal("0")
            )
    reserva_aberta = _flags_reserva(movs, dict(saldo_reservas))
    return {
        "pedido": _pedido_out(p),
        "versoes": versoes,
        "comparativo": _comparativo(
            p,
            versoes,
            movs,
            shopee=await _conta_shopee_do_pedido(session, p),
            tiktok=await _conta_tiktok_do_pedido(session, p, movs),
            amazon=await _conta_amazon_do_pedido(session, p, movs),
        ),
        "movimentos": [
            {**_movimento_out(m), "reserva_aberta": reserva_aberta.get(m.id)} for m in movs
        ],
        "frete_pago": await _fretes_pagos_do_pedido(session, p),
    }


async def _fretes_pagos_do_pedido(
    session: AsyncSession, p: ConciliacaoPedido
) -> list[dict[str, Any]]:
    """Envios do Melhor Envio casados com o pedido (bloco "Frete pago fora" do Raio-X)."""
    from app.models import ConciliacaoFretePago

    rows = (
        (
            await session.execute(
                select(ConciliacaoFretePago)
                .where(
                    ConciliacaoFretePago.integration_id == p.integration_id,
                    ConciliacaoFretePago.pedido_bling == p.pedido_bling,
                )
                .order_by(ConciliacaoFretePago.pago_em)
            )
        )
        .scalars()
        .all()
    )
    return [
        {
            "fonte": f.fonte,
            "servico": f.servico,
            "status": f.status,
            "rastreio": f.rastreio,
            "preco": _num(f.preco),
            "desconto": _num(f.desconto),
            "conferencia": _num(f.conferencia),
            "frete_pago": _num(f.frete_pago),
            "pago_em": f.pago_em.isoformat() if f.pago_em else None,
            "postado_em": f.postado_em.isoformat() if f.postado_em else None,
            "cancelado_em": f.cancelado_em.isoformat() if f.cancelado_em else None,
        }
        for f in rows
    ]


async def _conta_shopee_do_pedido(
    session: AsyncSession, p: ConciliacaoPedido
) -> dict[str, Decimal] | None:
    """A conta da Shopee (como a tela deles) a partir do escrow guardado pela Margem."""
    from app.models import MarketplaceOrderFinancial
    from app.services.conciliacao_calculo import conta_shopee

    if getattr(p.platform, "value", str(p.platform)) != "shopee" or not p.numeroloja:
        return None
    fin = (
        await session.execute(
            select(MarketplaceOrderFinancial).where(
                MarketplaceOrderFinancial.integration_id == p.integration_id,
                MarketplaceOrderFinancial.external_order_id == p.numeroloja,
            )
        )
    ).scalar_one_or_none()
    return conta_shopee(fin)


async def _conta_tiktok_do_pedido(
    session: AsyncSession, p: ConciliacaoPedido, movs: list[ConciliacaoMovimento]
) -> dict[str, Any] | None:
    """TikTok: previsto = as transações do pedido que a Margem guardou; real = as
    linhas do extrato diário que o painel leu (a mesma API, lida em horas diferentes)."""
    from app.models import MarketplaceOrderFinancial
    from app.models.conciliacao import ORIGEM_TIKTOK
    from app.services.conciliacao_calculo import conta_tiktok, conta_tiktok_de_transacoes

    if getattr(p.platform, "value", str(p.platform)) != "tiktok":
        return None
    fin = None
    if p.numeroloja:
        fin = (
            await session.execute(
                select(MarketplaceOrderFinancial).where(
                    MarketplaceOrderFinancial.integration_id == p.integration_id,
                    MarketplaceOrderFinancial.external_order_id == p.numeroloja,
                )
            )
        ).scalar_one_or_none()
    previsto = conta_tiktok(fin)
    real = conta_tiktok_de_transacoes(
        [
            m.raw
            for m in movs
            if m.origem == ORIGEM_TIKTOK and m.tipo == "venda_liberada" and isinstance(m.raw, dict)
        ]
    )
    if previsto is None and real is None:
        return None
    return {"previsto": previsto, "real": real}


async def _conta_amazon_do_pedido(
    session: AsyncSession, p: ConciliacaoPedido, movs: list[ConciliacaoMovimento]
) -> dict[str, Any] | None:
    """Amazon: previsto = transações do pedido que a Margem guardou; real = as linhas
    das Transações que o painel leu (a mesma API, lida em horas diferentes)."""
    from app.models import MarketplaceOrderFinancial
    from app.models.conciliacao import ORIGEM_AMAZON
    from app.services.conciliacao_calculo import conta_amazon, conta_amazon_de_transacoes

    if getattr(p.platform, "value", str(p.platform)) != "amazon":
        return None
    fin = None
    if p.numeroloja:
        fin = (
            await session.execute(
                select(MarketplaceOrderFinancial).where(
                    MarketplaceOrderFinancial.integration_id == p.integration_id,
                    MarketplaceOrderFinancial.external_order_id == p.numeroloja,
                )
            )
        ).scalar_one_or_none()
    previsto = conta_amazon(fin)
    real = conta_amazon_de_transacoes(
        [
            m.raw
            for m in movs
            if m.origem == ORIGEM_AMAZON
            and m.tipo in ("venda_liberada", "frete")
            and isinstance(m.raw, dict)
        ]
    )
    if previsto is None and real is None:
        return None
    if previsto is None:
        # loja por arquivo (KFA, 16/10): a Margem não lê a Amazon; a venda é o que o
        # próprio arquivo diz do pedido
        previsto = real
    return {"previsto": previsto, "real": real}


def _comparativo(
    p: ConciliacaoPedido,
    versoes: dict[str, Any],
    movs: list[ConciliacaoMovimento],
    *,
    shopee: dict[str, Decimal] | None = None,
    tiktok: dict[str, Any] | None = None,
    amazon: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Previsto (venda) × Real (extrato + fatura à parte) × Diferença — a leitura de
    conciliação que o Vinicius aprovou em 15/10 no lugar das três colunas."""
    v, e = versoes["venda"], versoes["extrato"]
    tem_extrato = e.get("valor") is not None

    def n(x: Any) -> float | None:
        return float(x) if x is not None else None

    # Cobranças da fatura que NÃO foram descontadas na liberação (de verdade: a fatura
    # marca o Mercado Envios como "não descontado" mesmo quando o extrato descontou).
    a_parte = 0.0
    for m in movs:
        if m.origem != "ml_fatura" or not m.debito or m.descontado_da_venda is not False:
            continue
        if m.tipo == "fatura_frete" and (e.get("frete") or 0) > 0:
            continue
        if (
            m.tipo in ("fatura_tarifa_venda", "fatura_tarifa_cobranca")
            and (e.get("tarifa") or 0) > 0
        ):
            continue
        a_parte += float(m.debito)
    a_parte = round(a_parte, 2)
    liberado = n(p.liberado_valor)
    recebido = n(p.recebido)
    # "depois" = só estornos/ajustes depois do lançamento (frete cobrado depois, como a
    # etiqueta Easy Ship, faz parte da liquidação — ver saidas_depois no cálculo)
    # calculado das linhas (não de p.saidas_depois, que fica None no cancelado): débitos
    # depois do lançamento que não são frete nem reserva
    saidas = (
        round(
            sum(
                float(m.debito or 0)
                for m in movs
                if m.origem in ORIGENS_EXTRATO
                and p.liberado_em is not None
                and m.data > p.liberado_em
                and m.tipo != "frete"
                and m.tipo not in TIPOS_RESERVA_SET
            ),
            2,
        )
        if liberado is not None
        else 0.0
    )
    depois = round(-saidas, 2) if (recebido is not None and liberado is not None) else None
    # A linha final compara o que a plataforma LANÇOU com o previsto; o que ficou com a
    # loja depois de estornos vai numa linha própria (15/10, Vinicius: "759,53 em cima
    # de um x tá estranho" — comparar o previsto com o líquido pós-estorno não diz nada).
    real_liquido = round(recebido - a_parte + saidas, 2) if recebido is not None else None
    ficou = round(recebido - a_parte, 2) if recebido is not None else None

    def _fechar(linhas_: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if depois is not None and abs(depois) > 0.005:
            linhas_.append(
                {
                    "chave": "ficou",
                    "rotulo": "= Ficou com a loja (depois de estornos e ajustes)",
                    "previsto": None,
                    "real": ficou,
                    "dif": None,
                    "estilo": "bold",
                }
            )
        return linhas_

    def linha(chave, rotulo, previsto, real, *, compara=True, estilo="") -> dict[str, Any]:
        dif = (
            round(previsto - real, 2)
            if (compara and previsto is not None and real is not None)
            else None
        )
        return {
            "chave": chave,
            "rotulo": rotulo,
            "previsto": previsto,
            "real": real,
            "dif": dif,
            "estilo": estilo,
        }

    plataforma = getattr(p.platform, "value", str(p.platform))
    if plataforma == "amazon":
        # As linhas da tela Pagamentos › Transações da Amazon (15/10, seção 24):
        # produto, frete cobrado do cliente, tarifas (comissão, Amazon for All, retenção,
        # etiqueta Easy Ship), promoção. Previsto = Margem; real = linhas lidas.
        liberou_am = recebido is not None
        prev = (amazon or {}).get("previsto")
        real_am = (amazon or {}).get("real") if liberou_am else None

        def par_am(
            chave: str,
            rotulo: str,
            campo: str,
            *,
            sinal: int = 1,
            estilo: str = "",
            so_se_tem: bool = False,
        ) -> dict[str, Any] | None:
            pv = round(float(prev[campo]) * sinal, 2) if prev is not None else None
            rl = round(float(real_am[campo]) * sinal, 2) if real_am is not None else None
            if so_se_tem and not pv and not rl:
                return None
            return linha(chave, rotulo, pv, rl, estilo=estilo)

        if prev is not None or real_am is not None:
            candidatas = [
                par_am("produtos", "Produto (custo total do produto)", "produto", estilo="bold"),
                par_am(
                    "frete_cliente", "+ Frete cobrado do cliente", "frete_cliente", so_se_tem=True
                ),
                par_am("valor", "= Valor (o que o cliente pagou)", "valor", estilo="bold"),
                par_am(
                    "promocao",
                    "− Desconto promocional da loja",
                    "promocao",
                    sinal=-1,
                    so_se_tem=True,
                ),
                par_am("comissao", "− Comissão da Amazon", "comissao", sinal=-1),
                par_am("amazon_for_all", "− Amazon for All (1,5%)", "amazon_for_all", sinal=-1),
                par_am(
                    "retencao",
                    "− Retenção sobre o frete",
                    "retencao_frete",
                    sinal=-1,
                    so_se_tem=True,
                ),
                par_am(
                    "etiqueta", "− Etiqueta Easy Ship (DBA)", "etiqueta", sinal=-1, so_se_tem=True
                ),
                par_am(
                    "outras_tarifas", "− Outras tarifas", "outras_tarifas", sinal=-1, so_se_tem=True
                ),
                par_am("outros", "± Outros", "outros", so_se_tem=True),
                linha(
                    "depois",
                    "± Depois do lançamento (estornos, ajustes)",
                    0.0 if tem_extrato else None,
                    depois,
                    compara=False,
                    estilo="muted",
                ),
                linha(
                    "liquido",
                    "= Total do pedido (o que a Amazon paga)",
                    round(float(prev["liquido"]), 2) if prev is not None else n(v.get("liquido")),
                    real_liquido,
                    estilo="bold",
                ),
            ]
            linhas = [x for x in candidatas if x is not None]
        else:
            linhas = [
                linha("produtos", "Produto", n(v.get("valor")), None),
                linha("tarifa", "− Tarifas da Amazon", n(v.get("tarifa")), None),
                linha("frete", "± Frete (etiqueta / frete cobrado)", n(v.get("frete")), None),
                linha(
                    "depois",
                    "± Depois do lançamento (estornos, ajustes)",
                    0.0 if tem_extrato else None,
                    depois,
                    compara=False,
                    estilo="muted",
                ),
                linha(
                    "liquido", "= Total do pedido", n(v.get("liquido")), real_liquido, estilo="bold"
                ),
            ]
        return {
            "linhas": _fechar(linhas),
            "real_rotulo": "transações da Amazon",
            "comparar": p.situacao != "cancelado",
            "nota": (
                "venda desfeita: as colunas mostram o que entrou e o que saiu"
                if p.situacao == "cancelado"
                else (
                    None
                    if liberou_am
                    else "a Amazon lança o pedido 3 a 4 dias depois do envio; até lá não há valor"
                )
            ),
            "rodape": (
                "as mesmas colunas de Pagamentos › Transações da Amazon · o valor fica preso na "
                "reserva até ~7 dias depois da entrega e sai no repasse de 14 dias"
            ),
            "frete_liquido_loja": None,
        }
    if plataforma == "tiktok":
        # A TikTok fecha a conta do pedido no extrato diário em que liquida, campo a
        # campo (15/10, Mini). Previsto = as transações que a Margem guardou; real = as
        # linhas do extrato que o painel leu. Até liquidar, só a estimativa da Central.
        liberou_tt = recebido is not None
        prev = (tiktok or {}).get("previsto")
        real_tt = (tiktok or {}).get("real") if liberou_tt else None

        def par(
            chave: str,
            rotulo: str,
            campo: str,
            *,
            sinal: int = 1,
            estilo: str = "",
            so_se_tem: bool = False,
        ) -> dict[str, Any] | None:
            pv = round(float(prev[campo]) * sinal, 2) if prev is not None else None
            rl = round(float(real_tt[campo]) * sinal, 2) if real_tt is not None else None
            if so_se_tem and not pv and not rl:
                return None
            return linha(chave, rotulo, pv, rl, estilo=estilo)

        if prev is not None or real_tt is not None:
            candidatas = [
                par("vendas", "Vendas brutas (preço de tabela)", "vendas_brutas", estilo="muted"),
                par("desconto", "− Desconto da loja", "desconto_loja", sinal=-1, estilo="muted"),
                par("produtos", "= Produtos após desconto da loja", "apos_desconto", estilo="bold"),
                par(
                    "subsidio",
                    "± Subsídio da TikTok e outros na receita",
                    "outros_receita",
                    estilo="muted",
                ),
                par("receita", "= Receita", "receita", estilo="bold"),
                par("comissao", "− Comissão da TikTok", "comissao", sinal=-1),
                par("taxa", "− Taxa de transação", "taxa_transacao", sinal=-1),
                par("afiliados", "− Afiliados", "afiliados", sinal=-1),
                par("outras_taxas", "− Outras taxas", "outras_taxas", sinal=-1, so_se_tem=True),
                par("frete", "± Frete (parte da loja)", "frete", sinal=-1),
                par("ajustes", "± Ajustes da TikTok", "ajustes", so_se_tem=True),
                par("outros", "± Outros", "outros", so_se_tem=True),
                linha(
                    "depois",
                    "± Depois da liquidação (estornos, ajustes)",
                    0.0 if tem_extrato else None,
                    depois,
                    compara=False,
                    estilo="muted",
                ),
                linha(
                    "liquido",
                    "= Liquidação do pedido",
                    round(float(prev["liquido"]), 2) if prev is not None else n(v.get("liquido")),
                    real_liquido,
                    estilo="bold",
                ),
            ]
            linhas = [x for x in candidatas if x is not None]
        else:
            linhas = [
                linha("produtos", "Produtos (após desconto da loja)", n(v.get("valor")), None),
                linha("tarifa", "− Taxas da TikTok", n(v.get("tarifa")), None),
                linha("frete", "± Frete (parte da loja)", n(v.get("frete")), None),
                linha(
                    "depois",
                    "± Depois da liquidação (estornos, ajustes)",
                    0.0 if tem_extrato else None,
                    depois,
                    compara=False,
                    estilo="muted",
                ),
                linha(
                    "liquido",
                    "= Liquidação do pedido",
                    n(v.get("liquido")),
                    real_liquido,
                    estilo="bold",
                ),
            ]
        return {
            "linhas": _fechar(linhas),
            "real_rotulo": "extrato diário da TikTok",
            "comparar": p.situacao != "cancelado",
            "nota": (
                "venda desfeita: as colunas mostram o que entrou e o que saiu"
                if p.situacao == "cancelado"
                else (
                    None
                    if liberou_tt
                    else (
                        "a TikTok fecha a conta do pedido no extrato do dia em que liquida; "
                        "até lá só a estimativa da Central do Vendedor"
                    )
                )
            ),
            "rodape": (
                "as linhas do extrato da TikTok por pedido (Finanças › Extratos) · "
                '"Liquidação" real é o que entrou no extrato diário e foi repassado pro banco'
            ),
            "frete_liquido_loja": None,
        }
    if plataforma == "shopee":
        # A carteira só traz o líquido. A abertura REAL (tarifa, frete, cupom) é o
        # escrow que a Shopee fecha na liberação — o robô relê na hora de liberar, e
        # é isso que a Margem guarda. Até liberar, só existe o previsto.
        liberou = recebido is not None

        def real(x: Any) -> float | None:
            return n(x) if liberou else None

        if shopee is not None:
            # As mesmas linhas de "Informações de Pagamento › Exibir Detalhes de Renda"
            # da Shopee, pra comparar linha a linha (15/10, pedido 301438).
            sh = {k: float(x) for k, x in shopee.items()}
            linhas = [
                linha("preco", "Preço de venda", sh["preco"], real(sh["preco"])),
                linha(
                    "subsidio",
                    "− Subsídio da Shopee no produto",
                    sh["subsidio"],
                    real(sh["subsidio"]),
                    estilo="muted",
                ),
                linha(
                    "produtos",
                    "= Subtotal dos produtos",
                    sh["subtotal_produtos"],
                    real(sh["subtotal_produtos"]),
                    estilo="bold",
                ),
                linha(
                    "frete",
                    "± Subtotal de frete (parte da loja)",
                    sh["frete_loja"],
                    real(sh["frete_loja"]),
                ),
                linha("cupom", "− Cupons e descontos da loja", sh["cupons"], real(sh["cupons"])),
                linha(
                    "tarifa",
                    "− Taxas e encargos (comissão, serviço, transação)",
                    sh["taxas"],
                    real(sh["taxas"]),
                ),
                linha("afiliados", "− Afiliados", sh["afiliados"], real(sh["afiliados"])),
                linha(
                    "depois",
                    "± Depois da liberação (estornos, ajustes)",
                    0.0 if tem_extrato else None,
                    depois,
                    compara=False,
                    estilo="muted",
                ),
                linha("liquido", "= Renda do pedido", sh["renda"], real_liquido, estilo="bold"),
            ]
            return {
                "linhas": _fechar(linhas),
                "real_rotulo": "escrow na liberação + carteira",
                "comparar": p.situacao != "cancelado",
                "nota": (
                    "venda desfeita: as colunas mostram o que entrou e o que saiu"
                    if p.situacao == "cancelado"
                    else (
                        None
                        if liberou
                        else (
                            "a Shopee fecha a conta da venda na liberação; "
                            "até lá só existe o previsto"
                        )
                    )
                ),
                "rodape": (
                    'as mesmas linhas de "Informações de Pagamento › Detalhes de Renda" '
                    "da Shopee · "
                    '"Renda do pedido" real é o que a carteira pagou'
                ),
                "frete_liquido_loja": None,
            }
        linhas = [
            linha("produtos", "Produtos (preço da venda)", n(v.get("valor")), real(v.get("valor"))),
            linha(
                "tarifa",
                "− Tarifas da Shopee (comissão, serviço, transação)",
                n(v.get("tarifa")),
                real(v.get("tarifa")),
            ),
            linha("frete", "− Frete (parte da loja)", n(v.get("frete")), real(v.get("frete"))),
            linha("cupom", "− Cupom / desconto da loja", n(v.get("cupom")), real(v.get("cupom"))),
            linha(
                "depois",
                "± Depois da liberação (estornos, ajustes)",
                0.0 if tem_extrato else None,
                depois,
                compara=False,
                estilo="muted",
            ),
            linha("liquido", "= Líquido", n(v.get("liquido")), real_liquido, estilo="bold"),
        ]
        return {
            "linhas": _fechar(linhas),
            "real_rotulo": "escrow na liberação + carteira",
            "comparar": p.situacao != "cancelado",
            "nota": (
                "venda desfeita: as colunas mostram o que entrou e o que saiu"
                if p.situacao == "cancelado"
                else (
                    None
                    if liberou
                    else "a Shopee fecha a conta da venda na liberação; até lá só existe o previsto"
                )
            ),
            "rodape": (
                "a carteira da Shopee só traz o líquido; a abertura real é o escrow que a "
                "Shopee fecha na liberação (o robô relê na hora) · o líquido real é o que a "
                "carteira pagou"
            ),
            "frete_liquido_loja": None,
        }
    fc = n(v.get("frete_comprador")) or 0.0
    linhas = [
        linha("produtos", "Produtos (preço da venda)", n(v.get("valor")), n(e.get("produtos"))),
        linha(
            "frete_comprador",
            "+ Frete pago pelo comprador",
            fc,
            n(e.get("frete_comprador")),
            compara=False,
            estilo="muted",
        ),
        linha(
            "total",
            "= Total pago pelo comprador",
            n(v.get("bruto")),
            n(e.get("valor")),
            estilo="bold",
        ),
        linha("tarifa", "− Tarifa de venda", n(v.get("tarifa")), n(e.get("tarifa"))),
        linha(
            "frete",
            "− Frete Mercado Envios (inteiro)",
            n(v.get("frete_cobrado")),
            n(e.get("frete")),
        ),
        linha("parcelamento", "− Parcelamento", None, n(e.get("parcelamento")), compara=False),
        linha("cupom", "− Cupom da loja", n(v.get("cupom")), n(e.get("cupom"))),
        linha(
            "depois",
            "± Depois da liberação (estornos, mediações)",
            0.0 if tem_extrato else None,
            depois,
            compara=False,
            estilo="muted",
        ),
        linha(
            "a_parte",
            "− Cobranças à parte na fatura",
            None,
            a_parte if tem_extrato else None,
            compara=False,
            estilo="muted",
        ),
        linha("liquido", "= Líquido", n(v.get("liquido")), real_liquido, estilo="bold"),
    ]
    if fc <= 0.005:
        linhas = [x for x in linhas if x["chave"] not in ("frete_comprador", "total")]
    if plataforma != "ml":
        linhas = [x for x in linhas if x["chave"] != "a_parte"]
    fl = n(v.get("frete"))
    return {
        "linhas": _fechar(linhas),
        "real_rotulo": "extrato + fatura",
        "rodape": (
            (
                f"frete líquido pra loja {fl:.2f} (frete inteiro − o que o comprador pagou) · "
                if fl is not None
                else ""
            )
            + '"Real" é o extrato do Mercado Pago mais o que a fatura cobra à parte'
        ),
        "comparar": p.situacao != "cancelado",
        "nota": (
            "venda desfeita: as colunas mostram o que entrou e o que saiu, não há previsto a cobrar"
            if p.situacao == "cancelado"
            else None
        ),
        "frete_liquido_loja": n(v.get("frete")),
    }


class PedidoPatch(BaseModel):
    loja: str | None = None
    observacao: str | None = Field(default=None, max_length=2000)
    conferido: bool | None = None
    # "finalizado" (à mão, com motivo quando teve pendência) ou "reaberto".
    etapa: str | None = Field(default=None, pattern="^(finalizado|reaberto)$")
    etapa_motivo: str | None = Field(default=None, max_length=500)


@router.patch("/pedidos/{pedido}")
async def marcar_pedido(
    pedido: str,
    body: PedidoPatch,
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    alvo = await _loja_ou_404(session, body.loja, escrita=True)
    p = await _pedido_ou_404(session, alvo, pedido)
    agora = datetime.now(UTC)
    quem = (getattr(user, "name", None) or user.email or "?")[:120]
    if body.observacao is not None:
        p.observacao = body.observacao.strip() or None
    if body.conferido is not None:
        p.conferido_em = agora if body.conferido else None
        p.conferido_por = quem if body.conferido else None
    if body.etapa == "finalizado":
        if p.situacao not in ("bateu", "cancelado") and not (body.etapa_motivo or "").strip():
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail={
                    "code": "motivo_obrigatorio",
                    "mensagem": "pra finalizar com pendência, diga o motivo",
                },
            )
        p.etapa = "finalizado"
        p.etapa_em = agora
        p.etapa_por = quem
        p.etapa_motivo = (body.etapa_motivo or "").strip() or None
    elif body.etapa == "reaberto":
        p.etapa = "reaberto"
        p.etapa_em = agora
        p.etapa_por = quem
        p.etapa_motivo = (body.etapa_motivo or "").strip() or "reaberto à mão"
    await session.commit()
    await session.refresh(p)
    return {"pedido": _pedido_out(p)}


_COLUNAS_XLSX = [
    ("Data", "data_pedido"),
    ("Pedido Bling", "pedido_bling"),
    ("Venda ML", "numeroloja"),
    ("Pack", "pack_id"),
    ("Loja", "conta"),
    ("Bling total", "bling_total"),
    ("Venda valor", "venda_valor"),
    ("Tarifa", "venda_tarifa"),
    ("Frete", "venda_frete"),
    ("Frete pago (Melhor Envio)", "venda_frete_pago"),
    ("Extras (cupom, parcelamento, outros)", "venda_extras"),
    ("Cupom", "venda_cupom"),
    ("Afiliados", "venda_afiliados"),
    ("Estorno", "venda_estorno"),
    ("Líquido previsto", "venda_liquido"),
    ("Fatura tarifa", "fatura_tarifa"),
    ("Fatura parcelamento", "fatura_parcelamento"),
    ("Fatura frete", "fatura_frete"),
    ("Fatura bonificação", "fatura_bonificacao"),
    ("Liberado em", "liberado_em"),
    ("Liberado", "liberado_valor"),
    ("Saídas depois", "saidas_depois"),
    ("Recebido", "recebido"),
    ("Dif. previsto × recebido", "dif_previsto_recebido"),
    ("Dif. Bling × venda", "dif_bling_venda"),
    ("Custo", "custo"),
    ("Reembolso", "reembolso"),
    ("Sobrou", "sobrou"),
    ("Margem", "margem"),
    ("Situação", "situacao_label"),
    ("Motivo", "motivo"),
    ("Status Bling", "situacao_bling"),
    ("Etapa", "etapa_label"),
    ("Etapa em", "etapa_em"),
    ("Etapa por", "etapa_por"),
    ("Conferido por", "conferido_por"),
    ("Observação", "observacao"),
]


# 16/10, revisão: célula de texto do Excel nunca vira fórmula (texto começando com =, +,
# -, @ ia como fórmula) e caractere de controle (a openpyxl recusa com
# IllegalCharacterError → 500) sai antes de gravar. As duas exportações passam por aqui.
_PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")
_CONTROLE_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # o mesmo conjunto da openpyxl


def _celula_excel(v: Any) -> Any:
    return _CONTROLE_RE.sub("", v) if isinstance(v, str) else v


def _linha_excel(ws: Any, valores: list[Any]) -> None:
    valores = [_celula_excel(v) for v in valores]
    ws.append(valores)
    if not any(isinstance(v, str) and v.startswith(_PREFIXOS_FORMULA) for v in valores):
        return
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith(_PREFIXOS_FORMULA):
            cell.data_type = "s"  # texto, não fórmula
            cell.quotePrefix = True  # o Excel mostra o texto como está, sem interpretar


@router.post("/export.xlsx")
async def exportar_xlsx(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    de: date | None = None,
    ate: date | None = None,
    data: Annotated[str, Query(pattern="^(pedido|liberacao)$")] = "pedido",
    situacao: str | None = None,
    etapa: str | None = None,
    q: str | None = None,
) -> StreamingResponse:
    from openpyxl import Workbook

    alvo = await _loja_ou_404(session, loja)
    conds = _filtros(alvo, de=de, ate=ate, data=data, situacao=situacao, etapa=etapa, q=q)
    rows = (
        (
            await session.execute(
                select(ConciliacaoPedido)
                .where(*conds)
                .order_by(ConciliacaoPedido.data_pedido.desc().nullslast())
                .limit(20000)
            )
        )
        .scalars()
        .all()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliação"
    _linha_excel(ws, [c for c, _ in _COLUNAS_XLSX])
    for p in rows:
        out = _pedido_out(p)
        linha = []
        for _, campo in _COLUNAS_XLSX:
            v = out.get(campo)
            if isinstance(v, str) and len(v) >= 19 and v[4] == "-" and "T" in v:
                try:
                    v = datetime.fromisoformat(v).astimezone(_BRT).strftime("%d/%m/%Y %H:%M")
                except ValueError:
                    pass
            linha.append(v)
        _linha_excel(ws, linha)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"conciliacao_{alvo.conta}_{datetime.now(_BRT):%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


# --- Extrato do dia: movimentos, vincular/classificar, sugestões, fechar o dia -----
# Passo 2c (14/10/2026, "Ok" do Vinicius): a tela de trabalho da manhã. Cada
# linha do extrato de ontem precisa de um destino — um pedido ou uma
# classificação — e o dia só fecha com 100%. Prova dos nove no rodapé: saldo
# inicial + créditos − débitos = saldo final = saldo informado pelo Mercado Pago.

from app.models import BlingOrder, ConciliacaoDia, ConciliacaoSync  # noqa: E402
from app.services.conciliacao_calculo import recalcular_loja  # noqa: E402
from app.services.conciliacao_extrato import TIPOS_RESERVA as TIPOS_RESERVA_SET  # noqa: E402
from app.services.conciliacao_extrato import bling_store_ids_da_loja  # noqa: E402

CLASSIFICACOES = (
    "saque",
    "fatura_ml",
    "publicidade",
    "anterior_ao_inicio",
    "reserva",
    "venda_desfeita",
    "ajuste",
    "outro",
)
CLASSIFICACAO_LABELS = {
    "saque": "Saque",
    "fatura_ml": "Fatura do ML",
    "publicidade": "Publicidade",
    "anterior_ao_inicio": "Pedido anterior ao início",
    "reserva": "Reserva (vai e volta)",
    "venda_desfeita": "Venda desfeita (sem pedido no Bling)",
    "ajuste": "Ajuste",
    "outro": "Outro",
}
TIPO_LABELS = {
    "venda_liberada": "Venda liberada",
    "mediacao": "Mediação (reclamação perdida)",
    "mediacao_cancelada": "Mediação cancelada",
    "estorno": "Estorno",
    "chargeback": "Chargeback",
    "disputa": "Disputa",
    "reserva_disputa": "Reserva por reclamação",
    "reserva_estorno": "Reserva pra estorno",
    "reserva_frete_devolucao": "Reserva do frete de devolução",
    "reserva_saque": "Reserva pro saque",
    "reserva_pagamento": "Reserva de pagamento",
    "reserva_outra": "Reserva (outra)",
    "fatura_cobrada": "Fatura do ML cobrada do saldo",
    "cobranca_divida": "Segurado pra pagar a fatura (dívida)",
    "saque": "Saque",
    "frete": "Frete",
    "frete_cancelado": "Frete cancelado",
    "frete_devolucao": "Frete de devolução",
    "antecipacao": "Antecipação",
    "saque_confirmado": "Saque confirmado",
    "saque_cancelado": "Saque cancelado",
    "ajuste_shopee": "Ajuste da Shopee",
    "ajuste_tiktok": "Ajuste da TikTok",
    "ajuste_amazon": "Ajuste da Amazon",
    "publicidade": "Publicidade (Ads)",
    "credito": "Crédito",
    "imposto": "Imposto",
    "outro": "Outro",
}
_SITUACOES_DIVERGENTES = {
    "recebido_a_menos",
    "recebido_a_mais",
    "bling_diferente",
    "estornado_depois",
    "vencido",
    "em_disputa",
}


def _status_movimento(m: ConciliacaoMovimento, pedido: ConciliacaoPedido | None) -> str:
    if m.classificacao == "anterior_ao_inicio":
        return "anterior"
    if m.classificacao:
        return "classificado"
    if not m.pedido_bling:
        return "pendente"
    if m.vinculado_por:
        return "manual"
    if (
        pedido is not None
        and m.tipo == "venda_liberada"
        and pedido.situacao in _SITUACOES_DIVERGENTES
    ):
        return "divergente"
    # 19/10 (revisão): a reserva que prende o dinheiro do pedido "em disputa" também é a
    # linha divergente do dia
    if pedido is not None and m.tipo in TIPOS_RESERVA_SET and pedido.situacao == "em_disputa":
        return "divergente"
    return "automatico"


# reserva pro saque é o saque em andamento (classificada "saque"): não é dinheiro retido
_RESERVAS_RETIDAS = tuple(t for t in TIPOS_RESERVA_SET if t != "reserva_saque")


def _flags_reserva(
    movs: list[ConciliacaoMovimento], saldo_por_pagamento: dict[str, Decimal]
) -> dict[Any, bool]:
    """Por linha de reserva: True = aberta (saiu e não voltou, ou sem nº de pagamento pra
    provar o par) → a tela mostra; False = par fechado (soma zero) → pode esconder."""
    out: dict[Any, bool] = {}
    for m in movs:
        if m.tipo not in TIPOS_RESERVA_SET:
            continue
        if m.tipo == "reserva_saque":
            out[m.id] = False
            continue
        pg = str(m.pagamento_id or "")
        out[m.id] = (not pg) or abs(saldo_por_pagamento.get(pg, Decimal("0"))) > Decimal("0.05")
    return out


async def _reservas_abertas(
    session: AsyncSession, loja: ConciliacaoLoja, movs: list[ConciliacaoMovimento]
) -> dict[Any, bool]:
    """19/10 (revisão): a tela escondia TODA reserva como "vão e voltam, somam zero" —
    injox (ml) 15/10 tinha uma única linha pendente (reserva +367,48 sem dono) e ela
    estava escondida. Por nº de pagamento, soma crédito − débito de todas as reservas da
    loja (qualquer dia): par fechado pode ser escondido; sobra ≠ 0 é dinheiro retido."""
    mov = ConciliacaoMovimento
    pagamentos = {
        str(m.pagamento_id) for m in movs if m.tipo in _RESERVAS_RETIDAS and m.pagamento_id
    }
    saldo: dict[str, Decimal] = {}
    if pagamentos:
        rows = (
            await session.execute(
                select(mov.pagamento_id, func.sum(mov.credito - mov.debito))
                .where(
                    mov.integration_id == loja.integration_id,
                    mov.origem.in_(ORIGENS_EXTRATO),
                    mov.tipo.in_(_RESERVAS_RETIDAS),
                    mov.pagamento_id.in_(pagamentos),
                )
                .group_by(mov.pagamento_id)
            )
        ).all()
        saldo = {str(pg): Decimal(str(v or 0)) for pg, v in rows}
    return _flags_reserva(movs, saldo)


_TOL_SALDO = Decimal("0.01")
# tamanho máximo do empate de horário que a busca completa percorre (acima disso, o
# encadeamento guloso de antes — as rajadas de 60 linhas da injox)
_MAX_EMPATE_BUSCA = 8


def _d0(v: Decimal | None) -> Decimal:
    return v if v is not None else Decimal("0")


def _encadeia(
    anterior: Decimal, m: ConciliacaoMovimento, reserva_divida: Decimal
) -> tuple[bool, Decimal]:
    """A linha `m` cabe depois de um saldo `anterior`? Devolve (cabe, quanto do débito foi
    pago com o "segurado pra dívida" e por isso NÃO saiu do saldo disponível).

    19/10 (revisão): a fatura do ML cobrada do saldo (`fatura_cobrada`) vem com o valor
    inteiro no NET_DEBIT mesmo quando parte (ou tudo) foi pago com o dinheiro que o
    Mercado Pago já tinha segurado (`reserve_for_debt_payment`) — o BALANCE só cai pela
    parte que saiu do disponível. kia (ml) 16/10 18:52: +362,32 · segurado −362,32
    (saldo 0) · fatura −362,32 (saldo continua 0) — o painel somava os dois débitos e
    acusava "saldo não bate +362,32" sem linha pra classificar. A parte paga com o
    segurado só é aceita até o que foi segurado no dia (não "explica" linha faltando)."""
    esperado = anterior + _d0(m.credito) - _d0(m.debito)
    if abs(esperado - m.saldo_apos) <= _TOL_SALDO:
        return True, Decimal("0")
    if m.tipo == "fatura_cobrada" and _d0(m.debito) > 0:
        consumo = m.saldo_apos - esperado
        teto = min(_d0(m.debito), reserva_divida)
        if consumo > _TOL_SALDO and consumo <= teto + _TOL_SALDO:
            return True, min(consumo, teto)
    return False, Decimal("0")


def _reserva_apos(m: ConciliacaoMovimento, reserva_divida: Decimal, consumo: Decimal) -> Decimal:
    if m.tipo == "cobranca_divida":
        reserva_divida += _d0(m.debito) - _d0(m.credito)
    return max(reserva_divida - consumo, Decimal("0"))


def _ordem_do_empate(
    grupo: list[ConciliacaoMovimento], anterior: Decimal | None, reserva_divida: Decimal
) -> tuple[list[ConciliacaoMovimento], list[Decimal]] | None:
    """Busca completa da ordem de um empate de horário: toda linha encadeia (a fatura
    paga com o segurado pode encadear com o consumo), e entre as ordens válidas vale a
    que menos "explica" com o segurado. None quando nenhuma ordem fecha."""
    melhor: list | None = None

    def dfs(ordem, consumos, saldo, reserva, restantes, total) -> None:
        nonlocal melhor
        if melhor is not None and total >= melhor[2]:
            return
        if not restantes:
            melhor = [list(ordem), list(consumos), total]
            return
        for i, m in enumerate(restantes):
            cabe, consumo = _encadeia(saldo, m, reserva)
            if not cabe:
                continue
            dfs(
                [*ordem, m],
                [*consumos, consumo],
                m.saldo_apos,
                _reserva_apos(m, reserva, consumo),
                restantes[:i] + restantes[i + 1 :],
                total + consumo,
            )

    if anterior is not None:
        dfs([], [], anterior, reserva_divida, list(grupo), Decimal("0"))
    else:
        # sem âncora: cada linha pode ser a primeira (o saldo de antes dela é o implícito)
        for i, m in enumerate(grupo):
            dfs(
                [m],
                [Decimal("0")],
                m.saldo_apos,
                _reserva_apos(m, reserva_divida, Decimal("0")),
                grupo[:i] + grupo[i + 1 :],
                Decimal("0"),
            )
    if melhor is None:
        return None
    return melhor[0], melhor[1]


def _ordena_dia(
    movs: list[ConciliacaoMovimento], saldo_anterior: Decimal | None = None
) -> list[ConciliacaoMovimento]:
    """Ordem do dia. O relatório não traz nº de linha; em empate de horário, a
    ordem certa é a que encadeia o saldo (saldo da linha = saldo da anterior +
    crédito − débito). `saldo_anterior` (fechamento do dia anterior) ancora o
    primeiro empate do dia — sem ele, pagamento + reserva no mesmo segundo formam
    um ciclo (0 → 362,15 → 0) e a ordem saía trocada (injox 15/10).

    Cada linha sai com `pago_com_reserva` (parte do débito da fatura do ML paga com o
    segurado pra dívida, que não saiu do saldo disponível — ver `_encadeia`)."""
    movs = sorted(movs, key=lambda m: (m.data, m.id))
    for m in movs:
        m.pago_com_reserva = Decimal("0")  # type: ignore[attr-defined]
    out: list[ConciliacaoMovimento] = []
    reserva_divida = Decimal("0")  # o que o MP segurou pra dívida hoje e ainda não devolveu
    i = 0
    while i < len(movs):
        grupo = [movs[i]]
        j = i + 1
        while j < len(movs) and movs[j].data == movs[i].data:
            grupo.append(movs[j])
            j += 1
        anterior = (
            out[-1].saldo_apos
            if out and out[-1].saldo_apos is not None
            else (saldo_anterior if not out else None)
        )
        com_saldo = all(m.saldo_apos is not None for m in grupo)
        resolvido = None
        if com_saldo and len(grupo) <= _MAX_EMPATE_BUSCA:
            resolvido = _ordem_do_empate(grupo, anterior, reserva_divida)
        if resolvido is not None:
            ordenado, consumos = resolvido
            for m, consumo in zip(ordenado, consumos, strict=True):
                m.pago_com_reserva = consumo  # type: ignore[attr-defined]
                reserva_divida = _reserva_apos(m, reserva_divida, consumo)
            grupo = ordenado
        elif len(grupo) > 1 and com_saldo:
            # empate grande ou sem ordem que feche: o encadeamento guloso de antes
            resto = list(grupo)
            ordenado = []
            while resto:
                prox = None
                if anterior is not None:
                    prox = next(
                        (m for m in resto if _encadeia(anterior, m, reserva_divida)[0]), None
                    )
                if prox is None:
                    # sem âncora: o primeiro é o que não é "seguinte" de nenhum outro
                    prox = next(
                        (
                            m
                            for m in resto
                            if not any(
                                abs((o.saldo_apos + m.credito - m.debito) - m.saldo_apos)
                                <= _TOL_SALDO
                                for o in resto
                                if o is not m
                            )
                        ),
                        resto[0],
                    )
                consumo = (
                    _encadeia(anterior, prox, reserva_divida)[1]
                    if anterior is not None
                    else Decimal("0")
                )
                prox.pago_com_reserva = consumo  # type: ignore[attr-defined]
                reserva_divida = _reserva_apos(prox, reserva_divida, consumo)
                ordenado.append(prox)
                resto.remove(prox)
                anterior = prox.saldo_apos
            grupo = ordenado
        else:
            for m in grupo:
                if m.saldo_apos is not None and anterior is not None:
                    _cabe, consumo = _encadeia(anterior, m, reserva_divida)
                    m.pago_com_reserva = consumo  # type: ignore[attr-defined]
                    reserva_divida = _reserva_apos(m, reserva_divida, consumo)
                else:
                    reserva_divida = _reserva_apos(m, reserva_divida, Decimal("0"))
                if m.saldo_apos is not None:
                    anterior = m.saldo_apos
        out.extend(grupo)
        i = j
    return out


def _totais_do_dia(
    movs: list[ConciliacaoMovimento], saldo_anterior: Decimal | None = None
) -> dict[str, Any]:
    creditos = sum((m.credito or Decimal("0") for m in movs), Decimal("0"))
    debitos = sum((m.debito or Decimal("0") for m in movs), Decimal("0"))
    # fatura do ML paga com o "segurado pra dívida": está nos débitos (é cobrança de
    # verdade) mas não saiu do saldo disponível — entra como linha própria do rodapé
    pago_com_reserva = sum(
        (getattr(m, "pago_com_reserva", None) or Decimal("0") for m in movs), Decimal("0")
    )
    primeiro = next((m for m in movs if m.saldo_apos is not None), None)
    ultimo = next((m for m in reversed(movs) if m.saldo_apos is not None), None)
    # saldo inicial = fechamento do dia anterior quando existe (16/10); senão, o que a
    # primeira linha do dia diz
    if saldo_anterior is not None and movs:
        saldo_inicial: Decimal | None = saldo_anterior
    else:
        saldo_inicial = (
            (primeiro.saldo_apos - primeiro.credito + primeiro.debito) if primeiro else None
        )
    saldo_final = (
        (saldo_inicial + creditos - debitos + pago_com_reserva)
        if saldo_inicial is not None
        else None
    )
    saldo_informado = ultimo.saldo_apos if ultimo else None
    return {
        "linhas": len(movs),
        "saldo_inicial": _num(saldo_inicial),
        "creditos": _num(creditos),
        "debitos": _num(debitos),
        "fatura_paga_com_reserva": _num(pago_com_reserva),
        "saldo_final": _num(saldo_final),
        "saldo_informado": _num(saldo_informado),
        "diferenca": _num(saldo_informado - saldo_final)
        if (saldo_informado is not None and saldo_final is not None)
        else None,
    }


async def _saldo_anterior(
    session: AsyncSession, loja: ConciliacaoLoja, dia: date
) -> Decimal | None:
    """Saldo informado logo antes do dia = saldo inicial de verdade do dia (prova dos
    nove independente da ordem das linhas do dia).

    Linhas no mesmo segundo têm ordem ambígua (kia (ml) 10/10: +7,04 / −7,04 / +7,04 às
    13:45:48; rajadas de 60 linhas em poucos segundos). Regra exata: âncora na última
    linha cujo horário é ÚNICO (saldo sem dúvida) + soma de tudo que veio depois dela.
    """
    if getattr(loja.platform, "value", str(loja.platform)) == "tiktok":
        # cada extrato da TikTok é um caixa próprio: começa em 0 e o repasse zera.
        # Extrato ainda não pago termina o dia no total liquidado (16/10).
        return Decimal("0")
    ini = datetime.combine(dia, datetime.min.time(), tzinfo=_BRT)
    mov = ConciliacaoMovimento
    ultimo = (
        (
            await session.execute(
                select(mov)
                .where(
                    mov.integration_id == loja.integration_id,
                    mov.origem.in_(ORIGENS_EXTRATO),
                    mov.data < ini,
                )
                .order_by(mov.data.desc())
                .limit(400)
            )
        )
        .scalars()
        .all()
    )
    if not ultimo:
        return None
    por_instante: dict[datetime, int] = {}
    for m in ultimo:
        por_instante[m.data] = por_instante.get(m.data, 0) + 1
    soma_depois = Decimal("0")
    for i, m in enumerate(ultimo):
        # a mais velha da janela pode ter gêmeas fora dela: não serve de âncora
        if por_instante[m.data] == 1 and m.saldo_apos is not None and i < len(ultimo) - 1:
            return Decimal(m.saldo_apos) + soma_depois
        soma_depois += (m.credito or Decimal("0")) - (m.debito or Decimal("0"))
    # tudo empatado: ordena o último dia pelo encadeamento (como antes)
    com_saldo = [m for m in ultimo if m.saldo_apos is not None]
    if not com_saldo:
        return None
    ultimo_dia = com_saldo[0].data.astimezone(_BRT).date()
    do_dia = [m for m in com_saldo if m.data.astimezone(_BRT).date() == ultimo_dia]
    ordenado = _ordena_dia(do_dia)
    return ordenado[-1].saldo_apos if ordenado else None


async def _movimentos_do_dia(
    session: AsyncSession,
    loja: ConciliacaoLoja,
    dia: date,
    saldo_anterior: Decimal | None = None,
) -> list[ConciliacaoMovimento]:
    ini = datetime.combine(dia, datetime.min.time(), tzinfo=_BRT)
    movs = (
        (
            await session.execute(
                select(ConciliacaoMovimento).where(
                    ConciliacaoMovimento.integration_id == loja.integration_id,
                    ConciliacaoMovimento.origem.in_(ORIGENS_EXTRATO),
                    ConciliacaoMovimento.data >= ini,
                    ConciliacaoMovimento.data < ini + timedelta(days=1),
                )
            )
        )
        .scalars()
        .all()
    )
    return _ordena_dia(list(movs), saldo_anterior)


def _ontem() -> date:
    return (datetime.now(_BRT) - timedelta(days=1)).date()


@router.get("/extrato")
async def extrato_do_dia(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    dia: date | None = None,
    mostrar: Annotated[str, Query(pattern="^(todos|pendentes|divergentes|sem_pedido)$")] = "todos",
    q: str | None = None,
) -> dict[str, Any]:
    alvo = await _loja_ou_404(session, loja)
    dia = dia or _ontem()
    saldo_ant = await _saldo_anterior(session, alvo, dia)
    movs = await _movimentos_do_dia(session, alvo, dia, saldo_ant)
    pedidos = {
        p.pedido_bling: p
        for p in (
            await session.execute(
                select(ConciliacaoPedido).where(
                    ConciliacaoPedido.integration_id == alvo.integration_id,
                    ConciliacaoPedido.pedido_bling.in_(
                        {m.pedido_bling for m in movs if m.pedido_bling} or {"-"}
                    ),
                )
            )
        ).scalars()
    }
    reserva_aberta = await _reservas_abertas(session, alvo, movs)
    linhas: list[dict[str, Any]] = []
    contagem: dict[str, int] = {
        "automatico": 0,
        "manual": 0,
        "classificado": 0,
        "anterior": 0,
        "pendente": 0,
        "divergente": 0,
    }
    for m in movs:
        st = _status_movimento(m, pedidos.get(m.pedido_bling or ""))
        contagem[st] += 1
        if mostrar == "pendentes" and st != "pendente":
            continue
        if mostrar == "divergentes" and st != "divergente":
            continue
        if mostrar == "sem_pedido" and m.pedido_bling:
            continue
        if q:
            termo = q.strip().lower()
            if (
                termo
                not in " ".join(
                    str(x or "")
                    for x in (
                        m.pedido_bling,
                        m.pedido_marketplace,
                        m.pack_id,
                        m.pagamento_id,
                        m.descricao,
                        m.sku,
                    )
                ).lower()
            ):
                continue
        ped = pedidos.get(m.pedido_bling or "")
        linhas.append(
            {
                **_movimento_out(m),
                "tipo_label": TIPO_LABELS.get(m.tipo, m.tipo),
                "status": st,
                "vinculado_por": m.vinculado_por,
                "vinculado_em": m.vinculado_em.isoformat() if m.vinculado_em else None,
                "classificacao_label": CLASSIFICACAO_LABELS.get(
                    m.classificacao or "", m.classificacao
                ),
                "pedido_situacao": ped.situacao if ped else None,
                "pedido_situacao_label": SITUACAO_LABELS.get(ped.situacao, ped.situacao)
                if ped
                else None,
                "pedido_etapa": ped.etapa if ped else None,
                # reserva: True = retida (aparece sempre); False = par fechado (pode esconder)
                "reserva_aberta": reserva_aberta.get(m.id),
                # fatura do ML: parte do débito paga com o segurado pra dívida (19/10)
                "pago_com_reserva": _num(getattr(m, "pago_com_reserva", None) or Decimal("0"))
                if m.tipo == "fatura_cobrada"
                else None,
            }
        )
    fechado = (
        await session.execute(
            select(ConciliacaoDia).where(
                ConciliacaoDia.integration_id == alvo.integration_id, ConciliacaoDia.dia == dia
            )
        )
    ).scalar_one_or_none()
    return {
        "loja": alvo.conta,
        "dia": dia.isoformat(),
        "movimentos": linhas,
        "totais": _totais_do_dia(movs, saldo_ant),
        "contagem": contagem,
        "resolvidos": len(movs) - contagem["pendente"],
        "fechado": {
            "em": fechado.fechado_em.isoformat(),
            "por": fechado.fechado_por,
            "motivo": fechado.fechado_motivo,
        }
        if fechado and fechado.fechado_em
        else None,
        "classificacoes": [{"key": k, "label": v} for k, v in CLASSIFICACAO_LABELS.items()],
    }


@router.get("/dias")
async def listar_dias(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    de: date | None = None,
    ate: date | None = None,
) -> dict[str, Any]:
    """Um resumo por dia (linhas, pendentes, fechado?) pra navegar e pro Resumo."""
    alvo = await _loja_ou_404(session, loja)
    ate = ate or _ontem()
    de = de or max(alvo.inicio, ate - timedelta(days=30))
    ini = datetime.combine(de, datetime.min.time(), tzinfo=_BRT)
    fim = datetime.combine(ate, datetime.min.time(), tzinfo=_BRT) + timedelta(days=1)
    dia_col = func.date(func.timezone(_FUSO_SQL, ConciliacaoMovimento.data))
    pend = func.sum(
        sa_case(
            (
                ConciliacaoMovimento.pedido_bling.is_(None)
                & ConciliacaoMovimento.classificacao.is_(None),
                1,
            ),
            else_=0,
        )
    )
    rows = (
        await session.execute(
            select(
                dia_col,
                func.count(),
                pend,
                func.sum(ConciliacaoMovimento.credito),
                func.sum(ConciliacaoMovimento.debito),
            )
            .where(
                ConciliacaoMovimento.integration_id == alvo.integration_id,
                ConciliacaoMovimento.origem.in_(ORIGENS_EXTRATO),
                ConciliacaoMovimento.data >= ini,
                ConciliacaoMovimento.data < fim,
            )
            .group_by(dia_col)
            .order_by(dia_col)
        )
    ).all()
    fechados = {
        d.dia: d
        for d in (
            await session.execute(
                select(ConciliacaoDia).where(
                    ConciliacaoDia.integration_id == alvo.integration_id,
                    ConciliacaoDia.dia >= de,
                    ConciliacaoDia.dia <= ate,
                )
            )
        ).scalars()
    }
    saida = []
    for d, n, p, c, deb in rows:
        f = fechados.get(d)
        saida.append(
            {
                "dia": d.isoformat(),
                "linhas": n,
                "pendentes": int(p or 0),
                "creditos": _num(c or 0),
                "debitos": _num(deb or 0),
                "fechado_em": f.fechado_em.isoformat() if f and f.fechado_em else None,
                "fechado_por": f.fechado_por if f else None,
            }
        )
    return {"loja": alvo.conta, "de": de.isoformat(), "ate": ate.isoformat(), "dias": saida}


class MovimentoPatch(BaseModel):
    loja: str | None = None
    # Vincular a um pedido do Bling (número), classificar sem pedido, ou desvincular.
    pedido_bling: str | None = Field(default=None, max_length=20)
    classificacao: str | None = Field(default=None, max_length=40)
    desvincular: bool = False


@router.get("/movimentos/{mov_id}/sugestoes")
async def sugestoes_de_vinculo(
    mov_id: UUID,
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
) -> dict[str, Any]:
    """Pedidos candidatos pra uma linha sem dono: mesmo pagamento, mesmo nº de
    venda/pack, ou mesmo valor por perto da data."""
    m = await session.get(ConciliacaoMovimento, mov_id)
    if m is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "movimento_nao_encontrado"})
    loja = (
        await session.execute(
            select(ConciliacaoLoja).where(ConciliacaoLoja.integration_id == m.integration_id)
        )
    ).scalar_one_or_none()
    lojas_bling = (
        await bling_store_ids_da_loja(session, await session.get(Integration, m.integration_id))
        if loja
        else set()
    )
    sugestoes: list[dict[str, Any]] = []
    vistos: set[str] = set()

    def _add(
        numero: str, numeroloja: str | None, data_pedido: Any, total: Any, motivo: str
    ) -> None:
        if numero in vistos:
            return
        vistos.add(numero)
        sugestoes.append(
            {
                "pedido_bling": numero,
                "numeroloja": numeroloja,
                "data_pedido": data_pedido.isoformat() if data_pedido else None,
                "total": _num(total),
                "motivo": motivo,
            }
        )

    if m.pagamento_id:
        irmaos = (
            await session.execute(
                select(ConciliacaoMovimento.pedido_bling, ConciliacaoMovimento.pedido_marketplace)
                .where(
                    ConciliacaoMovimento.integration_id == m.integration_id,
                    ConciliacaoMovimento.pagamento_id == m.pagamento_id,
                    ConciliacaoMovimento.pedido_bling.isnot(None),
                )
                .distinct()
            )
        ).all()
        for numero, numeroloja in irmaos:
            _add(str(numero), numeroloja, None, None, "mesmo pagamento")
    chaves = {x for x in (m.pedido_marketplace, m.pack_id) if x}
    if chaves:
        rows = (
            await session.execute(
                select(BlingOrder.numero, BlingOrder.numeroloja, BlingOrder.data, BlingOrder.total)
                .where(BlingOrder.numeroloja.in_(chaves), BlingOrder.loja.in_(lojas_bling or {"-"}))
                .distinct()
            )
        ).all()
        for numero, numeroloja, data_pedido, total in rows:
            _add(str(numero), numeroloja, data_pedido, total, "mesmo nº de venda")
    valor = (
        m.bruto.copy_abs()
        if m.bruto
        else ((m.credito or Decimal("0")) - (m.debito or Decimal("0"))).copy_abs()
    )
    if valor and valor > 0 and lojas_bling:
        rows = (
            await session.execute(
                select(BlingOrder.numero, BlingOrder.numeroloja, BlingOrder.data, BlingOrder.total)
                .where(
                    BlingOrder.loja.in_(lojas_bling),
                    BlingOrder.total.between(valor - Decimal("0.05"), valor + Decimal("0.05")),
                    BlingOrder.data >= m.data - timedelta(days=45),
                    BlingOrder.data <= m.data + timedelta(days=2),
                )
                .distinct()
                .limit(8)
            )
        ).all()
        for numero, numeroloja, data_pedido, total in rows:
            _add(str(numero), numeroloja, data_pedido, total, "mesmo valor")
    return {"movimento": _movimento_out(m), "sugestoes": sugestoes[:10]}


@router.patch("/movimentos/{mov_id}")
async def marcar_movimento(
    mov_id: UUID,
    body: MovimentoPatch,
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    m = await session.get(ConciliacaoMovimento, mov_id)
    if m is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "movimento_nao_encontrado"})
    loja = (
        await session.execute(
            select(ConciliacaoLoja).where(ConciliacaoLoja.integration_id == m.integration_id)
        )
    ).scalar_one_or_none()
    if loja is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "loja_nao_ligada"})
    await _travar_loja(session, m.integration_id)
    # a linha foi lida antes da trava: a rodada pode ter apagado/trocado (2ª rodada)
    try:
        await session.refresh(m)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail={"code": "movimento_nao_encontrado"}
        ) from exc
    agora = datetime.now(UTC)
    quem = (getattr(user, "name", None) or user.email or "?")[:120]
    # Dia fechado não se mexe: reabre primeiro (15/10, Vinicius).
    dia_mov = m.data.astimezone(_BRT).date()
    fechado = (
        await session.execute(
            select(ConciliacaoDia).where(
                ConciliacaoDia.integration_id == m.integration_id,
                ConciliacaoDia.dia == dia_mov,
                ConciliacaoDia.fechado_em.isnot(None),
            )
        )
    ).scalar_one_or_none()
    if fechado is not None:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={
                "code": "dia_fechado",
                "dia": dia_mov.isoformat(),
                "por": fechado.fechado_por,
                "em": fechado.fechado_em.isoformat() if fechado.fechado_em else None,
            },
        )
    if body.desvincular:
        m.pedido_bling = None
        m.bling_id = None
        m.classificacao = None
        m.vinculado_por = quem
        m.vinculado_em = agora
    elif body.pedido_bling:
        lojas_bling = await bling_store_ids_da_loja(
            session, await session.get(Integration, m.integration_id)
        )
        alvo = (
            await session.execute(
                select(BlingOrder.numero, BlingOrder.bling_id)
                .where(
                    BlingOrder.numero == body.pedido_bling.strip(),
                    BlingOrder.loja.in_(lojas_bling or {"-"}),
                )
                .limit(1)
            )
        ).first()
        if alvo is None:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND,
                detail={
                    "code": "pedido_bling_nao_encontrado",
                    "pedido": body.pedido_bling,
                    "loja": loja.conta,
                },
            )
        m.pedido_bling, m.bling_id = str(alvo[0]), alvo[1]
        m.classificacao = None
        m.vinculado_por = quem
        m.vinculado_em = agora
    elif body.classificacao:
        if body.classificacao not in CLASSIFICACOES:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY, detail={"code": "classificacao_invalida"}
            )
        m.classificacao = body.classificacao
        m.pedido_bling = None
        m.bling_id = None
        m.vinculado_por = quem
        m.vinculado_em = agora
    else:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, detail={"code": "nada_a_fazer"})
    await session.flush()
    await recalcular_loja(session, loja)
    await session.commit()
    await session.refresh(m)
    return {"movimento": _movimento_out(m), "status": _status_movimento(m, None)}


@router.post("/dias/{dia}/fechar")
async def fechar_dia(
    dia: date,
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
    loja: str | None = None,
    motivo: str | None = None,
) -> dict[str, Any]:
    if dia >= datetime.now(_BRT).date():
        # 16/10 (revisão, aprovado): hoje ainda recebe lançamentos; o botão fechava
        # hoje e até dias futuros
        raise HTTPException(
            422, detail={"code": "so_ate_ontem", "mensagem": "só dá pra fechar até ontem"}
        )
    alvo = await _loja_ou_404(session, loja, escrita=True)
    await _travar_loja(session, alvo.integration_id)  # fechar × vincular × robô (2ª rodada)
    saldo_ant = await _saldo_anterior(session, alvo, dia)
    movs = await _movimentos_do_dia(session, alvo, dia, saldo_ant)
    pendentes = [m for m in movs if not m.pedido_bling and not m.classificacao]
    if any(aguardando_repasse(m) for m in movs):
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={
                "code": "aguardando_repasse",
                "mensagem": (
                    "a TikTok ainda não pagou este extrato; o robô fecha o dia quando pagar"
                ),
            },
        )
    if pendentes:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={
                "code": "dia_com_pendencias",
                "pendentes": len(pendentes),
                "mensagem": "toda linha precisa de um destino",
            },
        )
    totais = _totais_do_dia(movs, saldo_ant)
    motivo = (motivo or "").strip() or None
    dif = totais["diferenca"]
    if dif is not None and abs(dif) > 0.05 and not motivo:
        # 16/10 (revisão, aprovado): o robô só fecha com diferença ≤ 0,05; à mão, com o
        # saldo diferente, só com motivo (fica gravado no dia)
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail={
                "code": "saldo_nao_bate",
                "diferenca": dif,
                "mensagem": (
                    f"o saldo não bate ({dif:+.2f}); pra fechar assim mesmo, informe o motivo"
                ),
            },
        )
    agora = datetime.now(UTC)
    quem = (getattr(user, "name", None) or user.email or "?")[:120]
    row = (
        await session.execute(
            select(ConciliacaoDia).where(
                ConciliacaoDia.integration_id == alvo.integration_id, ConciliacaoDia.dia == dia
            )
        )
    ).scalar_one_or_none()
    if row is None:
        row = ConciliacaoDia(integration_id=alvo.integration_id, dia=dia)
        session.add(row)
    for k in ("saldo_inicial", "creditos", "debitos", "saldo_final", "saldo_informado"):
        setattr(row, k, Decimal(str(totais[k])) if totais[k] is not None else None)
    row.linhas = totais["linhas"]
    row.pendentes = 0
    row.fechado_em = agora
    row.fechado_por = quem
    row.fechado_motivo = motivo if (dif is not None and abs(dif) > 0.05) else None
    row.reaberto_em = row.reaberto_por = row.reaberto_motivo = None
    await session.commit()
    return {
        "dia": dia.isoformat(),
        "fechado": {"em": agora.isoformat(), "por": quem, "motivo": row.fechado_motivo},
        "totais": totais,
    }


@router.post("/dias/{dia}/reabrir")
async def reabrir_dia(
    dia: date,
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
    loja: str | None = None,
    motivo: str | None = None,
) -> dict[str, Any]:
    alvo = await _loja_ou_404(session, loja, escrita=True)
    await _travar_loja(session, alvo.integration_id)
    row = (
        await session.execute(
            select(ConciliacaoDia).where(
                ConciliacaoDia.integration_id == alvo.integration_id, ConciliacaoDia.dia == dia
            )
        )
    ).scalar_one_or_none()
    quem = (getattr(user, "name", None) or user.email or "?")[:120]
    reaberto = None
    if row is not None and row.fechado_em is not None:
        # 16/10 (revisão, aprovado): fica gravado quem reabriu — o robô não refecha
        # o que a pessoa reabriu
        row.fechado_em = None
        row.fechado_por = None
        row.fechado_motivo = None
        row.reaberto_em = datetime.now(UTC)
        row.reaberto_por = quem
        row.reaberto_motivo = (motivo or "").strip() or "reaberto à mão"
        reaberto = {"em": row.reaberto_em.isoformat(), "por": quem, "motivo": row.reaberto_motivo}
        await session.commit()
    return {"dia": dia.isoformat(), "fechado": None, "reaberto": reaberto}


# --- Resumo (tela de entrada) e Fechamento do mês ----------------------------------
# Passo 2d (14/10/2026, "Ok" do Vinicius). Resumo: uma linha por loja ligada com o
# dia de ontem (fechou? pendentes? saldo bate?), as pendências dos pedidos e o mês
# até ontem; "Onde agir hoje" é a mesma lista que o vigia vai mandar no Threema.
# Fechamento: o mês por loja — vendas (pela data do pedido), caixa (pela data do
# movimento, com "ajustes de meses anteriores"), a fatura do ML (resumo × detalhes)
# e os dias fechados.

_ETAPAS_PENDENTES = ("pendencia", "reaberto")
_ATRASO_EXTRATO = timedelta(hours=36)
# fuso como literal: com parâmetro ligado o Postgres não acha timezone(varchar, timestamptz)
_FUSO_SQL = literal_column("'America/Sao_Paulo'")
_ATRASO_UPLOAD = timedelta(days=7)  # loja por arquivo: alguém envia o CSV de vez em quando


def _mes_range(mes: str | None) -> tuple[date, date]:
    """'AAAA-MM' → (primeiro dia, primeiro dia do mês seguinte)."""
    hoje = datetime.now(_BRT).date()
    try:
        ini = date.fromisoformat(f"{mes}-01") if mes else hoje.replace(day=1)
    except ValueError as exc:
        raise HTTPException(422, "mês inválido") from exc
    fim = (
        ini.replace(year=ini.year + 1, month=1)
        if ini.month == 12
        else ini.replace(month=ini.month + 1)
    )
    return ini, fim


async def _somas_pedidos(
    session: AsyncSession, loja: ConciliacaoLoja, ini: date, fim: date
) -> dict[str, Any]:
    """Totais dos pedidos da loja vendidos em [ini, fim) — os mesmos cartões da aba Pedidos."""
    ped = ConciliacaoPedido
    conds = [
        ped.integration_id == loja.integration_id,
        ped.data_pedido >= datetime.combine(ini, datetime.min.time(), tzinfo=_BRT),
        ped.data_pedido < datetime.combine(fim, datetime.min.time(), tzinfo=_BRT),
    ]
    s = (
        await session.execute(
            select(
                func.count(),
                func.sum(ped.venda_valor),
                func.sum(ped.venda_tarifa),
                func.sum(ped.venda_frete),
                func.sum(ped.fatura_parcelamento),
                func.sum(ped.venda_cupom),
                func.sum(ped.venda_liquido),
                func.sum(ped.recebido),
                func.sum(ped.saidas_depois),
                func.sum(ped.custo),
                func.sum(ped.reembolso),
                func.sum(ped.sobrou),
                func.count(
                    sa_case((ped.liberado_em.isnot(None) & (ped.situacao != "cancelado"), 1))
                ),
                func.count(
                    sa_case((ped.liberado_em.isnot(None) & (ped.situacao == "cancelado"), 1))
                ),
            ).where(*conds)
        )
    ).one()
    por_situacao = dict(
        (
            await session.execute(
                select(ped.situacao, func.count()).where(*conds).group_by(ped.situacao)
            )
        ).all()
    )
    por_etapa = dict(
        (
            await session.execute(select(ped.etapa, func.count()).where(*conds).group_by(ped.etapa))
        ).all()
    )
    custo = s[9] or Decimal("0")
    sobrou = s[11] or Decimal("0")
    return {
        "pedidos": s[0],
        "vendido": _num(s[1] or 0),
        "tarifa": _num(s[2] or 0),
        "frete": _num(s[3] or 0),
        "parcelamento": _num(s[4] or 0),
        "cupom": _num(s[5] or 0),
        "previsto": _num(s[6] or 0),
        "recebido": _num(s[7] or 0),
        "estornos": _num(s[8] or 0),
        "custo": _num(custo),
        "reembolso": _num(s[10] or 0),
        "sobrou": _num(sobrou),
        "margem": _num((sobrou / custo).quantize(Decimal("0.0001"))) if custo > 0 else None,
        "liberados": s[12],
        "liberados_cancelados": s[13],
        "finalizados": por_etapa.get("finalizado", 0),
        "por_situacao": por_situacao,
        "por_etapa": por_etapa,
    }


async def _ajustes_anteriores(
    session: AsyncSession, loja: ConciliacaoLoja, ini: date, fim: date
) -> tuple[float, int]:
    """Dinheiro do extrato em [ini, fim) que é de pedido de ANTES de `ini` (ou
    "anterior ao início"): o que o mês carrega de meses passados. (valor, linhas)"""
    mov = ConciliacaoMovimento
    ped = ConciliacaoPedido
    ini_dt = datetime.combine(ini, datetime.min.time(), tzinfo=_BRT)
    fim_dt = datetime.combine(fim, datetime.min.time(), tzinfo=_BRT)
    valor, n = (
        await session.execute(
            select(func.sum(mov.credito - mov.debito), func.count())
            .select_from(mov)
            .outerjoin(
                ped,
                (ped.integration_id == mov.integration_id) & (ped.pedido_bling == mov.pedido_bling),
            )
            .where(
                mov.integration_id == loja.integration_id,
                mov.origem.in_(ORIGENS_EXTRATO),
                mov.data >= ini_dt,
                mov.data < fim_dt,
                mov.tipo != "saque",
                mov.tipo.not_in(tuple(TIPOS_RESERVA_SET)),
                or_(mov.classificacao == "anterior_ao_inicio", ped.data_pedido < ini_dt),
            )
        )
    ).one()
    return float(valor or 0), int(n or 0)


async def _dinheiro_da_loja(session: AsyncSession, loja: ConciliacaoLoja) -> dict[str, Any]:
    """Onde o dinheiro está agora (bloco "Dinheiro nas plataformas", 15/10): na conta
    (último saldo informado pela plataforma), preso a liberar (líquido previsto das
    vendas ainda não liberadas) e reservado em disputa (só ML: reservas abertas)."""
    mov = ConciliacaoMovimento
    ultima = (
        await session.execute(
            select(mov.data)
            .where(
                mov.integration_id == loja.integration_id,
                mov.origem.in_(ORIGENS_EXTRATO),
                mov.saldo_apos.isnot(None),
            )
            .order_by(mov.data.desc())
            .limit(1)
        )
    ).scalar_one_or_none()
    na_conta = None
    saldo_em = None
    if ultima is not None:
        # 16/10, revisão: com a âncora do dia anterior a ordem das linhas do mesmo
        # segundo é a certa — e a última linha (o saldo de agora) também
        dia_ultimo = ultima.astimezone(_BRT).date()
        saldo_ant = await _saldo_anterior(session, loja, dia_ultimo)
        movs = await _movimentos_do_dia(session, loja, dia_ultimo, saldo_ant)
        na_conta = _totais_do_dia(movs, saldo_ant)["saldo_informado"]
        saldo_em = ultima.isoformat()
    ped = ConciliacaoPedido
    a_liberar, n_liberar = (
        await session.execute(
            select(func.sum(ped.venda_liquido), func.count()).where(
                ped.integration_id == loja.integration_id,
                ped.situacao.in_(("aguardando", "vencido")),
                ped.venda_liquido.isnot(None),
            )
        )
    ).one()
    # 19/10 (revisão): toda reserva retida conta (reclamação, estorno, frete de
    # devolução, pagamento) — só a reserva pro saque fica fora (é o saque em andamento)
    reservado = (
        await session.execute(
            select(func.sum(mov.debito - mov.credito)).where(
                mov.integration_id == loja.integration_id,
                mov.origem.in_(ORIGENS_EXTRATO),
                mov.tipo.in_(_RESERVAS_RETIDAS),
            )
        )
    ).scalar_one()
    plataforma = getattr(loja.platform, "value", str(loja.platform))
    a_receber = None
    nota = None
    if plataforma == "tiktok":
        # a TikTok não guarda saldo: o extrato vai pro banco; o que ela ainda não pagou
        # (extrato PROCESSING/FAILED) é "a receber"
        na_conta = saldo_em = None
        mov_tt = ConciliacaoMovimento
        v = (
            await session.execute(
                select(func.sum(mov_tt.credito - mov_tt.debito)).where(
                    mov_tt.integration_id == loja.integration_id,
                    mov_tt.origem == ORIGEM_TIKTOK,
                    mov_tt.raw["_payment_status"].astext.in_(AGUARDANDO),
                )
            )
        ).scalar()
        a_receber = _num(v or 0)
    preso = None
    if plataforma == "amazon":
        # reserva por data de entrega: lançamentos ainda DEFERRED. Na conta = liberado
        # (saldo total − preso); a liberar = o que está preso + o que ainda não lançou
        preso_v, n_preso = (
            await session.execute(
                select(func.sum(mov.credito - mov.debito), func.count()).where(
                    mov.integration_id == loja.integration_id,
                    mov.origem.in_(ORIGENS_EXTRATO),
                    mov.tipo.in_(("venda_liberada", "estorno")),
                    mov.sub_tipo == "DEFERRED",
                )
            )
        ).one()
        preso = max(_num(preso_v or 0) or 0.0, 0.0)
        if na_conta is not None:
            na_conta = round(na_conta - preso, 2)
        # pedido cujo dinheiro está preso já entra pelo "preso": não soma de novo pela
        # venda (16/10 à noite: kia mostrava 151 mil, o certo ~75 mil)
        presos_sq = select(mov.pedido_bling).where(
            mov.integration_id == loja.integration_id,
            mov.sub_tipo == "DEFERRED",
            mov.pedido_bling.isnot(None),
        )
        a_liberar, n_liberar = (
            await session.execute(
                select(func.sum(ped.venda_liquido), func.count()).where(
                    ped.integration_id == loja.integration_id,
                    ped.situacao.in_(("aguardando", "vencido")),
                    ped.venda_liquido.isnot(None),
                    ped.pedido_bling.notin_(presos_sq),
                )
            )
        ).one()
        a_liberar = (a_liberar or 0) + Decimal(str(preso))
        n_liberar = (n_liberar or 0) + int(n_preso or 0)
        if loja.modo_extrato == "upload":
            # o arquivo de transações não traz os repasses pro banco: o saldo não fecha
            na_conta = saldo_em = None
            nota = "por arquivo: repasses pro banco não entram"
    return {
        "na_conta": na_conta,
        "saldo_em": saldo_em,
        "nota": nota,
        "a_liberar": _num(a_liberar or 0),
        "a_liberar_pedidos": int(n_liberar or 0),
        "preso": preso,
        "a_receber": a_receber,
        # reserva aberta antes do início e solta depois deixaria negativo: zero é o piso
        "reservado": max(_num(reservado or 0) or 0.0, 0.0) if plataforma == "ml" else None,
    }


def _soma_dict(destino: dict[str, Any], origem: dict[str, Any]) -> None:
    for k, v in origem.items():
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            destino[k] = (destino.get(k) or 0) + v


@router.get("/resumo")
async def resumo_geral(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    dia: date | None = None,
) -> dict[str, Any]:
    return await montar_resumo(session, dia)


def _fatura_problema(loja: ConciliacaoLoja, plataforma: str, agora: datetime) -> str | None:
    """ML: a última rodada da fatura falhou (erro guardado na loja) ou a fatura nunca
    foi lida. Texto pronto pra tela; None quando está tudo certo."""
    if plataforma != "ml":
        return None
    erro = loja.ultimo_erro or ""
    if erro.startswith("fatura:"):
        return f"fatura não lida · {erro[len('fatura:') :].strip()[:90]}"
    return None


def _janela_dia(dia: date) -> tuple[datetime, datetime]:
    ini = datetime.combine(dia, datetime.min.time(), tzinfo=_BRT)
    return ini, ini + timedelta(days=1)


_EVENTOS_DIA = (
    # chave, coluna da data, condição
    ("estornados", "ultimo_movimento_em", ("situacao", "estornado_depois")),
    ("a_menos", "liberado_em", ("situacao", "recebido_a_menos")),
    ("falta_frete", "liberado_em", ("situacao", "falta_frete")),
    ("reabertos", "etapa_em", ("etapa", "reaberto")),
    ("pendencia", "etapa_em", ("etapa", "pendencia")),
    ("vencidos", "previsao_liberacao", ("situacao", "vencido")),
    ("em_disputa", "ultimo_movimento_em", ("situacao", "em_disputa")),
)
_ROTULO_EVENTO = {
    "estornados": "estornado(s) depois",
    "a_menos": "recebido(s) a menos",
    "falta_frete": "sem frete do Melhor Envio",
    "reabertos": "reaberto(s)",
    "pendencia": "pendência(s) aberta(s)",
    "vencidos": "vencido(s)",
    "em_disputa": "em disputa (dinheiro retido)",
}


async def _eventos_do_dia(
    session: AsyncSession, loja: ConciliacaoLoja, dia: date
) -> dict[str, int]:
    """O que aconteceu com os pedidos da loja NAQUELE dia (16/10, Vinicius: a obs do
    dia é só do dia): estorno lançado, liberado a menos, reaberto, pendência aberta,
    vencido. Pendência de outro dia não aparece no dia errado."""
    ini, fim = _janela_dia(dia)
    ped = ConciliacaoPedido
    colunas = []
    for _chave, col, (campo, valor) in _EVENTOS_DIA:
        c = getattr(ped, col)
        colunas.append(
            func.sum(sa_case((sa_and(getattr(ped, campo) == valor, c >= ini, c < fim), 1), else_=0))
        )
    row = (
        await session.execute(select(*colunas).where(ped.integration_id == loja.integration_id))
    ).one()
    return {chave: int(v or 0) for (chave, _c, _cond), v in zip(_EVENTOS_DIA, row, strict=True)}


def _obs_de_eventos(ev: dict[str, int]) -> list[str]:
    return [f"{ev[k]} {_ROTULO_EVENTO[k]}" for k, _c, _cond in _EVENTOS_DIA if ev.get(k)]


async def _resumo_dos_dias(
    session: AsyncSession, lojas: list[ConciliacaoLoja], ate: date, *, n: int = 15
) -> list[dict[str, Any]]:
    """Bloco "Dias" do Resumo (16/10): um dia por linha, com quantas lojas fecharam,
    quantas não tiveram movimento, quantas esperam o repasse da TikTok e quantas ainda
    têm pendência (com o nome e o motivo), mais a obs do que aconteceu no dia."""
    de = ate - timedelta(days=n - 1)
    ini_dt, _ = _janela_dia(de)
    _, fim_dt = _janela_dia(ate)
    agora = datetime.now(UTC)
    mov = ConciliacaoMovimento
    dia_local = func.date(func.timezone(_FUSO_SQL, mov.data))
    linhas_por: dict[tuple[Any, date], tuple[int, int, int, int]] = {}
    for iid, d, total, pend, aguard, falhou in (
        await session.execute(
            select(
                mov.integration_id,
                dia_local,
                func.count(),
                func.sum(
                    sa_case(
                        (sa_and(mov.pedido_bling.is_(None), mov.classificacao.is_(None)), 1),
                        else_=0,
                    )
                ),
                func.sum(sa_case((mov.raw["_payment_status"].astext.in_(AGUARDANDO), 1), else_=0)),
                func.sum(
                    sa_case(
                        (mov.raw["_payment_status"].astext.in_(["FAILED", "SEM_REPASSE"]), 1),
                        else_=0,
                    )
                ),
            )
            .where(mov.origem.in_(ORIGENS_EXTRATO), mov.data >= ini_dt, mov.data < fim_dt)
            .group_by(mov.integration_id, dia_local)
        )
    ).all():
        linhas_por[(iid, d)] = (
            int(total or 0),
            int(pend or 0),
            int(aguard or 0),
            int(falhou or 0),
        )
    fechados = {
        (iid, d)
        for iid, d in (
            await session.execute(
                select(ConciliacaoDia.integration_id, ConciliacaoDia.dia).where(
                    ConciliacaoDia.dia >= de,
                    ConciliacaoDia.dia <= ate,
                    ConciliacaoDia.fechado_em.isnot(None),
                )
            )
        ).all()
    }
    conta_por_id = {loja.integration_id: loja.conta for loja in lojas}
    ped = ConciliacaoPedido
    eventos: dict[date, dict[str, list[tuple[str, int]]]] = {}
    for chave, col, (campo, valor) in _EVENTOS_DIA:
        c = getattr(ped, col)
        for iid, d, qtd in (
            await session.execute(
                select(
                    ped.integration_id,
                    func.date(func.timezone(_FUSO_SQL, c)),
                    func.count(),
                )
                .where(getattr(ped, campo) == valor, c >= ini_dt, c < fim_dt)
                .group_by(ped.integration_id, func.date(func.timezone(_FUSO_SQL, c)))
            )
        ).all():
            nome = conta_por_id.get(iid)
            if nome:
                eventos.setdefault(d, {}).setdefault(chave, []).append((nome, int(qtd or 0)))
    saida: list[dict[str, Any]] = []
    d = ate
    while d >= de:
        fech = sem_mov = aguardando = 0
        pendentes: list[dict[str, str]] = []
        for loja in lojas:
            if d < loja.inicio:
                continue
            if (loja.integration_id, d) in fechados:
                fech += 1
                continue
            linhas = linhas_por.get((loja.integration_id, d))
            if not linhas:
                limite = _ATRASO_UPLOAD if loja.modo_extrato == "upload" else _ATRASO_EXTRATO
                atrasado = loja.ultimo_extrato_em is None or agora - loja.ultimo_extrato_em > limite
                plat = getattr(loja.platform, "value", str(loja.platform))
                if atrasado:
                    pendentes.append({"conta": loja.conta, "motivo": "extrato não lido"})
                elif _fatura_problema(loja, plat, agora):
                    pendentes.append({"conta": loja.conta, "motivo": "fatura não lida"})
                else:
                    sem_mov += 1
                continue
            total, pend, aguard, falhou = linhas
            if pend:
                pendentes.append({"conta": loja.conta, "motivo": f"{pend} linha(s) sem dono"})
                continue
            saldo_ant = await _saldo_anterior(session, loja, d)
            movs = await _movimentos_do_dia(session, loja, d, saldo_ant)
            totais = _totais_do_dia(movs, saldo_ant)
            dif = totais["diferenca"]
            if dif is not None and abs(dif) > 0.05:
                pendentes.append({"conta": loja.conta, "motivo": f"saldo não bate ({dif:+.2f})"})
                continue
            if falhou:
                # 16/10, revisão: repasse FAILED da TikTok não é dia fechado — é pendência
                pendentes.append(
                    {"conta": loja.conta, "motivo": "repasse falhou, a TikTok tenta de novo"}
                )
                continue
            if aguard:
                aguardando += 1
                continue
            pendentes.append({"conta": loja.conta, "motivo": "pronto pra fechar"})
        obs: list[str] = []
        for chave, _c, _cond in _EVENTOS_DIA:
            lista = eventos.get(d, {}).get(chave) or []
            if lista:
                qtd = sum(q for _n, q in lista)
                nomes = [nm for nm, _q in lista]
                extra = f" (+{len(nomes) - 3})" if len(nomes) > 3 else ""
                obs.append(f"{qtd} {_ROTULO_EVENTO[chave]} ({', '.join(nomes[:3])}{extra})")
        saida.append(
            {
                "dia": d.isoformat(),
                # sem movimento (lido) e aguardando repasse contam como fechados; os dois
                # números ao lado são só o detalhe
                "fechados": fech + sem_mov + aguardando,
                "sem_movimento": sem_mov,
                "aguardando_repasse": aguardando,
                "com_pendencia": len(pendentes),
                "pendentes": pendentes,
                "obs": obs,
            }
        )
        d -= timedelta(days=1)
    return saida


async def montar_resumo(session: AsyncSession, dia: date | None = None) -> dict[str, Any]:
    """O Resumo (tela de entrada) e a lista "Onde agir hoje" — a mesma lista que o
    Vigia de Conciliação registra na Ouvidoria e manda no Threema (15/10)."""
    dia = dia or _ontem()
    agora = datetime.now(UTC)
    mes_ini = dia.replace(day=1)
    lojas_out: list[dict[str, Any]] = []
    agir: list[dict[str, Any]] = []
    por_plataforma: dict[str, dict[str, Any]] = {}
    empresa: dict[str, Any] = {"lojas": 0}
    lojas = await lojas_ativas(session)
    for loja in lojas:
        plataforma = getattr(loja.platform, "value", str(loja.platform))
        saldo_ant = await _saldo_anterior(session, loja, dia)
        movs = await _movimentos_do_dia(session, loja, dia, saldo_ant)
        pedidos_do_dia = {
            p.pedido_bling: p
            for p in (
                await session.execute(
                    select(ConciliacaoPedido).where(
                        ConciliacaoPedido.integration_id == loja.integration_id,
                        ConciliacaoPedido.pedido_bling.in_(
                            {m.pedido_bling for m in movs if m.pedido_bling} or {"-"}
                        ),
                    )
                )
            ).scalars()
        }
        contagem: dict[str, int] = {
            "automatico": 0,
            "manual": 0,
            "classificado": 0,
            "anterior": 0,
            "pendente": 0,
            "divergente": 0,
        }
        for m in movs:
            contagem[_status_movimento(m, pedidos_do_dia.get(m.pedido_bling or ""))] += 1
        totais = _totais_do_dia(movs, saldo_ant)
        fechado = (
            await session.execute(
                select(ConciliacaoDia).where(
                    ConciliacaoDia.integration_id == loja.integration_id, ConciliacaoDia.dia == dia
                )
            )
        ).scalar_one_or_none()
        ped = ConciliacaoPedido
        base = [ped.integration_id == loja.integration_id]
        por_etapa = dict(
            (
                await session.execute(
                    select(ped.etapa, func.count()).where(*base).group_by(ped.etapa)
                )
            ).all()
        )
        # 16/10 (revisão): só o que ainda está em aberto — pedido finalizado com motivo
        # continuava cobrando no "Onde agir" (kia 302142)
        por_situacao = dict(
            (
                await session.execute(
                    select(ped.situacao, func.count())
                    .where(*base, ped.etapa.in_(_ETAPAS_PENDENTES))
                    .group_by(ped.situacao)
                )
            ).all()
        )
        mes = await _somas_pedidos(session, loja, mes_ini, dia + timedelta(days=1))
        mes["ajustes_meses_anteriores"], mes["ajustes_qtd"] = await _ajustes_anteriores(
            session, loja, mes_ini, dia + timedelta(days=1)
        )
        limite_atraso = _ATRASO_UPLOAD if loja.modo_extrato == "upload" else _ATRASO_EXTRATO
        extrato_atrasado = dia >= loja.inicio and (
            loja.ultimo_extrato_em is None or agora - loja.ultimo_extrato_em > limite_atraso
        )
        saldo_bate = totais["diferenca"] is not None and abs(totais["diferenca"]) <= 0.05
        sem_movimento = totais["linhas"] == 0
        nao_pagos = [m for m in movs if aguardando_repasse(m)]
        repasse_falhou = any(
            str((m.raw or {}).get("_payment_status") or "").upper() in ("FAILED", "SEM_REPASSE")
            for m in nao_pagos
        )
        # 16/10 (Vinicius): fatura do ML com erro não pode ficar escondida atrás de um
        # "fechado" verde — a loja aparece "fatura não lida" até a leitura passar
        fatura_problema = _fatura_problema(loja, plataforma, agora)
        if fatura_problema and not (fechado and fechado.fechado_em):
            # dia já fechado continua "fechado" (2ª rodada); o aviso fica em problemas
            dia_status, dia_texto = "fatura_nao_lida", "fatura não lida"
        elif fechado and fechado.fechado_em:
            dia_status, dia_texto = "fechado", f"fechado por {fechado.fechado_por}"
        elif sem_movimento and extrato_atrasado:
            dia_status, dia_texto = "extrato_nao_lido", "extrato não lido"
        elif sem_movimento:
            dia_status, dia_texto = "sem_movimento", "sem movimento"
        elif contagem["pendente"] == 0 and saldo_bate and repasse_falhou:
            dia_status, dia_texto = "repasse_falhou", "repasse falhou, a TikTok tenta de novo"
        elif contagem["pendente"] == 0 and saldo_bate and nao_pagos:
            dia_status, dia_texto = (
                "aguardando_repasse",
                "liquidado, aguardando o repasse da TikTok",
            )
        elif contagem["pendente"] == 0 and saldo_bate:
            dia_status, dia_texto = "pronto", "100% casado, saldo bate — pode fechar"
        else:
            dia_status, dia_texto = "aberto", "não fechou"
        reabertos = por_etapa.get("reaberto", 0)
        pendencias = por_etapa.get("pendencia", 0)
        vencidos = por_situacao.get("vencido", 0)
        a_menos = por_situacao.get("recebido_a_menos", 0)
        a_mais = por_situacao.get("recebido_a_mais", 0)
        sem_venda = por_situacao.get("sem_venda", 0)
        fora_da_loja = por_situacao.get("fora_da_loja", 0)
        sem_frete = por_situacao.get("falta_frete", 0)
        estornados = por_situacao.get("estornado_depois", 0)
        em_disputa = por_situacao.get("em_disputa", 0)
        problemas: list[str] = []
        if contagem["pendente"]:
            problemas.append(f"{contagem['pendente']} linha(s) sem dono")
        if totais["diferenca"] is not None and not saldo_bate:
            problemas.append(f"saldo não bate ({totais['diferenca']:+.2f})")
        if extrato_atrasado:
            problemas.append("extrato não lido")
        if fatura_problema:
            problemas.append(fatura_problema)
        if reabertos:
            problemas.append(f"{reabertos} reaberto(s)")
        if vencidos:
            problemas.append(f"{vencidos} sem liberação vencido(s)")
        if a_menos:
            problemas.append(f"{a_menos} recebido(s) a menos")
        if a_mais:
            problemas.append(f"{a_mais} recebido(s) a mais")
        if sem_venda:
            problemas.append(f"{sem_venda} venda(s) não lida(s) com dinheiro")
        if fora_da_loja:
            problemas.append(f"{fora_da_loja} pedido(s) que saíram da loja")
        if estornados:
            problemas.append(f"{estornados} estornado(s) depois")
        if em_disputa:
            problemas.append(f"{em_disputa} em disputa (dinheiro retido)")
        if loja.ultimo_erro:
            problemas.append(f"erro: {loja.ultimo_erro[:80]}")
        # Só o dia (16/10, Vinicius): o que trava o dia e o que aconteceu nele
        problemas_dia: list[str] = []
        if contagem["pendente"]:
            problemas_dia.append(f"{contagem['pendente']} linha(s) sem dono")
        if totais["diferenca"] is not None and not saldo_bate:
            problemas_dia.append(f"saldo não bate ({totais['diferenca']:+.2f})")
        if extrato_atrasado:
            problemas_dia.append("extrato não lido")
        if fatura_problema:
            problemas_dia.append(fatura_problema)
        if loja.ultimo_erro and not (loja.ultimo_erro.startswith("fatura:") and fatura_problema):
            problemas_dia.append(f"erro: {loja.ultimo_erro[:80]}")
        if fechado and not fechado.fechado_em and fechado.reaberto_motivo:
            problemas_dia.append(
                f"dia reaberto por {fechado.reaberto_por}: {fechado.reaberto_motivo}"
            )
        eventos_dia = await _eventos_do_dia(session, loja, dia)
        obs_dia = _obs_de_eventos(eventos_dia)
        dinheiro = await _dinheiro_da_loja(session, loja)
        linha = {
            "conta": loja.conta,
            "platform": plataforma,
            "inicio": loja.inicio.isoformat(),
            "modo_extrato": loja.modo_extrato,
            "dinheiro": dinheiro,
            "dia": {
                "status": dia_status,
                "texto": dia_texto,
                "linhas": totais["linhas"],
                "pendentes": contagem["pendente"],
                "divergentes": contagem["divergente"],
                "saldo_bate": saldo_bate if totais["linhas"] else None,
                "diferenca": totais["diferenca"],
                "fechado_em": fechado.fechado_em.isoformat()
                if fechado and fechado.fechado_em
                else None,
                "fechado_por": fechado.fechado_por if fechado else None,
            },
            "pedidos": {
                "pendencia": pendencias,
                "reabertos": reabertos,
                "aguardando": por_etapa.get("aguardando", 0),
                "finalizados": por_etapa.get("finalizado", 0),
                "conciliados": por_etapa.get("conciliado", 0),
                "vencidos": vencidos,
                "recebido_a_menos": a_menos,
                "recebido_a_mais": a_mais,
                "sem_venda": sem_venda,
                "fora_da_loja": fora_da_loja,
                "estornados": estornados,
                "em_disputa": em_disputa,
            },
            "leituras": {
                "extrato_em": loja.ultimo_extrato_em.isoformat()
                if loja.ultimo_extrato_em
                else None,
                "fatura_em": loja.ultimo_fatura_em.isoformat() if loja.ultimo_fatura_em else None,
                "recalculo_em": loja.ultimo_recalculo_em.isoformat()
                if loja.ultimo_recalculo_em
                else None,
                "extrato_atrasado": extrato_atrasado,
                "erro": loja.ultimo_erro,
            },
            "mes": mes,
            "eventos_dia": eventos_dia,
            "obs_dia": obs_dia,
            "problemas_dia": problemas_dia,
            "situacao": " · ".join(problemas)
            if problemas
            else ("dia fechado" if dia_status == "fechado" else "tudo certo"),
            "gravidade": "bad"
            if (contagem["pendente"] or extrato_atrasado or (totais["linhas"] and not saldo_bate))
            else ("warn" if problemas else "ok"),
        }
        lojas_out.append(linha)
        # Onde agir hoje
        if contagem["pendente"] or (totais["linhas"] and not saldo_bate):
            agir.append(
                {
                    "gravidade": "bad",
                    "loja": loja.conta,
                    "titulo": " · ".join(
                        [loja.conta]
                        + (
                            [f"{contagem['pendente']} linha(s) sem dono"]
                            if contagem["pendente"]
                            else []
                        )
                        + (
                            []
                            if saldo_bate or totais["diferenca"] is None
                            else [f"saldo não bate ({totais['diferenca']:+.2f})"]
                        )
                    ),
                    "detalhe": (
                        f"extrato de {dia:%d/%m}: {totais['linhas']} movimentos, "
                        f"{contagem['pendente']} pendentes"
                    ),
                    "acao": "extrato",
                    "dia": dia.isoformat(),
                }
            )
        if fatura_problema:
            agir.append(
                {
                    "gravidade": "bad",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · fatura do ML não lida",
                    "detalhe": fatura_problema,
                    "acao": "sincronizar",
                    "dia": dia.isoformat(),
                }
            )
        if extrato_atrasado:
            agir.append(
                {
                    "gravidade": "bad",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · extrato não lido",
                    "detalhe": (
                        f"última leitura {loja.ultimo_extrato_em.astimezone(_BRT):%d/%m %H:%M}"
                        if loja.ultimo_extrato_em
                        else "nunca lido"
                    )
                    + (f" · {loja.ultimo_erro[:80]}" if loja.ultimo_erro else ""),
                    "acao": "sincronizar",
                    "dia": dia.isoformat(),
                }
            )
        if reabertos:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {reabertos} pedido(s) reaberto(s)",
                    "detalhe": "chegou movimento novo depois de finalizado",
                    "acao": "pedidos",
                    "filtro": {"etapa": "reaberto"},
                }
            )
        if vencidos:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {vencidos} venda(s) sem liberação, prazo vencido",
                    "detalhe": "abrir chamado de repasse",
                    "acao": "pedidos",
                    "filtro": {"situacao": "vencido", "etapa": "pendencia,reaberto"},
                }
            )
        if sem_frete:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": (
                        f"{loja.conta} · {sem_frete} envio(s) próprio(s) "
                        "sem o frete do Melhor Envio"
                    ),
                    "detalhe": "o Melhor Envio ainda não trouxe o envio (ou o token está fora)",
                    "acao": "pedidos",
                    "filtro": {"situacao": "falta_frete", "etapa": "pendencia,reaberto"},
                }
            )
        if a_menos:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {a_menos} recebido(s) a menos",
                    "detalhe": "conferir o motivo e finalizar com motivo, ou abrir chamado",
                    "acao": "pedidos",
                    "filtro": {"situacao": "recebido_a_menos", "etapa": "pendencia,reaberto"},
                }
            )
        if estornados:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {estornados} estornado(s) depois de liberar",
                    "detalhe": "mediação, estorno ou chargeback depois do dinheiro entrar",
                    "acao": "pedidos",
                    "filtro": {"situacao": "estornado_depois", "etapa": "pendencia,reaberto"},
                }
            )
        if em_disputa:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {em_disputa} em disputa (dinheiro retido)",
                    "detalhe": (
                        "liberou, mas a plataforma reteve parte em reserva por reclamação/estorno; "
                        "acompanhar o chamado até a reserva voltar (ou virar mediação)"
                    ),
                    "acao": "pedidos",
                    "filtro": {"situacao": "em_disputa", "etapa": "pendencia,reaberto"},
                }
            )
        # 16/10 (revisão): "recebido a mais", "venda não lida com dinheiro" e "saiu da
        # loja" não geravam item — 8 pendências da Poofy (ml) invisíveis
        if a_mais:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {a_mais} recebido(s) a mais",
                    "detalhe": "entrou mais que o previsto: conferir a venda, finalizar com motivo",
                    "acao": "pedidos",
                    "filtro": {"situacao": "recebido_a_mais", "etapa": "pendencia,reaberto"},
                }
            )
        if sem_venda:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {sem_venda} venda(s) não lida(s) com dinheiro",
                    "detalhe": "entrou dinheiro, mas a Margem ainda não leu a venda",
                    "acao": "pedidos",
                    "filtro": {"situacao": "sem_venda", "etapa": "pendencia,reaberto"},
                }
            )
        if fora_da_loja:
            agir.append(
                {
                    "gravidade": "warn",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · {fora_da_loja} pedido(s) que saíram da loja",
                    "detalhe": "apagado ou mudou de loja no Bling: finalizar ou religar as linhas",
                    "acao": "pedidos",
                    "filtro": {"situacao": "fora_da_loja", "etapa": "pendencia,reaberto"},
                }
            )
        if dia_status == "pronto":
            agir.append(
                {
                    "gravidade": "ok",
                    "loja": loja.conta,
                    "titulo": f"{loja.conta} · dia {dia:%d/%m} pronto pra fechar",
                    "detalhe": "100% das linhas com destino e saldo batendo",
                    "acao": "extrato",
                    "dia": dia.isoformat(),
                }
            )
        plat = por_plataforma.setdefault(plataforma, {"lojas": 0})
        plat["lojas"] += 1
        for bloco_d in (plat, empresa):
            d = bloco_d.setdefault(
                "dinheiro",
                {"na_conta": 0.0, "a_liberar": 0.0, "a_liberar_pedidos": 0, "reservado": 0.0},
            )
            d["na_conta"] = round(d["na_conta"] + (dinheiro["na_conta"] or 0), 2)
            d["a_liberar"] = round(d["a_liberar"] + (dinheiro["a_liberar"] or 0), 2)
            d["a_liberar_pedidos"] += dinheiro["a_liberar_pedidos"]
            d["reservado"] = round(d["reservado"] + (dinheiro["reservado"] or 0), 2)
        _soma_dict(
            plat, {k: v for k, v in mes.items() if k not in ("margem", "por_situacao", "por_etapa")}
        )
        empresa["lojas"] += 1
        _soma_dict(
            empresa,
            {k: v for k, v in mes.items() if k not in ("margem", "por_situacao", "por_etapa")},
        )
    for bloco in (*por_plataforma.values(), empresa):
        custo = bloco.get("custo") or 0
        bloco["margem"] = round((bloco.get("sobrou") or 0) / custo, 4) if custo else None
    ordem = {"bad": 0, "warn": 1, "ok": 2}
    agir.sort(key=lambda x: (ordem.get(x["gravidade"], 9), x["loja"]))
    return {
        "dia": dia.isoformat(),
        "mes": mes_ini.isoformat()[:7],
        "lojas": lojas_out,
        # a lista fica fixa nos últimos 15 dias até ontem (clicar num dia só troca a
        # tabela de cima; antes cada clique empurrava a lista pro passado — 16/10)
        "dias": await _resumo_dos_dias(session, lojas, max(dia, _ontem())),
        "agir": agir,
        "por_plataforma": por_plataforma,
        "empresa": empresa,
        "cartoes": {
            "lojas": len(lojas_out),
            # 16/10 (Vinicius): sem movimento com extrato lido e "aguardando o repasse"
            # da TikTok contam como dia fechado (verde, com obs)
            "dias_fechados": sum(
                1
                for x in lojas_out
                if x["dia"]["status"] in ("fechado", "sem_movimento", "aguardando_repasse")
            ),
            "dias_prontos": sum(1 for x in lojas_out if x["dia"]["status"] == "pronto"),
            "pendentes": sum(x["dia"]["pendentes"] for x in lojas_out),
            "reabertos": sum(x["pedidos"]["reabertos"] for x in lojas_out),
            "vencidos": sum(x["pedidos"]["vencidos"] for x in lojas_out),
            "extrato_atrasado": sum(
                1
                for x in lojas_out
                if x["leituras"]["extrato_atrasado"] or x["dia"]["status"] == "fatura_nao_lida"
            ),
            "saldo_bate": sum(1 for x in lojas_out if x["dia"]["saldo_bate"]),
            "com_movimento": sum(1 for x in lojas_out if x["dia"]["linhas"]),
        },
    }


_TIPOS_CAIXA = {
    "venda_liberada": "liberado",
    "mediacao": "mediacoes",
    "mediacao_cancelada": "mediacoes_canceladas",
    "estorno": "estornos",
    "chargeback": "chargebacks",
    "frete": "frete",
    "frete_cancelado": "frete",
    "frete_devolucao": "frete",
    "saque": "saques",
}


@router.get("/fechamento")
async def fechamento_do_mes(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    mes: Annotated[str | None, Query(pattern=r"^\d{4}-\d{2}$")] = None,
) -> dict[str, Any]:
    alvo = await _loja_ou_404(session, loja)
    ini, fim_mes = _mes_range(mes)
    # 16/10, revisão: o mês corrente para em ontem (hoje ainda está mudando) — vendas,
    # caixa e dias no MESMO corte; antes o caixa ia até hoje contra o saldo de ontem
    fim = max(ini, min(fim_mes, _ontem() + timedelta(days=1)))
    ini_dt = datetime.combine(ini, datetime.min.time(), tzinfo=_BRT)
    fim_dt = datetime.combine(fim, datetime.min.time(), tzinfo=_BRT)

    vendas = await _somas_pedidos(session, alvo, ini, fim)

    # Caixa: pela data do movimento, reservas de fora, com "ajustes de meses anteriores".
    mov = ConciliacaoMovimento
    pedd = ConciliacaoPedido
    rows = (
        await session.execute(
            select(
                mov.tipo,
                mov.classificacao,
                pedd.data_pedido,
                pedd.situacao,
                func.count(),
                func.sum(mov.credito),
                func.sum(mov.debito),
            )
            .select_from(mov)
            .outerjoin(
                pedd,
                (pedd.integration_id == mov.integration_id)
                & (pedd.pedido_bling == mov.pedido_bling),
            )
            .where(
                mov.integration_id == alvo.integration_id,
                mov.origem.in_(ORIGENS_EXTRATO),
                mov.data >= ini_dt,
                mov.data < fim_dt,
            )
            .group_by(mov.tipo, mov.classificacao, pedd.data_pedido, pedd.situacao)
        )
    ).all()
    caixa: dict[str, Any] = {
        "liberado": 0.0,
        "mediacoes": 0.0,
        "mediacoes_canceladas": 0.0,
        "estornos": 0.0,
        "chargebacks": 0.0,
        "frete": 0.0,
        "saques": 0.0,
        "outros": 0.0,
        "reservas_liquido": 0.0,
        "ajustes_meses_anteriores": 0.0,
        "ajustes_qtd": 0,
        # Cancelado por reclamação: o ML libera a venda e tira no mesmo dia
        # (liberação + mediação + estorno da tarifa = zero). Fica à parte pra
        # não parecer dinheiro que entrou.
        "cancelados_entrou": 0.0,
        "cancelados_saiu": 0.0,
        "cancelados_liquido": 0.0,
        "cancelados_linhas": 0,
        "linhas": 0,
    }
    for tipo, classificacao, data_pedido, situacao_ped, n, c, d in rows:
        c = float(c or 0)
        d = float(d or 0)
        caixa["linhas"] += n
        if tipo in TIPOS_RESERVA_SET:
            caixa["reservas_liquido"] += c - d
            continue
        anterior = classificacao == "anterior_ao_inicio" or (
            data_pedido is not None and data_pedido < ini_dt
        )
        if anterior and tipo != "saque":
            caixa["ajustes_meses_anteriores"] += c - d
            caixa["ajustes_qtd"] += n
            continue
        if situacao_ped == "cancelado" and tipo != "saque":
            caixa["cancelados_entrou"] += c
            caixa["cancelados_saiu"] += d
            caixa["cancelados_liquido"] += c - d
            caixa["cancelados_linhas"] += n
            continue
        chave = _TIPOS_CAIXA.get(tipo, "outros")
        caixa[chave] += c - d
    caixa = {k: (round(v, 2) if isinstance(v, float) else v) for k, v in caixa.items()}
    # O que entrou de verdade no mês (antes de tirar): vendas liberadas, menos o que
    # saiu por mediação/estorno/chargeback/frete, mais ajustes de meses anteriores.
    caixa["resultado_caixa"] = round(
        caixa["liberado"]
        + caixa["mediacoes"]
        + caixa["mediacoes_canceladas"]
        + caixa["estornos"]
        + caixa["chargebacks"]
        + caixa["frete"]
        + caixa["outros"]
        + caixa["cancelados_liquido"]
        + caixa["ajustes_meses_anteriores"],
        2,
    )
    # Tudo somado (com reservas e saques) tem que ser saldo final − saldo inicial.
    caixa["movimento_liquido"] = round(
        caixa["resultado_caixa"] + caixa["reservas_liquido"] + caixa["saques"], 2
    )

    # Dias do mês: linhas, pendentes, fechado — e o saldo do primeiro e do último dia.
    dias_resp = await listar_dias(
        session, _user, loja=alvo.conta, de=ini, ate=fim - timedelta(days=1)
    )
    dias = dias_resp["dias"]
    saldo_inicial = saldo_final = None
    if dias:
        # 16/10, revisão: a mesma âncora do Extrato (`_saldo_anterior`) — o saldo inicial
        # é o fechamento da véspera e a ordem das linhas do mesmo segundo é a certa
        d_ini = date.fromisoformat(dias[0]["dia"])
        d_fim = date.fromisoformat(dias[-1]["dia"])
        anc_ini = await _saldo_anterior(session, alvo, d_ini)
        anc_fim = await _saldo_anterior(session, alvo, d_fim)
        primeiro = await _movimentos_do_dia(session, alvo, d_ini, anc_ini)
        ultimo = await _movimentos_do_dia(session, alvo, d_fim, anc_fim)
        saldo_inicial = _totais_do_dia(primeiro, anc_ini)["saldo_inicial"]
        saldo_final = _totais_do_dia(ultimo, anc_fim)["saldo_informado"]

    # Fatura do ML: resumo do período (chave AAAA-MM-01) × soma dos detalhes lidos.
    chave = ini.strftime("%Y-%m-01")
    sync = (
        await session.execute(
            select(ConciliacaoSync).where(
                ConciliacaoSync.integration_id == alvo.integration_id,
                ConciliacaoSync.origem == "ml_fatura",
                ConciliacaoSync.chave == f"{chave}|resumo",
            )
        )
    ).scalar_one_or_none()
    resumo_ml = ((sync.detalhe or {}).get("resumo") if sync else None) or {}
    periodo = resumo_ml.get("period") or {}
    bill = resumo_ml.get("bill_includes") or {}
    fatura: dict[str, Any] = {
        "periodo": {
            "chave": chave,
            "de": periodo.get("date_from"),
            "ate": periodo.get("date_to"),
            "expira": periodo.get("expiration_date"),
        },
        "resumo_cobrancas": [
            {"tipo": c.get("type"), "label": c.get("label"), "valor": _num(c.get("amount"))}
            for c in (bill.get("charges") or [])
            if isinstance(c, dict)
        ],
        "resumo_bonificacoes": [
            {"tipo": c.get("type"), "label": c.get("label"), "valor": _num(c.get("amount"))}
            for c in (bill.get("bonuses") or [])
            if isinstance(c, dict)
        ],
        "resumo_total": _num(bill.get("total_amount")),
        "pago": resumo_ml.get("payment_collected") or {},
        "detalhes_por_sub_tipo": [],
        "detalhes_cobrancas": 0.0,
        "detalhes_bonificacoes": 0.0,
        "bonificacoes_mesma_fatura": 0.0,
        "bonificacoes_de_faturas_anteriores": 0.0,
        "bonificacoes_anteriores_linhas": 0,
        "lido_em": sync.lido_em.isoformat() if sync and sync.lido_em else None,
    }
    if periodo.get("date_from") and periodo.get("date_to"):
        p_ini = datetime.combine(
            date.fromisoformat(periodo["date_from"]), datetime.min.time(), tzinfo=_BRT
        )
        p_fim = datetime.combine(
            date.fromisoformat(periodo["date_to"]), datetime.min.time(), tzinfo=_BRT
        ) + timedelta(days=1)
        det = (
            await session.execute(
                select(
                    mov.sub_tipo,
                    mov.tipo,
                    func.count(),
                    func.sum(mov.debito),
                    func.sum(mov.credito),
                )
                .where(
                    mov.integration_id == alvo.integration_id,
                    mov.origem == "ml_fatura",
                    mov.data >= p_ini,
                    mov.data < p_fim,
                )
                .group_by(mov.sub_tipo, mov.tipo)
                .order_by(func.sum(mov.debito).desc())
            )
        ).all()
        fatura["detalhes_por_sub_tipo"] = [
            {
                "sub_tipo": st,
                "tipo": t,
                "linhas": n,
                "cobrado": _num(d or 0),
                "bonificado": _num(c or 0),
            }
            for st, t, n, d, c in det
        ]
        fatura["detalhes_cobrancas"] = round(sum(float(d or 0) for _, _, _, d, _ in det), 2)
        fatura["detalhes_bonificacoes"] = round(sum(float(c or 0) for _, _, _, _, c in det), 2)
        # Bonificação: o resumo do ML só conta a que devolve cobrança da MESMA fatura
        # (confirmado 15/10 na jlas2: 193 linhas = 11.131,25 = resumo). A que devolve
        # cobrança de fatura anterior é crédito real, mas fica fora do resumo.
        datas_cobranca = dict(
            (
                await session.execute(
                    select(mov.id_externo, mov.data).where(
                        mov.integration_id == alvo.integration_id,
                        mov.origem == "ml_fatura",
                        mov.debito > 0,
                    )
                )
            ).all()
        )
        bonus = (
            await session.execute(
                select(mov.credito, mov.raw).where(
                    mov.integration_id == alvo.integration_id,
                    mov.origem == "ml_fatura",
                    mov.credito > 0,
                    mov.data >= p_ini,
                    mov.data < p_fim,
                )
            )
        ).all()
        mesma = anteriores = 0.0
        n_ant = 0
        for cred, raw in bonus:
            orig = str(((raw or {}).get("charge_info") or {}).get("charge_bonified_id") or "")
            dt_orig = datas_cobranca.get(orig)
            if dt_orig is not None and p_ini <= dt_orig < p_fim:
                mesma += float(cred or 0)
            else:
                anteriores += float(cred or 0)
                n_ant += 1
        fatura["bonificacoes_mesma_fatura"] = round(mesma, 2)
        fatura["bonificacoes_de_faturas_anteriores"] = round(anteriores, 2)
        fatura["bonificacoes_anteriores_linhas"] = n_ant
    soma_resumo = round(sum(x["valor"] or 0 for x in fatura["resumo_cobrancas"]), 2)
    soma_resumo_bonus = round(sum(abs(x["valor"] or 0) for x in fatura["resumo_bonificacoes"]), 2)
    fatura["bate_bonificacoes"] = (
        (abs(soma_resumo_bonus - fatura.get("bonificacoes_mesma_fatura", 0.0)) <= 1.0)
        if fatura["resumo_bonificacoes"]
        else None
    )
    fatura["bate"] = (
        (abs(soma_resumo - fatura["detalhes_cobrancas"]) <= 1.0)
        if fatura["resumo_cobrancas"]
        else None
    )

    bate = (
        abs(saldo_inicial + caixa["movimento_liquido"] - saldo_final) <= 0.05
        if (saldo_inicial is not None and saldo_final is not None)
        else None
    )
    return {
        "loja": alvo.conta,
        "mes": ini.isoformat()[:7],
        "de": ini.isoformat(),
        "ate": (fim_mes - timedelta(days=1)).isoformat(),
        # até onde as contas foram (mês corrente: ontem)
        "ate_corte": (fim - timedelta(days=1)).isoformat(),
        "vendas": vendas,
        "caixa": {
            **caixa,
            "saldo_inicial": saldo_inicial,
            "saldo_final": saldo_final,
            "bate": bate,
        },
        "dias": dias,
        "dias_fechados": sum(1 for d in dias if d["fechado_em"]),
        "fatura": fatura,
    }


@router.post("/fechamento.xlsx")
async def fechamento_xlsx(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
    loja: str | None = None,
    mes: Annotated[str | None, Query(pattern=r"^\d{4}-\d{2}$")] = None,
) -> StreamingResponse:
    """O Fechamento do mês em Excel: Vendas, Caixa, Fatura e Dias (mesmos números da tela)."""
    from openpyxl import Workbook

    f = await fechamento_do_mes(session, _user, loja=loja, mes=mes)
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    _linha_excel(ws, [f"Vendas de {f['mes']} · {f['loja']} · pela data do pedido", ""])
    v = f["vendas"]
    for rotulo, chave in (
        ("Pedidos", "pedidos"),
        ("Vendido (preço da venda)", "vendido"),
        ("Tarifa de venda", "tarifa"),
        ("Frete (parte nossa)", "frete"),
        ("Parcelamento (fatura)", "parcelamento"),
        ("Cupons", "cupom"),
        ("Líquido previsto", "previsto"),
        ("Recebido", "recebido"),
        ("Liberados (e ficou)", "liberados"),
        ("Liberados e devolvidos (cancelados)", "liberados_cancelados"),
        ("Saiu depois (mediação, estorno)", "estornos"),
        ("Custo dos produtos", "custo"),
        ("Reembolsos", "reembolso"),
        ("Sobrou (pedido sem liberação entra pelo previsto)", "sobrou"),
        ("Margem sobre o custo", "margem"),
        ("Finalizados", "finalizados"),
    ):
        _linha_excel(ws, [rotulo, v.get(chave)])
    _linha_excel(ws, [])
    _linha_excel(ws, ["Por situação", "pedidos"])
    for k, n in sorted(v["por_situacao"].items()):
        _linha_excel(ws, [k, n])
    _linha_excel(ws, [])
    _linha_excel(ws, ["Por etapa", "pedidos"])
    for k, n in sorted(v["por_etapa"].items()):
        _linha_excel(ws, [k, n])

    ws = wb.create_sheet("Caixa")
    c = f["caixa"]
    _linha_excel(ws, [f"Caixa de {f['mes']} · pela data do movimento ({c['linhas']} linhas)", ""])
    for rotulo, chave in (
        ("Saldo inicial", "saldo_inicial"),
        ("Vendas liberadas", "liberado"),
        ("Mediações (reclamação perdida)", "mediacoes"),
        ("Mediações canceladas", "mediacoes_canceladas"),
        ("Estornos", "estornos"),
        ("Chargebacks", "chargebacks"),
        ("Frete (cobrado / devolvido)", "frete"),
        ("Outros", "outros"),
        ("Cancelados por reclamação: entrou", "cancelados_entrou"),
        ("Cancelados por reclamação: saiu", "cancelados_saiu"),
        ("Cancelados por reclamação: líquido", "cancelados_liquido"),
        ("Ajustes de meses anteriores", "ajustes_meses_anteriores"),
        ("Ajustes de meses anteriores (linhas)", "ajustes_qtd"),
        ("= Resultado de caixa do mês", "resultado_caixa"),
        ("Reservas (líquido)", "reservas_liquido"),
        ("Saques", "saques"),
        ("= Movimento líquido", "movimento_liquido"),
        ("Saldo final informado", "saldo_final"),
        ("Prova dos nove bate?", "bate"),
    ):
        _linha_excel(ws, [rotulo, c.get(chave)])

    ws = wb.create_sheet("Fatura ML")
    fa = f["fatura"]
    per = fa["periodo"]
    _linha_excel(
        ws, [f"Fatura do ML · período {per['chave']} ({per['de']} a {per['ate']})", "", ""]
    )
    _linha_excel(ws, ["O que o ML diz (resumo)", "tipo", "valor"])
    for x in fa["resumo_cobrancas"]:
        _linha_excel(ws, [x["label"], x["tipo"], -(x["valor"] or 0)])
    for x in fa["resumo_bonificacoes"]:
        _linha_excel(ws, [x["label"], x["tipo"], abs(x["valor"] or 0)])
    _linha_excel(ws, ["Total da fatura (resumo)", "", fa["resumo_total"]])
    _linha_excel(ws, [])
    _linha_excel(ws, ["O que lemos linha a linha", "sub-tipo", "linhas", "cobrado", "bonificado"])
    for d in fa["detalhes_por_sub_tipo"]:
        _linha_excel(ws, [d["tipo"], d["sub_tipo"], d["linhas"], d["cobrado"], d["bonificado"]])
    _linha_excel(ws, ["Cobranças: detalhes", "", "", fa["detalhes_cobrancas"], ""])
    _linha_excel(ws, ["Bonificações: detalhes", "", "", "", fa["detalhes_bonificacoes"]])
    _linha_excel(
        ws,
        ["  das quais devolvem cobrança desta fatura", "", "", "", fa["bonificacoes_mesma_fatura"]],
    )
    _linha_excel(
        ws,
        [
            "  das quais devolvem cobrança de fatura anterior",
            "",
            fa["bonificacoes_anteriores_linhas"],
            "",
            fa["bonificacoes_de_faturas_anteriores"],
        ],
    )
    _linha_excel(ws, ["Resumo = detalhes (cobranças)?", "", fa["bate"]])
    _linha_excel(
        ws, ["Resumo = detalhes (bonificações da mesma fatura)?", "", fa["bate_bonificacoes"]]
    )

    ws = wb.create_sheet("Dias")
    _linha_excel(
        ws, ["Dia", "Linhas", "Sem dono", "Créditos", "Débitos", "Fechado em", "Fechado por"]
    )
    for d in f["dias"]:
        _linha_excel(
            ws,
            [
                d["dia"],
                d["linhas"],
                d["pendentes"],
                d["creditos"],
                d["debitos"],
                d["fechado_em"],
                d["fechado_por"],
            ],
        )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"fechamento_{f['loja']}_{f['mes']}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


# --- Sondagens (só leitura) ---------------------------------------------------------
# 15/10/2026, madrugada ("Faz o 6 ... pode fazer a leitura começar a sondar"):
# (a) por que as bonificações do resumo do ML não batem com a soma dos detalhes;
# (b) o que a API da Shopee dá de extrato/repasse, numa conta ativa qualquer.


def _soma_por_sub(detalhes: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for det in detalhes:
        ci = det.get("charge_info") or {}
        sub = str(ci.get("detail_sub_type") or "?")
        x = out.setdefault(sub, {"linhas": 0, "valor": 0.0})
        x["linhas"] += 1
        x["valor"] = round(x["valor"] + float(ci.get("detail_amount") or 0), 2)
    return out


@router.get("/sonda/bonus")
async def sonda_bonus(
    conta: Annotated[str, Query(min_length=1)],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    periodo: Annotated[str | None, Query(pattern=r"^\d{4}-\d{2}-01$")] = None,
    so_banco: bool = False,
) -> dict[str, Any]:
    """Bonificações do período: o resumo do ML × os detalhes do ML (lidos agora) × o que
    está gravado. Duas chamadas ao ML (summary/details e a página de BONUS). Só leitura.
    `so_banco=true` não chama o ML: só cruza o gravado (bonificação × cobrança original)."""
    from app.services.conciliacao_fatura import ler_pagina, ler_resumo

    chave = periodo or datetime.now(_BRT).date().strftime("%Y-%m-01")
    if so_banco:
        return await _bonus_so_banco(session, conta, chave)
    cliente = await _cliente_da_conta(session, conta)
    # O balde de 5/min do ML é dividido com o robô da Margem (:15/:45): no 429,
    # espera um minuto e tenta de novo, como a leitura da fatura faz.
    resumo: dict[str, Any] = {}
    for tentativa in range(4):
        try:
            resumo = await ler_resumo(cliente, chave)
            break
        except RuntimeError as e:
            if "429" not in str(e) or tentativa == 3:
                raise
            await asyncio.sleep(65)
    bill = resumo.get("bill_includes") or {}
    per = resumo.get("period") or {}
    detalhes: list[dict[str, Any]] = []
    from_id = 0
    for _ in range(3):
        pagina = await ler_pagina(cliente, chave, "BONUS", from_id, 1000)
        results = [x for x in (pagina.get("results") or []) if isinstance(x, dict)]
        detalhes.extend(results)
        from_id = int(pagina.get("last_id") or 0)
        if len(results) < 1000 or not from_id:
            break
    ml_por_sub = _soma_por_sub(detalhes)
    resumo_por_tipo = {
        str(b.get("type")): round(abs(float(b.get("amount") or 0)), 2)
        for b in (bill.get("bonuses") or [])
        if isinstance(b, dict)
    }
    # campos que podem explicar a diferença
    campos: dict[str, Counter] = defaultdict(Counter)
    meses: Counter = Counter()
    for det in detalhes:
        ci = det.get("charge_info") or {}
        for k in ("document_type", "status", "debited_from_operation", "charge_bonified_id"):
            v = ci.get(k)
            campos[k][
                ("tem" if v not in (None, "") else "vazio") if k == "charge_bonified_id" else str(v)
            ] += 1
        meses[str(ci.get("creation_date_time") or "")[:7]] += 1
    # o que está gravado no mesmo período (por data)
    gravado: dict[str, dict[str, Any]] = {}
    ids_gravados: set[str] = set()
    if per.get("date_from") and per.get("date_to"):
        p_ini = datetime.combine(
            date.fromisoformat(per["date_from"]), datetime.min.time(), tzinfo=_BRT
        )
        p_fim = datetime.combine(
            date.fromisoformat(per["date_to"]), datetime.min.time(), tzinfo=_BRT
        ) + timedelta(days=1)
        loja = await _loja_ou_404(session, conta)
        mov = ConciliacaoMovimento
        linhas = (
            await session.execute(
                select(mov.id_externo, mov.sub_tipo, mov.credito, mov.raw).where(
                    mov.integration_id == loja.integration_id,
                    mov.origem == "ml_fatura",
                    mov.credito > 0,
                    mov.data >= p_ini,
                    mov.data < p_fim,
                )
            )
        ).all()
        for id_ext, sub, cred, raw in linhas:
            ids_gravados.add(str(id_ext))
            x = gravado.setdefault(
                str(sub or "?"), {"linhas": 0, "valor": 0.0, "periodos": Counter()}
            )
            x["linhas"] += 1
            x["valor"] = round(x["valor"] + float(cred or 0), 2)
            x["periodos"][str((raw or {}).get("_periodo") or "?")] += 1
        for x in gravado.values():
            x["periodos"] = dict(x["periodos"])
    ids_ml = {str((d.get("charge_info") or {}).get("detail_id")) for d in detalhes}
    so_no_ml = [
        d
        for d in detalhes
        if str((d.get("charge_info") or {}).get("detail_id")) not in ids_gravados
    ]
    so_gravado = sorted(ids_gravados - ids_ml)
    # amostras do sub-tipo com maior diferença resumo × detalhes do ML
    difs = {
        sub: round(ml_por_sub.get(sub, {}).get("valor", 0.0) - resumo_por_tipo.get(sub, 0.0), 2)
        for sub in set(ml_por_sub) | set(resumo_por_tipo)
    }
    pior = max(difs, key=lambda k: abs(difs[k])) if difs else None
    amostras = [
        _recorta(d)
        for d in detalhes
        if str((d.get("charge_info") or {}).get("detail_sub_type")) == pior
    ][:3]
    saida = {
        "conta": conta,
        "periodo": {"chave": chave, **per},
        "resumo_ml": {
            "bonificacoes": resumo_por_tipo,
            "total_bonificacoes": round(sum(resumo_por_tipo.values()), 2),
            "total_fatura": bill.get("total_amount"),
        },
        "detalhes_ml_agora": {
            "linhas": len(detalhes),
            "por_sub_tipo": ml_por_sub,
            "total": round(sum(x["valor"] for x in ml_por_sub.values()), 2),
            "campos": {k: dict(v) for k, v in campos.items()},
            "por_mes_de_criacao": dict(meses),
        },
        "gravado_no_periodo_por_data": {
            "linhas": sum(x["linhas"] for x in gravado.values()),
            "por_sub_tipo": gravado,
            "total": round(sum(x["valor"] for x in gravado.values()), 2),
        },
        "diferenca_por_sub_tipo (detalhes ML − resumo ML)": difs,
        "so_no_ml_agora": {
            "linhas": len(so_no_ml),
            "valor": round(
                sum(
                    float((d.get("charge_info") or {}).get("detail_amount") or 0) for d in so_no_ml
                ),
                2,
            ),
            "por_sub_tipo": _soma_por_sub(so_no_ml),
        },
        "so_gravado": {"linhas": len(so_gravado), "ids": so_gravado[:20]},
        "amostras_do_pior_sub_tipo": {"sub_tipo": pior, "itens": amostras},
    }
    logger.info("conciliacao_sonda_bonus", conta=conta, periodo=chave, linhas_ml=len(detalhes))
    return saida


async def _bonus_so_banco(session: AsyncSession, conta: str, chave: str) -> dict[str, Any]:
    """Hipótese: o resumo do ML só conta a bonificação cuja cobrança original está no
    MESMO período da fatura. Cruza cada bonificação gravada (charge_bonified_id) com a
    cobrança gravada e soma por 'cobrança no período' / 'de período anterior' / 'não achada'."""
    loja = await _loja_ou_404(session, conta)
    sync = (
        await session.execute(
            select(ConciliacaoSync).where(
                ConciliacaoSync.integration_id == loja.integration_id,
                ConciliacaoSync.origem == "ml_fatura",
                ConciliacaoSync.chave == f"{chave}|resumo",
            )
        )
    ).scalar_one_or_none()
    resumo_ml = ((sync.detalhe or {}).get("resumo") if sync else None) or {}
    per = resumo_ml.get("period") or {}
    if not (per.get("date_from") and per.get("date_to")):
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "resumo_nao_lido"})
    p_ini = datetime.combine(date.fromisoformat(per["date_from"]), datetime.min.time(), tzinfo=_BRT)
    p_fim = datetime.combine(
        date.fromisoformat(per["date_to"]), datetime.min.time(), tzinfo=_BRT
    ) + timedelta(days=1)
    mov = ConciliacaoMovimento
    linhas = (
        await session.execute(
            select(mov.id_externo, mov.sub_tipo, mov.credito, mov.debito, mov.data, mov.raw).where(
                mov.integration_id == loja.integration_id, mov.origem == "ml_fatura"
            )
        )
    ).all()
    por_id = {str(i): (sub, float(c or 0), float(d or 0), dt) for i, sub, c, d, dt, _ in linhas}
    classes: dict[str, dict[str, Any]] = {}
    por_sub: dict[str, dict[str, float]] = {}
    exemplos: dict[str, list[dict[str, Any]]] = {}
    for id_ext, sub, cred, _deb, dt, raw in linhas:
        if float(cred or 0) <= 0 or not (p_ini <= dt < p_fim):
            continue
        ci = (raw or {}).get("charge_info") or {}
        orig = str(ci.get("charge_bonified_id") or "")
        o = por_id.get(orig)
        if o is None:
            classe = "cobranca_nao_gravada"
        elif p_ini <= o[3] < p_fim:
            classe = "cobranca_no_mesmo_periodo"
        elif o[3] < p_ini:
            classe = "cobranca_de_periodo_anterior"
        else:
            classe = "cobranca_de_periodo_posterior"
        x = classes.setdefault(classe, {"linhas": 0, "valor": 0.0})
        x["linhas"] += 1
        x["valor"] = round(x["valor"] + float(cred), 2)
        ps = por_sub.setdefault(str(sub), {})
        ps[classe] = round(ps.get(classe, 0.0) + float(cred), 2)
        if len(exemplos.setdefault(classe, [])) < 3:
            exemplos[classe].append(
                {
                    "bonificacao": id_ext,
                    "sub_tipo": sub,
                    "valor": float(cred),
                    "data": dt.isoformat(),
                    "cobranca": orig,
                    "cobranca_data": o[3].isoformat() if o else None,
                    "cobranca_valor": o[2] if o else None,
                }
            )
    bill = resumo_ml.get("bill_includes") or {}
    resumo_por_tipo = {
        str(b.get("type")): round(abs(float(b.get("amount") or 0)), 2)
        for b in (bill.get("bonuses") or [])
        if isinstance(b, dict)
    }
    return {
        "conta": conta,
        "periodo": {"chave": chave, **per},
        "resumo_ml": {
            "bonificacoes": resumo_por_tipo,
            "total": round(sum(resumo_por_tipo.values()), 2),
        },
        "gravado_por_classe": classes,
        "gravado_por_sub_tipo_e_classe": por_sub,
        "exemplos": exemplos,
    }


def _resumo_carteira(lista: list[dict[str, Any]]) -> dict[str, Any]:
    """Resumo das transações da carteira Shopee: por tipo, fluxo, com/sem pedido."""
    por_tipo: dict[str, dict[str, Any]] = {}
    fluxo: Counter = Counter()
    com_pedido = 0
    datas: list[int] = []
    for t in lista:
        tipo = str(t.get("transaction_type") or "?")
        x = por_tipo.setdefault(tipo, {"linhas": 0, "valor": 0.0})
        x["linhas"] += 1
        x["valor"] = round(x["valor"] + float(t.get("amount") or 0), 2)
        fluxo[str(t.get("money_flow") or "?")] += 1
        if t.get("order_sn"):
            com_pedido += 1
        if isinstance(t.get("create_time"), int):
            datas.append(t["create_time"])
    return {
        "linhas": len(lista),
        "por_tipo": por_tipo,
        "por_fluxo": dict(fluxo),
        "com_pedido": com_pedido,
        "de": datetime.fromtimestamp(min(datas), tz=_BRT).isoformat() if datas else None,
        "ate": datetime.fromtimestamp(max(datas), tz=_BRT).isoformat() if datas else None,
        "campos": sorted({k for t in lista for k in t}) if lista else [],
    }


@router.get("/sonda/shopee")
async def sonda_shopee(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    conta: Annotated[
        str | None, Query(description="nome da loja Shopee; sem valor = a primeira ativa")
    ] = None,
    dias: Annotated[int, Query(ge=1, le=31)] = 7,
) -> dict[str, Any]:
    """O que a API da Shopee dá de dinheiro: carteira (extrato), liberações (escrow),
    detalhe de uma liberação e repasses (payout). Só leitura, numa conta ativa."""
    from app.models.integration import IntegrationPlatform

    q = select(Integration).where(
        Integration.platform == IntegrationPlatform.SHOPEE,
        Integration.archived_at.is_(None),
        Integration.status == "active",
    )
    if conta:
        q = q.where(func.lower(func.trim(Integration.name)) == conta.strip().lower())
    integ = (
        await session.execute(q.order_by(Integration.created_at).limit(1))
    ).scalar_one_or_none()
    if integ is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail={"code": "conta_shopee_nao_encontrada", "conta": conta},
        )
    cliente = await cliente_da_integracao(integ)
    agora = int(time.time())
    t_de = agora - dias * 86400
    chamadas: list[dict[str, Any]] = []

    async def chama(
        nome: str, metodo: str, path: str, *, params=None, json=None
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        """Uma chamada registrada: status, tempo, erro da Shopee e uma amostra do corpo.
        Devolve também o corpo inteiro, pra resumir listas sem cortar."""
        t0 = time.monotonic()
        item: dict[str, Any] = {"nome": f"{nome} ({metodo})", "url": path}
        try:
            r = await cliente._request(metodo, path, params=params, json=json)
        except Exception as e:  # noqa: BLE001 — a sondagem registra e segue
            item.update(ok=False, erro=f"{type(e).__name__}: {str(e)[:300]}")
            item["ms"] = int((time.monotonic() - t0) * 1000)
            chamadas.append(item)
            return item, {}
        item.update(status=r.status_code, ms=int((time.monotonic() - t0) * 1000))
        try:
            body = r.json() or {}
        except ValueError:
            body = {"texto": (r.text or "")[:_MAX_TEXTO]}
        if not isinstance(body, dict):
            body = {"corpo": body}
        item["ok"] = r.status_code == 200 and not body.get("error")
        if body.get("error"):
            item["erro_shopee"] = f"{body.get('error')}: {body.get('message')}"
        item["amostra"] = _recorta(body)
        chamadas.append(item)
        return item, body

    def resposta(body: dict[str, Any]) -> dict[str, Any]:
        r = body.get("response")
        return r if isinstance(r, dict) else {}

    await chama("loja", "GET", "/api/v2/shop/get_shop_info")

    # 1) carteira = o extrato (POST; se a Shopee reclamar do método, GET)
    janela = {"page_no": 0, "page_size": 100, "create_time_from": t_de, "create_time_to": agora}
    item, body = await chama(
        "carteira", "POST", "/api/v2/payment/get_wallet_transaction_list", json=janela
    )
    if item.get("status") in (404, 405) or "method" in str(item.get("erro_shopee", "")).lower():
        item, body = await chama(
            "carteira", "GET", "/api/v2/payment/get_wallet_transaction_list", params=janela
        )
    carteira: dict[str, Any] = {}
    lista = resposta(body).get("transaction_list")
    if isinstance(lista, list):
        carteira = _resumo_carteira([t for t in lista if isinstance(t, dict)])
        carteira["mais_paginas"] = resposta(body).get("more")

    # 2) liberações (escrow) no período e o detalhe de uma
    item, body = await chama(
        "liberacoes",
        "GET",
        "/api/v2/payment/get_escrow_list",
        params={
            "release_time_from": t_de,
            "release_time_to": agora,
            "page_size": 100,
            "page_no": 1,
        },
    )
    esc = resposta(body).get("escrow_list")
    liberacoes: dict[str, Any] = {}
    primeiro_sn = None
    if isinstance(esc, list):
        itens = [e for e in esc if isinstance(e, dict)]
        liberacoes = {
            "linhas": len(itens),
            "valor": round(sum(float(e.get("payout_amount") or 0) for e in itens), 2),
            "mais_paginas": resposta(body).get("more"),
            "campos": sorted({k for e in itens for k in e}),
        }
        primeiro_sn = itens[0].get("order_sn") if itens else None
    if primeiro_sn:
        await chama(
            "liberacao_detalhe",
            "GET",
            "/api/v2/payment/get_escrow_detail",
            params={"order_sn": primeiro_sn},
        )

    # 3) repasses pro banco (as duas rotas que a Shopee documenta)
    # get_payout_info pede `cursor` ("" na primeira página; depois o next_cursor da resposta)
    rep = {"page_size": 100, "page_no": 0, "payout_time_from": t_de, "payout_time_to": agora}
    await chama(
        "repasses_info",
        "POST",
        "/api/v2/payment/get_payout_info",
        json={"page_size": 100, "cursor": "", "payout_time_from": t_de, "payout_time_to": agora},
    )
    await chama(
        "repasses_detalhe",
        "GET",
        "/api/v2/payment/get_payout_detail",
        params={**rep, "page_no": 1},
    )

    logger.info(
        "conciliacao_sonda_shopee",
        conta=integ.name,
        chamadas=[(c["nome"], c.get("status"), c.get("erro")) for c in chamadas],
    )
    return {
        "conta": integ.name,
        "integration_id": str(integ.id),
        "janela": {
            "de": datetime.fromtimestamp(t_de, tz=_BRT).isoformat(),
            "ate": datetime.fromtimestamp(agora, tz=_BRT).isoformat(),
        },
        "carteira": carteira,
        "liberacoes": liberacoes,
        "chamadas": chamadas,
    }


_CAMPOS_ESCROW = (
    "escrow_amount",
    "escrow_amount_after_adjustment",
    "total_adjustment_amount",
    "buyer_total_amount",
    "cost_of_goods_sold",
    "order_selling_price",
    "commission_fee",
    "net_commission_fee",
    "service_fee",
    "net_service_fee",
    "seller_transaction_fee",
    "campaign_fee",
    "ams_commission_fee",
    "order_ams_commission_fee",
    "actual_shipping_fee",
    "buyer_paid_shipping_fee",
    "final_shipping_fee",
    "estimated_shipping_fee",
    "shopee_shipping_rebate",
    "seller_shipping_discount",
    "shipping_seller_protection_fee_amount",
    "delivery_seller_protection_fee_premium_amount",
    "voucher_from_seller",
    "voucher_from_shopee",
    "seller_coin_cash_back",
    "seller_discount",
    "seller_return_refund",
    "drc_adjustable_refund",
    "reverse_shipping_fee",
    "return_handling_fee",
    "seller_lost_compensation",
    "pix_discount",
    "seller_product_rebate",
)


@router.get("/sonda/shopee-escrow")
async def sonda_shopee_escrow(
    conta: Annotated[str, Query(min_length=1)],
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    pedidos: Annotated[str | None, Query(description="nºs do Bling separados por vírgula")] = None,
    situacao: str | None = "recebido_a_menos",
    vivo: bool = True,
) -> dict[str, Any]:
    """Por que a carteira pagou diferente do líquido da Margem? Lado a lado: o escrow
    gravado pela Margem, o escrow como a Shopee diz AGORA (1 chamada por pedido) e a
    carteira; e os campos do escrow que valem exatamente a diferença. Só leitura."""
    from app.models import MarketplaceOrderFinancial

    loja = await _loja_ou_404(session, conta)
    q = select(ConciliacaoPedido).where(ConciliacaoPedido.integration_id == loja.integration_id)
    if pedidos:
        q = q.where(ConciliacaoPedido.pedido_bling.in_([x.strip() for x in pedidos.split(",")]))
    elif situacao:
        q = q.where(ConciliacaoPedido.situacao.in_(situacao.split(",")))
    rows = (
        (await session.execute(q.order_by(ConciliacaoPedido.data_pedido.desc()).limit(10)))
        .scalars()
        .all()
    )
    cliente = None
    if vivo:
        integ = await session.get(Integration, loja.integration_id)
        cliente = await cliente_da_integracao(integ) if integ else None

    def recorte(income: dict[str, Any]) -> dict[str, Any]:
        return {k: income.get(k) for k in _CAMPOS_ESCROW if income.get(k) not in (None, 0, 0.0, "")}

    saida = []
    for p in rows:
        fin = (
            await session.execute(
                select(MarketplaceOrderFinancial).where(
                    MarketplaceOrderFinancial.integration_id == loja.integration_id,
                    MarketplaceOrderFinancial.external_order_id == p.numeroloja,
                )
            )
        ).scalar_one_or_none()
        raw = fin.raw if (fin and isinstance(fin.raw, dict)) else {}
        gravado = (
            ((raw.get("escrow") or {}).get("order_income") or {})
            if isinstance(raw.get("escrow"), dict)
            else {}
        )
        agora: dict[str, Any] = {}
        erro_vivo = None
        if cliente and p.numeroloja:
            try:
                r = await cliente._request(
                    "GET", "/api/v2/payment/get_escrow_detail", params={"order_sn": p.numeroloja}
                )
                body = r.json() or {}
                if body.get("error"):
                    erro_vivo = f"{body.get('error')}: {body.get('message')}"
                agora = (
                    ((body.get("response") or {}).get("order_income") or {})
                    if isinstance(body.get("response"), dict)
                    else {}
                )
            except Exception as e:  # noqa: BLE001
                erro_vivo = f"{type(e).__name__}: {str(e)[:200]}"
        carteira = (
            (
                await session.execute(
                    select(ConciliacaoMovimento)
                    .where(
                        ConciliacaoMovimento.integration_id == loja.integration_id,
                        ConciliacaoMovimento.origem == "shopee_carteira",
                        ConciliacaoMovimento.pedido_marketplace == p.numeroloja,
                    )
                    .order_by(ConciliacaoMovimento.data)
                )
            )
            .scalars()
            .all()
        )
        liquido = _num(p.venda_liquido)
        recebido = _num(p.recebido)
        dif = (
            round(liquido - recebido, 2) if (liquido is not None and recebido is not None) else None
        )
        pistas = []
        for origem_nome, income in (("agora", agora), ("gravado", gravado)):
            for k, v in income.items():
                if (
                    isinstance(v, int | float)
                    and dif is not None
                    and abs(abs(v) - abs(dif)) < 0.011
                    and v
                ):
                    pistas.append(f"{origem_nome}.{k} = {v}")
        mudou = (
            {
                k: {"gravado": gravado.get(k), "agora": agora.get(k)}
                for k in _CAMPOS_ESCROW
                if gravado.get(k) != agora.get(k) and (gravado.get(k) or agora.get(k))
            }
            if agora
            else {}
        )
        saida.append(
            {
                "pedido_bling": p.pedido_bling,
                "order_sn": p.numeroloja,
                "situacao": p.situacao,
                "margem": {
                    "valor": _num(p.venda_valor),
                    "tarifa": _num(p.venda_tarifa),
                    "frete": _num(p.venda_frete),
                    "cupom": _num(p.venda_cupom),
                    "liquido": liquido,
                    "lido_em": fin.updated_at.isoformat()
                    if fin and getattr(fin, "updated_at", None)
                    else None,
                },
                "carteira": [
                    {
                        "data": m.data.isoformat(),
                        "tipo": m.tipo,
                        "credito": _num(m.credito),
                        "debito": _num(m.debito),
                    }
                    for m in carteira
                ],
                "recebido": recebido,
                "diferenca": dif,
                "pistas": pistas,
                "escrow_gravado": recorte(gravado),
                "escrow_agora": recorte(agora),
                "o_que_mudou": mudou,
                "erro_vivo": erro_vivo,
            }
        )
    logger.info("conciliacao_sonda_shopee_escrow", conta=conta, pedidos=len(saida))
    return {"conta": conta, "pedidos": saida}


def _epoch_para_iso(v: Any) -> str | None:
    try:
        return datetime.fromtimestamp(int(v), tz=_BRT).isoformat() if v else None
    except (TypeError, ValueError, OSError):
        return None


@router.get("/sonda/tiktok")
async def sonda_tiktok(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    conta: Annotated[
        str | None, Query(description="nome da loja TikTok; sem valor = a que mais vende")
    ] = None,
    dias: Annotated[int, Query(ge=1, le=90)] = 30,
) -> dict[str, Any]:
    """O que a API financeira da TikTok dá: extratos (statements), repasses (payments)
    e as transações de um extrato (pedido a pedido). Só leitura. Sem `conta`, pega a
    loja TikTok com mais vendas lidas pela Margem nos últimos 30 dias (15/10, Vinicius:
    "pode pegar a que mais vende e sondar")."""
    from app.models import MarketplaceOrderFinancial
    from app.models.integration import IntegrationPlatform

    agora_dt = datetime.now(UTC)
    base_q = select(Integration).where(
        Integration.platform == IntegrationPlatform.TIKTOK,
        Integration.archived_at.is_(None),
        Integration.status == "active",
    )
    integ = None
    if conta:
        candidatas = list(
            (
                await session.execute(
                    base_q.where(func.lower(func.trim(Integration.name)) == conta.strip().lower())
                )
            ).scalars()
        )
        if len(candidatas) > 1:
            # 16/10, revisão: duas integrações ativas com o mesmo nome davam 500
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                detail={
                    "code": "conta_tiktok_duplicada",
                    "conta": conta,
                    "integration_ids": [str(i.id) for i in candidatas],
                    "mensagem": "há mais de uma integração TikTok ativa com esse nome",
                },
            )
        integ = candidatas[0] if candidatas else None
    else:
        fin = MarketplaceOrderFinancial
        ranking = (
            await session.execute(
                select(fin.integration_id, func.count())
                .where(
                    fin.platform == IntegrationPlatform.TIKTOK,
                    fin.created_at >= agora_dt - timedelta(days=30),
                )
                .group_by(fin.integration_id)
                .order_by(func.count().desc())
            )
        ).all()
        ativas = {i.id: i for i in (await session.execute(base_q)).scalars()}
        for iid, _n in ranking:
            if iid in ativas:
                integ = ativas[iid]
                break
        if integ is None and ativas:
            integ = sorted(ativas.values(), key=lambda i: i.created_at or agora_dt)[0]
    if integ is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail={"code": "conta_tiktok_nao_encontrada", "conta": conta},
        )
    cliente = await cliente_da_integracao(integ)
    agora = int(time.time())
    t_de = agora - dias * 86400
    chamadas: list[dict[str, Any]] = []

    async def chama(nome: str, path: str, params: dict[str, str] | None = None) -> dict[str, Any]:
        t0 = time.monotonic()
        item: dict[str, Any] = {"nome": nome, "url": path, "params": params}
        try:
            body = await cliente._get(path, params)
        except Exception as e:  # noqa: BLE001 — a sondagem registra e segue
            item.update(ok=False, erro=f"{type(e).__name__}: {str(e)[:300]}")
            item["ms"] = int((time.monotonic() - t0) * 1000)
            chamadas.append(item)
            return {}
        item["ms"] = int((time.monotonic() - t0) * 1000)
        body = body if isinstance(body, dict) else {"corpo": body}
        item["code"] = body.get("code")
        item["ok"] = body.get("code") == 0
        if not item["ok"]:
            item["mensagem"] = str(body.get("message") or "")[:300]
        item["amostra"] = _recorta(body.get("data") if "data" in body else body)
        chamadas.append(item)
        return body.get("data") if isinstance(body.get("data"), dict) else {}

    # 1) extratos (statements) do período
    dados = await chama(
        "extratos",
        "/finance/202309/statements",
        {
            "statement_time_ge": str(t_de),
            # a API filtra por dia e exclui o dia do `lt`: amanhã pra ver o de hoje
            "statement_time_lt": str(agora + 86400),
            "page_size": "50",
            "sort_field": "statement_time",
            "sort_order": "DESC",
        },
    )
    statements = [x for x in (dados.get("statements") or []) if isinstance(x, dict)]
    resumo_ext: dict[str, Any] = {}
    if statements:
        status_c: Counter = Counter(str(x.get("payment_status") or "?") for x in statements)
        resumo_ext = {
            "linhas": len(statements),
            "mais_paginas": bool(dados.get("next_page_token")),
            "por_status": dict(status_c),
            "liquidado_total": round(
                sum(float(x.get("settlement_amount") or 0) for x in statements), 2
            ),
            "de": _epoch_para_iso(min(int(x.get("statement_time") or 0) for x in statements)),
            "ate": _epoch_para_iso(max(int(x.get("statement_time") or 0) for x in statements)),
            "campos": sorted({k for x in statements for k in x}),
        }

    # 2) transações do extrato mais recente (pedido a pedido)
    transacoes_resumo: dict[str, Any] = {}
    if statements:
        sid = statements[0].get("id")
        dados_t = await chama(
            "transacoes_do_extrato",
            f"/finance/202309/statements/{sid}/statement_transactions",
            {"page_size": "100", "sort_field": "order_create_time", "sort_order": "DESC"},
        )
        txs = [x for x in (dados_t.get("statement_transactions") or []) if isinstance(x, dict)]
        if txs:
            tipos: Counter = Counter(str(x.get("type") or "?") for x in txs)
            transacoes_resumo = {
                "extrato": sid,
                "linhas": len(txs),
                "mais_paginas": bool(dados_t.get("next_page_token")),
                "por_tipo": dict(tipos),
                "com_pedido": sum(1 for x in txs if x.get("order_id")),
                "liquidado_total": round(
                    sum(float(x.get("settlement_amount") or 0) for x in txs), 2
                ),
                "campos": sorted({k for x in txs for k in x}),
                "cabecalho_do_extrato": {
                    k: v for k, v in dados_t.items() if k != "statement_transactions"
                },
            }
            pedido = next((x.get("order_id") for x in txs if x.get("order_id")), None)
            if pedido:
                await chama(
                    "transacoes_do_pedido",
                    f"/finance/202309/orders/{pedido}/statement_transactions",
                )

    # 3) repasses pro banco (payments)
    dados_p = await chama(
        "repasses",
        "/finance/202309/payments",
        {
            "create_time_ge": str(t_de),
            "create_time_lt": str(agora + 86400),
            "page_size": "50",
            "sort_field": "create_time",
            "sort_order": "DESC",
        },
    )
    payments = [x for x in (dados_p.get("payments") or []) if isinstance(x, dict)]
    resumo_rep: dict[str, Any] = {}
    if payments:
        st: Counter = Counter(str(x.get("status") or "?") for x in payments)

        def _valor(x: dict[str, Any]) -> float:
            a = x.get("amount")
            if isinstance(a, dict):
                return float(a.get("value") or 0)
            return float(a or 0)

        resumo_rep = {
            "linhas": len(payments),
            "mais_paginas": bool(dados_p.get("next_page_token")),
            "por_status": dict(st),
            "total": round(sum(_valor(x) for x in payments), 2),
            "campos": sorted({k for x in payments for k in x}),
        }

    logger.info(
        "conciliacao_sonda_tiktok",
        conta=integ.name,
        chamadas=[(c["nome"], c.get("code"), c.get("erro") or c.get("mensagem")) for c in chamadas],
    )
    return {
        "conta": integ.name,
        "integration_id": str(integ.id),
        "janela": {"de": _epoch_para_iso(t_de), "ate": _epoch_para_iso(agora)},
        "extratos": resumo_ext,
        "transacoes": transacoes_resumo,
        "repasses": resumo_rep,
        "chamadas": chamadas,
    }


@router.get("/sonda/tiktok/espera")
async def sonda_tiktok_espera(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    conta: Annotated[str, Query(description="loja TikTok ligada no painel")],
    dias: Annotated[int, Query(ge=1, le=180)] = 90,
) -> dict[str, Any]:
    """TikTok "Em espera" (pedidos ainda não liquidados, `orders/unsettled`) × nosso
    "A liberar" (estimativa dos pedidos aguardando desde o início da loja). Só leitura.
    15/10, Vinicius: "aqui mostra em espera 47.634,93, o nosso 48.928 — por que?"."""
    loja = await _loja_ou_404(session, conta)
    integ = await session.get(Integration, loja.integration_id)
    if integ is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail={"code": "integracao_nao_encontrada"})
    cliente = await cliente_da_integracao(integ)
    agora = int(time.time())

    def _d(v: Any) -> Decimal | None:
        if v in (None, ""):
            return None
        try:
            return Decimal(str(v))
        except (InvalidOperation, ValueError):
            return None

    inicio_dt = datetime.combine(loja.inicio, datetime.min.time(), tzinfo=_BRT)

    # 1) o que a TikTok diz que está em espera, pedido a pedido
    tiktok: dict[str, dict[str, Any]] = {}
    ajustes = Decimal("0")
    ajustes_qtd = 0
    paginas = 0
    token: str | None = None
    for _ in range(20):
        body = await cliente.get_unsettled_orders(
            search_time_ge=agora - dias * 86400, search_time_lt=agora + 86400, page_token=token
        )
        paginas += 1
        if not isinstance(body, dict) or body.get("code") not in (0, "0", None):
            raise HTTPException(
                status.HTTP_502_BAD_GATEWAY,
                detail={
                    "code": "tiktok_unsettled",
                    "mensagem": str((body or {}).get("message"))[:300],
                },
            )
        data = body.get("data") if isinstance(body.get("data"), dict) else {}
        for t in data.get("transactions") or []:
            if not isinstance(t, dict):
                continue
            est = _d(t.get("est_settlement_amount"))
            if t.get("adjustment_id") or not t.get("order_id"):
                ajustes += est or Decimal("0")
                ajustes_qtd += 1
                continue
            oid = str(t.get("order_id")).strip()
            slot = tiktok.setdefault(
                oid,
                {
                    "est": Decimal("0"),
                    "criado": _epoch_para_iso(int(t.get("order_create_time") or 0))
                    if t.get("order_create_time")
                    else None,
                    "criado_epoch": int(t.get("order_create_time") or 0),
                    "motivo": t.get("unsettled_reason"),
                },
            )
            slot["est"] += est or Decimal("0")
        token = data.get("next_page_token") or None
        if not token:
            break

    # 2) o que o painel conta como "a liberar" (aguardando/vencido) e o resto dos pedidos
    ped = ConciliacaoPedido
    nossos = {
        p.numeroloja: p
        for p in (
            await session.execute(select(ped).where(ped.integration_id == loja.integration_id))
        ).scalars()
        if p.numeroloja
    }
    a_liberar = {n: p for n, p in nossos.items() if p.situacao in ("aguardando", "vencido")}

    def soma(ids: set[str], de: dict[str, Any], chave: str) -> float:
        total = Decimal("0")
        for i in ids:
            v = de[i][chave] if isinstance(de[i], dict) else getattr(de[i], chave, None)
            total += _d(v) or Decimal("0")
        return _num(total) or 0.0

    ambos = set(tiktok) & set(a_liberar)
    so_tiktok = set(tiktok) - set(a_liberar)
    so_nosso = set(a_liberar) - set(tiktok)
    antes_inicio = {
        i
        for i in so_tiktok
        if tiktok[i]["criado_epoch"] and tiktok[i]["criado_epoch"] < inicio_dt.timestamp()
    }
    so_tiktok_outros = so_tiktok - antes_inicio
    por_situacao: Counter = Counter(
        nossos[i].situacao if i in nossos else "fora_do_painel" for i in so_tiktok_outros
    )
    return {
        "conta": loja.conta,
        "inicio": loja.inicio.isoformat(),
        "paginas": paginas,
        "tiktok_em_espera": {
            "pedidos": len(tiktok),
            "total": soma(set(tiktok), tiktok, "est"),
            "ajustes": _num(ajustes),
            "ajustes_qtd": ajustes_qtd,
        },
        "nosso_a_liberar": {
            "pedidos": len(a_liberar),
            "total": soma(set(a_liberar), a_liberar, "venda_liquido"),
        },
        "nos_dois": {
            "pedidos": len(ambos),
            "tiktok": soma(ambos, tiktok, "est"),
            "nosso": soma(ambos, a_liberar, "venda_liquido"),
        },
        "so_na_tiktok": {
            "antes_do_inicio": {
                "pedidos": len(antes_inicio),
                "total": soma(antes_inicio, tiktok, "est"),
                "exemplos": sorted(antes_inicio)[:10],
            },
            "desde_o_inicio": {
                "pedidos": len(so_tiktok_outros),
                "total": soma(so_tiktok_outros, tiktok, "est"),
                "por_situacao_no_painel": dict(por_situacao),
                "exemplos": [
                    {
                        "pedido": i,
                        "est": _num(tiktok[i]["est"]),
                        "criado": tiktok[i]["criado"],
                        "motivo": tiktok[i]["motivo"],
                        "no_painel": nossos[i].situacao if i in nossos else None,
                        "pedido_bling": nossos[i].pedido_bling if i in nossos else None,
                    }
                    for i in sorted(so_tiktok_outros)[:15]
                ],
            },
        },
        "so_no_painel": {
            "pedidos": len(so_nosso),
            "total": soma(so_nosso, a_liberar, "venda_liquido"),
            "exemplos": [
                {
                    "pedido": i,
                    "pedido_bling": a_liberar[i].pedido_bling,
                    "nosso": _num(a_liberar[i].venda_liquido),
                    "status_bling": a_liberar[i].situacao_bling,
                    "venda_status": a_liberar[i].venda_status,
                }
                for i in sorted(so_nosso)[:15]
            ],
        },
    }


def _soma_quebras(no: Any, acc: dict[str, Decimal], prefixo: str = "") -> None:
    """Soma `breakdownAmount` por `breakdownType` (recursivo) das transações da Amazon."""
    if isinstance(no, dict):
        tipo = no.get("breakdownType")
        amt = (no.get("breakdownAmount") or {}).get("currencyAmount")
        if tipo and amt is not None:
            try:
                acc[f"{prefixo}{tipo}"] += Decimal(str(amt))
            except (InvalidOperation, ValueError):
                pass
        for b in no.get("breakdowns") or []:
            _soma_quebras(b, acc, f"{prefixo}{tipo} › " if tipo else prefixo)
    elif isinstance(no, list):
        for x in no:
            _soma_quebras(x, acc, prefixo)


def _resumo_relatorio_repasse(texto: str) -> dict[str, Any]:
    """Resume o settlement report (TSV V2) sem dados pessoais: cabeçalho da
    liquidação, linhas por tipo e as somas por tipo de valor."""
    linhas = [ln for ln in texto.splitlines() if ln.strip()]
    if not linhas:
        return {"linhas": 0}
    cab = linhas[0].split("\t")
    idx = {c.strip(): i for i, c in enumerate(cab)}

    def col(partes: list[str], nome: str) -> str:
        i = idx.get(nome)
        return partes[i].strip() if (i is not None and i < len(partes)) else ""

    def dec(v: str) -> Decimal:
        try:
            return Decimal(v.replace(",", "."))
        except (InvalidOperation, ValueError, AttributeError):
            return Decimal("0")

    corpo = [ln.split("\t") for ln in linhas[1:]]
    liquidacao = corpo[0] if corpo else []
    tipos: Counter = Counter()
    por_valor: Counter = Counter()
    somas: dict[str, Decimal] = defaultdict(Decimal)
    canais: Counter = Counter()
    pedidos: set[str] = set()
    total_linhas = Decimal("0")
    for partes in corpo[1:]:
        tipo = col(partes, "transaction-type") or "?"
        tipos[tipo] += 1
        chave = f"{tipo} · {col(partes, 'amount-type')} · {col(partes, 'amount-description')}"
        por_valor[chave] += 1
        valor = dec(col(partes, "amount"))
        somas[chave] += valor
        total_linhas += valor
        if col(partes, "fulfillment-id"):
            canais[col(partes, "fulfillment-id")] += 1
        if col(partes, "order-id"):
            pedidos.add(col(partes, "order-id"))
    top = sorted(somas.items(), key=lambda kv: abs(kv[1]), reverse=True)[:25]
    return {
        "colunas": cab,
        "linhas": len(corpo),
        "liquidacao": {
            "settlement_id": col(liquidacao, "settlement-id"),
            "inicio": col(liquidacao, "settlement-start-date"),
            "fim": col(liquidacao, "settlement-end-date"),
            "deposito": col(liquidacao, "deposit-date"),
            "total": col(liquidacao, "total-amount"),
            "moeda": col(liquidacao, "currency"),
        },
        "soma_das_linhas": _num(total_linhas),
        "por_tipo": dict(tipos),
        "canais": dict(canais),
        "pedidos_distintos": len(pedidos),
        "somas": [{"chave": k, "linhas": por_valor[k], "total": _num(v)} for k, v in top],
    }


@router.get("/sonda/amazon")
async def sonda_amazon(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    conta: Annotated[str, Query(description="nome da conta Amazon (kfa, kia, poofy…)")],
    dias: Annotated[int, Query(ge=1, le=120)] = 60,
    baixar: bool = True,
) -> dict[str, Any]:
    """O que a Amazon dá pra conciliação, por conta — só leitura (15/10, Vinicius:
    "pode rodar"): transações do período e grupos de repasse (Finances), relatórios
    de repasse (settlement, Reports API) com o mais recente baixado e resumido, o que
    a Margem já leu, pedidos por tipo de envio desde 01/10 e o Melhor Envio pelo token."""
    from app.models import Logistica, MarketplaceOrderFinancial
    from app.models.integration import IntegrationPlatform
    from app.services.marketplaces.amazon import _decode_amazon_report
    from app.services.melhor_envio import MelhorEnvioClient

    integ = (
        (
            await session.execute(
                select(Integration).where(
                    Integration.platform == IntegrationPlatform.AMAZON,
                    Integration.archived_at.is_(None),
                    func.lower(func.trim(Integration.name)) == conta.strip().lower(),
                )
            )
        )
        .scalars()
        .first()
    )
    if integ is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail={"code": "conta_amazon_nao_encontrada", "conta": conta},
        )
    cliente = await cliente_da_integracao(integ)
    agora = datetime.now(UTC)
    de = agora - timedelta(days=dias)
    chamadas: list[dict[str, Any]] = []

    async def chama(nome: str, path: str, params: dict[str, Any] | None = None) -> Any:
        t0 = time.monotonic()
        item: dict[str, Any] = {"nome": nome, "url": path, "params": params}
        try:
            r = await cliente._request("GET", path, params=params)
        except Exception as e:  # noqa: BLE001 — a sondagem registra e segue
            item.update(ok=False, erro=f"{type(e).__name__}: {str(e)[:300]}")
            chamadas.append(item)
            return None
        item["ms"] = int((time.monotonic() - t0) * 1000)
        item["status"] = r.status_code
        try:
            body = r.json()
        except ValueError:
            body = {"texto": r.text[:300]}
        item["ok"] = r.status_code == 200
        if not item["ok"]:
            item["mensagem"] = str(body)[:300]
        chamadas.append(item)
        return body if item["ok"] else None

    # 1) transações do período (Finances 2024-06-19) — o que a Margem lê por pedido
    tx_body = await chama(
        "transacoes",
        "/finances/2024-06-19/transactions",
        {"postedAfter": de.isoformat(), "marketplaceId": cliente.marketplace_id},
    )
    transacoes: dict[str, Any] = {}
    if isinstance(tx_body, dict):
        payload = tx_body.get("payload") if isinstance(tx_body.get("payload"), dict) else tx_body
        txs = [t for t in (payload.get("transactions") or []) if isinstance(t, dict)]
        token = payload.get("nextToken")
        for _ in range(4):
            if not token:
                break
            mais = await chama(
                "transacoes_pagina",
                "/finances/2024-06-19/transactions",
                {"postedAfter": de.isoformat(), "nextToken": token},
            )
            if not isinstance(mais, dict):
                break
            pl = mais.get("payload") if isinstance(mais.get("payload"), dict) else mais
            txs.extend(t for t in (pl.get("transactions") or []) if isinstance(t, dict))
            token = pl.get("nextToken")
        tx_body = {"transactions": txs, "nextToken": token}
        tipos: Counter = Counter(str(t.get("transactionType") or "?") for t in txs)
        quebras: Counter = Counter()
        total = Decimal("0")
        for t in txs:
            amt = t.get("totalAmount") or {}
            try:
                total += Decimal(str(amt.get("currencyAmount") or 0))
            except (InvalidOperation, ValueError):
                pass
            for it in t.get("items") or []:
                for b in (it.get("breakdowns") or []) if isinstance(it, dict) else []:
                    if isinstance(b, dict):
                        quebras[str(b.get("breakdownType") or "?")] += 1
                        for b2 in b.get("breakdowns") or []:
                            if isinstance(b2, dict):
                                quebras[
                                    f"{b.get('breakdownType')} › {b2.get('breakdownType')}"
                                ] += 1
        somas_tipo: dict[str, Decimal] = defaultdict(Decimal)
        somas_quebra: dict[str, Decimal] = defaultdict(Decimal)
        repasses_banco: list[dict[str, Any]] = []
        por_status: Counter = Counter()
        for t in txs:
            amt = t.get("totalAmount") or {}
            try:
                v = Decimal(str(amt.get("currencyAmount") or 0))
            except (InvalidOperation, ValueError):
                v = Decimal("0")
            tipo_t = str(t.get("transactionType") or "?")
            somas_tipo[tipo_t] += v
            por_status[str(t.get("transactionStatus") or "?")] += 1
            if tipo_t in ("Transfer", "DebtRecovery"):
                repasses_banco.append(
                    {
                        "tipo": tipo_t,
                        "data": str(t.get("postedDate") or "")[:10],
                        "valor": _num(v),
                        "descricao": str(t.get("description") or "")[:60],
                        "status": t.get("transactionStatus"),
                    }
                )
            for it in t.get("items") or []:
                if isinstance(it, dict):
                    _soma_quebras(it.get("breakdowns"), somas_quebra)
        transacoes = {
            "linhas": len(txs),
            "mais_paginas": bool(tx_body.get("nextToken")),
            "por_tipo": dict(tipos),
            "por_status": dict(por_status),
            "soma_por_tipo": {k: _num(v) for k, v in somas_tipo.items()},
            "soma_por_quebra": {
                k: _num(v)
                for k, v in sorted(somas_quebra.items(), key=lambda kv: abs(kv[1]), reverse=True)[
                    :25
                ]
            },
            "repasses_banco": sorted(repasses_banco, key=lambda x: x["data"])[-12:],
            "total": _num(total),
            "campos": sorted({k for t in txs for k in t}),
            "quebras": dict(quebras.most_common(30)),
            "datas": sorted({str(t.get("postedDate") or "")[:10] for t in txs})[-5:],
        }
    # 1b) o que ainda está preso (DEFERRED) × liberado (RELEASED) — reserva por data de entrega
    por_situacao: dict[str, Any] = {}
    for situ in ("DEFERRED", "RELEASED"):
        b = await chama(
            f"transacoes_{situ.lower()}",
            "/finances/2024-06-19/transactions",
            {
                "postedAfter": de.isoformat(),
                "marketplaceId": cliente.marketplace_id,
                "transactionStatus": situ,
            },
        )
        if isinstance(b, dict):
            pl = b.get("payload") if isinstance(b.get("payload"), dict) else b
            lst = [t for t in (pl.get("transactions") or []) if isinstance(t, dict)]
            soma = Decimal("0")
            tipos_s: Counter = Counter()
            for t in lst:
                tipos_s[str(t.get("transactionType") or "?")] += 1
                try:
                    soma += Decimal(str((t.get("totalAmount") or {}).get("currencyAmount") or 0))
                except (InvalidOperation, ValueError):
                    pass
            por_situacao[situ] = {
                "linhas": len(lst),
                "mais_paginas": bool(pl.get("nextToken")),
                "por_tipo": dict(tipos_s),
                "soma": _num(soma),
            }
    # 2) grupos de repasse (Finances v0): um grupo = um período de repasse pro banco
    g_body = await chama(
        "grupos_repasse",
        "/finances/v0/financialEventGroups",
        {"FinancialEventGroupStartedAfter": de.isoformat(), "MaxResultsPerPage": 100},
    )
    grupos: list[dict[str, Any]] = []
    if isinstance(g_body, dict):
        lista = ((g_body.get("payload") or {}).get("FinancialEventGroupList")) or []
        for g in lista:
            if not isinstance(g, dict):
                continue
            grupos.append(
                {
                    "id": g.get("FinancialEventGroupId"),
                    "status": g.get("ProcessingStatus"),
                    "repasse": g.get("FundTransferStatus"),
                    "inicio": g.get("FinancialEventGroupStart"),
                    "fim": g.get("FinancialEventGroupEnd"),
                    "repassado_em": g.get("FundTransferDate"),
                    "total": (g.get("OriginalTotal") or {}).get("CurrencyAmount"),
                    "saldo_inicial": (g.get("BeginningBalance") or {}).get("CurrencyAmount"),
                }
            )
    # 3) relatórios de repasse (settlement) — a Amazon gera sozinha; só listar/baixar
    r_body = await chama(
        "relatorios_repasse",
        "/reports/2021-06-30/reports",
        {
            "reportTypes": "GET_V2_SETTLEMENT_REPORT_DATA_FLAT_FILE_V2",
            "createdSince": de.isoformat(),
            "pageSize": 100,
        },
    )
    relatorios: list[dict[str, Any]] = []
    if isinstance(r_body, dict):
        for rep in r_body.get("reports") or []:
            if isinstance(rep, dict):
                relatorios.append(
                    {
                        "id": rep.get("reportId"),
                        "criado": rep.get("createdTime"),
                        "inicio": rep.get("dataStartTime"),
                        "fim": rep.get("dataEndTime"),
                        "status": rep.get("processingStatus"),
                        "documento": rep.get("reportDocumentId"),
                    }
                )
    relatorio_baixado: dict[str, Any] = {}
    relatorios_baixados: list[dict[str, Any]] = []
    prontos = [x for x in relatorios if x["status"] == "DONE" and x["documento"]]
    if baixar and prontos:
        ordem = sorted(prontos, key=lambda x: str(x["criado"]), reverse=True)[:12]
        for i, alvo in enumerate(ordem):
            d_body = await chama(
                "documento_repasse", f"/reports/2021-06-30/documents/{alvo['documento']}"
            )
            if not (isinstance(d_body, dict) and d_body.get("url")):
                continue
            try:
                async with httpx.AsyncClient(timeout=60.0) as c:
                    dl = await c.get(d_body["url"])
                bruto = dl.content
                if d_body.get("compressionAlgorithm") == "GZIP":
                    bruto = gzip.decompress(bruto)
                resumo = _resumo_relatorio_repasse(_decode_amazon_report(bruto))
            except Exception as e:  # noqa: BLE001
                relatorios_baixados.append(
                    {"relatorio": alvo["id"], "erro": f"{type(e).__name__}: {str(e)[:200]}"}
                )
                continue
            if i == 0:
                relatorio_baixado = {"relatorio": alvo["id"], "http": dl.status_code, **resumo}
            liq = resumo.get("liquidacao") or {}
            relatorios_baixados.append(
                {
                    "relatorio": alvo["id"],
                    "criado": alvo["criado"],
                    "settlement_id": liq.get("settlement_id"),
                    "periodo": [liq.get("inicio"), liq.get("fim")],
                    "deposito": liq.get("deposito"),
                    "total": liq.get("total"),
                    "soma_das_linhas": resumo.get("soma_das_linhas"),
                    "linhas": resumo.get("linhas"),
                    "pedidos": resumo.get("pedidos_distintos"),
                    "por_tipo": resumo.get("por_tipo"),
                    "canais": resumo.get("canais"),
                    "chaves": [
                        f"{x['chave']} = {x['total']}" for x in (resumo.get("somas") or [])[:8]
                    ],
                }
            )
    # 4) o que a Margem já leu desta conta
    fin = MarketplaceOrderFinancial
    margem = {
        str(st): int(n)
        for st, n in (
            await session.execute(
                select(fin.status, func.count())
                .where(fin.integration_id == integ.id)
                .group_by(fin.status)
            )
        ).all()
    }
    # 5) pedidos desde 01/10 por tipo de envio (DBA × próprio), pela Logística
    lojas_bling = await bling_store_ids_da_loja(session, integ)
    inicio = datetime(2026, 10, 1, tzinfo=_BRT)
    numeros = set(
        (
            await session.execute(
                select(BlingOrder.numero).where(
                    BlingOrder.loja.in_(lojas_bling or {"-"}), BlingOrder.data >= inicio
                )
            )
        ).scalars()
    )
    canais: Counter = Counter()
    servicos: Counter = Counter()
    canal_por_numero: dict[str, str] = {}
    if numeros:
        for numero, canal, servico in (
            await session.execute(
                select(
                    Logistica.pedido_bling, Logistica.amazon_canal, Logistica.servico_envio
                ).where(Logistica.pedido_bling.in_(numeros))
            )
        ).all():
            canais[str(canal or "sem_canal")] += 1
            servicos[str(servico or "?")] += 1
            canal_por_numero[str(numero)] = str(canal or "sem_canal")
    # o que a Margem tem desses pedidos: status, líquido e os nomes das quebras por canal
    numero_por_externo = {
        str(nl): str(n)
        for n, nl in (
            await session.execute(
                select(BlingOrder.numero, BlingOrder.numeroloja).where(
                    BlingOrder.loja.in_(lojas_bling or {"-"}), BlingOrder.data >= inicio
                )
            )
        ).all()
        if nl
    }
    margem_out: dict[str, Any] = {
        "status": Counter(),
        "liquido_posted": Decimal("0"),
        "amostras": {},
    }

    def _quebras(no: Any, acc: dict[str, Decimal], prefixo: str = "") -> None:
        if isinstance(no, dict):
            tipo = no.get("breakdownType")
            amt = (no.get("breakdownAmount") or {}).get("currencyAmount")
            if tipo and amt is not None:
                try:
                    acc[f"{prefixo}{tipo}"] += Decimal(str(amt))
                except (InvalidOperation, ValueError):
                    pass
            for b in no.get("breakdowns") or []:
                _quebras(b, acc, f"{prefixo}{tipo} › " if tipo else prefixo)
        elif isinstance(no, list):
            for x in no:
                _quebras(x, acc, prefixo)

    if numero_por_externo:
        rows = (
            (
                await session.execute(
                    select(fin).where(
                        fin.integration_id == integ.id,
                        fin.external_order_id.in_(set(numero_por_externo)),
                    )
                )
            )
            .scalars()
            .all()
        )
        for f in rows:
            margem_out["status"][str(f.status)] += 1
            if f.status == "posted" and f.net_amount is not None:
                margem_out["liquido_posted"] += Decimal(str(f.net_amount))
            canal = canal_por_numero.get(numero_por_externo.get(str(f.external_order_id), ""), "?")
            if (
                f.status == "posted"
                and canal not in margem_out["amostras"]
                and isinstance(f.raw, dict)
            ):
                acc: dict[str, Decimal] = defaultdict(Decimal)
                corpo = f.raw.get("transactions")
                pl = (
                    corpo.get("payload")
                    if isinstance(corpo, dict) and isinstance(corpo.get("payload"), dict)
                    else corpo
                )
                txl = (pl or {}).get("transactions") if isinstance(pl, dict) else None
                tipos_tx: Counter = Counter()
                for t in txl or []:
                    if isinstance(t, dict):
                        tipos_tx[str(t.get("transactionType"))] += 1
                        for it in t.get("items") or []:
                            if isinstance(it, dict):
                                _quebras(it.get("breakdowns"), acc)
                margem_out["amostras"][canal] = {
                    "pedido": numero_por_externo.get(str(f.external_order_id)),
                    "liquido": _num(Decimal(str(f.net_amount)))
                    if f.net_amount is not None
                    else None,
                    "transacoes": dict(tipos_tx),
                    "quebras": {k: _num(v) for k, v in acc.items()},
                }
    margem_out["status"] = dict(margem_out["status"])
    margem_out["liquido_posted"] = _num(margem_out["liquido_posted"])
    # 6) Melhor Envio pelo token do DaVinci: dá pra ler os envios e o valor pago?
    me: dict[str, Any] = {}
    try:
        cli_me = MelhorEnvioClient()
        if not cli_me.token:
            me = {"ok": False, "erro": "sem token"}
        else:
            headers = {
                "Authorization": f"Bearer {cli_me.token}",
                "Accept": "application/json",
                "User-Agent": "DaVinci Conciliacao (contato@davinci)",
            }
            async with httpx.AsyncClient(timeout=30) as c:
                r1 = await c.get(f"{cli_me.base_url}/api/v2/me", headers=headers)
                r2 = await c.get(
                    f"{cli_me.base_url}/api/v2/me/orders", headers=headers, params={"page": 1}
                )
            me = {"conta_http": r1.status_code, "envios_http": r2.status_code}
            if r1.status_code == 200:
                j1 = r1.json() or {}
                me["conta"] = {k: j1.get(k) for k in ("firstname", "company") if k in j1}
            if r2.status_code == 200:
                j2 = r2.json() or {}
                dados = j2.get("data") if isinstance(j2, dict) else j2
                dados = dados if isinstance(dados, list) else []
                me["envios_na_pagina"] = len(dados)
                me["total"] = j2.get("total") if isinstance(j2, dict) else None
                if dados and isinstance(dados[0], dict):
                    e0 = dados[0]
                    me["campos"] = sorted(e0.keys())
                    me["exemplo"] = {
                        "price": e0.get("price"),
                        "discount": e0.get("discount"),
                        "status": e0.get("status"),
                        "tags": [
                            t.get("tag") for t in (e0.get("tags") or []) if isinstance(t, dict)
                        ],
                        "conciliation": e0.get("conciliation"),
                        "tem_payment": "payment" in e0,
                    }
            else:
                me["mensagem"] = r2.text[:200]
    except Exception as e:  # noqa: BLE001
        me = {"ok": False, "erro": f"{type(e).__name__}: {str(e)[:200]}"}
    logger.info(
        "conciliacao_sonda_amazon",
        conta=integ.name,
        chamadas=[
            (c["nome"], c.get("status"), c.get("erro") or c.get("mensagem")) for c in chamadas
        ],
    )
    return {
        "conta": integ.name,
        "integration_id": str(integ.id),
        "janela": {"de": de.isoformat(), "ate": agora.isoformat()},
        "chamadas": chamadas,
        "transacoes": transacoes,
        "transacoes_por_situacao": por_situacao,
        "grupos_repasse": grupos[:20],
        "relatorios_repasse": relatorios[:20],
        "relatorio_baixado": relatorio_baixado,
        "relatorios_baixados": relatorios_baixados,
        "margem_financeiro": margem,
        "margem_desde_01_10": margem_out,
        "pedidos_desde_01_10": {
            "pedidos": len(numeros),
            "canais": dict(canais),
            "servicos": dict(servicos.most_common(8)),
        },
        "melhor_envio": me,
    }


# ---------------------------------------------------------------------------
# Token do Melhor Envio colado pela operação (15/10, Vinicius: "me mostra onde eu colo")
# ---------------------------------------------------------------------------
def _resumo_token_me(token: str | None, detalhe: dict[str, Any]) -> dict[str, Any]:
    from app.config import get_settings

    tem_painel = bool(token)
    tem_servidor = bool((get_settings().melhor_envio_token or "").strip())
    return {
        "tem_token": tem_painel or tem_servidor,
        "origem": "painel" if tem_painel else ("servidor" if tem_servidor else None),
        "atualizado_em": detalhe.get("atualizado_em"),
        "atualizado_por": detalhe.get("atualizado_por"),
        "conta": detalhe.get("conta"),
        "testado_em": detalhe.get("testado_em"),
        "aviso": (
            None
            if tem_painel
            else (
                "o token do servidor está vencido (401): cole um novo aqui"
                if tem_servidor
                else "sem token: cole o token gerado no Melhor Envio"
            )
        ),
    }


@router.get("/melhor-envio/token")
async def melhor_envio_token_status(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
) -> dict[str, Any]:
    from app.services.melhor_envio import token_do_painel

    token, detalhe = await token_do_painel(session)
    return _resumo_token_me(token, detalhe)


@router.put("/melhor-envio/token")
async def melhor_envio_token_salvar(
    body: Annotated[dict[str, Any], Body()],
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    """Confere o token no Melhor Envio (GET /me) e guarda cifrado. O token nunca
    volta na resposta nem vai pro log."""
    from app.services.melhor_envio import (
        guardar_token_do_painel,
        ler_conta,
        nome_da_conta,
        token_do_painel,
    )

    token = str(body.get("token") or "").strip()
    if len(token) < 20:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail={"code": "token_vazio", "mensagem": "cole o token inteiro"},
        )
    http, corpo = await ler_conta(token)
    if http != 200:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail={
                "code": "token_recusado",
                "mensagem": (
                    f"o Melhor Envio recusou o token (HTTP {http}): "
                    f"{str(corpo.get('message') or corpo)[:120]}"
                ),
            },
        )
    agora = datetime.now(UTC).isoformat()
    await guardar_token_do_painel(
        session,
        token,
        por=(user.name or user.email),
        detalhe={"conta": nome_da_conta(corpo), "testado_em": agora},
    )
    await session.commit()
    logger.info("conciliacao_melhor_envio_token_salvo", por=user.email, conta=nome_da_conta(corpo))
    novo, detalhe = await token_do_painel(session)
    return _resumo_token_me(novo, detalhe)


@router.post("/melhor-envio/token/testar")
async def melhor_envio_token_testar(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM_PAINEL)],
) -> dict[str, Any]:
    """Com o token guardado: quem é a conta e quantos envios a lista devolve."""
    from app.services.melhor_envio import ler_conta, ler_envios, nome_da_conta, token_do_painel

    token, _ = await token_do_painel(session)
    if not token:
        from app.config import get_settings

        token = (get_settings().melhor_envio_token or "").strip()
    if not token:
        return {"ok": False, "mensagem": "sem token"}
    http, corpo = await ler_conta(token)
    if http != 200:
        return {"ok": False, "http": http, "mensagem": str(corpo.get("message") or corpo)[:160]}
    http2, envios = await ler_envios(token)
    dados = envios.get("data") if isinstance(envios.get("data"), list) else []
    return {
        "ok": True,
        "conta": nome_da_conta(corpo),
        "envios_http": http2,
        "envios_total": envios.get("total") if http2 == 200 else None,
        "envios_na_pagina": len(dados),
        "campos": sorted(dados[0].keys()) if dados and isinstance(dados[0], dict) else [],
    }


@router.delete("/melhor-envio/token")
async def melhor_envio_token_apagar(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    from app.services.melhor_envio import apagar_token_do_painel, token_do_painel

    apagou = await apagar_token_do_painel(session)
    await session.commit()
    logger.info("conciliacao_melhor_envio_token_apagado", por=user.email, apagou=apagou)
    token, detalhe = await token_do_painel(session)
    return _resumo_token_me(token, detalhe)


@router.get("/sonda/melhor-envio")
async def sonda_melhor_envio(
    session: Annotated[AsyncSession, Depends(get_session)],
    _user: Annotated[User, Depends(_PERM)],
    q: Annotated[
        str | None, Query(description="busca do Melhor Envio (nº do pedido, rastreio…)")
    ] = None,
    page: Annotated[int, Query(ge=1, le=50)] = 1,
) -> dict[str, Any]:
    """Como o Melhor Envio devolve um envio (só o que interessa pra conciliação; sem
    nome/endereço do cliente): preço pago, desconto, conferência de peso, etiquetas
    (tags) e campos de referência. Só leitura (15/10)."""
    from app.services.melhor_envio import ler_envios, token_do_painel

    token, _ = await token_do_painel(session)
    if not token:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail={"code": "sem_token_melhor_envio"})
    http, corpo = await ler_envios(token, page=page, q=q)
    dados = corpo.get("data") if isinstance(corpo.get("data"), list) else []

    def compacto(e: dict[str, Any]) -> dict[str, Any]:
        servico = e.get("service") if isinstance(e.get("service"), dict) else {}
        tags = e.get("tags") if isinstance(e.get("tags"), list) else []
        conc = e.get("conciliation")
        return {
            "id": e.get("id"),
            "protocol": e.get("protocol"),
            "status": e.get("status"),
            "servico": servico.get("name"),
            "price": e.get("price"),
            "discount": e.get("discount"),
            "insurance_value": e.get("insurance_value"),
            "billed_weight": e.get("billed_weight"),
            "paid_at": e.get("paid_at"),
            "posted_at": e.get("posted_at"),
            "delivered_at": e.get("delivered_at"),
            "canceled_at": e.get("canceled_at"),
            "tracking": e.get("tracking"),
            "self_tracking": e.get("self_tracking"),
            "tags": tags[:8],
            "additional_info_chaves": sorted(e["additional_info"].keys())
            if isinstance(e.get("additional_info"), dict)
            else e.get("additional_info"),
            "conciliation": conc if not isinstance(conc, list) else conc[:3],
            "freight_arrangements": e.get("freight_arrangements"),
            "products": [
                {k: p.get(k) for k in ("name", "quantity", "unitary_value") if k in p}
                for p in (e.get("products") or [])[:3]
                if isinstance(p, dict)
            ],
            "campos": sorted(e.keys()),
        }

    return {
        "http": http,
        "q": q,
        "total": corpo.get("total"),
        "pagina": page,
        "envios": [compacto(e) for e in dados[:5] if isinstance(e, dict)],
    }


def _numero_ou_422(body: dict[str, Any], campo: str, padrao: float) -> float:
    """Número ≥ 0 do corpo (vazio ou 0 = padrão, como antes); texto que não é número
    responde 422 em vez de 500 (16/10, revisão)."""
    bruto = body.get(campo)
    if not bruto:
        return padrao
    try:
        valor = float(bruto)
    except (TypeError, ValueError):
        valor = math.nan
    if not math.isfinite(valor) or valor < 0:
        raise HTTPException(
            422,
            detail={
                "code": "parametro_invalido",
                "campo": campo,
                "mensagem": f"{campo} precisa ser um número maior ou igual a zero",
            },
        )
    return valor


@router.post("/carga-inicial")
async def carga_inicial(
    body: Annotated[dict[str, Any], Body()],
    _user: Annotated[User, Depends(_PERM_PAINEL_EDIT)],
) -> dict[str, Any]:
    """Primeira carga das lojas recém-ligadas, em fila (15/10): um job por loja; as
    do ML espaçadas (`intervalo_min`, padrão 12) por causa do limite da fatura; as
    outras quase juntas. `contas` e `plataformas` filtram; `so_sem_carga` (padrão
    true) pega só loja que nunca recalculou."""
    from app.services.conciliacao_jobs import enfileirar_rodadas

    contas = {str(c).strip().lower() for c in (body.get("contas") or []) if str(c).strip()}
    plataformas = tuple(str(p) for p in (body.get("plataformas") or []) if p)
    intervalo_ml = _numero_ou_422(body, "intervalo_min", 12.0)
    intervalo_outras = _numero_ou_422(body, "intervalo_outras_seg", 30.0)
    plano = await enfileirar_rodadas(
        await get_arq_pool(),
        fatura=True,
        extrato=True,
        recalculo=True,
        tag="carga",
        intervalo_ml_min=intervalo_ml,
        intervalo_outras_seg=intervalo_outras,
        so_plataformas=plataformas or None,
        contas=contas or None,
        so_sem_carga=bool(body.get("so_sem_carga", True)),
    )
    return {"lojas": len(plano), "plano": plano}
