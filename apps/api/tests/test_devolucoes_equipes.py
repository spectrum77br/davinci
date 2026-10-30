"""Devoluções × equipe de vendas (30/10/2026).

Eduardo, 30/10: três equipes — 101 Londres, 102 Israel, 103 Marrocos — e quem
é de uma equipe só enxerga as lojas dela, em todas as telas. Em Devoluções a
LINHA é da equipe pelo pedido (`bling_orders.loja` entre as lojas dela) ou
pela conta escrita nela, na régua plataforma + nome ("Shopee Inova" é da
Marrocos, "ML Inova" da Israel); o PEDIDO do Bling (Acompanhamento/Fraude,
busca do pedido, atualizar clientes) é da equipe pela loja, como no Controle
de Estoque. Admin e quem não tem equipe não mudam; equipe sem loja cadastrada
não vê nada (nunca "tudo").

Cobre: lista + cards, busca, exportação, Acompanhamento/Fraude, busca do
pedido para lançar, linha única (ler, salvar, excluir, estoque, senha,
anexos), lançar e mover linha para pedido/conta de outra equipe, ações por
pedido (rastreio, vídeo, situação no Bling) e o "atualizar clientes".
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models import (
    BlingOrder,
    DevolucaoAnexo,
    DevolucaoRastreio,
    Devolution,
    SituacaoBling,
    StoreInfo,
    User,
    UserRole,
    UserStatus,
)
from app.routers.devolutions import acompanhamento_rows
from app.services import devolution_stock_return as dsr

pytestmark = pytest.mark.asyncio

PERM = {"devolucoes": {"view": True, "edit": True, "delete": True}}

LOJA_LONDRES = "9101"  # Shopee vortan (101)
LOJA_ISRAEL = "9102"  # Shopee atv (102)
LOJA_ISRAEL_ML = "9103"  # ML inova (102)
LOJA_MARROCOS = "9104"  # Shopee inova (103)
LOJA_SEM_CADASTRO = "9109"  # loja do Bling que nenhuma equipe tem

# O apelido de cada loja (stores.apelido_override): é o que a busca do pedido
# devolve como `conta` e o que a linha de devolução guarda.
APELIDOS = {
    LOJA_LONDRES: ("Shopee Vortan", "shopee"),
    LOJA_ISRAEL: ("Shopee ATV", "shopee"),
    LOJA_ISRAEL_ML: ("ML Inova", "ml"),
    LOJA_MARROCOS: ("Shopee Inova", "shopee"),
}

# Pedido → loja; todos em Aguardando Devolução (83957).
PEDIDOS = {
    "880001": LOJA_LONDRES,
    "880002": LOJA_ISRAEL,
    "880003": LOJA_ISRAEL_ML,
    "880004": LOJA_MARROCOS,
    "880009": LOJA_SEM_CADASTRO,
}

# As linhas de devolução: (chave, pedido, conta, reembolso).
LINHAS = (
    ("lon", "880001", "Shopee Vortan", True),
    ("lon_sem_pedido", None, "Shopee Vortan", False),  # Londres pela conta
    ("isr", "880002", "Shopee ATV", True),
    ("isr_ml", "880003", "ML Inova", False),
    # Pedido fora do espelho: Israel pela conta (ML inova) — e NÃO a Marrocos,
    # que tem a Shopee inova (o nome sozinho vazava entre equipes).
    ("isr_fora_espelho", "889001", "ML Inova", False),
    # Sem plataforma na conta: vale o nome que é só da equipe (atv)...
    ("isr_nome_puro", None, "atv", False),
    # ...e o nome repetido em equipes diferentes não é de ninguém.
    ("inova_puro", None, "Inova", False),
    ("mar", "880004", "Shopee Inova", False),
    ("sem_cadastro", "880009", "Loja 9109", False),
)
TODAS = {chave for chave, *_ in LINHAS}

VE = {
    "admin": TODAS,
    "galpao": TODAS,  # sem equipe: vê tudo, como sempre
    "londres": {"lon", "lon_sem_pedido"},
    "israel": {"isr", "isr_ml", "isr_fora_espelho", "isr_nome_puro"},
    "marrocos": {"mar"},
    "equipe_sem_loja": set(),
}
PEDIDOS_QUE_VE = {
    "admin": set(PEDIDOS),
    "galpao": set(PEDIDOS),
    "londres": {"880001"},
    "israel": {"880002", "880003"},
    "marrocos": {"880004"},
    "equipe_sem_loja": set(),
}


async def _user(
    db: AsyncSession, nome: str, *, role: UserRole = UserRole.USER, equipes: list[int] | None = None
) -> User:
    email = f"{nome}-{uuid.uuid4().hex[:6]}@davinci-test.com"
    u = User(
        open_id=f"email:{email}",
        email=email,
        name=nome,
        role=role,
        status=UserStatus.ACTIVE,
        permissions=PERM,
        sales_teams=equipes,
    )
    db.add(u)
    await db.commit()
    await db.refresh(u)
    return u


async def _criar_view(db: AsyncSession) -> None:
    """A vw_devolucoes de produção (0126) por cima do espelho semeado; o
    apelido/plataforma da loja vem de uma tabela de valores (no lugar de
    `stores`)."""
    schema = get_settings().database_schema
    valores = ", ".join(
        f"('{loja}', '{apelido}', '{plat}')" for loja, (apelido, plat) in APELIDOS.items()
    )
    await db.execute(text(f'DROP VIEW IF EXISTS "{schema}".vw_devolucoes'))
    await db.execute(
        text(
            f"""
            CREATE VIEW "{schema}".vw_devolucoes AS
            SELECT
                bo.id AS bling_order_item_id,
                bo.bling_id,
                bo.numero AS pedido_bling,
                bo.numeroloja AS pedido_marketplace,
                bo.data,
                bo.situacao,
                NULL::text AS situacao_nome,
                bo.loja AS bling_loja_id,
                s.apelido AS loja_nome,
                s.plat AS plataforma_bling,
                bo.item_codigo AS sku,
                bo.item_descricao AS produto,
                bo.item_quantidade AS quantidade,
                bo.nome_destinatario,
                bo.cep_destino,
                bo.endereco_destino,
                bo.numero_destino,
                bo.complemento_destino,
                bo.bairro_destino,
                bo.cidade_destino,
                bo.uf_destino
            FROM "{schema}".bling_orders bo
            LEFT JOIN (VALUES {valores}) AS s(loja, apelido, plat) ON s.loja = bo.loja
            WHERE bo.situacao IN ('83957', '83960', '83961', '83966', '84677', '545902')
            """  # noqa: S608
        )
    )
    await db.commit()


async def _derrubar_view(db: AsyncSession) -> None:
    schema = get_settings().database_schema
    await db.execute(text(f'DROP VIEW IF EXISTS "{schema}".vw_devolucoes'))
    await db.commit()


@pytest_asyncio.fixture
async def cenario(db: AsyncSession):
    admin = await _user(db, "admin", role=UserRole.ADMIN)
    db.add_all(
        [
            StoreInfo(
                user_id=admin.id,
                platform="shopee",
                account_name="vortan",
                bling_store_id=LOJA_LONDRES,
                sales_team=101,
            ),
            StoreInfo(
                user_id=admin.id,
                platform="shopee",
                account_name="atv",
                bling_store_id=LOJA_ISRAEL,
                sales_team=102,
            ),
            # "mercadolivre" no cadastro, "ML" na frente da conta: a mesma loja.
            StoreInfo(
                user_id=admin.id,
                platform="mercadolivre",
                account_name="inova",
                bling_store_id=LOJA_ISRAEL_ML,
                sales_team=102,
            ),
            StoreInfo(
                user_id=admin.id,
                platform="shopee",
                account_name="inova",
                bling_store_id=LOJA_MARROCOS,
                sales_team=103,
            ),
            SituacaoBling(id=83957, nome="Aguardando Devolução"),
            SituacaoBling(id=83960, nome="Problemas"),
        ]
    )
    agora = datetime.now(UTC)
    for numero, loja in PEDIDOS.items():
        db.add(
            BlingOrder(
                bling_id=int(numero),
                numero=numero,
                numeroloja=f"MK{numero}",
                loja=loja,
                situacao="83957",
                item_index=0,
                item_codigo=f"x{numero[-3:]}.ci",
                item_descricao=f"Produto {numero}",
                item_quantidade=1,
                data=agora - timedelta(days=5),
                aguardando_devolucao_data=(agora - timedelta(days=2)).date(),
            )
        )
    ids: dict[str, uuid.UUID] = {}
    for chave, pedido, conta, reembolso in LINHAS:
        dev = Devolution(
            pedido_bling=pedido,
            conta=conta,
            sku="x000.ci",
            reembolso=reembolso,
            data=agora - timedelta(days=1),
        )
        db.add(dev)
        await db.flush()
        ids[chave] = dev.id
    await db.commit()
    usuarios = {
        "admin": admin,
        "galpao": await _user(db, "galpao"),
        "londres": await _user(db, "londres", equipes=[101]),
        "israel": await _user(db, "israel", equipes=[102]),
        "marrocos": await _user(db, "marrocos", equipes=[103]),
        # Equipe sem loja cadastrada (bug de cadastro): não vê nada.
        "equipe_sem_loja": await _user(db, "sem-loja", equipes=[104]),
    }
    await _criar_view(db)
    yield {"u": usuarios, "ids": ids}
    await _derrubar_view(db)


def _chaves(cenario: dict, items: list[dict]) -> set[str]:
    por_id = {str(v): k for k, v in cenario["ids"].items()}
    return {por_id.get(i["id"], i["id"]) for i in items}


async def _existe(db: AsyncSession, dev_id) -> bool:
    await db.commit()  # transação nova: enxerga o que a API gravou
    n = (
        await db.execute(
            select(func.count()).select_from(Devolution).where(Devolution.id == dev_id)
        )
    ).scalar_one()
    return bool(n)


# ── Lista, cards, busca e exportação ─────────────────────────────────────


@pytest.mark.parametrize("quem", sorted(VE))
async def test_lista_e_cards_so_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str
):
    auth_as(cenario["u"][quem])
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert r.status_code == 200, r.text
    body = r.json()
    esperado = VE[quem]
    assert _chaves(cenario, body["items"]) == esperado
    # Os cards contam o MESMO recorte da lista.
    assert body["total"] == len(esperado)
    linhas = [linha for linha in LINHAS if linha[0] in esperado]
    com_pedido = {p for _, p, _, _ in linhas if p}
    sem_pedido = sum(1 for _, p, _, _ in linhas if not p)
    assert body["total_pedidos"] == len(com_pedido) + sem_pedido
    assert body["reembolso_itens"] == sum(1 for *_, reemb in linhas if reemb)


async def test_busca_e_filtro_seguem_dentro_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    """A busca procura dentro da equipe: "Inova" para a Israel traz só as
    linhas do ML inova, não a Shopee inova da Marrocos nem o nome solto."""
    auth_as(cenario["u"]["israel"])
    r = await client.get("/api/devolutions", params={"search": "Inova"})
    assert r.status_code == 200
    assert _chaves(cenario, r.json()["items"]) == {"isr_ml", "isr_fora_espelho"}
    r = await client.get("/api/devolutions", params={"reembolso": "true"})
    assert _chaves(cenario, r.json()["items"]) == {"isr"}

    auth_as(cenario["u"]["admin"])
    r = await client.get("/api/devolutions", params={"search": "Inova"})
    assert _chaves(cenario, r.json()["items"]) == {
        "isr_ml",
        "isr_fora_espelho",
        "inova_puro",
        "mar",
    }


@pytest.mark.parametrize("quem", ["admin", "londres", "marrocos", "equipe_sem_loja"])
async def test_exportacao_so_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str
):
    auth_as(cenario["u"][quem])
    r = await client.get("/api/devolutions/export.xlsx")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    linhas = list(ws.iter_rows(values_only=True))
    cab = list(linhas[0])
    contas = sorted(str(lin[cab.index("Conta")]) for lin in linhas[1:])
    esperado = sorted(conta for chave, _, conta, _ in LINHAS if chave in VE[quem])
    assert contas == esperado


# ── Acompanhamento / Fraude ──────────────────────────────────────────────


@pytest.mark.parametrize("quem", sorted(PEDIDOS_QUE_VE))
async def test_acompanhamento_e_fraude_so_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str
):
    """As abas Acompanhamento e Fraude (e os contadores delas) saem da mesma
    lista: o pedido da Israel movido para Fraude só aparece para a Israel."""
    db.add(DevolucaoRastreio(pedido_bling="880002", fila_manual="fraude"))
    await db.commit()
    auth_as(cenario["u"][quem])
    r = await client.get("/api/devolutions/acompanhamento")
    assert r.status_code == 200, r.text
    body = r.json()
    pedidos = {i["pedido_bling"] for i in body["items"]}
    assert pedidos == PEDIDOS_QUE_VE[quem]
    assert body["total_pedidos"] == len(PEDIDOS_QUE_VE[quem])
    fraude = {i["pedido_bling"] for i in body["items"] if i["fila"] == "fraude"}
    assert fraude == ({"880002"} & PEDIDOS_QUE_VE[quem])


async def test_informar_continua_com_todos_os_pedidos(db: AsyncSession, cenario):
    """O relatório Informar chama a mesma consulta sem cerca: nada muda."""
    todos = await acompanhamento_rows(db)
    assert {r["pedido_bling"] for r in todos} == set(PEDIDOS)
    assert await acompanhamento_rows(db, lojas=set()) == []
    so_londres = await acompanhamento_rows(db, lojas={LOJA_LONDRES})
    assert {r["pedido_bling"] for r in so_londres} == {"880001"}


# ── Busca do pedido para lançar ──────────────────────────────────────────


@pytest.mark.parametrize(
    ("quem", "pedido", "acha"),
    [
        ("londres", "880001", True),
        ("londres", "880002", False),
        ("israel", "880002", True),
        ("israel", "880004", False),
        ("marrocos", "880004", True),
        ("equipe_sem_loja", "880001", False),
        ("admin", "880002", True),
        ("galpao", "880009", True),
    ],
)
async def test_busca_do_pedido_so_acha_pedido_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str, pedido: str, acha: bool
):
    auth_as(cenario["u"][quem])
    r = await client.get("/api/devolutions/order-lookup", params={"pedido": pedido})
    assert r.status_code == 200, r.text
    rows = r.json()
    if acha:
        assert rows and {x["pedido_bling"] for x in rows} == {pedido}
        assert rows[0]["conta"] == APELIDOS.get(PEDIDOS[pedido], ("Loja 9109",))[0]
    else:
        assert rows == []


# ── Linha de outra equipe: ler, salvar, excluir, estoque, senha, anexos ──


async def test_linha_de_outra_equipe_e_recusada(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    ids = cenario["ids"]
    alvo = ids["isr"]
    anexo = DevolucaoAnexo(
        devolution_id=alvo,
        filename="foto.jpg",
        content_type="image/jpeg",
        size_bytes=3,
        blob=b"abc",
        created_by=cenario["u"]["admin"].id,
    )
    db.add(anexo)
    await db.commit()
    auth_as(cenario["u"]["londres"])

    chamadas = [
        ("GET", f"/api/devolutions/{alvo}", {}),
        ("PATCH", f"/api/devolutions/{alvo}", {"json": {"observacao": "mexi"}}),
        ("POST", f"/api/devolutions/{alvo}/estoque/tentar-de-novo", {"json": {}}),
        ("POST", f"/api/devolutions/{alvo}/estoque/ja-entrou", {"json": {"sku": "x002.ci"}}),
        ("POST", f"/api/devolutions/{alvo}/senha/pedir-de-novo", {}),
        (
            "POST",
            f"/api/devolutions/{alvo}/anexos",
            {"files": {"file": ("f.jpg", b"\xff\xd8\xff", "image/jpeg")}},
        ),
        ("GET", f"/api/devolutions/anexos/{anexo.id}", {}),
        ("DELETE", f"/api/devolutions/anexos/{anexo.id}", {}),
        ("DELETE", f"/api/devolutions/{alvo}", {}),
    ]
    for metodo, url, kw in chamadas:
        r = await client.request(metodo, url, **kw)
        assert r.status_code == 403, (metodo, url, r.status_code, r.text)
        detalhe = r.json()["detail"]
        assert detalhe["code"] == "pedido_de_outra_equipe"
        assert detalhe["message"]
        # Pelo id da linha, o número do pedido de outra equipe não volta.
        assert "pedidos" not in detalhe
        assert "880002" not in r.text

    # Nada mudou: a linha, o texto e o anexo continuam lá.
    assert await _existe(db, alvo)
    obs = (
        await db.execute(select(Devolution.observacao).where(Devolution.id == alvo))
    ).scalar_one()
    assert obs is None
    n_anexos = (
        await db.execute(
            select(func.count())
            .select_from(DevolucaoAnexo)
            .where(DevolucaoAnexo.devolution_id == alvo)
        )
    ).scalar_one()
    assert n_anexos == 1

    # Linha que não existe continua 404 (não vira 403).
    r = await client.get(f"/api/devolutions/{uuid.uuid4()}")
    assert r.status_code == 404


async def test_linha_da_propria_equipe_continua_funcionando(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    ids = cenario["ids"]
    auth_as(cenario["u"]["londres"])
    for chave in ("lon", "lon_sem_pedido"):
        r = await client.get(f"/api/devolutions/{ids[chave]}")
        assert r.status_code == 200, r.text
    r = await client.patch(f"/api/devolutions/{ids['lon']}", json={"observacao": "conferido"})
    assert r.status_code == 200, r.text
    assert r.json()["observacao"] == "conferido"
    # Sem pendência de estoque: passa da cerca e cai na regra de sempre (409).
    r = await client.post(f"/api/devolutions/{ids['lon']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "estoque_nao_pendente"
    r = await client.post(
        f"/api/devolutions/{ids['lon']}/anexos",
        files={"file": ("f.jpg", b"\xff\xd8\xff", "image/jpeg")},
    )
    assert r.status_code == 201, r.text
    anexo_id = r.json()["anexos"][0]["id"]
    r = await client.get(f"/api/devolutions/anexos/{anexo_id}")
    assert r.status_code == 200
    r = await client.delete(f"/api/devolutions/anexos/{anexo_id}")
    assert r.status_code == 204

    # Israel enxerga e mexe na linha que é dela pela conta (pedido fora do espelho).
    auth_as(cenario["u"]["israel"])
    r = await client.patch(f"/api/devolutions/{ids['isr_fora_espelho']}", json={"observacao": "ok"})
    assert r.status_code == 200, r.text


@pytest.mark.parametrize("quem", ["admin", "galpao"])
async def test_admin_e_sem_equipe_mexem_em_qualquer_linha(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str
):
    auth_as(cenario["u"][quem])
    for chave in ("isr", "mar", "sem_cadastro", "inova_puro"):
        r = await client.get(f"/api/devolutions/{cenario['ids'][chave]}")
        assert r.status_code == 200, (chave, r.text)
    r = await client.patch(
        f"/api/devolutions/{cenario['ids']['mar']}", json={"observacao": "admin"}
    )
    assert r.status_code == 200


async def test_equipe_sem_loja_nao_mexe_em_nada(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    auth_as(cenario["u"]["equipe_sem_loja"])
    for chave in ("lon", "isr", "sem_cadastro", "inova_puro"):
        r = await client.get(f"/api/devolutions/{cenario['ids'][chave]}")
        assert r.status_code == 403, chave


# ── Lançar / mover linha para pedido ou conta de outra equipe ────────────


@pytest.mark.parametrize(
    ("pedido", "conta"),
    [
        ("880002", "Shopee Vortan"),  # a conta da própria equipe não salva
        ("880002", "Shopee ATV"),
        ("880009", "Shopee Vortan"),  # loja do Bling sem equipe
        ("889002", "Shopee ATV"),  # fora do espelho: manda a conta
        (None, "Shopee ATV"),
        (None, "Inova"),  # nome repetido entre equipes
    ],
)
async def test_lancar_pedido_de_outra_equipe_e_recusado(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, pedido: str | None, conta: str
):
    auth_as(cenario["u"]["londres"])
    antes = (await db.execute(select(func.count()).select_from(Devolution))).scalar_one()
    r = await client.post(
        "/api/devolutions", json={"pedido_bling": pedido, "conta": conta, "sku": "x002.ci"}
    )
    assert r.status_code == 403, r.text
    assert r.json()["detail"]["code"] == "pedido_de_outra_equipe"
    await db.commit()
    depois = (await db.execute(select(func.count()).select_from(Devolution))).scalar_one()
    assert depois == antes


@pytest.mark.parametrize(
    ("quem", "pedido", "conta"),
    [
        ("londres", "880001", "Shopee Vortan"),
        # O apelido da busca nem sempre bate com o cadastro: a loja decide.
        ("londres", "880001", "Loja 9101"),
        ("londres", "889003", "Shopee Vortan"),  # fora do espelho, conta da equipe
        ("londres", None, "Shopee - Vortan"),
        ("israel", "880003", "ML Inova"),
        ("admin", "880002", "Shopee ATV"),
        ("galpao", "880009", "Loja 9109"),
    ],
)
async def test_lancar_pedido_da_propria_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, quem: str, pedido, conta: str
):
    auth_as(cenario["u"][quem])
    r = await client.post(
        "/api/devolutions", json={"pedido_bling": pedido, "conta": conta, "sku": "x001.ci"}
    )
    assert r.status_code == 201, r.text
    # E a linha nova aparece na lista de quem lançou.
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert r.json()["total"] == len(VE[quem]) + 1


async def test_mover_linha_para_outra_equipe_e_recusado(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    ids = cenario["ids"]
    auth_as(cenario["u"]["londres"])
    for body in (
        {"pedido_bling": "880002"},  # pedido de outra equipe
        {"conta": "Shopee ATV"},  # conta de outra equipe no pedido da equipe
        {"pedido_bling": None, "conta": "Shopee Inova"},
    ):
        r = await client.patch(f"/api/devolutions/{ids['lon']}", json=body)
        assert r.status_code == 403, (body, r.text)
        assert r.json()["detail"]["code"] == "pedido_de_outra_equipe"
    await db.commit()
    row = (
        await db.execute(
            select(Devolution.pedido_bling, Devolution.conta).where(Devolution.id == ids["lon"])
        )
    ).one()
    assert tuple(row) == ("880001", "Shopee Vortan")
    # Trocar por outra grafia da conta da própria equipe passa.
    r = await client.patch(f"/api/devolutions/{ids['lon']}", json={"conta": "Shopee - Vortan"})
    assert r.status_code == 200, r.text


# ── Ações por pedido: rastreio, vídeo, situação no Bling ─────────────────


async def test_acoes_por_pedido_de_outra_equipe_sao_recusadas(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    # O admin pede o vídeo do pedido da Israel (pra existir o que cancelar).
    auth_as(cenario["u"]["admin"])
    r = await client.post("/api/devolutions/acompanhamento/880002/video/solicitar", json={})
    assert r.status_code == 200, r.text

    auth_as(cenario["u"]["londres"])
    chamadas = [
        ("PATCH", "/api/devolutions/acompanhamento/880002", {"json": {"observacao": "x"}}),
        ("POST", "/api/devolutions/acompanhamento/880002/video/solicitar", {"json": {}}),
        ("DELETE", "/api/devolutions/acompanhamento/880002/video", {}),
        ("POST", "/api/devolutions/pedido/880002/situacao-bling", {"json": {"situacao_id": 83960}}),
        # Pedido que nem existe: não dá pra provar que é da equipe.
        ("PATCH", "/api/devolutions/acompanhamento/999999", {"json": {"observacao": "x"}}),
    ]
    for metodo, url, kw in chamadas:
        r = await client.request(metodo, url, **kw)
        assert r.status_code == 403, (metodo, url, r.status_code, r.text)
        detalhe = r.json()["detail"]
        assert detalhe["code"] == "pedido_de_outra_equipe"
        assert len(detalhe["pedidos"]) == 1

    await db.commit()
    rast = (
        await db.execute(
            select(DevolucaoRastreio).where(DevolucaoRastreio.pedido_bling == "880002")
        )
    ).scalar_one()
    assert rast.observacao is None
    assert rast.video_solicitado_em is not None  # o cancelamento não passou
    assert (await db.get(DevolucaoRastreio, "999999")) is None


async def test_acoes_por_pedido_da_propria_equipe_continuam(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    auth_as(cenario["u"]["londres"])
    r = await client.patch(
        "/api/devolutions/acompanhamento/880001", json={"observacao": "chega amanhã"}
    )
    assert r.status_code == 200, r.text
    r = await client.post("/api/devolutions/acompanhamento/880001/video/solicitar", json={})
    assert r.status_code == 200, r.text
    assert r.json()["video_status"] == "pendente"
    r = await client.delete("/api/devolutions/acompanhamento/880001/video")
    assert r.status_code == 200, r.text
    # Situação: passa da cerca e para no Bling (sem integração no teste).
    r = await client.post(
        "/api/devolutions/pedido/880001/situacao-bling", json={"situacao_id": 83960}
    )
    assert r.status_code == 503
    assert r.json()["detail"]["code"] == "bling_sem_integracao"

    # Admin: pedido que não existe continua 404.
    auth_as(cenario["u"]["admin"])
    r = await client.patch("/api/devolutions/acompanhamento/999999", json={"observacao": "x"})
    assert r.status_code == 404


# ── Atualizar clientes (endereços do Bling) ──────────────────────────────


class _BlingFalso:
    def __init__(self) -> None:
        self.pedidos: list[int] = []

    async def get_order(self, bling_id: int) -> dict:
        self.pedidos.append(int(bling_id))
        return {}


@pytest.mark.parametrize(
    ("quem", "pedidos"),
    [
        ("londres", {"880001"}),
        ("israel", {"880002", "880003"}),
        ("equipe_sem_loja", set()),
        ("admin", set(PEDIDOS)),
        ("galpao", set(PEDIDOS)),
    ],
)
async def test_atualizar_clientes_so_da_equipe(
    client: AsyncClient, db: AsyncSession, auth_as, cenario, monkeypatch, quem: str, pedidos
):
    fake = _BlingFalso()

    async def _cliente(_session):
        return fake

    monkeypatch.setattr(dsr, "_get_bling_client", _cliente)
    auth_as(cenario["u"][quem])
    r = await client.post("/api/devolutions/backfill-addresses")
    assert r.status_code == 200, r.text
    assert r.json()["processed"] == len(pedidos)
    assert {str(p) for p in fake.pedidos} == pedidos


# ── Pedido no espelho: só a loja decide (revisão de 30/10) ───────────────


def _ids(body: dict) -> set[str]:
    return {i["id"] for i in body["items"]}


async def test_conta_de_outra_equipe_em_pedido_da_equipe_nao_abre_para_a_outra(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    """A Londres lança o pedido 880001 (loja dela) escrevendo a conta da
    Israel ("Shopee ATV"). O lançamento passa — a loja do pedido é da
    Londres —, mas a conta escrita não abre a linha nem o pedido para a
    Israel: com o pedido no espelho, só a loja dele decide. Antes a Israel
    via e salvava a linha, e mexia no rastreio, no vídeo e na situação do
    pedido da Londres no Bling."""
    auth_as(cenario["u"]["londres"])
    r = await client.post(
        "/api/devolutions",
        json={"pedido_bling": "880001", "conta": "Shopee ATV", "sku": "x001.ci"},
    )
    assert r.status_code == 201, r.text
    novo = r.json()["id"]
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert novo in _ids(r.json())

    auth_as(cenario["u"]["israel"])
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert r.status_code == 200, r.text
    assert _chaves(cenario, r.json()["items"]) == VE["israel"]
    assert r.json()["total"] == len(VE["israel"])
    for busca in ("ATV", "880001"):
        r = await client.get("/api/devolutions", params={"search": busca})
        assert novo not in _ids(r.json()), busca
    r = await client.get("/api/devolutions/export.xlsx")
    assert r.status_code == 200
    linhas = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert len(linhas) - 1 == len(VE["israel"])

    # A linha (pelo id): 403 sem contar de que pedido ela é.
    for metodo, url, kw in (
        ("GET", f"/api/devolutions/{novo}", {}),
        ("PATCH", f"/api/devolutions/{novo}", {"json": {"observacao": "israel mexeu"}}),
        ("DELETE", f"/api/devolutions/{novo}", {}),
    ):
        r = await client.request(metodo, url, **kw)
        assert r.status_code == 403, (metodo, url, r.status_code, r.text)
        assert r.json()["detail"]["code"] == "pedido_de_outra_equipe"
        assert "880001" not in r.text
    # O pedido da Londres: 403 antes do Bling (sem integração seria 503).
    for metodo, url, kw in (
        ("PATCH", "/api/devolutions/acompanhamento/880001", {"json": {"observacao": "israel"}}),
        ("POST", "/api/devolutions/acompanhamento/880001/video/solicitar", {"json": {}}),
        ("POST", "/api/devolutions/pedido/880001/situacao-bling", {"json": {"situacao_id": 83960}}),
    ):
        r = await client.request(metodo, url, **kw)
        assert r.status_code == 403, (metodo, url, r.status_code, r.text)
        detalhe = r.json()["detail"]
        assert detalhe["code"] == "pedido_de_outra_equipe"
        assert detalhe["pedidos"] == ["880001"]

    # Nada mudou: a linha segue lá sem o recado, o pedido sem rastreio/vídeo.
    assert await _existe(db, uuid.UUID(novo))
    obs = (
        await db.execute(select(Devolution.observacao).where(Devolution.id == uuid.UUID(novo)))
    ).scalar_one()
    assert obs is None
    rast = await db.get(DevolucaoRastreio, "880001")
    assert rast is None or (rast.observacao is None and rast.video_solicitado_em is None)
    # E na Acompanhamento da Israel o pedido da Londres não aparece.
    r = await client.get("/api/devolutions/acompanhamento")
    assert {i["pedido_bling"] for i in r.json()["items"]} == PEDIDOS_QUE_VE["israel"]

    # A Londres segue mexendo na linha (o pedido é dela).
    auth_as(cenario["u"]["londres"])
    r = await client.patch(f"/api/devolutions/{novo}", json={"observacao": "conferido"})
    assert r.status_code == 200, r.text


async def test_pedido_fora_do_espelho_continua_decidido_pela_conta(
    client: AsyncClient, db: AsyncSession, auth_as, cenario
):
    """O outro lado da régua: pedido que não está no espelho (ou está sem
    loja) é da equipe pela conta da linha — na lista, na linha e nas ações
    por pedido."""
    # 880010: no espelho, mas sem loja — conta como fora dele.
    db.add(
        BlingOrder(
            bling_id=880010,
            numero="880010",
            numeroloja="MK880010",
            loja=None,
            situacao="83957",
            item_index=0,
            item_codigo="x010.ci",
            item_quantidade=1,
            data=datetime.now(UTC) - timedelta(days=3),
        )
    )
    await db.commit()

    auth_as(cenario["u"]["londres"])
    novos: dict[str, str] = {}
    for pedido in ("889100", "880010"):
        r = await client.post(
            "/api/devolutions",
            json={"pedido_bling": pedido, "conta": "Shopee Vortan", "sku": "x001.ci"},
        )
        assert r.status_code == 201, (pedido, r.text)
        novos[pedido] = r.json()["id"]
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert set(novos.values()) <= _ids(r.json())
    assert r.json()["total"] == len(VE["londres"]) + 2
    # Ações por pedido passam da cerca: o fora do espelho para no 404, o sem loja salva.
    r = await client.patch("/api/devolutions/acompanhamento/889100", json={"observacao": "x"})
    assert r.status_code == 404, r.text
    r = await client.patch(
        "/api/devolutions/acompanhamento/880010", json={"observacao": "a caminho"}
    )
    assert r.status_code == 200, r.text
    # O pedido fora do espelho da Israel (889001, "ML Inova") não é da Londres.
    r = await client.patch("/api/devolutions/acompanhamento/889001", json={"observacao": "x"})
    assert r.status_code == 403, r.text

    auth_as(cenario["u"]["israel"])
    r = await client.get("/api/devolutions", params={"limit": 500})
    assert _chaves(cenario, r.json()["items"]) == VE["israel"]
    for novo in novos.values():
        r = await client.patch(f"/api/devolutions/{novo}", json={"observacao": "israel"})
        assert r.status_code == 403, r.text
    for pedido in novos:
        r = await client.patch(
            f"/api/devolutions/acompanhamento/{pedido}", json={"observacao": "israel"}
        )
        assert r.status_code == 403, (pedido, r.text)
        assert r.json()["detail"]["pedidos"] == [pedido]
    # E o fora do espelho dela passa da cerca (404: não há pedido para salvar).
    r = await client.patch("/api/devolutions/acompanhamento/889001", json={"observacao": "x"})
    assert r.status_code == 404, r.text

    await db.commit()
    rast = await db.get(DevolucaoRastreio, "880010")
    assert rast is not None and rast.observacao == "a caminho"
