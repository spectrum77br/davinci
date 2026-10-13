"""Filtro "tag destino" da aba Lançamentos (Vinicius 13/10).

`tag` é a tag do SKU vendido (origem). O item pode voltar ao estoque em outra
tag — saiu `.sp`, entrou `.sa` — e esse destino é o SKU que o Bling recebeu
(`estoque_mov_sku`, o selo verde da tela). `tag_destino` filtra por ele, com a
regra do Controle de Estoque; entrada estornada não conta.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.asyncio


def _perm() -> dict:
    return {"devolucoes": {"view": True, "edit": False, "delete": False}}


async def _seed(db: AsyncSession) -> None:
    await db.execute(
        text(
            """
            INSERT INTO devolutions
              (pedido_bling, conta, sku, tag, condicao_produto,
               estoque_mov_sku, estoque_mov_revertido_at)
            VALUES
              ('1001', 'ML', 'x001.sp', 'sp',   'Usado', 'x001.sa', NULL),   -- sp → sa
              ('1002', 'ML', 'x002.sp', 'sp',   'Novo',  'x002.sp', NULL),   -- sp → sp
              ('1003', 'ML', 'x003.sp', 'sp',   'Usado', 'x003.sa', now()),  -- estornado
              ('1004', 'ML', 'x004.sp', 'sp',   'Usado', NULL,      NULL),   -- não voltou
              ('1005', 'ML', 'b031.24', 'mala', 'Usado', 'z0361.mala', NULL),
              ('1006', 'ML', 'x006.sa', 'sa',   'Usado', 'z0400',   NULL)    -- z sem sufixo → us
            """
        )
    )
    await db.commit()


async def _pedidos(client, **params) -> set[str]:
    response = await client.get("/api/devolutions", params=params)
    assert response.status_code == 200
    return {item["pedido_bling"] for item in response.json()["items"]}


async def test_tag_destino_filtra_pelo_sku_que_entrou_no_estoque(
    client, db: AsyncSession, make_user, auth_as
):
    auth_as(await make_user(permissions=_perm()))
    await _seed(db)

    assert await _pedidos(client, tag_destino="sa") == {"1001"}
    assert await _pedidos(client, tag_destino="sp") == {"1002"}
    assert await _pedidos(client, tag_destino="mala") == {"1005"}
    assert await _pedidos(client, tag_destino="us") == {"1006"}
    assert await _pedidos(client, tag_destino="all") == {
        "1001", "1002", "1003", "1004", "1005", "1006",
    }


async def test_origem_e_destino_juntos_acham_quem_mudou_de_estoque(
    client, db: AsyncSession, make_user, auth_as
):
    auth_as(await make_user(permissions=_perm()))
    await _seed(db)

    assert await _pedidos(client, tag="sp", tag_destino="sa") == {"1001"}
    assert await _pedidos(client, tag="sa", tag_destino="sa") == set()


async def test_export_traz_coluna_tag_destino(
    client, db: AsyncSession, make_user, auth_as
):
    auth_as(await make_user(permissions=_perm()))
    await _seed(db)

    response = await client.get("/api/devolutions/export.xlsx", params={"tag": "sp"})

    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content)).active
    header = [c.value for c in ws[1]]
    col_pedido = header.index("Pedido Bling")
    col_destino = header.index("Tag destino")
    destino = {
        row[col_pedido]: row[col_destino] for row in ws.iter_rows(min_row=2, values_only=True)
    }
    # openpyxl lê célula vazia como None.
    assert destino == {"1001": "sa", "1002": "sp", "1003": None, "1004": None}
