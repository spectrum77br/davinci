"""Conferência — Mercado Livre e Amazon (07/10/2026): a mesma fila, o mesmo
relatório, coleta no SERVIDOR com um coletor FALSO.

O que fica travado:
  • perfil (plataformas.py): estados de cada plataforma, título e nome de
    arquivo iguais aos da tela (o JSON compartilhado com lib/conferencia.ts);
  • cálculo: métrica com estado fica sem número em linha, total e Geral (mesmo
    que a coleta mande), o relatório leva `plataforma` e `estados`, notas da
    plataforma; eletro de item pela marca do coletor, mas o mapa do DaVinci
    (SKU) ganha; o % do ML é só o Ads;
  • arquivos: o texto do estado no lugar do "—" (planilha, Excel em cinza,
    HTML com a etiqueta, CSV, MD), título e nome com a plataforma; Threema;
  • fila: uma rodada coletando POR plataforma; o lease do Mac só entrega
    Shopee; `lease_servidor` só ML/Amazon e da execução pedida; rodada parada
    para re-enfileirar;
  • job do servidor: ok, parcial (seção esperada faltando), erro (ColetorError,
    exceção, tempo, formato torto, sem coletor, conta apagada), fecha o
    relatório com o eletro pelo SKU do DaVinci;
  • rotas: ?plataforma= na lista, no "Gerar agora" (enfileira o job) e nas
    contas; integração da conta (mesma plataforma, uma vez só, ativa só
    ligada); a lista de integrações;
  • worker: agenda do ML/Amazon só com a chave; o job; o varredor re-enfileira.

Números INVENTADOS com a forma do ColetaDados (a do dry-run de 07/10/2026).
"""

from __future__ import annotations

import asyncio
import json
import sys
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app import worker, worker_pool
from app.config import get_settings
from app.models import (
    ConferenciaShopeeConta,
    ConferenciaShopeeExecucao,
    Integration,
    IntegrationPlatform,
    Product,
    ProductCategory,
    User,
    UserRole,
    UserStatus,
)
from app.routers import marketing_conferencia as mc
from app.services.conferencia_shopee import (
    calculo,
    coletores,
    fila,
    periodos,
    plataformas,
    saida,
    servidor,
)
from app.services.conferencia_shopee import threema_aviso as ta
from tests.test_conferencia_shopee_calculo import GERADO
from tests.test_conferencia_shopee_calculo import coleta as coleta_calc
from tests.test_conferencia_shopee_calculo import dados as dados_shopee
from tests.test_conferencia_shopee_calculo import execucao as execucao_calc
from tests.test_conferencia_shopee_calculo import semana as semana_shopee
from tests.test_conferencia_shopee_fila import AGORA, coletas_de, dados_loja, semear_contas

API = "/api/marketing/conferencia-shopee"
SEMANAS = periodos.semanas("semanal", AGORA.date())
_FIXTURE = Path(__file__).resolve().parents[3] / "apps" / "web" / "tests" / (
    "conferencia-arredondamento.json"
)


# ───────────────────────────────────────────────────────────── números


def semana_servidor(
    s: dict,
    *,
    vendas: float = 1000.0,
    ads: bool = True,
    ads_itens: bool = True,
    eletro_marca: bool | None = None,
    afiliados: bool = False,
) -> dict:
    """Uma semana do ColetaDados do servidor: Vendas pelos itens do Bling (SKU)
    e Ads do ML com os anúncios (MLB…) já com o SKU do vínculo."""
    item_eletro = {"item_id": "uaf001", "sku": "uaf001", "nome": "Air Fryer 4L",
                   "valor": 300.0, "pedidos": 3}
    item_cel = {"item_id": "dg053.sp", "sku": "dg053.sp", "nome": "Capinha",
                "valor": vendas - 300.0, "pedidos": 7}
    if eletro_marca is not None:
        item_eletro["eletro"] = eletro_marca
    out: dict = {
        "inicio": s["inicio"],
        "fim": s["fim"],
        "vendas": {"valor": vendas, "pedidos": 10},
        "vendas_itens": [item_cel, item_eletro],
        "avisos": [],
    }
    if ads:
        out["ads"] = {"impressoes": 5000, "cliques": 100, "gasto": 50.0, "vendas": 400.0,
                      "pedidos": 4}
        if ads_itens:
            out["ads_itens"] = [
                {"item_id": "MLB1", "sku": "uaf001", "nome": "Air Fryer", "impressoes": 2000,
                 "cliques": 40, "gasto": 20.0, "vendas": 200.0, "pedidos": 2},
                {"item_id": "MLB2", "sku": "dg053.sp", "nome": "Capinha", "impressoes": 3000,
                 "cliques": 60, "gasto": 30.0, "vendas": 200.0, "pedidos": 2},
            ]
    if afiliados:  # coletor com defeito mandando o que não devia
        out["afiliados"] = {"vendas": 999.0, "comissao": 99.0, "pedidos": 9, "cliques": 90}
    return out


def numeros(semanas: list[dict], **kw) -> dict:
    return {"versao": 1, "semanas": [semana_servidor(s, **kw) for s in semanas], "avisos": []}


def _grupo(rel: dict, chave: str) -> dict:
    return next(g for g in rel["grupos"] if g["chave"] == chave)


def _linha(rel: dict, grupo: str, conta: str) -> dict:
    return next(lin for lin in _grupo(rel, grupo)["linhas"] if lin["conta"] == conta)


def rel_servidor(plataforma: str, **kw) -> dict:
    return calculo.montar_relatorio(
        execucao_calc(plataforma=plataforma),
        [
            coleta_calc("Kia", "celular", numeros(SEMANAS, **kw), uid=None),
            coleta_calc("KFA", "mala", numeros(SEMANAS, **kw), uid=None),
        ],
        gerado_em=GERADO,
    )


# ───────────────────────────────────────────────────────────── perfil


def test_perfil_das_plataformas():
    assert plataformas.PLATAFORMAS == ("shopee", "ml", "amazon")
    assert plataformas.SERVIDOR == ("ml", "amazon")
    assert plataformas.estados("shopee") == {}
    ml = plataformas.estados("ml")
    assert {k for k, v in ml.items() if v == "nao_coletado"} == {
        "vendas_afiliados", "impressoes_afiliados", "conversao_afiliados", "invest_afiliados",
        "cliques_afiliados", "pedidos_afiliados",
    }
    assert ml["saldo_ads"] == "nao_se_aplica" and "pct" not in ml
    amazon = plataformas.estados("amazon")
    assert amazon["vendas_afiliados"] == "nao_se_aplica"
    assert {k for k, v in amazon.items() if v == "aguardando_acesso"} == {
        "vendas_ads", "impressoes", "conversao_ads", "invest_ads", "cliques_ads", "pedidos_ads",
        "pct",
    }
    # Toda chave de estado é uma métrica ou a linha sem métrica da planilha.
    for est in plataformas.ESTADOS.values():
        assert set(est) <= {*calculo.CHAVES, plataformas.IMPRESSOES_AFILIADOS}
        assert set(est.values()) <= set(plataformas.ESTADOS_VALIDOS)
    assert plataformas.valida(None) == "shopee" and plataformas.valida(" ML ") == "ml"
    with pytest.raises(ValueError):
        plataformas.valida("tiktok")
    assert plataformas.de({}) == "shopee" and plataformas.de({"plataforma": "xx"}) == "shopee"


def test_textos_iguais_aos_da_tela():
    """Os MESMOS casos que tests/conferencia-lib.cjs confere na lib da tela."""
    if not _FIXTURE.exists():
        pytest.skip("web fora do checkout")
    casos = json.loads(_FIXTURE.read_text(encoding="utf-8"))
    for chave, rot in casos["plataformas"]:
        assert plataformas.rotulo(chave) == rot
    for estado, texto in casos["estados"]:
        assert saida.texto_estado(estado) == texto
    for plataforma, inicio, fim, titulo in casos["titulos"]:
        rel = {"plataforma": plataforma, "semanas": [{"inicio": inicio, "fim": fim}]}
        assert saida.titulo(rel) == titulo
    for plataforma, inicio, fim, ext, nome in casos["nomes_arquivo"]:
        rel = {"plataforma": plataforma, "semanas": [{"inicio": inicio, "fim": fim}]}
        assert saida.nome_arquivo(rel, ext) == nome


# ───────────────────────────────────────────────────────────── cálculo


def test_ml_afiliados_nao_coletado_mesmo_que_a_coleta_mande():
    rel = rel_servidor("ml", afiliados=True)
    assert rel["plataforma"] == "ml"
    assert rel["estados"] == plataformas.estados("ml")
    kia = _linha(rel, "celular", "Kia")["semanas"][0]
    # Vendas 1000 − eletro 300; Ads 100 cliques − 40 do air fryer.
    assert (kia["vendas"], kia["cliques_ads"], kia["pedidos_ads"]) == (700.0, 60, 2)
    for c in ("vendas_afiliados", "invest_afiliados", "cliques_afiliados", "pedidos_afiliados",
              "conversao_afiliados", "saldo_ads"):
        assert kia[c] is None, c
        assert _grupo(rel, "celular")["total"]["semanas"][0][c] is None
        assert rel["geral"]["semanas"][0][c] is None
    # % do ML = só Ads: 30 ÷ 700.
    assert kia["pct"] == 4.29
    assert kia["conversao_ads"] == 3.33  # 2 ÷ 60
    eletro = _linha(rel, "eletro", "Kia")["semanas"][0]
    assert (eletro["vendas"], eletro["invest_ads"], eletro["conversao_ads"]) == (300.0, 20.0, 5.0)
    assert eletro["vendas_afiliados"] is None
    geral = rel["geral"]["semanas"][0]
    assert (geral["vendas"], geral["invest_ads"]) == (2000.0, 100.0)
    assert any("Mercado Livre" in n for n in rel["notas"])
    assert not any("Central do Vendedor da Shopee" in n for n in rel["notas"])


def test_amazon_ads_aguardando_e_afiliados_nao_se_aplica():
    rel = rel_servidor("amazon")
    kfa = _linha(rel, "mala", "KFA")["semanas"][0]
    assert kfa["vendas"] == 1000.0
    for c in ("vendas_ads", "impressoes", "invest_ads", "cliques_ads", "pedidos_ads",
              "conversao_ads", "pct", "vendas_afiliados", "invest_afiliados"):
        assert kfa[c] is None, c
        assert rel["geral"]["semanas"][0][c] is None
    assert rel["estados"]["pct"] == "aguardando_acesso"
    assert _linha(rel, "eletro", "Kia")["semanas"][0]["vendas"] == 300.0
    assert any("aguardando acesso" in n for n in rel["notas"])
    # Ads que veio mesmo assim (sem os anúncios por item) não vira aviso.
    rel = rel_servidor("amazon", ads_itens=False)
    assert _linha(rel, "celular", "Kia")["avisos"] == []
    assert rel["nao_atribuido_ads"] == []


def test_parcial_traz_o_trecho_comparado_na_nota_da_plataforma():
    rel = calculo.montar_relatorio(
        execucao_calc(plataforma="ml", tipo="parcial"),
        [coleta_calc("Kia", "celular", numeros(SEMANAS), uid=None)],
        gerado_em=GERADO,
    )
    assert rel["notas"][-1].startswith("Semana parcial:")


def test_marca_do_coletor_vale_e_o_mapa_do_davinci_ganha():
    # Marca do coletor: o air fryer NÃO é eletro (sem mapa) → a venda fica em
    # Celular, apesar do título; o anúncio dele (sem marca) segue o título.
    rel = rel_servidor("ml", eletro_marca=False)
    assert _linha(rel, "celular", "Kia")["semanas"][0]["vendas"] == 1000.0
    assert _linha(rel, "eletro", "Kia")["semanas"][0]["invest_ads"] == 20.0
    with_mapa = calculo.montar_relatorio(
        execucao_calc(plataforma="ml"),
        [coleta_calc("Kia", "celular", numeros(SEMANAS, eletro_marca=False), chave="c1",
                     uid=None)],
        mapa_davinci={"c1": {"uaf001": True, "MLB1": True}},
        gerado_em=GERADO,
    )
    assert _linha(with_mapa, "eletro", "Kia")["semanas"][0]["vendas"] == 300.0
    # Sem mapa e sem marca: o título decide ("Air Fryer").
    sem_nada = rel_servidor("ml")
    assert _linha(sem_nada, "eletro", "Kia")["semanas"][0]["vendas"] == 300.0


def test_sem_anuncios_por_item_o_ads_fica_em_celular_com_aviso():
    rel = rel_servidor("ml", ads_itens=False)
    kia = _linha(rel, "celular", "Kia")
    assert kia["semanas"][0]["invest_ads"] == 50.0
    assert any("sem os itens de Ads" in a for a in kia["avisos"])
    assert _linha(rel, "eletro", "Kia")["semanas"][0]["invest_ads"] is None


def test_shopee_continua_igual():
    rel = calculo.montar_relatorio(
        execucao_calc(),
        [coleta_calc("Inova", "mala", dados_shopee([semana_shopee(i) for i in range(4)]))],
        gerado_em=GERADO,
    )
    assert (rel["plataforma"], rel["estados"]) == ("shopee", {})
    assert _linha(rel, "mala", "Inova")["semanas"][0]["vendas_afiliados"] == 1000.0
    assert list(rel["notas"][: len(calculo.NOTAS_FIXAS)]) == list(calculo.NOTAS_FIXAS)


# ───────────────────────────────────────────────────────────── arquivos


def test_planilha_e_arquivos_com_o_texto_do_estado():
    rel = rel_servidor("ml")
    p = saida.planilha(rel)
    por_linha = {(lin["categoria"], lin["sub"]): lin for lin in p["linhas"]}
    imp_af = por_linha[("Impressões", "afiliados")]
    assert imp_af["estado"] == "nao_coletado"
    assert imp_af["variacoes"][0]["texto"] == "não coletado"
    assert por_linha[("Vendas", "afiliados")]["estado"] == "nao_coletado"
    assert por_linha[("Vendas", "Ads")]["estado"] is None
    assert por_linha[("Resumo", "Vendas no período")]["valores"][-1][-1] == 2000.0

    assert saida.titulo(rel).startswith("Conferência Mercado Livre — ")
    assert saida.nome_arquivo(rel, "xlsx") == "conferencia-ml-2026-09-28_2026-10-04.xlsx"

    ws = load_workbook(saida.excel(rel)).active
    assert ws["A1"].value.startswith("Conferência Mercado Livre")
    textos = [c.value for row in ws.iter_rows() for c in row]
    # As 4 linhas de afiliados × (4 semanas + Variação) × 4 grupos.
    assert textos.count("não coletado") == 4 * (5 * 4)
    celula = next(c for row in ws.iter_rows(min_row=6, max_row=6) for c in row
                  if c.value == "não coletado")
    assert celula.font.italic and celula.font.color.rgb.endswith("808080")

    html = saida.html(rel)
    assert "Conferência Mercado Livre" in html
    assert html.count('<span class="estado-pill">não coletado</span>') == 4 * 20
    csv = saida.csv(rel).decode("utf-8-sig")
    cab, primeira = csv.splitlines()[:2]
    col = cab.split(";").index("Vendas afiliados")
    assert primeira.split(";")[col] == "não coletado"
    assert primeira.split(";")[cab.split(";").index("Saldo Ads")] == "não se aplica"
    md = saida.markdown(rel)
    assert "| Vendas | afiliados | não coletado |" in md


def test_amazon_no_excel_e_no_threema():
    rel = rel_servidor("amazon")
    ws = load_workbook(saida.excel(rel)).active
    valores = [c.value for row in ws.iter_rows() for c in row]
    assert "aguardando acesso" in valores and "não se aplica" in valores
    msg = ta.texto(rel, "https://x/excel", ta.link_davinci("abc", "amazon"))
    assert msg.startswith("📊 Conferência Amazon — 28/09 a 04/10/2026")
    assert "invest. aguardando acesso · % s/ vendas aguardando acesso" in msg
    assert "conf=amazon" in msg
    assert ta.link_davinci("abc", "shopee").endswith("/marketing?aba=conferencia&execucao=abc")


def test_relatorio_antigo_sem_plataforma_e_shopee():
    rel = {"semanas": [{"inicio": "2026-09-28", "fim": "2026-10-04"}], "grupos": [],
           "geral": {}}
    assert saida.titulo(rel) == "Conferência Shopee — 28/09 a 04/10/2026"
    assert all(lin["estado"] is None for lin in saida.planilha(rel)["linhas"])


# ───────────────────────────────────────────────────────────── banco


@pytest.fixture
async def dono(db) -> User:
    u = User(open_id="email:conf-plat@davinci-test.com", email="conf-plat@davinci-test.com",
             role=UserRole.ADMIN, status=UserStatus.ACTIVE)
    db.add(u)
    await db.commit()
    return u


async def _integracao(db, dono: User, plataforma: str, nome: str) -> Integration:
    i = Integration(user_id=dono.id, platform=IntegrationPlatform(plataforma), name=nome,
                    credentials=b"x")
    db.add(i)
    await db.commit()
    return i


async def contas_servidor(
    db, dono: User, plataforma: str = "ml"
) -> dict[str, ConferenciaShopeeConta]:
    kia = await _integracao(db, dono, plataforma, "kia")
    kfa = await _integracao(db, dono, plataforma, "kfa")
    contas = {
        "Kia": ConferenciaShopeeConta(plataforma=plataforma, nome="Kia", grupo="celular",
                                      integration_id=kia.id, bling_loja_id="204713113",
                                      conta_key="kia", ordem=0),
        "KFA": ConferenciaShopeeConta(plataforma=plataforma, nome="KFA", grupo="mala",
                                      integration_id=kfa.id, bling_loja_id="204438129",
                                      conta_key="kfa", ordem=1),
        "Nexus": ConferenciaShopeeConta(plataforma=plataforma, nome="Nexus", grupo="celular",
                                        ativo=False, conta_key="nexus"),
    }
    db.add_all(contas.values())
    await db.commit()
    return contas


@pytest.fixture
def coletor_falso():
    """Registra um coletor do ML que devolve `numeros(semanas)`, ou o que o
    teste pôr em `respostas[nome]` (exceção, corrotina lenta, dict)."""
    chamadas: list[coletores.ContaColeta] = []
    respostas: dict = {}

    async def collect(conta, semanas, *, session):
        chamadas.append(conta)
        assert all(isinstance(s, coletores.Semana) for s in semanas)
        r = respostas.get(conta.nome)
        if isinstance(r, BaseException):
            raise r
        if callable(r):
            return await r()
        if r is not None:
            return r
        return numeros([s.chave() for s in semanas])

    coletores.registrar("ml", collect)
    yield chamadas, respostas
    coletores.registrar("ml", None)


@pytest.fixture
def fila_arq(monkeypatch):
    """O pool do arq de mentira: guarda o que foi enfileirado."""
    jobs: list[tuple] = []

    class _Pool:
        async def enqueue_job(self, nome, *args, **kw):
            jobs.append((nome, args, kw))

    async def _pool():
        return _Pool()

    monkeypatch.setattr(worker_pool, "get_arq_pool", _pool)
    return jobs


async def test_uma_rodada_coletando_por_plataforma(db, dono):
    await semear_contas(db)
    await contas_servidor(db, dono, "ml")
    shopee = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA)
    await db.commit()
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    assert (shopee.plataforma, ml.plataforma) == ("shopee", "ml")
    assert ml.semanas == shopee.semanas and ml.prazo == shopee.prazo
    # Só as ativas DO ML, sem perfil do AdsPower.
    coletas = await coletas_de(db, ml.id)
    assert [(c.nome, c.adspower_user_id) for c in coletas] == [("Kia", None), ("KFA", None)]
    with pytest.raises(fila.FilaError, match="conferencia_em_andamento"):
        await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.rollback()
    with pytest.raises(fila.FilaError, match="conferencia_sem_contas"):
        await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="amazon")
    await db.rollback()
    with pytest.raises(fila.FilaError, match="plataforma_invalida"):
        await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="tiktok")


async def test_lease_do_mac_so_shopee_e_o_do_servidor_so_ml_amazon(db, dono):
    await contas_servidor(db, dono, "ml")
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    job, _ = await fila.lease(db, "mac", AGORA)
    assert job is None  # o executor do Mac nunca recebe coleta do ML
    await semear_contas(db, (("k-ana", "Ana", "mala", True, 0),))
    shopee = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA)
    await db.commit()
    job, _ = await fila.lease(db, "mac", AGORA)
    await db.commit()
    assert job["execucao_id"] == str(shopee.id)
    # O servidor: só da execução pedida, e não pega a da Shopee.
    coleta, _ = await fila.lease_servidor(db, shopee.id, AGORA)
    assert coleta is None
    coleta, _ = await fila.lease_servidor(db, ml.id, AGORA)
    await db.commit()
    assert (coleta.nome, coleta.status, coleta.agente, coleta.tentativas) == (
        "Kia", "coletando", "servidor", 1,
    )


async def test_rodada_do_servidor_parada_volta_para_a_fila(db, dono):
    await contas_servidor(db, dono, "ml")
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    # Recém-criada: o job da criação ainda vai pegar.
    assert await fila.execucoes_servidor_paradas(db, AGORA + timedelta(minutes=1)) == []
    depois = AGORA + timedelta(minutes=10)
    assert await fila.execucoes_servidor_paradas(db, depois) == [ml.id]
    # Alguém coletando agora: não está parada.
    await fila.lease_servidor(db, ml.id, depois)
    await db.commit()
    assert await fila.execucoes_servidor_paradas(db, depois + timedelta(minutes=5)) == []
    # Largada há mais de 20 min: parada de novo.
    assert await fila.execucoes_servidor_paradas(db, depois + timedelta(minutes=25)) == [ml.id]


async def test_job_do_servidor_coleta_fecha_e_classifica_pelo_sku(db, dono, coletor_falso):
    chamadas, _ = coletor_falso
    contas = await contas_servidor(db, dono, "ml")
    # Produto do DaVinci: categoria do Bling "Eletro" num SKU que não começa
    # com u — o SKU do item manda (eletro_por_sku), não o título.
    db.add(ProductCategory(bling_category_id=777, name="Eletro"))
    db.add(Product(user_id=dono.id, sku="DG053.SP", name="Capinha", category="777"))
    await db.commit()
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()

    resumo = await servidor.coletar_execucao(ml.id, relogio=lambda: AGORA)
    assert resumo == {"execucao": str(ml.id), "contas": 2, "fechou": True}
    assert [c.nome for c in chamadas] == ["Kia", "KFA"]
    assert chamadas[0].integration_id == contas["Kia"].integration_id
    assert (chamadas[0].bling_loja_id, chamadas[0].plataforma, chamadas[0].grupo) == (
        "204713113", "ml", "celular",
    )
    await db.refresh(ml)
    assert ml.status == "pronto"
    kia, kfa = await coletas_de(db, ml.id)
    assert (kia.status, kfa.status) == ("ok", "ok")
    assert kia.dados["coletado_em"] == "2026-10-06T16:30:00Z"
    rel = ml.relatorio
    assert rel["plataforma"] == "ml"
    # dg053.sp é eletro pelo DaVinci (categoria 777 = Eletro); uaf001 pelo SKU.
    eletro = _linha(rel, "eletro", "Kia")["semanas"][0]
    assert eletro["vendas"] == 1000.0 and eletro["invest_ads"] == 50.0
    assert _linha(rel, "celular", "Kia")["semanas"][0]["vendas"] == 0.0
    # Rodar de novo: nada a coletar.
    assert (await servidor.coletar_execucao(ml.id, relogio=lambda: AGORA))["motivo"] == (
        "nada a coletar"
    )


async def test_job_do_servidor_parcial_e_erros_por_conta(db, dono, coletor_falso):
    _, respostas = coletor_falso
    await contas_servidor(db, dono, "ml")
    for nome, grupo in (("Lenta", "celular"), ("Torta", "celular"), ("Pifou", "mala"),
                        ("Curta", "mala")):
        db.add(ConferenciaShopeeConta(plataforma="ml", nome=nome, grupo=grupo, conta_key=nome))
    await db.commit()
    respostas["Kia"] = numeros(SEMANAS, ads=False)  # Ads sem permissão
    respostas["KFA"] = coletores.ColetorError("integração sem o Ads ligado")

    async def lenta():
        await asyncio.sleep(1)

    respostas["Lenta"] = lenta
    respostas["Torta"] = {"versao": 1, "semanas": [{"inicio": "x"}]}
    respostas["Pifou"] = RuntimeError("bug")
    # Só 3 das 4 semanas: parcial.
    respostas["Curta"] = {"versao": 1, "semanas": [semana_servidor(s) for s in SEMANAS[:3]]}
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    await servidor.coletar_execucao(ml.id, relogio=lambda: AGORA, limite_por_conta_s=0.2)
    por_nome = {c.nome: c for c in await coletas_de(db, ml.id)}
    assert (por_nome["Kia"].status, por_nome["Kia"].dados is not None) == ("parcial", True)
    assert (por_nome["KFA"].status, por_nome["KFA"].erro) == (
        "erro", "integração sem o Ads ligado",
    )
    assert por_nome["Lenta"].status == "erro"
    assert por_nome["Lenta"].erro.startswith(servidor.ERRO_TEMPO)
    assert (por_nome["Torta"].status, por_nome["Torta"].erro) == (
        "erro", fila.ERRO_DADOS_INVALIDOS,
    )
    assert (por_nome["Pifou"].status, por_nome["Pifou"].erro) == (
        "erro", "falha na coleta (RuntimeError)",
    )
    assert (por_nome["Curta"].status, por_nome["Curta"].erro) == ("parcial", None)
    await db.refresh(ml)
    assert ml.status == "pronto"
    rel = ml.relatorio
    assert {d["conta"] for d in rel["contas_sem_dados"]} == {"KFA", "Lenta", "Torta", "Pifou"}
    kia = _linha(rel, "celular", "Kia")["semanas"][0]
    assert kia["vendas"] == 700.0 and kia["invest_ads"] is None


async def test_job_sem_coletor_e_conta_apagada(db, dono, monkeypatch):
    await contas_servidor(db, dono, "amazon")
    am = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="amazon")
    await db.commit()
    kia, _kfa = await coletas_de(db, am.id)
    await db.delete(await db.get(ConferenciaShopeeConta, kia.conta_id))
    await db.commit()
    coletores.registrar("amazon", None)
    # Sem o módulo do coletor (None no sys.modules = o import falha).
    monkeypatch.setitem(sys.modules, f"{coletores.__name__}.amazon", None)
    await servidor.coletar_execucao(am.id, relogio=lambda: AGORA)
    por_nome = {c.nome: c for c in await coletas_de(db, am.id)}
    assert (por_nome["Kia"].status, por_nome["Kia"].erro) == ("erro", servidor.ERRO_SEM_CONTA)
    assert por_nome["KFA"].status == "erro"
    assert "sem coletor" in por_nome["KFA"].erro


async def test_job_ignora_shopee_e_rodada_cancelada(db, dono, coletor_falso):
    chamadas, _ = coletor_falso
    await semear_contas(db, (("k-ana", "Ana", "mala", True, 0),))
    shopee = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA)
    await db.commit()
    r = await servidor.coletar_execucao(shopee.id, relogio=lambda: AGORA)
    assert r["motivo"] == "nada a coletar" and chamadas == []
    await contas_servidor(db, dono, "ml")
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    await fila.cancelar(db, ml, AGORA)
    await db.commit()
    r = await servidor.coletar_execucao(ml.id, relogio=lambda: AGORA)
    assert r["motivo"] == "nada a coletar" and chamadas == []


async def test_eletro_por_sku(db, dono):
    from app.services.conferencia_shopee import classificacao

    db.add(ProductCategory(bling_category_id=778, name="Eletro Kit"))
    db.add(Product(user_id=dono.id, sku="kit.cafe", name="Cafeteira", category="778"))
    db.add(Product(user_id=dono.id, sku="ua001", name="Fone", category="Celular"))
    await db.commit()
    r = await classificacao.eletro_por_sku(db, ["KIT.CAFE ", "ua001", "uzz999", "dg1", None, ""])
    assert r == {"kit.cafe": True, "ua001": True, "uzz999": True, "dg1": False}
    mapa = await classificacao.carregar_mapa_por_sku(
        db, {"c1": {"MLB9": "kit.cafe", "dg1": "dg1"}, "c2": {}}
    )
    assert mapa == {"c1": {"MLB9": True, "dg1": False}, "c2": {}}


# ───────────────────────────────────────────────────────────── rotas


@pytest.fixture(scope="module", autouse=True)
def _monta_router():
    from app.main import app

    if not any(getattr(r, "path", "").startswith(API) for r in app.routes):
        app.include_router(mc.router)


@pytest.fixture
async def admin(make_user, auth_as, monkeypatch):
    monkeypatch.setattr(mc, "_agora", lambda: AGORA)
    monkeypatch.setattr(get_settings(), "conferencia_shopee_threema", False)
    u = await make_user(role=UserRole.ADMIN)
    auth_as(u)
    return u


async def test_rotas_por_plataforma(client, db, dono, admin, fila_arq):
    await semear_contas(db, (("k-ana", "Ana", "mala", True, 0),))
    await contas_servidor(db, dono, "ml")
    r = await client.post(f"{API}/execucoes?plataforma=ml", json={"tipo": "semanal"})
    assert r.status_code == 200, r.text
    ex_ml = r.json()["execucao"]
    assert ex_ml["plataforma"] == "ml"
    assert [c["nome"] for c in r.json()["coletas"]] == ["Kia", "KFA"]
    ((nome, args, kw),) = fila_arq
    assert (nome, args) == (servidor.JOB, (ex_ml["id"],))
    assert kw["_job_id"].startswith(f"{servidor.JOB}:{ex_ml['id']}:")
    # A Shopee não depende da do ML (e o corpo também escolhe a plataforma).
    r = await client.post(f"{API}/execucoes", json={"tipo": "semanal", "plataforma": "shopee"})
    assert r.status_code == 200 and r.json()["execucao"]["plataforma"] == "shopee"
    assert len(fila_arq) == 1  # Shopee é do executor do Mac: nada no worker
    r = await client.post(f"{API}/execucoes?plataforma=ml", json={"tipo": "semanal"})
    assert r.status_code == 409
    assert (await client.post(f"{API}/execucoes?plataforma=x", json={"tipo": "semanal"})
            ).status_code == 422

    lista_ml = (await client.get(f"{API}/execucoes?plataforma=ml")).json()
    assert [e["id"] for e in lista_ml] == [ex_ml["id"]]
    assert [e["plataforma"] for e in (await client.get(f"{API}/execucoes")).json()] == ["shopee"]

    contas = (await client.get(f"{API}/contas?plataforma=ml")).json()
    # Ordem da lista: (ordem, nome) — Kia e Nexus são 0, KFA é 1.
    assert [(c["nome"], c["integracao_nome"], c["bling_loja_id"]) for c in contas] == [
        ("Kia", "kia", "204713113"), ("Nexus", None, None), ("KFA", "kfa", "204438129"),
    ]
    assert [c["nome"] for c in (await client.get(f"{API}/contas")).json()] == ["Ana"]


async def test_redis_fora_nao_derruba_o_gerar_agora(client, db, dono, admin, monkeypatch):
    await contas_servidor(db, dono, "amazon")

    async def _sem_redis():
        raise ConnectionError("redis fora")

    monkeypatch.setattr(worker_pool, "get_arq_pool", _sem_redis)
    r = await client.post(f"{API}/execucoes?plataforma=amazon", json={"tipo": "parcial"})
    assert r.status_code == 200, r.text
    assert r.json()["execucao"]["status"] == "coletando"


async def test_editar_conta_do_ml_liga_integracao(client, db, dono, admin):
    contas = await contas_servidor(db, dono, "ml")
    nexus = await _integracao(db, dono, "ml", "nexus")
    da_amazon = await _integracao(db, dono, "amazon", "nexus")
    url = f"{API}/contas/{contas['Nexus'].id}"
    # Ativar sem as ligações: não.
    r = await client.put(url, json={"ativo": True})
    assert (r.status_code, r.json()["detail"]["code"]) == (422, "conta_sem_integracao")
    r = await client.put(url, json={"integration_id": str(da_amazon.id)})
    assert (r.status_code, r.json()["detail"]["code"]) == (422, "integracao_invalida")
    r = await client.put(url, json={"integration_id": str(uuid.uuid4())})
    assert r.json()["detail"]["code"] == "integracao_invalida"
    r = await client.put(url, json={"integration_id": str(contas["Kia"].integration_id)})
    assert (r.status_code, r.json()["detail"]["code"]) == (409, "integracao_em_uso")
    r = await client.put(url, json={"integration_id": str(nexus.id), "ativo": True})
    assert (r.status_code, r.json()["detail"]["code"]) == (422, "conta_sem_loja_bling")
    r = await client.put(url, json={"bling_loja_id": "12a"})
    assert r.status_code == 422
    r = await client.put(url, json={"integration_id": str(nexus.id), "bling_loja_id": " 206064394 ",
                                    "ativo": True})
    assert r.status_code == 200, r.text
    assert (r.json()["integracao_nome"], r.json()["bling_loja_id"], r.json()["ativo"]) == (
        "nexus", "206064394", True,
    )
    # Desligar a integração de uma conta ativa: não.
    r = await client.put(url, json={"integration_id": None})
    assert r.json()["detail"]["code"] == "conta_sem_integracao"
    # Shopee não tem integração na lista.
    await semear_contas(db, (("k-ana", "Ana", "mala", True, 0),))
    ana = (await db.execute(select(ConferenciaShopeeConta).where(
        ConferenciaShopeeConta.nome == "Ana"))).scalar_one()
    r = await client.put(f"{API}/contas/{ana.id}", json={"bling_loja_id": "1"})
    assert r.json()["detail"]["code"] == "campo_so_ml_amazon"

    integ = (await client.get(f"{API}/contas/integracoes?plataforma=ml")).json()
    assert [i["nome"] for i in integ] == ["kfa", "kia", "nexus"]
    assert (await client.get(f"{API}/contas/integracoes?plataforma=shopee")).status_code == 422


async def test_loja_do_bling_uma_vez_por_plataforma(client, db, dono, admin):
    """A mesma loja do Bling em duas contas do ML contaria as mesmas vendas
    duas vezes (Mala/Celular/Eletro e Geral): 409 `loja_bling_em_uso`. Na
    Amazon a mesma loja pode (outra plataforma); a própria conta pode
    regravar a dela."""
    contas = await contas_servidor(db, dono, "ml")
    nexus = contas["Nexus"]
    r = await client.put(f"{API}/contas/{nexus.id}", json={"bling_loja_id": "204438129"})
    assert (r.status_code, r.json()["detail"]) == (
        409, {"code": "loja_bling_em_uso", "conta": "KFA"},
    )
    await db.refresh(nexus)
    assert nexus.bling_loja_id is None
    r = await client.put(f"{API}/contas/{contas['KFA'].id}", json={"bling_loja_id": "204438129"})
    assert r.status_code == 200, r.text
    amazon = ConferenciaShopeeConta(plataforma="amazon", nome="KFA", grupo="mala",
                                    conta_key="kfa", ativo=False)
    db.add(amazon)
    await db.commit()
    r = await client.put(f"{API}/contas/{amazon.id}", json={"bling_loja_id": "204438129"})
    assert r.status_code == 200, r.text
    # O único parcial é o cinto (dois salvando ao mesmo tempo).
    db.add(ConferenciaShopeeConta(plataforma="ml", nome="X", grupo="mala",
                                  bling_loja_id="204713113", ativo=False))
    with pytest.raises(IntegrityError, match="uq_conferencia_shopee_conta_plataforma_bling_loja"):
        await db.commit()
    await db.rollback()


async def test_executor_do_mac_nao_ve_rodada_do_ml(client, db, dono, admin, monkeypatch):
    monkeypatch.setattr(get_settings(), "marketing_agent_token", "tok-plat")
    await contas_servidor(db, dono, "ml")
    await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    r = await client.post(f"{API}/agent/lease", json={}, headers={"X-Agent-Token": "tok-plat"})
    assert r.json() == {"job": None}


async def test_recalcular_rodada_do_ml_e_baixar_arquivo(client, db, dono, admin, coletor_falso):
    await contas_servidor(db, dono, "ml")
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    await servidor.coletar_execucao(ml.id, relogio=lambda: AGORA)
    r = await client.post(f"{API}/execucoes/{ml.id}/recalcular")
    assert r.status_code == 200 and r.json()["relatorio"]["plataforma"] == "ml"
    r = await client.get(f"{API}/execucoes/{ml.id}/arquivo/csv")
    assert r.status_code == 200
    assert 'filename="conferencia-ml-2026-09-28_2026-10-04.csv"' in r.headers[
        "content-disposition"
    ]


# ───────────────────────────────────────────────────────────── worker


@pytest.fixture
def liga(monkeypatch):
    def _liga(nome: str, valor) -> None:
        for cfg in {id(c): c for c in (get_settings(), worker._settings)}.values():
            monkeypatch.setattr(cfg, nome, valor)

    _liga("enable_marketing", True)
    _liga("conferencia_ml_cron", False)
    _liga("conferencia_amazon_cron", False)
    _liga("conferencia_shopee_threema", False)
    return _liga


async def _execucoes(db) -> list[ConferenciaShopeeExecucao]:
    return list((await db.execute(select(ConferenciaShopeeExecucao))).scalars())


async def test_agenda_do_ml_e_da_amazon_so_com_a_chave(db, dono, liga, fila_arq, monkeypatch):
    await contas_servidor(db, dono, "ml")
    monkeypatch.setattr(periodos, "tipo_da_agenda", lambda dia: "semanal")
    await worker.conferencia_ml_agenda({})
    await worker.conferencia_amazon_agenda({})
    assert await _execucoes(db) == [] and fila_arq == []

    liga("conferencia_ml_cron", True)
    liga("conferencia_amazon_cron", True)
    await worker.conferencia_ml_agenda({})
    (ex,) = await _execucoes(db)
    assert (ex.plataforma, ex.origem) == ("ml", "agenda")
    assert [j[0] for j in fila_arq] == [servidor.JOB]
    # Amazon sem conta ativa: pula sem erro; ML já coletando: pula.
    await worker.conferencia_amazon_agenda({})
    await worker.conferencia_ml_agenda({})
    assert len(await _execucoes(db)) == 1 and len(fila_arq) == 1


async def test_job_do_worker_e_varredor_reenfileira(db, dono, liga, fila_arq, coletor_falso):
    await contas_servidor(db, dono, "ml")
    antes = datetime.now(UTC) - timedelta(minutes=30)
    ex = await fila.criar_execucao(db, "semanal", "manual", None, antes, plataforma="ml")
    await db.commit()
    await worker.conferencia_shopee_varrer({})
    assert [(j[0], j[1]) for j in fila_arq] == [(servidor.JOB, (str(ex.id),))]
    r = await worker.conferencia_coletar_servidor({}, str(ex.id))
    assert r["fechou"] is True
    await db.refresh(ex)
    assert ex.status == "pronto"
    await worker.conferencia_shopee_varrer({})
    assert len(fila_arq) == 1


async def test_saldo_de_ads_nao_se_grava_sem_perfil(db, dono):
    """Um coletor que mandasse saldo_ads (não deveria) não quebra: o histórico
    de saldo é por perfil do AdsPower, e o ML não tem."""
    await contas_servidor(db, dono, "ml")
    ml = await fila.criar_execucao(db, "semanal", "agenda", None, AGORA, plataforma="ml")
    await db.commit()
    coleta, _ = await fila.lease_servidor(db, ml.id, AGORA)
    await db.commit()
    d = dados_loja(ml.semanas)
    c = await fila.registrar_resultado(db, coleta.id, "ok", None, d, AGORA)
    await db.commit()
    assert c.status == "ok"
