"""Reembolso por equipe — as rotas que ficavam fora do recorte (30/10/2026,
revisão das equipes Londres/Israel/Marrocos): a planilha (export.xlsx), a
busca e o custo do pedido e a edição/exclusão por id. Quem é de equipe só
mexe no que enxerga na lista; admin e quem não tem equipe, como antes."""

from __future__ import annotations

import io
import uuid
from collections.abc import Callable

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import BlingOrder, Refund, StoreInfo, User, UserRole, UserStatus

pytestmark = pytest.mark.asyncio

PERM = {"reembolso": {"view": True, "edit": True, "delete": True}}


def _usuario(role: UserRole, equipes: list[int] | None) -> User:
    email = f"reemb-eq-{uuid.uuid4().hex[:6]}@davinci-test.com"
    return User(
        open_id=f"email:{email}",
        email=email,
        role=role,
        status=UserStatus.ACTIVE,
        permissions=PERM,
        sales_teams=equipes,
    )


async def _como(db: AsyncSession, auth_as, equipes: list[int] | None, role=UserRole.USER) -> User:
    u = _usuario(role, equipes)
    db.add(u)
    await db.commit()
    auth_as(u)
    return u


@pytest_asyncio.fixture
async def dados(db: AsyncSession) -> dict[str, Refund]:
    """Duas lojas de equipes diferentes, um pedido e um reembolso de cada."""
    dono = _usuario(UserRole.ADMIN, None)
    db.add(dono)
    await db.flush()
    db.add_all(
        [
            StoreInfo(
                user_id=dono.id,
                platform="shopee",
                account_name="vortan",
                bling_store_id="b-lon",
                sales_team=101,
            ),
            StoreInfo(
                user_id=dono.id,
                platform="shopee",
                account_name="atv",
                bling_store_id="b-isr",
                sales_team=102,
            ),
            BlingOrder(numero="900101", numeroloja="MKT-LON", loja="b-lon"),
            BlingOrder(numero="900102", numeroloja="MKT-ISR", loja="b-isr"),
        ]
    )
    lon = Refund(pedido_bling="900101", conta="Shopee Vortan", tipo="Cliente", reembolso=-30)
    isr = Refund(pedido_bling="900102", conta="Shopee ATV", tipo="Cliente", reembolso=-40)
    db.add_all([lon, isr])
    await db.commit()
    return {"londres": lon, "israel": isr}


async def _pedidos_da_planilha(client: AsyncClient) -> set[str]:
    r = await client.get("/api/refunds/export.xlsx")
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    cab = [c.value for c in ws[1]]
    col = next(i for i, v in enumerate(cab) if v and "Bling" in str(v))
    return {str(row[col].value) for row in ws.iter_rows(min_row=2) if row[col].value}


async def test_planilha_so_da_equipe(db, client, auth_as: Callable, dados):
    await _como(db, auth_as, [101])
    assert await _pedidos_da_planilha(client) == {"900101"}
    await _como(db, auth_as, None)  # sem equipe: tudo, como antes
    assert await _pedidos_da_planilha(client) == {"900101", "900102"}
    await _como(db, auth_as, [101], role=UserRole.ADMIN)
    assert await _pedidos_da_planilha(client) == {"900101", "900102"}


async def test_busca_e_custo_do_pedido_de_outra_equipe(db, client, auth_as: Callable, dados):
    await _como(db, auth_as, [101])
    r = await client.get("/api/refunds/order-lookup", params={"pedido": "900102"})
    assert r.status_code == 403 and r.json()["detail"]["code"] == "pedido_de_outra_equipe"
    r = await client.get("/api/refunds/order-lookup", params={"pedido": "MKT-ISR"})
    assert r.status_code == 403
    r = await client.get(
        "/api/refunds/order-cost", params={"pedido_bling": "900102", "conta": "atv"}
    )
    assert r.status_code == 403
    # O da própria equipe passa da trava (a consulta da conciliação, que vem
    # depois, é uma view que o banco de teste não tem — por isso só a trava).
    from app.deps.team_scope import resolve_team_scope
    from app.routers.refunds import _pedido_de_outra_equipe

    eu = await _como(db, auth_as, [101])
    escopo = await resolve_team_scope(db, eu)
    assert await _pedido_de_outra_equipe(db, escopo, "900101") is False
    assert await _pedido_de_outra_equipe(db, escopo, "MKT-LON") is False
    assert await _pedido_de_outra_equipe(db, escopo, "nao-existe") is False
    assert await _pedido_de_outra_equipe(db, escopo, "MKT-ISR") is True
    sem_equipe = await resolve_team_scope(db, await _como(db, auth_as, None))
    assert await _pedido_de_outra_equipe(db, sem_equipe, "900102") is False


async def test_editar_e_apagar_so_o_que_enxerga(db, client, auth_as: Callable, dados):
    await _como(db, auth_as, [101])
    isr, lon = dados["israel"], dados["londres"]
    r = await client.patch(f"/api/refunds/{isr.id}", json={"observacao": "x"})
    assert r.status_code == 404
    r = await client.delete(f"/api/refunds/{isr.id}")
    assert r.status_code == 404
    r = await client.patch(f"/api/refunds/{lon.id}", json={"observacao": "conferido"})
    assert r.status_code == 200, r.text
    await _como(db, auth_as, None)
    assert (
        await client.patch(f"/api/refunds/{isr.id}", json={"observacao": "y"})
    ).status_code == 200
