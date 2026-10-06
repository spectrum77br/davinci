"""Botão "Excel" do painel Logística: a tela manda as linhas que está mostrando
(já filtradas e formatadas) e o servidor devolve o .xlsx igual."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services import logistica_export_xlsx

_COLUNAS = [
    {"titulo": "Data", "tipo": "data"},
    {"titulo": "Pedido Bling"},
    {"titulo": "Pedido Marketplace"},
    {"titulo": "Produto"},
    {"titulo": "Status Plataforma"},
]


def _corpo(linhas, aba="Amazon Envio próprio"):
    return {"aba": aba, "colunas": _COLUNAS, "linhas": linhas}


async def test_export_devolve_as_linhas_da_tela(client, make_user, auth_as):
    auth_as(await make_user(permissions={"logistica": {"view": True}}))

    r = await client.post(
        "/api/logistica/export.xlsx",
        json=_corpo(
            [
                ["2026-10-05", "301324", "2000018792949872", "Uranyx A17 ×2\nCapinha", "Cancelado | Fechada"],
                [None, "296819", None, None, "=HYPERLINK(\"x\")"],
            ]
        ),
    )

    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    ws = load_workbook(BytesIO(r.content)).active
    assert ws.title == "Amazon Envio próprio"
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0] == ("Data", "Pedido Bling", "Pedido Marketplace", "Produto", "Status Plataforma")
    assert linhas[1] == (
        datetime(2026, 10, 5),
        "301324",
        # Texto: número de pedido grande não pode virar 2,00002E+15.
        "2000018792949872",
        "Uranyx A17 ×2\nCapinha",
        "Cancelado | Fechada",
    )
    assert ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["D2"].alignment.wrap_text
    # Vazio na tela = célula vazia; texto com "=" não vira fórmula.
    assert linhas[2] == (None, "296819", None, None, '=HYPERLINK("x")')
    assert ws["E3"].data_type == "s"
    assert ws.freeze_panes == "A2"


async def test_export_sem_linhas_so_cabecalho(client, make_user, auth_as):
    auth_as(await make_user(permissions={"logistica": {"view": True}}))

    r = await client.post("/api/logistica/export.xlsx", json=_corpo([], aba="TikTok"))

    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    assert ws.title == "TikTok"
    assert ws.max_row == 1


async def test_export_exige_ver_logistica(client, make_user, auth_as):
    auth_as(await make_user(permissions={}))

    r = await client.post("/api/logistica/export.xlsx", json=_corpo([]))

    assert r.status_code == 403


def test_nome_aba_respeita_o_excel():
    assert logistica_export_xlsx.nome_aba("Amazon [DBA]: a/b") == "Amazon  DBA   a b"
    assert len(logistica_export_xlsx.nome_aba("x" * 50)) == 31
    assert logistica_export_xlsx.nome_aba("") == "Logística"


def test_data_invalida_fica_como_texto():
    conteudo = logistica_export_xlsx.montar_xlsx(
        "ML", [("Data", "data")], [["2026-13-40"], ["ontem"]]
    )
    ws = load_workbook(BytesIO(conteudo)).active
    assert [c.value for c in ws["A"]] == ["Data", "2026-13-40", "ontem"]
