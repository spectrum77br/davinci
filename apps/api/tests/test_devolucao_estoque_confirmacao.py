"""Entrada no estoque do Bling guardada na linha (alembic 0421).

Caso 300588 (15/10/2026): a mala voltou como Novo, o Bling recusou a entrada
(limite de chamadas), o aviso de erro sumiu em 6 s e a "Data devolvido
estoque" ficou na linha como se tivesse entrado — e o pedido foi para
Resolvido. Casos 300334/298826 (mesmo dia): o Bling GRAVOU a entrada e
respondeu 504; o DaVinci tratou como falha, sem registro (sem estorno).
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from uuid import uuid4

import httpx
import pytest
from sqlalchemy import select

from app.models import Devolution, StockMovement
from app.routers import devolutions as devolutions_router
from app.services import devolution_stock_return as dsr
from app.services.marketplaces.bling import BlingCloudflareError

pytestmark = pytest.mark.asyncio

PID = 16377501606


def _perms() -> dict:
    return {"devolucoes": {"view": True, "edit": True, "delete": True}}


ESTOQUES_URL = "https://api.bling.com.br/Api/v3/estoques"


def _http_error(code: int, url: str = ESTOQUES_URL) -> httpx.HTTPStatusError:
    req = httpx.Request("POST", url)
    resp = httpx.Response(code, request=req, json={"error": {"description": f"erro {code}"}})
    return httpx.HTTPStatusError(f"{code}", request=req, response=resp)


class _FakeBling:
    """Só o que o caminho "destino existente" usa. `falhas` = exceções a
    levantar nas próximas chamadas de entrada, em ordem."""

    def __init__(self, *falhas: Exception):
        self.falhas = list(falhas)
        self.entradas: list[dict] = []

    async def update_stock_by_id(self, pid, *, qty, operation, observacao="", custo=None, **_kw):
        self.entradas.append(
            {"pid": pid, "qty": qty, "op": operation, "obs": observacao, "em": datetime.now(UTC)}
        )
        if self.falhas:
            raise self.falhas.pop(0)
        return {}


def _wire(monkeypatch, fake: _FakeBling) -> list:
    async def _client(_session):
        return fake

    async def _pid(_session, _client, _sku):
        return PID

    situacoes: list[str] = []

    async def _situacao(_session, pedido, **_kw):
        situacoes.append(pedido)

    monkeypatch.setattr(dsr, "_get_bling_client", _client)
    monkeypatch.setattr(dsr, "_resolve_product_id", _pid)
    monkeypatch.setattr(devolutions_router, "apply_order_situacao", _situacao)
    monkeypatch.setattr(dsr, "apply_order_situacao", _situacao)
    return situacoes


def _body(**kw) -> dict:
    body = {
        "pedido_bling": f"30{uuid4().hex[:4]}",
        "conta": "Shopee KFA",
        "sku": "x020.sp",
        "produtos": "Produto teste",
        "custo_produto": 18.9,
        "condicao_produto": "Novo",
        "devolver_estoque": True,
        "estoque_destino_sku": "x020.sp",
        "quantidade": 1,
    }
    body.update(kw)
    return body


async def _linha(db, dev_id) -> Devolution:
    # populate_existing (e não expire_all): expirar derrubaria o usuário do
    # auth_as, que vive na mesma sessão.
    return (
        await db.execute(
            select(Devolution)
            .where(Devolution.id == dev_id)
            .execution_options(populate_existing=True)
        )
    ).scalar_one()


# ── POST: o resultado do Bling fica na linha ─────────────────────────────


async def test_bling_recusou_por_limite_fica_falhou_sem_movimento(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling(BlingCloudflareError("status=429 cf_html=False"))
    situacoes = _wire(monkeypatch, fake)

    r = await client.post("/api/devolutions", json=_body())
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["estoque_status"] == "falhou"
    assert "limite de chamadas" in out["estoque_erro"]
    assert out["bling_stock_result"]["action"] == "error"
    row = await _linha(db, out["id"])
    assert row.estoque_mov_bling_id is None
    assert row.estoque_tentativa_at is not None
    # A data continua sendo o pedido — mas a situação não pode virar Resolvido:
    # o item com a entrada pendente não conta como resolvido.
    # O pedido fica (ou volta) em Aguardando Devolução — nunca Resolvido.
    assert dsr._order_situacao_target([row]) == dsr.SITUACAO_AGUARDANDO_DEVOLUCAO
    assert situacoes == [out["pedido_bling"]]


async def test_bling_respondeu_504_fica_conferindo_com_o_produto(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(_http_error(504)))

    r = await client.post("/api/devolutions", json=_body())
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["estoque_status"] == "conferindo"
    assert out["bling_stock_result"]["action"] == "unconfirmed"
    row = await _linha(db, out["id"])
    assert row.estoque_conferir_bling_id == PID
    assert row.estoque_conferir_sku == "x020.sp"
    assert row.estoque_conferir_action == "entry_existing"
    assert row.estoque_mov_bling_id is None


async def test_timeout_de_leitura_tambem_e_incerto_mas_conexao_nao():
    req = httpx.Request("POST", "https://api.bling.com.br/Api/v3/estoques")
    assert dsr._resposta_incerta(httpx.ReadTimeout("t", request=req))
    # Erro de OUTRA chamada (renovar o token antes do POST) não é incerto.
    assert not dsr._resposta_incerta(_http_error(502, "https://www.bling.com.br/Api/v3/oauth/token"))
    assert not dsr._resposta_incerta(httpx.ReadTimeout("t"))
    assert dsr._resposta_incerta(_http_error(500))
    assert dsr._resposta_incerta(_http_error(502))
    assert not dsr._resposta_incerta(_http_error(503))
    assert not dsr._resposta_incerta(_http_error(400))
    assert not dsr._resposta_incerta(httpx.ConnectError("x"))
    assert not dsr._resposta_incerta(BlingCloudflareError("status=429 cf_html=False"))


async def test_entrada_ok_fica_ok_com_movimento(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling())

    r = await client.post("/api/devolutions", json=_body())
    out = r.json()
    assert out["estoque_status"] == "ok"
    assert out["estoque_erro"] is None
    row = await _linha(db, out["id"])
    assert row.estoque_mov_bling_id == PID
    assert dsr._order_situacao_target([row]) == dsr.SITUACAO_RESOLVIDO


async def test_linha_antiga_com_movimento_sai_como_ok():
    out = devolutions_router.DevolutionOut.model_validate(
        SimpleNamespace(
            id=uuid4(), estoque_status=None, estoque_mov_bling_id=123,
            reembolso=False, devolver_estoque=True, manutencao=False,
            created_at=datetime.now(UTC), updated_at=datetime.now(UTC), conta="x",
        ),
        from_attributes=True,
    )
    assert out.estoque_status == "ok"


# ── Situação do pedido ────────────────────────────────────────────────────


def _r(
    condicao: str, status: str | None = None, mov: int | None = None, tentou: bool = True
) -> SimpleNamespace:
    return SimpleNamespace(
        condicao_produto=condicao, manutencao_destino=None,
        estoque_status=status, estoque_mov_bling_id=mov,
        estoque_tentativa_at=datetime.now(UTC) if tentou else None,
    )


async def test_pedido_nao_vai_para_resolvido_com_entrada_pendente():
    aguardando = dsr.SITUACAO_AGUARDANDO_DEVOLUCAO
    for st in ("falhou", "conferindo", "sem_confirmacao"):
        assert dsr._order_situacao_target([_r("Novo", st)]) == aguardando
        # Pedido de 2 linhas: a 1ª entrou (já mandou Resolvido), a 2ª falhou →
        # volta para Aguardando Devolução.
        assert dsr._order_situacao_target([_r("Novo", "ok"), _r("Usado", st)]) == aguardando
        # Trocado não força Resolvido se a entrada dele, ou de outra linha, não entrou.
        assert dsr._order_situacao_target([_r("Trocado", st)]) == aguardando
        assert dsr._order_situacao_target([_r("Trocado", "ok"), _r("Novo", st)]) == aguardando
    # Linha antiga do backfill (nunca tentada por este código): não mexe no
    # pedido, só deixa de contar como resolvida.
    assert dsr._order_situacao_target([_r("Novo", "sem_confirmacao", tentou=False)]) is None
    # Extraviado continua mandando, com ou sem pendência.
    assert dsr._order_situacao_target([_r("Extraviado"), _r("Novo", "falhou")]) == (
        dsr.SITUACAO_EXTRAVIADO
    )
    assert dsr._order_situacao_target([_r("Novo", None, mov=1)]) == dsr.SITUACAO_RESOLVIDO


# ── Tentar de novo ────────────────────────────────────────────────────────


async def test_tentar_de_novo_lanca_e_fica_ok(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling(BlingCloudflareError("status=429 cf_html=False"))
    situacoes = _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"

    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"
    assert len(fake.entradas) == 2
    row = await _linha(db, out["id"])
    assert row.estoque_mov_bling_id == PID
    # Confirmada → recalcula a situação do pedido.
    assert situacoes[-1] == out["pedido_bling"]

    # Já entrou: um 2º clique não lança de novo.
    r2 = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r2.status_code == 409
    assert r2.json()["detail"]["code"] == "estoque_nao_pendente"
    assert len(fake.entradas) == 2


async def test_tentar_de_novo_sem_confirmacao_exige_conferir_no_bling(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    _wire(monkeypatch, fake)
    row = Devolution(
        pedido_bling="298503", conta="Shopee Inova", sku="a073", condicao_produto="Novo",
        quantidade=1, devolver_estoque=True, estoque_destino_sku="a073",
        data_devolvido_estoque=datetime.now(UTC) - timedelta(days=9),
        estoque_status="sem_confirmacao", estoque_erro="Lançamento sem confirmação do Bling",
    )
    db.add(row)
    await db.commit()

    r = await client.post(f"/api/devolutions/{row.id}/estoque/tentar-de-novo", json={})
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == "estoque_conferir_extrato"
    assert fake.entradas == []

    r = await client.post(
        f"/api/devolutions/{row.id}/estoque/tentar-de-novo",
        json={"confirmo_conferi_no_bling": True},
    )
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"
    assert len(fake.entradas) == 1


async def test_conferindo_trava_mudar_estoque_excluir_e_tentar(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling(_http_error(504))
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "conferindo"

    r = await client.patch(f"/api/devolutions/{out['id']}", json={"condicao_produto": "Usado"})
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "estoque_conferindo"
    # Campo que não mexe no estoque passa.
    r = await client.patch(f"/api/devolutions/{out['id']}", json={"observacao": "ok"})
    assert r.status_code == 200, r.text

    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 409
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 409
    assert len(fake.entradas) == 1


async def test_mudar_para_extraviado_limpa_a_falha(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(BlingCloudflareError("status=429 cf_html=False")))
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"

    r = await client.patch(
        f"/api/devolutions/{out['id']}",
        json={
            "condicao_produto": "Extraviado",
            "devolver_estoque": False,
            "link_abertura": "nao ha",
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] is None
    row = await _linha(db, out["id"])
    assert row.estoque_erro is None
    assert dsr._order_situacao_target([row]) == dsr.SITUACAO_EXTRAVIADO


async def test_get_de_uma_linha(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(_http_error(504)))
    out = (await client.post("/api/devolutions", json=_body())).json()
    r = await client.get(f"/api/devolutions/{out['id']}")
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "conferindo"
    # As rotas fixas continuam casando antes da rota por id.
    r = await client.get("/api/devolutions/situacoes-bling")
    assert r.status_code == 200


# ── Conferência do worker ─────────────────────────────────────────────────


def _pid_novo() -> int:
    # stock_movements não é limpo entre testes: cada teste usa um produto só dele.
    return 10**10 + uuid4().int % 10**9


async def _conferindo(db, *, t0: datetime, pid: int | None, qty: int = 1) -> Devolution:
    row = Devolution(
        pedido_bling=f"30{uuid4().hex[:4]}", conta="Shopee ATV", sku="a001.sp",
        condicao_produto="Novo", quantidade=qty, devolver_estoque=True,
        estoque_destino_sku="a001.pi", data_devolvido_estoque=t0,
        estoque_status="conferindo", estoque_tentativa_at=t0,
        estoque_conferir_bling_id=pid, estoque_conferir_sku="a001.pi",
        estoque_conferir_action="entry_existing",
    )
    db.add(row)
    await db.commit()
    return row


async def _entrada(db, *, quando: datetime, pid: int, qty: int = 1, tipo: str = "E"):
    db.add(
        StockMovement(
            bling_product_id=pid, sku="a001.pi", date=quando, tipo=tipo, quantidade=qty,
        )
    )
    await db.commit()


async def test_conferencia_confirma_quando_a_entrada_aparece(db, monkeypatch):
    situacoes = _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    row = await _conferindo(db, t0=t0, pid=pid)
    await _entrada(db, quando=t0 + timedelta(seconds=31), pid=pid)

    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 1
    row = await _linha(db, row.id)
    assert row.estoque_status == "ok"
    assert row.estoque_mov_bling_id == pid
    assert row.estoque_mov_sku == "a001.pi"
    assert row.estoque_mov_action == "entry_existing"
    assert row.estoque_mov_ref is not None
    assert row.estoque_conferir_bling_id is None
    assert situacoes == [row.pedido_bling]


async def test_conferencia_nao_usa_a_entrada_de_outra_devolucao(db, monkeypatch):
    """Duas devoluções do mesmo SKU no mesmo minuto: a que entrou normal (ok)
    é dona da única entrada que apareceu — a incerta não pode pegá-la."""
    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    outra = Devolution(
        pedido_bling="301535", conta="Shopee ATV", sku="a001.sp", condicao_produto="Novo",
        quantidade=1, devolver_estoque=True, data_devolvido_estoque=t0,
        estoque_status="ok", estoque_tentativa_at=t0 - timedelta(seconds=20),
        estoque_mov_bling_id=pid, estoque_mov_sku="a001.pi", estoque_mov_qty=1,
    )
    db.add(outra)
    await db.commit()
    row = await _conferindo(db, t0=t0, pid=pid)
    await _entrada(db, quando=t0 - timedelta(seconds=2), pid=pid)

    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 0
    assert (await _linha(db, row.id)).estoque_status == "conferindo"

    # A segunda entrada aparece → agora sobra uma para ela.
    await _entrada(db, quando=t0 + timedelta(seconds=14), pid=pid)
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 1


async def test_conferencia_ignora_saida_quantidade_errada_e_entrada_antiga(db, monkeypatch):
    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    row = await _conferindo(db, t0=t0, pid=pid)
    await _entrada(db, quando=t0 + timedelta(seconds=5), tipo="S", pid=pid)
    await _entrada(db, quando=t0 + timedelta(seconds=5), qty=2, pid=pid)
    await _entrada(db, quando=t0 - timedelta(minutes=3), pid=pid)
    await _entrada(db, quando=t0 + timedelta(seconds=5), pid=999)

    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 0
    assert (await _linha(db, row.id)).estoque_status == "conferindo"


async def test_uma_entrada_confirma_uma_linha_so(db, monkeypatch):
    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    a = await _conferindo(db, t0=t0, pid=pid)
    b = await _conferindo(db, t0=t0 + timedelta(seconds=1), pid=pid)
    await _entrada(db, quando=t0 + timedelta(seconds=10), pid=pid)

    # Uma entrada só para duas linhas conferindo: não dá para saber de quem é
    # — nenhuma confirma (na dúvida, não confirma).
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 0
    # Chega a segunda: agora cada uma tem a sua.
    await _entrada(db, quando=t0 + timedelta(seconds=12), pid=pid)
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 2
    sts = {(await _linha(db, a.id)).estoque_status, (await _linha(db, b.id)).estoque_status}
    assert sts == {"ok"}
    refs = {(await _linha(db, a.id)).estoque_mov_ref, (await _linha(db, b.id)).estoque_mov_ref}
    assert len(refs) == 2


async def test_conferencia_desiste_em_10_min(db, monkeypatch):
    situacoes = _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=11)
    pid = _pid_novo()
    row = await _conferindo(db, t0=t0, pid=pid)

    res = await dsr.conferir_entradas_pendentes(db)
    assert res["sem_confirmacao"] == 1
    row = await _linha(db, row.id)
    assert row.estoque_status == "sem_confirmacao"
    assert "a001.pi" in row.estoque_erro
    assert row.estoque_mov_bling_id is None
    assert situacoes == []


async def test_conferindo_sem_produto_vira_sem_confirmacao(db, monkeypatch):
    """Tentar de novo que caiu no meio (processo morreu): linha conferindo sem
    produto — passados 10 min, a tela pede para conferir no extrato."""
    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=12)
    row = await _conferindo(db, t0=t0, pid=None)
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["sem_confirmacao"] == 1
    assert (await _linha(db, row.id)).estoque_status == "sem_confirmacao"


# ── Revisão de 15/10 ──────────────────────────────────────────────────────


async def test_conferencia_acha_a_entrada_que_chegou_antes_do_504(
    client, db, make_user, auth_as, monkeypatch
):
    """O Bling grava, o webhook chega em 1 s e o 504 só volta depois: a
    tentativa é a hora de ANTES de chamar o Bling, senão a entrada fica fora
    da janela da conferência (achado da revisão)."""
    auth_as(await make_user(permissions=_perms()))
    pid = _pid_novo()

    class _Lento(_FakeBling):
        async def update_stock_by_id(self, *a, **kw):
            self.entradas.append({"em": datetime.now(UTC)})
            await _entrada(db, quando=datetime.now(UTC) + timedelta(seconds=1), pid=pid)
            raise _http_error(504)

    fake = _Lento()
    _wire(monkeypatch, fake)

    async def _pid(_s, _c, _sku):
        return pid

    monkeypatch.setattr(dsr, "_resolve_product_id", _pid)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "conferindo"
    row = await _linha(db, out["id"])
    assert row.estoque_tentativa_at <= fake.entradas[0]["em"]
    # Simula o 504 voltando 30 s depois: a hora gravada continua a do envio.
    res = await dsr.conferir_entradas_pendentes(db, agora=datetime.now(UTC) + timedelta(seconds=30))
    assert res["confirmadas"] == 1
    assert (await _linha(db, out["id"])).estoque_status == "ok"


async def test_conferencia_nao_confirma_com_entrada_do_robo_de_prioridade(db, monkeypatch):
    from app.models import PrioridadeEstoqueMovimento

    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    row = await _conferindo(db, t0=t0, pid=pid)
    db.add(
        PrioridadeEstoqueMovimento(
            pedido_bling="999001", sku="a001.pi", bling_product_id=pid, operacao="E",
            quantidade=1, status="ok", lancado_at=t0 + timedelta(seconds=3),
        )
    )
    await db.commit()
    await _entrada(db, quando=t0 + timedelta(seconds=4), pid=pid)
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 0
    assert (await _linha(db, row.id)).estoque_status == "conferindo"


async def test_z_criado_fica_na_linha_e_a_nova_tentativa_usa_ele(db, monkeypatch):
    """Usado com tag nova: o z000N é criado e a entrada leva 429. A linha
    guarda o z como destino — o Tentar de novo lança NELE, sem criar outro."""
    criados: list[str] = []

    class _Z(_FakeBling):
        async def product_exists_by_sku(self, sku):
            return sku in criados

        async def create_product(self, *, sku, name, price=None, category_id=None):
            criados.append(sku)
            return {"id": 990000 + len(criados)}

        async def get_product(self, pid):
            return {"nome": "Mala", "categoria": None}

    fake = _Z(BlingCloudflareError("status=429 cf_html=False"))

    async def _client(_s):
        return fake

    async def _nada(*_a, **_k):
        return None

    monkeypatch.setattr(dsr, "_get_bling_client", _client)
    monkeypatch.setattr(dsr, "_ensure_local_product", _nada)
    monkeypatch.setattr(dsr, "_persist_cost", _nada)
    async def _pid(_s, _c, sku):
        # O z criado acha pela tabela local (o _ensure_local_product o grava).
        return 990000 + criados.index(sku) + 1 if sku in criados else None

    monkeypatch.setattr(dsr, "_resolve_product_id", _pid)
    row = Devolution(
        pedido_bling="304999", conta="Loja", sku="b007.24", produtos="Mala",
        condicao_produto="Usado", quantidade=1, devolver_estoque=True,
        estoque_nova_tag="mala", custo_produto=130,
    )
    db.add(row)
    await db.commit()

    sr = await dsr.return_product_to_bling_stock(db, row)
    dsr.registrar_resultado_estoque(row, sr, inicio=datetime.now(UTC))
    assert row.estoque_status == "falhou"
    assert row.estoque_destino_sku == criados[0]
    assert row.estoque_nova_tag is None

    sr2 = await dsr.return_product_to_bling_stock(db, row)
    assert sr2["ok"] and sr2["sku"] == criados[0], sr2
    assert len(criados) == 1  # nenhum z novo
    await db.commit()


async def test_patch_nao_relanca_linha_sem_confirmacao(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    _wire(monkeypatch, fake)
    row = Devolution(
        pedido_bling="298503", conta="Shopee Inova", sku="a073", condicao_produto="Novo",
        quantidade=1, devolver_estoque=True, estoque_destino_sku="a073",
        data_devolvido_estoque=datetime.now(UTC) - timedelta(days=9),
        estoque_status="sem_confirmacao",
    )
    db.add(row)
    await db.commit()
    for patch in ({"devolver_estoque": False}, {"condicao_produto": "Usado"}):
        r = await client.patch(f"/api/devolutions/{row.id}", json=patch)
        assert r.status_code == 409, patch
        assert r.json()["detail"]["code"] == "estoque_sem_confirmacao"
    assert fake.entradas == []


async def test_ja_entrou_registra_sem_lancar_e_recalcula_o_pedido(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    situacoes = _wire(monkeypatch, fake)

    async def _pid_local(_s, sku):
        return 16000000073 if sku == "a073" else None

    monkeypatch.setattr(devolutions_router, "_local_bling_product_id", _pid_local)
    row = Devolution(
        pedido_bling="298503", conta="Shopee Inova", sku="a073", condicao_produto="Novo",
        quantidade=1, devolver_estoque=True, estoque_destino_sku="a073",
        data_devolvido_estoque=datetime.now(UTC) - timedelta(days=9),
        estoque_status="sem_confirmacao",
    )
    db.add(row)
    await db.commit()

    r = await client.post(f"/api/devolutions/{row.id}/estoque/ja-entrou", json={"sku": "xyz"})
    assert r.status_code == 422
    r = await client.post(f"/api/devolutions/{row.id}/estoque/ja-entrou", json={"sku": "a073"})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"
    linha = await _linha(db, row.id)
    assert linha.estoque_mov_bling_id == 16000000073
    assert linha.estoque_mov_action == "entry_conferida_manual"
    assert fake.entradas == []  # nada lançado no Bling
    assert situacoes == ["298503"]
    # Só serve para "sem confirmação".
    r = await client.post(f"/api/devolutions/{row.id}/estoque/ja-entrou", json={"sku": "a073"})
    assert r.status_code == 409


async def test_estorno_que_falha_nao_relanca_e_tentar_de_novo_estorna_antes(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "ok" and len(fake.entradas) == 1

    # Troca para Usado: o estorno (S) leva 429 → não lança no destino novo.
    fake.falhas = [BlingCloudflareError("status=429 cf_html=False")]
    r = await client.patch(
        f"/api/devolutions/{out['id']}",
        json={"condicao_produto": "Usado", "estoque_destino_sku": "x020.us"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "falhou"
    assert r.json()["bling_stock_result"]["action"] == "reverse_error"
    assert [e["op"] for e in fake.entradas] == ["E", "S"]
    linha = await _linha(db, out["id"])
    assert linha.estoque_mov_bling_id == PID and linha.estoque_mov_revertido_at is None

    # Tentar de novo: estorna a antiga e só então lança a nova.
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"
    assert [e["op"] for e in fake.entradas] == ["E", "S", "S", "E"]


async def test_desligar_o_toggle_de_linha_falhou_recalcula_o_pedido(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    situacoes = _wire(monkeypatch, _FakeBling(BlingCloudflareError("status=429 cf_html=False")))
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"
    situacoes.clear()
    r = await client.patch(f"/api/devolutions/{out['id']}", json={"devolver_estoque": False})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] is None
    assert situacoes == [out["pedido_bling"]]


async def test_correcao_de_estoque_incerta_manda_conferir_o_extrato(
    client, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(_http_error(504)))
    r = await client.post(
        "/api/devolutions/stock-correction",
        json={"sku": "x020.sp", "condicao_produto": "Novo", "estoque_destino_sku": "x020.sp"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["action"] == "unconfirmed_correcao"
    assert "Confira o extrato" in r.json()["message"]


async def test_planilha_tem_a_coluna_entrada_no_bling(client, make_user, auth_as, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(BlingCloudflareError("status=429 cf_html=False")))
    out = (await client.post("/api/devolutions", json=_body())).json()
    r = await client.get("/api/devolutions/export.xlsx")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    cab = [c.value for c in ws[1]]
    col = cab.index("Entrada no Bling")
    pb = cab.index("Pedido Bling") if "Pedido Bling" in cab else None
    linhas = [
        [c.value for c in linha] for linha in ws.iter_rows(min_row=2)
        if pb is None or str(linha[pb].value) == out["pedido_bling"]
    ]
    assert any(str(linha[col]).startswith("NÃO entrou") for linha in linhas)


# ── Revisão 2 (15/10) ─────────────────────────────────────────────────────


async def test_linha_do_backfill_nao_trava_nem_destrava_o_pedido():
    antiga = _r("Novo", "sem_confirmacao", tentou=False)
    # Fica fora da conta: o pedido volta a Resolvido quando a irmã resolve.
    assert dsr._order_situacao_target([antiga, _r("Novo", "ok")]) == dsr.SITUACAO_RESOLVIDO
    assert dsr._order_situacao_target(
        [antiga, SimpleNamespace(condicao_produto="Manutenção", manutencao_destino=None,
                                 estoque_status=None, estoque_mov_bling_id=None,
                                 estoque_tentativa_at=None)]
    ) == dsr.SITUACAO_MANUTENCAO


async def test_devolucao_antiga_reprocessada_que_falha_nao_reabre_o_pedido():
    velha = _r("Novo", "falhou")
    velha.created_at = datetime.now(UTC) - timedelta(days=40)
    # A entrada foi pedida agora (Manutenção que voltou), mas a devolução é
    # antiga: não vai para Resolvido e também não volta para Aguardando Devolução.
    velha.data_devolvido_estoque = datetime.now(UTC)
    assert dsr._order_situacao_target([velha]) is None
    nova = _r("Novo", "falhou")
    nova.created_at = datetime.now(UTC) - timedelta(hours=2)
    assert dsr._order_situacao_target([nova]) == dsr.SITUACAO_AGUARDANDO_DEVOLUCAO


async def test_movimento_estornado_nao_conta_como_ok():
    row = SimpleNamespace(estoque_status=None, estoque_mov_bling_id=1,
                          estoque_mov_revertido_at=datetime.now(UTC))
    assert dsr.status_estoque(row) is None


async def test_estorno_incerto_fica_sem_confirmacao_e_ninguem_estorna_de_novo_sem_conferir(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "ok"

    # Troca para Usado: o estorno (S) leva 504 — pode ter saído.
    fake.falhas = [_http_error(504)]
    r = await client.patch(
        f"/api/devolutions/{out['id']}",
        json={"condicao_produto": "Usado", "estoque_destino_sku": "x020.us"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "sem_confirmacao"
    assert r.json()["bling_stock_result"]["action"] == "reverse_incerto"
    assert [e["op"] for e in fake.entradas] == ["E", "S"]

    # Excluir não estorna de novo enquanto ninguém conferiu.
    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "estoque_estorno_sem_confirmacao"
    # Tentar de novo exige conferir o extrato.
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 422
    assert [e["op"] for e in fake.entradas] == ["E", "S"]

    # A pessoa viu a saída no extrato: "Estorno já saiu" — falta a entrada nova.
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/ja-entrou", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "falhou"
    linha = await _linha(db, out["id"])
    assert linha.estoque_mov_revertido_at is not None
    assert [e["op"] for e in fake.entradas] == ["E", "S"]  # nada lançado

    # Tentar de novo agora só lança a entrada nova (sem estornar de novo).
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"
    assert [e["op"] for e in fake.entradas] == ["E", "S", "E"]


async def test_excluir_linha_pendente_recalcula_o_pedido(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    situacoes = _wire(monkeypatch, _FakeBling(BlingCloudflareError("status=429 cf_html=False")))
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"
    situacoes.clear()
    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 200, r.text
    assert situacoes == [out["pedido_bling"]]


async def test_patch_que_lanca_deixa_a_linha_conferindo_durante_o_bling(
    client, db, make_user, auth_as, monkeypatch
):
    """Durante a chamada ao Bling a linha já está "conferindo" (gravado): um
    Tentar de novo ou outro save no meio recebe 409, não lança 2×."""
    from app.db import SessionLocal

    auth_as(await make_user(permissions=_perms()))
    vistos: list[str | None] = []

    class _Espia(_FakeBling):
        alvo = None

        async def update_stock_by_id(self, pid, **kw):
            if self.alvo is None:
                return await super().update_stock_by_id(pid, **kw)
            async with SessionLocal() as outra:
                st = (
                    await outra.execute(
                        select(Devolution.estoque_status).where(Devolution.id == self.alvo)
                    )
                ).scalar_one()
            vistos.append(st)
            return await super().update_stock_by_id(pid, **kw)

    fake = _Espia(BlingCloudflareError("status=429 cf_html=False"))
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    fake.alvo = out["id"]
    vistos.clear()
    r = await client.patch(
        f"/api/devolutions/{out['id']}",
        json={"condicao_produto": "Usado", "estoque_destino_sku": "x020.us"},
    )
    assert r.status_code == 200, r.text
    assert vistos == ["conferindo"]
    assert r.json()["estoque_status"] == "ok"


async def test_ja_entrou_em_sku_diferente_do_enviado_pede_confirmacao(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling())

    async def _pid_local(_s, sku):
        return {"a001.pi": 16376311221, "a001.ci": 16376311000}.get(sku)

    monkeypatch.setattr(devolutions_router, "_local_bling_product_id", _pid_local)
    row = Devolution(
        pedido_bling="300334", conta="Shopee ATV", sku="a001.sp", condicao_produto="Novo",
        quantidade=1, devolver_estoque=True, estoque_destino_sku="a001.pi",
        data_devolvido_estoque=datetime.now(UTC), estoque_status="sem_confirmacao",
        estoque_tentativa_at=datetime.now(UTC), estoque_conferir_bling_id=16376311221,
        estoque_conferir_sku="a001.pi",
    )
    db.add(row)
    await db.commit()
    r = await client.post(f"/api/devolutions/{row.id}/estoque/ja-entrou", json={"sku": "a001.ci"})
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "estoque_sku_diferente"
    r = await client.post(
        f"/api/devolutions/{row.id}/estoque/ja-entrou",
        json={"sku": "a001.ci", "confirmo_outro_sku": True},
    )
    assert r.status_code == 200, r.text
    assert (await _linha(db, row.id)).estoque_mov_sku == "a001.ci"


# ── Revisão 3 (15/10) ─────────────────────────────────────────────────────


async def test_save_qualquer_nao_apaga_o_estorno_incerto(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    fake.falhas = [_http_error(504)]
    r = await client.patch(f"/api/devolutions/{out['id']}", json={"devolver_estoque": False})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "sem_confirmacao"
    # Editar a observação não pode apagar a dúvida do estorno.
    r = await client.patch(f"/api/devolutions/{out['id']}", json={"observacao": "olhei"})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "sem_confirmacao"
    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 409
    assert [e["op"] for e in fake.entradas] == ["E", "S"]


async def test_estorno_incerto_da_exclusao_conferido_so_libera_excluir(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling()
    situacoes = _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    fake.falhas = [_http_error(504)]
    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 502
    linha = await _linha(db, out["id"])
    assert linha.estoque_status == "sem_confirmacao"
    situacoes.clear()

    # "Estorno já saiu": não pede entrada nova e não mexe no pedido.
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/ja-entrou", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] is None
    assert situacoes == []
    # Excluir de novo: sem outra baixa.
    r = await client.delete(f"/api/devolutions/{out['id']}")
    assert r.status_code == 200, r.text
    assert [e["op"] for e in fake.entradas] == ["E", "S"]


# ── Checagem ampla (19/10) ────────────────────────────────────────────────


async def test_trocado_irmao_nao_resolve_com_pendencia_antiga():
    velha = _r("Novo", "falhou")
    velha.created_at = datetime.now(UTC) - timedelta(days=10)
    assert dsr._order_situacao_target([_r("Trocado", "ok"), velha]) is None


async def test_correcao_de_estoque_minutos_depois_nao_confirma_a_devolucao(db, monkeypatch):
    _wire(monkeypatch, _FakeBling())
    t0 = datetime.now(UTC) - timedelta(minutes=6)
    pid = _pid_novo()
    row = await _conferindo(db, t0=t0, pid=pid)
    # Uma entrada do mesmo SKU 5 min depois (correção à mão): não é desta tentativa.
    await _entrada(db, quando=t0 + timedelta(minutes=5), pid=pid)
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 0
    assert (await _linha(db, row.id)).estoque_status == "conferindo"


async def test_segunda_devolucao_do_mesmo_sku_segundos_depois_confirma(db, monkeypatch):
    _wire(monkeypatch, _FakeBling())
    t_a = datetime.now(UTC) - timedelta(minutes=1)
    pid = _pid_novo()
    db.add(
        Devolution(
            pedido_bling="301535", conta="Shopee ATV", sku="a001.sp", condicao_produto="Novo",
            quantidade=1, devolver_estoque=True, data_devolvido_estoque=t_a,
            estoque_status="ok", estoque_tentativa_at=t_a,
            estoque_mov_bling_id=pid, estoque_mov_sku="a001.pi", estoque_mov_qty=1,
        )
    )
    await db.commit()
    await _entrada(db, quando=t_a + timedelta(seconds=1), pid=pid)  # a da 1ª
    t_b = t_a + timedelta(seconds=10)
    row = await _conferindo(db, t0=t_b, pid=pid)
    await _entrada(db, quando=t_b + timedelta(seconds=1), pid=pid)  # a da 2ª
    res = await dsr.conferir_entradas_pendentes(db)
    assert res["confirmadas"] == 1
    assert (await _linha(db, row.id)).estoque_status == "ok"


# ── Tentativa automática do 429 (21/10, caso 297390) ──────────────────────


async def test_recusa_por_limite_fica_para_o_worker_tentar_de_novo(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    fake = _FakeBling(BlingCloudflareError("status=429 cf_html=False"))
    situacoes = _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"
    assert out["estoque_conferir_action"] == "auto_retry"
    assert "tenta de novo sozinho" in out["estoque_erro"]

    # Antes de 90 s o worker não mexe.
    res = await dsr.retentar_entradas_recusadas(db)
    assert res["tentadas"] == 0
    # Depois, tenta de novo sozinho — e agora o Bling aceita.
    situacoes.clear()
    res = await dsr.retentar_entradas_recusadas(db, agora=datetime.now(UTC) + timedelta(minutes=2))
    assert res == {"tentadas": 1, "entraram": 1}
    row = await _linha(db, out["id"])
    assert row.estoque_status == "ok" and row.estoque_mov_bling_id == PID
    assert row.estoque_auto_tentativas == 0  # zera quando a entrada entra
    assert situacoes == [out["pedido_bling"]]
    assert len(fake.entradas) == 2


async def test_worker_desiste_depois_de_5_recusas(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    recusa = BlingCloudflareError("status=429 cf_html=False")
    fake = _FakeBling(*([recusa] * 6))
    _wire(monkeypatch, fake)
    out = (await client.post("/api/devolutions", json=_body())).json()
    futuro = datetime.now(UTC)
    for _ in range(5):
        futuro += timedelta(minutes=30)  # passa da espera, que dobra a cada vez
        await dsr.retentar_entradas_recusadas(db, agora=futuro)
    row = await _linha(db, out["id"])
    assert row.estoque_status == "falhou"
    assert row.estoque_auto_tentativas == 5
    assert row.estoque_conferir_action is None
    assert "use Tentar de novo" in row.estoque_erro
    res = await dsr.retentar_entradas_recusadas(db, agora=futuro + timedelta(hours=2))
    assert res["tentadas"] == 0
    assert len(fake.entradas) == 6
    # A pessoa clica Tentar de novo: o contador zera e entra.
    r = await client.post(f"/api/devolutions/{out['id']}/estoque/tentar-de-novo", json={})
    assert r.status_code == 200, r.text
    assert r.json()["estoque_status"] == "ok"


async def test_recusa_que_nao_e_limite_nao_tenta_sozinho(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(_http_error(400)))
    out = (await client.post("/api/devolutions", json=_body())).json()
    assert out["estoque_status"] == "falhou"
    assert out["estoque_conferir_action"] is None
    res = await dsr.retentar_entradas_recusadas(db, agora=datetime.now(UTC) + timedelta(minutes=5))
    assert res["tentadas"] == 0


# ── Revisão da tentativa automática (21/10) ───────────────────────────────


async def test_espera_dobra_a_cada_tentativa(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    recusa = BlingCloudflareError("status=429 cf_html=False")
    _wire(monkeypatch, _FakeBling(*([recusa] * 3)))
    out = (await client.post("/api/devolutions", json=_body())).json()
    t0 = (await _linha(db, out["id"])).estoque_tentativa_at
    # 1ª espera 1,5 min
    assert (await dsr.retentar_entradas_recusadas(db, agora=t0 + timedelta(seconds=80)))["tentadas"] == 0
    assert (await dsr.retentar_entradas_recusadas(db, agora=t0 + timedelta(seconds=95)))["tentadas"] == 1
    t1 = (await _linha(db, out["id"])).estoque_tentativa_at
    # 2ª espera 3 min
    assert (await dsr.retentar_entradas_recusadas(db, agora=t1 + timedelta(seconds=170)))["tentadas"] == 0
    assert (await dsr.retentar_entradas_recusadas(db, agora=t1 + timedelta(seconds=185)))["tentadas"] == 1


async def test_tempestade_para_no_primeiro_429_do_tick(client, db, make_user, auth_as, monkeypatch):
    auth_as(await make_user(permissions=_perms()))
    recusa = BlingCloudflareError("status=429 cf_html=False")
    fake = _FakeBling(*([recusa] * 10))
    _wire(monkeypatch, fake)
    ids = [(await client.post("/api/devolutions", json=_body())).json()["id"] for _ in range(3)]
    antes = len(fake.entradas)
    res = await dsr.retentar_entradas_recusadas(db, agora=datetime.now(UTC) + timedelta(minutes=2))
    assert res["tentadas"] == 1
    assert len(fake.entradas) == antes + 1
    tentativas = sorted([(await _linha(db, i)).estoque_auto_tentativas for i in ids])
    assert tentativas == [0, 0, 1]


async def test_save_de_tela_velha_nao_desfaz_o_destino_da_linha_recusada(
    client, db, make_user, auth_as, monkeypatch
):
    auth_as(await make_user(permissions=_perms()))
    _wire(monkeypatch, _FakeBling(BlingCloudflareError("status=429 cf_html=False")))
    out = (await client.post("/api/devolutions", json=_body())).json()
    db.add(await _linha(db, out["id"]))
    linha = await _linha(db, out["id"])
    linha.estoque_destino_sku, linha.estoque_nova_tag = "z0001.mala", None
    await db.commit()
    r = await client.patch(
        f"/api/devolutions/{out['id']}",
        json={"observacao": "x", "estoque_destino_sku": None, "estoque_nova_tag": "mala"},
    )
    assert r.status_code == 200, r.text
    linha = await _linha(db, out["id"])
    assert linha.estoque_destino_sku == "z0001.mala" and linha.estoque_nova_tag is None


async def test_situacao_recusada_por_limite_fica_para_o_proximo_tick(db, monkeypatch):
    from app.redis_client import redis

    chamadas: list[str] = []

    async def _situacao(_s, pedido, **_kw):
        chamadas.append(pedido)
        if len(chamadas) == 1:
            return {"ok": False, "action": "situacao_transitoria"}
        return {"ok": True, "action": "situacao_patched"}

    monkeypatch.setattr(dsr, "apply_order_situacao", _situacao)
    await redis.delete(dsr.SITUACAO_PENDENTE_KEY)
    await dsr.aplicar_situacao_ou_guardar(db, "300999")
    assert await redis.sismember(dsr.SITUACAO_PENDENTE_KEY, "300999")
    assert await dsr.reaplicar_situacoes_pendentes(db) == 1
    assert not await redis.sismember(dsr.SITUACAO_PENDENTE_KEY, "300999")
    assert chamadas == ["300999", "300999"]


async def test_situacao_com_429_nao_faz_o_desvio():
    assert dsr._situacao_transitoria(BlingCloudflareError("status=429 cf_html=False"))
    assert dsr._situacao_transitoria(_http_error(502))
    assert not dsr._situacao_transitoria(_http_error(400))
    assert dsr._falha_sem_envio(RuntimeError("bling_cf_cooldown_active ttl_s=30"))
