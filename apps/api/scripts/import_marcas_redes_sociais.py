"""Importa a planilha `redes sociais.xlsx` em marcas + redes_sociais (+ o
padrão de e-mail SAC de cada marca).

Eduardo, 15/09/2026: as abas Cadastros › Marcas, Redes Sociais e E-mails
nascem da planilha que a operação mantinha à mão (com senha em texto puro —
por isso ela fica FORA do git, em data/seed/, e este script NUNCA imprime
senha).

Uso:
    uv run python -m scripts.import_marcas_redes_sociais "<xlsx>" [--dry-run]

Forma da planilha (as tabelas são achadas pelo CONTEÚDO do cabeçalho —
lower/strip/sem acento + apelidos tipo 'intagram' → instagram —, nunca por
número de linha fixo):

  • aba `marcas`: uma linha por marca — marca, inpi, usuario, senha, email,
    dominio br, dominio, dono dominio, validade, atuação, email, obs.
    'atuação' vira `classe`; 'validade' é do DOMÍNIO (`dominio_validade`);
    a 2ª coluna 'email' é um lembrete igual em todas as linhas ("colocar
    todos os email disponíveis") — NÃO vai pra obs, sai uma vez no resumo.
  • aba `r.social`: DUAS tabelas na mesma aba. A PIVOT (linha = marca na
    coluna A; fone, usuario — que na prática é o e-mail sac@… —, senha, um
    @ por plataforma, função, obs) e, depois de linhas em branco, a
    NORMALIZADA (Plataforma, Conta, E-mail, Fone, Senha, tipo, obs), que
    NÃO tem coluna de marca: cada linha casa com a pivot por e-mail, senão
    por fone, senão pelo @ daquela plataforma.

O que vai pra onde ("seguir bem a planilha", Eduardo 15/09/2026):

  • A linha da PIVOT é a MARCA: fone / usuario / senha viram
    marcas.sac_fone (só dígitos) / sac_email / sac_senha_enc — UMA credencial
    por marca, compartilhada pelas redes e pelo SAC. Célula da pivot vazia:
    se TODAS as linhas normalizadas da marca concordam num valor, ele é
    promovido pra marca (backfill); se divergem entre si, cada conta guarda
    o seu e sai um aviso pra alguém preencher a pivot.
  • Por (marca, plataforma) sai UMA linha em redes_sociais — união das
    colunas da pivot com @ preenchido e das plataformas que têm linha
    normalizada. O @ (`conta`) vem da pivot. e-mail/fone/senha da conta são
    OVERRIDE: só gravados quando DIFEREM do valor da marca (senão NULL =
    herda; facebook, que só existe na pivot, herda tudo). `usuario` guarda a
    'Conta' da normalizada quando difere do @ (instagram: login 'locagil' ×
    @locagiloficial). função, tipo (1º não vazio da normalizada) e obs da
    pivot são por MARCA — ficam em marcas.funcao/tipo/obs.
  • site: quando a marca não tem, "https://" + 1º domínio de dominio_br
    (a célula pode listar vários, separados por vírgula/ponto-e-vírgula/
    espaço) — só se passar na validação de URL.
  • Padrão de e-mail SAC: toda marca SEM nenhum padrão de contexto 'sac'
    ganha "Resposta padrão SAC" (assunto/corpo genéricos; o operador ajusta
    na aba E-mails). Pré-checado em memória — a 2ª execução não recria.
  • Verificação (verificacao_status da conta, whatsapp_verificacao_* da
    marca) e `ativo` são da tela: a importação nunca toca.

Idempotente: marca por slug (_slugify do nome), conta por (marca,
plataforma). Na reexecução só grava célula preenchida (nunca apaga o que
já está no banco) e a senha só é recifrada quando o texto mudou. Conta cujo
@ no banco difere do da planilha é OUTRA conta: aviso e nada gravado. obs é
anexada com dedupe (lower/strip/espaços colapsados). --dry-run mostra o
plano (marca / sac: fone|e-mail|senha sim|não; marca / plataforma / conta /
senha: sim (própria|herda da marca)|não; padrão SAC a criar) e não grava
nada. Erro numa linha não derruba a importação: sai no resumo só com o nº
da linha e os campos.
"""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, cast
from uuid import UUID, uuid4

import openpyxl
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import session_scope
from app.models import (
    MARCA_INPI_STATUS,
    REDES_SOCIAIS_PLATAFORMAS,
    Marca,
    MarcaEmailPadrao,
    RedeSocial,
)
from app.schemas.marcas import (
    InpiStatus,
    MarcaCreate,
    Plataforma,
    RedeSocialCreate,
    _digitos,
    _email,
    _handle,
    _site,
    _texto,
)
from app.schemas.segments import _slugify
from app.security.cipher import decrypt, encrypt

# Cabeçalho normalizado (lower/strip/sem acento) → nome interno.
_APELIDOS_CABECALHO = {
    "intagram": "instagram",  # (sic) na planilha
    "atuacao": "classe",  # Eduardo pediu pra renomear
    "dominio br": "dominio_br",
    "dono dominio": "dono_dominio",
    "e-mail": "email",
    "funcao": "funcao",
}

# Campos da marca gravados "só quando a célula está preenchida".
_CAMPOS_MARCA = (
    "nome",
    "usuario",
    "email",
    "dominio_br",
    "dominio",
    "dono_dominio",
    "dominio_validade",
    "classe",
)
_CAMPOS_REDE = ("usuario", "email", "fone")
# Credencial que a pivot dá por MARCA (sac_*) e a normalizada por conta
# (override). Rótulo só pros avisos/plano.
_CAMPOS_SAC = ("fone", "email", "senha")
_ROTULO_SAC = {"fone": "fone", "email": "e-mail", "senha": "senha"}

# Padrão de e-mail criado pra toda marca sem NENHUM padrão de contexto 'sac'
# (SPEC v3.2, Eduardo 15/09/2026). Nome fixo; o operador edita na aba
# E-mails. Só placeholders conhecidos (services/email_marca.PLACEHOLDERS).
PADRAO_SAC_CONTEXTO = "sac"
PADRAO_SAC_NOME = "Resposta padrão SAC"
PADRAO_SAC_ASSUNTO = "{{ marca }} — atendimento"
PADRAO_SAC_CORPO = (
    "Olá {{ cliente }},\n\n"
    "Recebemos sua mensagem e já estamos cuidando do seu atendimento.\n\n"
    "Qualquer dúvida, fale com a gente pelo WhatsApp {{ whatsapp }} ou responda este e-mail.\n\n"
    "Atenciosamente,\nEquipe {{ marca }}"
)


class LinhaInvalidaError(ValueError):
    """Linha que não dá pra importar. A mensagem carrega SÓ nomes de campo
    (vai pro resumo — nunca o valor da célula)."""


# ================================================================== resumo


@dataclass
class Resumo:
    dry_run: bool = False
    marcas_criadas: int = 0
    marcas_atualizadas: int = 0
    redes_criadas: int = 0
    redes_atualizadas: int = 0
    # Padrões de e-mail SAC criados (marcas que não tinham nenhum).
    padroes_criados: int = 0
    # Linhas da normalizada que não casaram com nenhuma marca da pivot.
    linhas_sem_marca: int = 0
    avisos: list[str] = field(default_factory=list)
    # Valores distintos da 2ª coluna 'email' da aba marcas (lembrete).
    lembretes: list[str] = field(default_factory=list)
    # Linhas do plano (dry-run): "marca X / senha: sim|não", "marca X / sac:
    # fone sim|não, e-mail sim|não, senha sim|não", "marca / plataforma /
    # conta / senha: sim (própria|herda da marca)|não" e "marca X / padrão
    # de e-mail SAC: criar". NUNCA o valor de uma senha.
    plano: list[str] = field(default_factory=list)

    def aviso(self, msg: str) -> None:
        self.avisos.append(msg)

    def imprimir(self) -> None:
        if self.dry_run:
            print("DRY-RUN — nada gravado. Plano:")
            for linha in self.plano:
                print(f"  {linha}")
        print(
            f"marcas: criadas={self.marcas_criadas} atualizadas={self.marcas_atualizadas} | "
            f"redes: criadas={self.redes_criadas} atualizadas={self.redes_atualizadas} | "
            f"padroes: criados={self.padroes_criados} | "
            f"avisos={len(self.avisos)}"
        )
        for a in self.avisos:
            print(f"  aviso: {a}")
        for lembrete in self.lembretes:
            print(f"lembrete da planilha: {lembrete}")


# ============================================================ normalização


def _sem_acento(s: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch))


def _cabecalho(v: Any) -> str:
    """Nome interno da coluna: lower/strip/sem acento/espaços colapsados +
    apelidos. Também serve pro valor da coluna 'Plataforma' da normalizada."""
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", _sem_acento(str(v)).strip().lower())
    return _APELIDOS_CABECALHO.get(s, s)


def _canon(s: str | None) -> str:
    """Chave de dedupe de texto livre: lower/strip/espaços colapsados."""
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _txt(v: Any) -> str | None:
    """Texto livre da planilha: strip; vazio ou '?' (= não sabe) vira None."""
    s = _texto(v)
    return None if s == "?" else s


def _mail(v: Any) -> str | None:
    s = _email(v)
    return None if s == "?" else s


def _fone(v: Any) -> str | None:
    """Fone int/str/float → só dígitos (célula numérica pode vir float)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _digitos(v)


def _senha_planilha(v: Any) -> str | None:
    """Senha: vazio/'?' → None. Sem strip (schemas/marcas._senha: espaço pode
    ser parte da senha) — só descarta célula que é SÓ espaço."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v)
    if not s.strip() or s.strip() == "?":
        return None
    return s


def _data(v: Any) -> date | None:
    """Validade: datetime (planilha) → date; aceita texto ISO ou dd/mm/aaaa."""
    if v is None or v == "" or v == "?":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise LinhaInvalidaError("validade")


def _site_do_dominio(dominio_br: str | None) -> str | None:
    """Site da marca = "https://" + 1º domínio de dominio_br (a célula pode
    listar vários: "a.com.br, b.com.br"). Domínio nu vai em minúsculas; se
    alguém já escreveu a URL inteira, fica como está. None se não há ou não
    valida (schemas/marcas._site — a mesma regra do PATCH da marca)."""
    tokens = [t for t in re.split(r"[,;\s]+", dominio_br or "") if t and t != "?"]
    if not tokens:
        return None
    dom = tokens[0]
    url = dom if dom.lower().startswith(("http://", "https://")) else f"https://{dom.lower()}"
    try:
        return _site(url)
    except ValueError:
        return None


def _mapeia_inpi(raw: str | None) -> tuple[InpiStatus | None, str | None]:
    """'ok' → registrado, 'aguardando' → aguardando, valor já do enum → ele
    mesmo; vazio → None (create usa o default nao_registrado; na reexecução
    não mexe); outro → nao_registrado + nota pra obs."""
    if raw is None:
        return None, None
    v = _sem_acento(raw).strip().lower()
    if v == "ok":
        return "registrado", None
    if v == "aguardando":
        return "aguardando", None
    if v in MARCA_INPI_STATUS:
        return cast(InpiStatus, v), None
    return "nao_registrado", f"inpi na planilha: {raw}"


def _campos_do_erro(e: ValidationError) -> str:
    # include_input=False: o resumo nunca pode carregar o valor da célula.
    campos = sorted({".".join(str(p) for p in err["loc"]) for err in e.errors(include_input=False)})
    return ", ".join(campos) or "?"


def _sn(tem: bool) -> str:
    return "sim" if tem else "não"


# ================================================================= tabelas


@dataclass
class Tabela:
    aba: str
    linha_cabecalho: int
    # nome interno → índices (0-based) das colunas; 'email' aparece 2x na
    # aba marcas (a 2ª é o lembrete).
    colunas: dict[str, list[int]]
    # (nº da linha na planilha, valores)
    linhas: list[tuple[int, list[Any]]]

    def valor(self, linha: list[Any], nome: str, n: int = 0) -> Any:
        idx = self.colunas.get(nome)
        if not idx or len(idx) <= n or idx[n] >= len(linha):
            return None
        return linha[idx[n]]


def _em_branco(row: list[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def _tipo_de_tabela(heads: list[str]) -> str | None:
    hs = {h for h in heads if h}
    if {"marca", "inpi"} <= hs:
        return "marcas"
    if {"fone", "usuario", "senha"} <= hs and hs & set(REDES_SOCIAIS_PLATAFORMAS):
        return "pivot"
    if heads and heads[0] == "plataforma" and "conta" in hs:
        return "normalizada"
    return None


def _acha_tabelas(wb: openpyxl.Workbook) -> dict[str, Tabela]:
    """Varre todas as abas procurando os três cabeçalhos pelo conteúdo. A
    tabela vai até a 1ª linha em branco (a normalizada, última da aba, vai
    até o fim pulando brancos). Fica a 1ª ocorrência de cada tipo."""
    achadas: dict[str, Tabela] = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        i = 0
        while i < len(rows):
            heads = [_cabecalho(v) for v in rows[i]]
            tipo = _tipo_de_tabela(heads)
            if tipo is None or tipo in achadas:
                i += 1
                continue
            colunas: dict[str, list[int]] = {}
            for j, h in enumerate(heads):
                if h:
                    colunas.setdefault(h, []).append(j)
            linhas: list[tuple[int, list[Any]]] = []
            j = i + 1
            while j < len(rows):
                if _em_branco(rows[j]):
                    if tipo != "normalizada":
                        break
                elif tipo == "normalizada" and _tipo_de_tabela([_cabecalho(v) for v in rows[j]]):
                    break  # outra tabela começou
                else:
                    linhas.append((j + 1, rows[j]))
                j += 1
            achadas[tipo] = Tabela(ws.title, i + 1, colunas, linhas)
            i = j
    return achadas


# ============================================================ linhas lidas


@dataclass
class _Pivot:
    linha: int
    nome: str
    slug: str
    email: str | None  # coluna 'usuario' da pivot (é o sac@…)
    fone: str | None
    senha: str | None
    handles: dict[str, str | None]  # plataforma → @ (None se vazio/'?')
    funcao: str | None
    obs: str | None


@dataclass
class _Normal:
    linha: int
    plataforma: str
    conta: str | None
    email: str | None
    fone: str | None
    senha: str | None
    tipo: str | None
    obs: str | None


@dataclass
class _Sac:
    """fone/e-mail/senha EFETIVOS da marca depois desta rodada (pivot →
    backfill → o que já estava no banco). Cada conta compara os seus com
    estes: igual → NULL (herda); diferente → override. `senha` é o texto em
    claro só pra comparar — nunca vai pro resumo/plano."""

    fone: str | None = None
    email: str | None = None
    senha: str | None = None


# =============================================================== importador


class _Importador:
    def __init__(self, session: AsyncSession, resumo: Resumo) -> None:
        self.session = session
        self.resumo = resumo
        self.marcas: dict[str, Marca] = {}  # slug → marca (banco + novas)
        self.redes: dict[UUID, list[RedeSocial]] = {}  # marca_id → contas
        # (plataforma, lower(conta)) → marca_id: espelha o índice único
        # uq_redes_sociais_plataforma_conta pra avisar em vez de estourar.
        self.contas_em_uso: dict[tuple[str, str], UUID] = {}
        # marca_id com ALGUM padrão de e-mail no contexto 'sac' (banco +
        # criados nesta rodada): pré-checagem em memória, nunca IntegrityError.
        self.marcas_com_padrao_sac: set[UUID] = set()
        self.marcas_novas: set[str] = set()
        self.marcas_alteradas: set[str] = set()

    async def carregar(self) -> None:
        for m in (await self.session.execute(select(Marca))).scalars():
            self.marcas[m.slug] = m
        for r in (await self.session.execute(select(RedeSocial))).scalars():
            self.redes.setdefault(r.marca_id, []).append(r)
            if r.conta:
                self.contas_em_uso[(r.plataforma, r.conta.lower())] = r.marca_id
        stmt = select(MarcaEmailPadrao.marca_id).where(
            MarcaEmailPadrao.contexto == PADRAO_SAC_CONTEXTO
        )
        for marca_id in (await self.session.execute(stmt)).scalars():
            self.marcas_com_padrao_sac.add(marca_id)

    # ------------------------------------------------------------ helpers

    def _marca(self, nome: str, slug: str) -> Marca:
        """Marca por slug; cria (nome/slug) se não existe. Não renomeia —
        o nome de exibição vem da aba marcas, a pivot só garante existir."""
        m = self.marcas.get(slug)
        if m is None:
            m = Marca(id=uuid4(), nome=nome, slug=slug, inpi_status="nao_registrado")
            self.session.add(m)
            self.marcas[slug] = m
            self.marcas_novas.add(slug)
        return m

    def _marca_mudou(self, m: Marca, mudou: bool) -> None:
        if mudou and m.slug not in self.marcas_novas:
            self.marcas_alteradas.add(m.slug)

    @staticmethod
    def _grava(obj: Any, campo: str, valor: Any) -> bool:
        """Só valor preenchido e diferente do atual — nunca apaga."""
        if valor is None or getattr(obj, campo) == valor:
            return False
        setattr(obj, campo, valor)
        return True

    @staticmethod
    def _grava_senha(obj: Any, senha: str | None, campo: str = "senha_enc") -> bool:
        """Cifra só quando há senha e ela mudou (a cifra tem nonce aleatório:
        recifrar a mesma senha viraria 'atualizada' em toda reexecução).
        `campo`: senha_enc (registro/conta) ou sac_senha_enc (marca)."""
        if not senha:
            return False
        atual = getattr(obj, campo)
        if atual:
            # Chave rotacionada/blob ruim: não compara, recifra.
            with contextlib.suppress(Exception):
                if decrypt(atual) == senha:
                    return False
        setattr(obj, campo, encrypt(senha))
        return True

    @staticmethod
    def _decifra(valor_enc: str | None) -> str | None:
        """Senha do banco em claro SÓ pra comparar com a da planilha (nunca
        sai daqui pro resumo). None se não há ou se a chave não abre o blob."""
        if not valor_enc:
            return None
        with contextlib.suppress(Exception):
            return decrypt(valor_enc)
        return None

    @staticmethod
    def _tem_linha(obs: str | None, texto: str) -> bool:
        alvo = _canon(texto)
        return any(_canon(linha) == alvo for linha in (obs or "").splitlines())

    def _anexa_obs(self, obj: Any, texto: str | None) -> bool:
        """Anexa uma linha à obs se ela ainda não está lá (dedupe canônico)."""
        if not texto or self._tem_linha(obj.obs, texto):
            return False
        obj.obs = f"{obj.obs}\n{texto}" if obj.obs else texto
        return True

    def _handle_ou_aviso(self, v: Any, *, linha: int, campo: str) -> str | None:
        try:
            return _handle(v)
        except ValueError:
            self.resumo.aviso(f"aba r.social linha {linha}: {campo} inválido — ignorado")
            return None

    # -------------------------------------------------------- aba marcas

    def importar_marcas(self, tab: Tabela) -> None:
        for linha, row in tab.linhas:
            try:
                self._importa_marca(tab, linha, row)
            except ValidationError as e:
                self.resumo.aviso(
                    f"aba {tab.aba} linha {linha}: inválida ({_campos_do_erro(e)}) — ignorada"
                )
            except LinhaInvalidaError as e:
                self.resumo.aviso(f"aba {tab.aba} linha {linha}: inválida ({e}) — ignorada")
            except Exception as e:
                # Só o tipo do erro: a mensagem poderia carregar o valor da célula.
                self.resumo.aviso(f"aba {tab.aba} linha {linha}: {type(e).__name__} — ignorada")

    def _importa_marca(self, tab: Tabela, linha: int, row: list[Any]) -> None:
        nome = _txt(tab.valor(row, "marca"))
        if not nome:
            raise LinhaInvalidaError("marca")
        slug = _slugify(nome)
        if not slug:
            raise LinhaInvalidaError("marca (slug)")
        inpi, nota_inpi = _mapeia_inpi(_txt(tab.valor(row, "inpi")))
        # 2ª coluna 'email' = lembrete igual em todas as linhas: não é dado
        # da marca — vai UMA vez pro resumo.
        lembrete = _txt(tab.valor(row, "email", 1))
        if lembrete and lembrete not in self.resumo.lembretes:
            self.resumo.lembretes.append(lembrete)
        dados = MarcaCreate(
            nome=nome,
            slug=slug,
            inpi_status=inpi or "nao_registrado",
            usuario=_txt(tab.valor(row, "usuario")),
            senha=_senha_planilha(tab.valor(row, "senha")),
            email=_mail(tab.valor(row, "email", 0)),
            dominio_br=_txt(tab.valor(row, "dominio_br")),
            dominio=_txt(tab.valor(row, "dominio")),
            dono_dominio=_txt(tab.valor(row, "dono_dominio")),
            dominio_validade=_data(tab.valor(row, "validade")),
            classe=_txt(tab.valor(row, "classe")),
            obs=_txt(tab.valor(row, "obs")),
        )
        m = self._marca(dados.nome, dados.slug or slug)
        mudou = False
        for campo in _CAMPOS_MARCA:
            mudou |= self._grava(m, campo, getattr(dados, campo))
        if inpi:
            mudou |= self._grava(m, "inpi_status", inpi)
        mudou |= self._anexa_obs(m, dados.obs)
        mudou |= self._anexa_obs(m, nota_inpi)
        mudou |= self._grava_senha(m, dados.senha)
        if not m.site:
            # Assinatura dos e-mails precisa do site: deriva do 1º domínio de
            # dominio_br (o que já está no banco conta). Nunca sobrescreve um
            # site posto na tela.
            mudou |= self._grava(m, "site", _site_do_dominio(m.dominio_br))
        self._marca_mudou(m, mudou)
        self.resumo.plano.append(f"marca {m.nome} / senha: {_sn(bool(dados.senha))}")

    # -------------------------------------------------------- aba r.social

    def ler_pivot(self, tab: Tabela) -> list[_Pivot]:
        # Coluna A não tem cabeçalho na planilha real; se alguém nomear
        # 'marca', vale ela.
        idx_marca = tab.colunas.get("marca", [0])[0]
        plataformas = [p for p in REDES_SOCIAIS_PLATAFORMAS if p in tab.colunas]
        pivots: list[_Pivot] = []
        for linha, row in tab.linhas:
            nome = _txt(row[idx_marca] if idx_marca < len(row) else None)
            slug = _slugify(nome) if nome else ""
            if not nome or not slug:
                self.resumo.aviso(f"aba {tab.aba} linha {linha}: sem marca na coluna A — ignorada")
                continue
            pivots.append(
                _Pivot(
                    linha=linha,
                    nome=nome,
                    slug=slug,
                    email=_mail(tab.valor(row, "usuario")),
                    fone=_fone(tab.valor(row, "fone")),
                    senha=_senha_planilha(tab.valor(row, "senha")),
                    handles={
                        p: self._handle_ou_aviso(tab.valor(row, p), linha=linha, campo=p)
                        for p in plataformas
                    },
                    funcao=_txt(tab.valor(row, "funcao")),
                    obs=_txt(tab.valor(row, "obs")),
                )
            )
        return pivots

    def ler_normalizada(self, tab: Tabela) -> list[_Normal]:
        normais: list[_Normal] = []
        for linha, row in tab.linhas:
            plat = _cabecalho(tab.valor(row, "plataforma"))
            if not plat:
                self.resumo.aviso(f"aba {tab.aba} linha {linha}: sem plataforma — ignorada")
                continue
            if plat not in REDES_SOCIAIS_PLATAFORMAS:
                self.resumo.aviso(
                    f"aba {tab.aba} linha {linha}: plataforma '{plat}' desconhecida — ignorada"
                )
                continue
            try:
                conta = _handle(tab.valor(row, "conta"))
            except ValueError:
                self.resumo.aviso(f"aba {tab.aba} linha {linha}: inválida (conta) — ignorada")
                continue
            normais.append(
                _Normal(
                    linha=linha,
                    plataforma=plat,
                    conta=conta,
                    email=_mail(tab.valor(row, "email")),
                    fone=_fone(tab.valor(row, "fone")),
                    senha=_senha_planilha(tab.valor(row, "senha")),
                    tipo=_txt(tab.valor(row, "tipo")),
                    obs=_txt(tab.valor(row, "obs")),
                )
            )
        return normais

    @staticmethod
    def _casa(n: _Normal, pivots: list[_Pivot]) -> _Pivot | None:
        """Linha normalizada → marca da pivot: e-mail == usuario da pivot;
        senão fone (dígitos); senão Conta == @ daquela plataforma."""
        if n.email:
            for p in pivots:
                if p.email == n.email:
                    return p
        if n.fone:
            for p in pivots:
                if p.fone == n.fone:
                    return p
        if n.conta:
            for p in pivots:
                h = p.handles.get(n.plataforma)
                if h and h.lower() == n.conta.lower():
                    return p
        return None

    def _resolve_sac(self, m: Marca, p: _Pivot, linhas: list[_Normal]) -> tuple[_Sac, bool]:
        """fone/e-mail/senha da MARCA (sac_*) a partir da linha da pivot.
        Célula vazia: se TODAS as linhas normalizadas da marca que TÊM o
        valor concordam (linha sem o valor não vota), ele sobe pra marca
        (backfill); se divergem entre si, fica por conta e sai aviso.
        Devolve os valores efetivos (o que a marca tem depois desta rodada,
        inclusive o que já estava no banco) pras contas decidirem entre
        override e herança — e se a marca mudou."""
        sac = _Sac()
        mudou = False
        da_planilha: dict[str, bool] = {}
        for campo in _CAMPOS_SAC:
            valor = getattr(p, campo)
            distintos = {getattr(n, campo) for n in linhas if getattr(n, campo)}
            if valor is None and len(distintos) == 1:
                valor = distintos.pop()
            elif valor is None and len(distintos) > 1:
                self.resumo.aviso(
                    f"{m.nome}: {_ROTULO_SAC[campo]} vazio na pivot e diferente entre as "
                    f"linhas normalizadas — fica por conta"
                )
            da_planilha[campo] = valor is not None
            if campo == "senha":
                mudou |= self._grava_senha(m, valor, "sac_senha_enc")
                sac.senha = valor or self._decifra(m.sac_senha_enc)
            else:
                mudou |= self._grava(m, f"sac_{campo}", valor)
                setattr(sac, campo, getattr(m, f"sac_{campo}"))
        self.resumo.plano.append(
            f"marca {m.nome} / sac: fone {_sn(da_planilha['fone'])}, "
            f"e-mail {_sn(da_planilha['email'])}, senha {_sn(da_planilha['senha'])}"
        )
        return sac, mudou

    def importar_redes(self, tab_pivot: Tabela, tab_normal: Tabela | None) -> None:
        pivots = self.ler_pivot(tab_pivot)
        normais = self.ler_normalizada(tab_normal) if tab_normal else []
        aba_normal = tab_normal.aba if tab_normal else tab_pivot.aba

        por_marca: dict[str, list[_Normal]] = {}
        for n in normais:
            p = self._casa(n, pivots)
            if p is None:
                self.resumo.linhas_sem_marca += 1
                self.resumo.aviso(f"aba {aba_normal} linha {n.linha} sem marca — ignorada")
                continue
            por_marca.setdefault(p.slug, []).append(n)

        for p in pivots:
            m = self._marca(p.nome, p.slug)
            linhas_marca = por_marca.get(p.slug, [])
            # função e obs são da linha da marca na pivot; tipo é o 1º não
            # vazio entre as linhas normalizadas da marca (é igual em todas).
            mudou = self._grava(m, "funcao", p.funcao)
            mudou |= self._anexa_obs(m, p.obs)
            mudou |= self._grava(m, "tipo", next((n.tipo for n in linhas_marca if n.tipo), None))
            sac, mudou_sac = self._resolve_sac(m, p, linhas_marca)
            self._marca_mudou(m, mudou or mudou_sac)

            por_plat: dict[str, list[_Normal]] = {}
            for n in linhas_marca:
                por_plat.setdefault(n.plataforma, []).append(n)
            # União: @ preenchido na pivot ∪ plataformas com linha normalizada.
            for plat in REDES_SOCIAIS_PLATAFORMAS:
                if not p.handles.get(plat) and plat not in por_plat:
                    continue
                candidatas = por_plat.get(plat, [])
                normal = candidatas[0] if candidatas else None
                for extra in candidatas[1:]:
                    self.resumo.aviso(
                        f"aba {tab_pivot.aba} linha {extra.linha}: {plat} repetido pra marca "
                        f"{m.nome} — ignorada (vale a linha {candidatas[0].linha})"
                    )
                linha_ref = normal.linha if normal else p.linha
                try:
                    self._importa_rede(m, p, plat, normal, sac)
                except ValidationError as e:
                    self.resumo.aviso(
                        f"aba {tab_pivot.aba} linha {linha_ref} ({plat}): inválida "
                        f"({_campos_do_erro(e)}) — ignorada"
                    )
                except Exception as e:
                    # Só o tipo do erro: a mensagem poderia carregar a senha.
                    self.resumo.aviso(
                        f"aba {tab_pivot.aba} linha {linha_ref} ({plat}): "
                        f"{type(e).__name__} — ignorada"
                    )

    def _importa_rede(
        self, m: Marca, p: _Pivot, plat: str, n: _Normal | None, sac: _Sac
    ) -> None:
        conta = p.handles.get(plat)
        # 'Conta' da normalizada só interessa quando é outro texto que o @
        # (login do instagram); igual ao @ é redundante.
        usuario = None
        if n and n.conta and (conta is None or n.conta.lower() != conta.lower()):
            usuario = n.conta
        # obs da normalizada vai pra conta, a não ser que seja a mesma da
        # pivot (que já está na marca) — não repete.
        obs = None
        if n and n.obs and _canon(n.obs) != _canon(p.obs) and not self._tem_linha(m.obs, n.obs):
            obs = n.obs
        # e-mail/fone/senha da conta são OVERRIDE: só quando a linha
        # normalizada traz valor DIFERENTE do da marca (sac_*). Igual ou
        # vazio → NULL = herda (facebook, que só existe na pivot, herda tudo).
        email = n.email if n and n.email and n.email != sac.email else None
        fone = n.fone if n and n.fone and n.fone != sac.fone else None
        senha = n.senha if n and n.senha and n.senha != sac.senha else None
        # Valida com o MESMO schema do POST (strip, sem @, e-mail lower,
        # fone só dígitos, limites) — erro sai com os nomes dos campos.
        dados = RedeSocialCreate(
            marca_id=m.id,
            plataforma=cast(Plataforma, plat),
            conta=conta,
            usuario=usuario,
            email=email,
            fone=fone,
            senha=senha,
            obs=obs,
        )
        rotulo = f"@{dados.conta}" if dados.conta else "(sem conta)"
        if dados.senha:
            senha_plano = "sim (própria)"
        elif sac.senha:
            senha_plano = "sim (herda da marca)"
        else:
            senha_plano = "não"
        self.resumo.plano.append(f"{m.nome} / {plat} / {rotulo} / senha: {senha_plano}")

        existentes = [r for r in self.redes.get(m.id, []) if r.plataforma == plat]
        if len(existentes) > 1:
            # Mais de uma conta da marca na plataforma (cadastradas na tela):
            # só atualiza a que tem o mesmo @ da planilha.
            if dados.conta:
                alvo = dados.conta.lower()
                existentes = [r for r in existentes if r.conta and r.conta.lower() == alvo]
            else:
                existentes = [r for r in existentes if r.conta is None]
            if len(existentes) != 1:
                self.resumo.aviso(
                    f"{m.nome}/{plat}: mais de uma conta no banco e nenhuma casa com a "
                    f"planilha ({rotulo}) — ignorada"
                )
                return

        chave = (plat, dados.conta.lower()) if dados.conta else None
        if chave and self.contas_em_uso.get(chave, m.id) != m.id:
            self.resumo.aviso(f"{m.nome}/{plat}: {rotulo} já pertence a outra marca — ignorada")
            return

        if not existentes:
            # verificacao_status fica no default do model (nao_solicitado):
            # o andamento do selo é registrado na tela, nunca pela planilha.
            r = RedeSocial(
                id=uuid4(),
                marca_id=m.id,
                plataforma=plat,
                conta=dados.conta,
                usuario=dados.usuario,
                email=dados.email,
                fone=dados.fone,
                senha_enc=encrypt(dados.senha) if dados.senha else None,
                obs=dados.obs,
                ativo=True,
            )
            self.session.add(r)
            self.redes.setdefault(m.id, []).append(r)
            if chave:
                self.contas_em_uso[chave] = m.id
            self.resumo.redes_criadas += 1
            return

        r = existentes[0]
        mudou = False
        if dados.conta and chave:
            if r.conta is None:
                # Linha "sem conta" ganha o @ (sai do índice de placeholder).
                r.conta = dados.conta
                mudou = True
                self.contas_em_uso[chave] = m.id
            elif r.conta.lower() != dados.conta.lower():
                # Conta diferente = OUTRA conta: não sobrescreve e-mail/fone/
                # senha (a equipe pode ter trocado a credencial na tela e a
                # planilha traria a antiga por cima, sem ninguém ver).
                self.resumo.aviso(
                    f"{m.nome}/{plat}: @ no banco (@{r.conta}) difere da planilha "
                    f"({rotulo}) — linha mantida como está, nada gravado"
                )
                return
        # Override vazio (= herda) nunca apaga um override que está no banco
        # (_grava ignora None): a regra é não zerar nada na reexecução.
        for campo in _CAMPOS_REDE:
            mudou |= self._grava(r, campo, getattr(dados, campo))
        mudou |= self._anexa_obs(r, dados.obs)
        mudou |= self._grava_senha(r, dados.senha)
        # verificacao_status/verificacao_obs/ativo: nunca tocados na
        # reexecução (são da tela).
        if mudou:
            self.resumo.redes_atualizadas += 1

    # ------------------------------------------------------ padrão SAC

    def criar_padroes_sac(self) -> None:
        """Toda marca (banco + planilha) SEM nenhum padrão de e-mail no
        contexto 'sac' ganha o "Resposta padrão SAC". Pré-checado em memória
        (marcas_com_padrao_sac vem do carregar() e recebe os criados aqui) —
        nunca depende de IntegrityError; a 2ª execução não cria nada."""
        for m in self.marcas.values():
            if m.id in self.marcas_com_padrao_sac:
                continue
            self.session.add(
                MarcaEmailPadrao(
                    id=uuid4(),
                    marca_id=m.id,
                    contexto=PADRAO_SAC_CONTEXTO,
                    nome=PADRAO_SAC_NOME,
                    assunto=PADRAO_SAC_ASSUNTO,
                    corpo=PADRAO_SAC_CORPO,
                    incluir_logo=True,
                    incluir_assinatura=True,
                    ativo=True,
                )
            )
            self.marcas_com_padrao_sac.add(m.id)
            self.resumo.padroes_criados += 1
            self.resumo.plano.append(f"marca {m.nome} / padrão de e-mail SAC: criar")

    def fechar(self) -> None:
        self.resumo.marcas_criadas = len(self.marcas_novas)
        self.resumo.marcas_atualizadas = len(self.marcas_alteradas)


# ==================================================================== API


async def _importar(session: AsyncSession, path: str, dry_run: bool) -> Resumo:
    resumo = Resumo(dry_run=dry_run)
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        tabelas = _acha_tabelas(wb)
    finally:
        wb.close()
    if "marcas" not in tabelas:
        resumo.aviso("tabela de marcas não encontrada (cabeçalho com 'marca' e 'inpi')")
    if "pivot" not in tabelas:
        resumo.aviso("tabela pivot de r.social não encontrada (fone/usuario/senha + plataformas)")
    if "normalizada" not in tabelas:
        resumo.aviso("tabela normalizada de r.social não encontrada (coluna A = 'Plataforma')")

    imp = _Importador(session, resumo)
    await imp.carregar()
    if "marcas" in tabelas:
        imp.importar_marcas(tabelas["marcas"])
    if "pivot" in tabelas:
        imp.importar_redes(tabelas["pivot"], tabelas.get("normalizada"))
    imp.criar_padroes_sac()
    imp.fechar()

    if dry_run:
        # Nada foi flushado (as consultas rodaram antes dos adds); o
        # rollback descarta os objetos pendentes.
        await session.rollback()
        return resumo
    try:
        await session.commit()
    except IntegrityError as e:
        await session.rollback()
        # A mensagem do driver traz valores da linha — no aviso só a constraint
        # (o asyncpg real fica em e.orig.__cause__).
        causa = getattr(e.orig, "__cause__", None) or e.orig
        constraint = getattr(causa, "constraint_name", None) or "?"
        resumo.aviso(f"gravação abortada por unicidade ({constraint}); nada foi salvo")
        raise
    return resumo


async def importar(
    path: str, *, dry_run: bool = False, session: AsyncSession | None = None
) -> Resumo:
    """Importa a planilha. `session` é pros testes (usam a sessão do
    fixture); sem ela abre app.db.session_scope()."""
    if session is not None:
        return await _importar(session, path, dry_run)
    async with session_scope() as s:
        return await _importar(s, path, dry_run)


def _args(argv: list[str] | None = None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(
        prog="python -m scripts.import_marcas_redes_sociais",
        description=(
            "Importa `redes sociais.xlsx` em marcas + redes_sociais + padrão de e-mail SAC "
            "(nunca imprime senha)."
        ),
    )
    ap.add_argument("xlsx", help="caminho da planilha")
    ap.add_argument("--dry-run", action="store_true", help="mostra o plano e não grava nada")
    return ap.parse_args(argv)


async def main(argv: list[str] | None = None) -> None:
    args = _args(argv)
    resumo = await importar(args.xlsx, dry_run=args.dry_run)
    resumo.imprimir()


if __name__ == "__main__":
    asyncio.run(main())
